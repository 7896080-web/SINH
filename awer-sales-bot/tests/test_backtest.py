"""Проверка плана на прошлых месяцах: план — по данным до начала месяца, факт — 1С или ведомости, без записи в базу."""
import io
from datetime import date, timedelta

import openpyxl
import pytest
from fastapi.testclient import TestClient

from awer_sales_bot.importers.onec_sales import import_file
from awer_sales_bot.planning import backtest
from awer_sales_bot.web import create_app
from test_flow import env, make_export  # noqa: F401

AUTH = ("admin", "pw")


def test_october_plan_matches_reference_and_uses_only_past(env):  # noqa: F811
    cfg, db = env
    before = backtest.run(db, cfg, "2026-10")
    assert before.store.plan == pytest.approx(5_605_452, abs=1)  # эталон октября (SPEC, test_engine)
    assert before.source == "нет данных" and before.rows == []

    days = [date(2026, 10, 1) + timedelta(days=i) for i in range(17)]
    lines_f, docs_f = make_export(cfg.onec.inbox, days, ["Продавец04", "Админ2"])
    import_file(docs_f, db, cfg)
    import_file(lines_f, db, cfg)
    r = backtest.run(db, cfg, "2026-10")
    assert r.store.plan == pytest.approx(before.store.plan)  # продажи октября не меняют план октября
    assert r.source == "1С" and r.coverage == pytest.approx(17 / 31)
    x = next(x for x in r.rows if x.name == "Продавец04")
    assert x.plan == pytest.approx(561_764, abs=1) and x.method == "personal"
    assert x.days == 17 and x.fte == pytest.approx(17 / cfg.norm_shifts)
    assert x.performance == pytest.approx((x.sales / x.fte) / x.plan)
    assert x.plan_worked == pytest.approx(x.plan * x.fte)
    assert not db.one("SELECT 1 FROM plans") and not db.one("SELECT 1 FROM results")  # ничего не записано


def test_payroll_months_and_page(env):  # noqa: F811
    cfg, db = env
    aug, sep = backtest.run(db, cfg, "2026-08"), backtest.run(db, cfg, "2026-09")
    assert aug.source == sep.source == "ведомости" and aug.rows and sep.rows
    assert all(x.plan for x in sep.rows if x.method == "personal")
    assert 0 <= sep.done <= len(sep.rows) and sep.rates > 0

    sid = db.seller_id("Продавец04", create=False)
    db.fire(sid, date(2026, 9, 30))  # уволенный после месяца — в проверке этого месяца он есть
    assert any(x.seller_id == sid for x in backtest.run(db, cfg, "2026-09").rows)

    web = TestClient(create_app(cfg, db))
    page = web.get("/check?frm=2026-08&to=2026-10", auth=AUTH).text
    assert "Проверка плана на прошлых месяцах" in page and "Продавец04" in page
    assert "<h2>2026-10</h2>" in page and "загрузите выгрузку 1С" in page
    assert web.get("/check?frm=мусор", auth=AUTH).status_code == 200  # неверный месяц — значения по умолчанию
    x = web.get("/check.xlsx?frm=2026-08&to=2026-09", auth=AUTH)
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Магазины", "2026-08", "2026-09"] and wb["2026-09"].max_row == len(sep.rows) + 1
