"""Админка (локальный веб): настройки месяца, ручные коэффициенты и планы, продавцы, итоги/ведомость.

Запуск: awer-bot web  →  http://127.0.0.1:8080  (логин admin, пароль из ADMIN_PASSWORD)
Все расчётные значения можно перекрыть вручную; пустое поле = расчётное значение.
"""
from __future__ import annotations

import html
import io
import os
import re
import secrets
from datetime import date, datetime, timedelta
from urllib.parse import quote, urlparse

import openpyxl
from openpyxl.styles import Font
from fastapi import Depends, FastAPI, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, PlainTextResponse, RedirectResponse, StreamingResponse
from fastapi.security import HTTPBasic, HTTPBasicCredentials

from . import botcfg, content, health, names
from . import tasks as tasks_mod
from . import crews as crews_mod
from .config import Settings
from .db import DB
from .messages import pct  # вниз до целого (аудит П1)
from .planning import backtest, service
from .planning.engine import METHOD_NAMES, month_add, rate_for, threshold_amounts

def _how_note(method: str) -> str:
    """Пометка у имени в «Личных планах»: «после перерыва» или «новичок»."""
    return f' <span class=note>{METHOD_NAMES[method]}</span>' if method in ("returning", "newcomer") else ""


def qr_svg(text: str) -> str:
    """QR-код ссылки на бота — SVG прямо в странице (без картинок и интернета)."""
    import qrcode
    import qrcode.image.svg

    img = qrcode.make(text, image_factory=qrcode.image.svg.SvgPathImage, box_size=10)
    return img.to_string(encoding="unicode")


CSS = """
body{font-family:Arial,sans-serif;margin:24px;color:#1f2328;background:#fff}
h1{font-size:20px}h2{font-size:16px;margin-top:28px}
table{border-collapse:collapse;margin:8px 0}td,th{border:1px solid #d0d7de;padding:5px 8px;font-size:13px}
th{background:#eef2f7;text-align:left}td.n{text-align:right;white-space:nowrap}
input[type=text],input[type=number]{width:110px;padding:3px;background:#fffbe6;border:1px solid #c8b26b}
.calc{color:#57606a}.man{font-weight:bold;color:#9a3412}.ok{background:#e6f4ea}
.btn{padding:6px 14px;margin:6px 6px 6px 0;cursor:pointer}
nav{display:flex;flex-wrap:wrap;gap:6px;padding-bottom:10px;border-bottom:1px solid #d0d7de}
nav a,.links a{display:inline-block;padding:6px 12px;border:1px solid #d0d7de;border-radius:6px;background:#f6f8fa;
color:#1f2328;text-decoration:none;font-size:13px}
nav a:hover,.links a:hover{background:#e7ecf2}nav a.cur{background:#0969da;border-color:#0969da;color:#fff}
.links{display:flex;flex-wrap:wrap;gap:6px}.links a.act{background:#dafbe1;border-color:#4ac26b;font-weight:bold}
.note{color:#57606a;font-size:12px}
.qr svg{width:260px;height:260px}.code{font-size:22px;letter-spacing:3px}
table.grid td,table.grid th{padding:2px 3px;text-align:center}table.grid td.nm{text-align:left;white-space:nowrap}
table.grid th.we,table.grid td.we{background:#f6f8fa}.off{color:#9a3412;font-weight:bold}
table.grid td.alien{background:#eceff3}table.grid td.swap{background:#fde68a}table.grid td.sh{font-size:11px;color:#57606a}
.health{font-size:12px;color:#57606a;margin:4px 0 0}.health.bad{color:#9a3412;font-weight:bold}
"""


def _rub(x):
    return "" if x is None else f"{x:,.0f}".replace(",", " ")


def _num(s: str | None) -> float | None:
    if s is None or str(s).strip() == "":
        return None
    return float(str(s).replace(" ", "").replace("\xa0", "").replace(",", ".").rstrip("%"))


MONTH = re.compile(r"\d{4}-(0[1-9]|1[0-2])")
ROLES = ("seller", "admin", "test")  # test — тестовая карточка: сообщения как у продавца, без ведомости и средних
ROLE_NAMES = {"seller": "продавец", "admin": "админ", "test": "тестовый"}
MIN_PLAN = 10_000  # ручной план на полный месяц — не меньше (аудит А1)


def create_app(cfg: Settings, db: DB, runner=None, max_client=None) -> FastAPI:
    """runner — поток бота в службе (bot.Runner; None, если админка запущена отдельно: awer-bot web);
    max_client — фабрика клиента MAX по токену (тесты подменяют)."""
    app = FastAPI(title="AWER — планы продавцов", docs_url=None, redoc_url=None, openapi_url=None)  # аудит А2
    security = HTTPBasic()
    password = os.environ.get("ADMIN_PASSWORD", "")

    @app.middleware("http")
    async def guard(request: Request, call_next):
        """Аудит А2: месяц в адресе — только ГГГГ-ММ; POST — только со страниц самой админки (чужая страница в браузере
        владельца не может подключить владельца бота, сменить токен или план). /api/upload защищён своим токеном."""
        parts = request.url.path.split("/")
        if len(parts) > 2 and parts[1] in ("month", "schedule") and not MONTH.fullmatch(parts[2]):
            return PlainTextResponse("Неверный месяц в адресе: нужен ГГГГ-ММ, например 2026-11", status_code=400)
        if request.method == "POST" and request.url.path != "/api/upload":
            src = request.headers.get("origin") or request.headers.get("referer")
            if src and urlparse(src).netloc != request.headers.get("host"):
                return PlainTextResponse("Запрос пришёл не со страницы админки — отклонён", status_code=403)
        return await call_next(request)

    def auth(c: HTTPBasicCredentials = Depends(security)):
        if not password:
            raise HTTPException(500, "Задайте ADMIN_PASSWORD в .env")
        if not (secrets.compare_digest(c.username, "admin") and secrets.compare_digest(c.password, password)):
            raise HTTPException(401, headers={"WWW-Authenticate": "Basic"})

    def page(title: str, body: str) -> HTMLResponse:
        cur = date.today().strftime("%Y-%m")
        items = [(f"/month/{cur}", "Настройки месяца"), ("/sellers", "Продавцы"), ("/bot", "Бот MAX"),
                 ("/crews", "Смены"), (f"/schedule/{cur}", "График работы"), (f"/month/{cur}/results", "Итоги и ведомость"),
                 ("/check", "Проверка плана"), ("/learn", "Обучение"), ("/tasks", "Задания")]
        here = lambda label: title.startswith(label.split()[0]) and not (  # noqa: E731
            label == "Настройки месяца" and title.startswith("Итоги"))
        nav = "<nav>" + "".join(f'<a href="{u}"{" class=cur" if here(t) else ""}>{t}</a>' for u, t in items) + "</nav>"
        title = html.escape(title)  # в заголовке бывает имя из 1С или часть адреса (аудит А2)
        return HTMLResponse(f"<!doctype html><html lang=ru><head><meta charset=utf-8><title>{title}</title>"
                            f"<style>{CSS}</style></head><body>{nav}{health.line(db, cfg)}<h1>{title}</h1>{body}</body></html>")

    @app.get("/ping")
    def ping():
        """Без пароля: установщик проверяет, что на порту отвечает именно эта программа (аудит У3)."""
        return PlainTextResponse("awer-sales-bot")

    @app.get("/", dependencies=[Depends(auth)])
    def root():
        return RedirectResponse(f"/month/{date.today():%Y-%m}")

    # ------------------------------------------------------------------ месяц
    @app.get("/month/{month}", dependencies=[Depends(auth)])
    def month_page(month: str, msg: str = ""):
        mp = service.compute(db, cfg, month)
        p, sp = mp.params, mp.store
        ms = db.one("SELECT * FROM month_settings WHERE month=?", month)
        y, mo = map(int, month.split("-"))
        last = date(y + mo // 12, mo % 12 + 1, 1) - timedelta(days=1)
        today = min(date.today() - timedelta(days=1), last)
        f = lambda v, d=3: "" if v is None else f"{v:.{d}f}"  # noqa: E731
        prev, nxt = month_add(month, -1), month_add(month, 1)
        b = [f'<p class=links><a href="/month/{prev}">← {prev}</a><a href="/month/{nxt}">{nxt} →</a></p>',
             f'<p class=off>{html.escape(msg)}</p>' if msg else "",
             f'<p class=note>План {"зафиксирован " + mp.frozen_at if mp.frozen_at else "НЕ зафиксирован"}. '
             'Ежедневные сообщения используют зафиксированный план. Жёлтые поля — ручная правка; пусто = расчёт.</p>',
             f'<form method=post action="/month/{month}/settings"><h2>Коэффициенты и ставки месяца</h2><table>',
             '<tr><th>Показатель</th><th>Расчёт</th><th>Вручную</th><th>Применяется</th></tr>',
             f'<tr><td>Коэффициент год к году (тренд, 3 мес. к прошлому году)</td><td class="n calc">{f(sp.trend_calc)}</td>'
             f'<td><input type=text name=trend_manual value="{f(mp.trend_manual) if mp.trend_manual is not None else ""}"></td>'
             f'<td class="n {"man" if mp.trend_manual is not None else ""}">{f(sp.trend)}</td></tr>',
             f'<tr><td>Шаг перевыполнения плана, %</td><td class="n calc">{cfg.params.step:.0%}</td>'
             f'<td><input type=text name=step value="{"" if not ms or ms["step"] is None else round(ms["step"]*100, 2)}"></td>'
             f'<td class=n>{p.step:.0%}</td></tr>',
             f'<tr><td>Надбавка к ставке за шаг, ₽</td><td class="n calc">{_rub(cfg.params.step_bonus)}</td>'
             f'<td><input type=text name=step_bonus value="{"" if not ms or ms["step_bonus"] is None else _rub(ms["step_bonus"])}"></td>'
             f'<td class=n>{_rub(p.step_bonus)}</td></tr>',
             f'<tr><td>Базовая ставка, ₽</td><td class="n calc">{_rub(cfg.params.base_rate)}</td>'
             f'<td><input type=text name=base_rate value="{"" if not ms or ms["base_rate"] is None else _rub(ms["base_rate"])}"></td>'
             f'<td class=n>{_rub(p.base_rate)}</td></tr>',
             f'<tr><td>Максимальная ставка (без админ. доплаты), ₽</td><td class="n calc">{_rub(cfg.params.max_rate)}</td>'
             f'<td><input type=text name=max_rate value="{"" if not ms or ms["max_rate"] is None else _rub(ms["max_rate"])}"></td>'
             f'<td class=n>{_rub(p.max_rate)}</td></tr>',
             f'<tr><td>Админ. доплата старшим, ₽ (фиксированная, схемой не меняется)</td><td class=n>{_rub(p.admin_bonus)}</td>'
             '<td class=note>config.yaml</td><td class=n>' + _rub(p.admin_bonus) + '</td></tr></table>',
             f'Комментарий: <input type=text style="width:420px" name=note value="{html.escape(ms["note"] or "") if ms else ""}">',
             '<br><button class=btn>Сохранить настройки</button></form>',
             '<h2>Магазин</h2><table>',
             f'<tr><td>Выторг этого месяца прошлого года</td><td class=n>{_rub(sp.last_year)}</td></tr>',
             f'<tr><td>Сезонность: месяц к предыдущему (прошлый год)</td><td class=n>{f(sp.seasonality, 2)}</td></tr>',
             f'<tr><td>План магазина = прошлый год × коэффициент</td><td class=n><b>{_rub(sp.plan)}</b></td></tr>',
             f'<tr><td>Средний месяц магазина за 3 последних</td><td class=n>{_rub(sp.recent_avg)}</td></tr>',
             f'<tr><td>Коэффициент месяца (план / средний месяц)</td><td class=n>{f(sp.coef, 3)}</td></tr>',
             f'<tr><td>План новичка (средний рядовой прошлого года × коэффициент)</td><td class=n>{_rub(sp.newcomer_plan)}</td></tr>',
             '</table>',
             f'<form method=post action="/month/{month}/plans"><h2>Личные планы (на полный месяц)</h2><table>',
             '<tr><th>Продавец</th><th>Роль</th><th>Обычный месяц (доля за год)</th><th>План расчёт</th><th>План вручную</th>'
             f'<th>План итог</th><th>+{p.step:.0%} → ставка</th><th>Факт с 1-го</th><th>Выполн. (за рабочий день)</th><th>Ставка сейчас</th></tr>']
        for r in mp.rows:
            st = service.status(db, cfg, r.seller_id, today) if today.strftime("%Y-%m") == month else None
            th = threshold_amounts(r.plan_final, p, 1)[1] if r.plan_final else None
            perf = st.performance if st else None
            b.append(
                f'<tr><td>{html.escape(r.name)}{_how_note(r.method)}</td><td>{ROLE_NAMES.get(r.role, "продавец")}</td>'
                f'<td class=n>{_rub(r.level)}</td><td class="n calc">{_rub(r.plan_calc)}</td>'
                f'<td><input type=text name="p{r.seller_id}" value="{_rub(r.plan_manual)}"></td>'
                f'<td class="n {"man" if r.plan_manual is not None else ""}"><b>{_rub(r.plan_final)}</b></td>'
                f'<td class=n>{_rub(th[0]) + " → " + _rub(rate_for(1, r.role, p)) if th else ""}</td>'
                f'<td class=n>{_rub(st.sales) if st else ""}</td>'
                f'<td class="n {"ok" if perf and perf >= 1 + p.step else ""}">{"" if perf is None else pct(perf)}</td>'
                f'<td class=n>{_rub(st.current_rate) if st else ""}</td></tr>')
        b += ['</table><button class=btn>Сохранить ручные планы</button><span class=note> — после фиксации ручной план '
              'применяется сразу и только к тому, кому вписан (например, новичку); остальным цели не меняются.</span></form>',
              f'<form method=post action="/month/{month}/freeze"><button class=btn>Зафиксировать план месяца</button>'
              '<span class=note> — перезаписывает зафиксированные планы; повторная фиксация в середине месяца меняет цели продавцам.</span></form>']
        return page(f"Настройки месяца {month}", "".join(b))

    @app.post("/month/{month}/settings", dependencies=[Depends(auth)])
    def save_settings(month: str, trend_manual: str = Form(""), step: str = Form(""), step_bonus: str = Form(""),
                      base_rate: str = Form(""), max_rate: str = Form(""), note: str = Form("")):
        try:
            tm, st, sb, br, mr = map(_num, (trend_manual, step, step_bonus, base_rate, max_rate))
        except ValueError:
            return _back(month, "Не сохранено: в числовом поле не число.")
        # аудит А1: «0,15» в поле шага в процентах — это 0,15 %, и продавец с 104 % получил бы максимум
        errs = [t for bad, t in (
            (tm is not None and not 0.3 <= tm <= 3, "коэффициент год к году — от 0,3 до 3 (например 0,95)"),
            (st is not None and not 1 <= st <= 100, "шаг — в процентах, от 1 до 100 (например 15, а не 0,15)"),
            (sb is not None and not 0 < sb <= 100_000, "надбавка за шаг — больше 0 и не больше 100 000 ₽"),
            (br is not None and not 1_000 <= br <= 1_000_000, "базовая ставка — от 1 000 до 1 000 000 ₽"),
            ((mr if mr is not None else cfg.params.max_rate) < (br if br is not None else cfg.params.base_rate),
             "максимальная ставка не может быть меньше базовой")) if bad]
        if errs:
            return _back(month, "Не сохранено: " + "; ".join(errs) + ".")
        service.save_month_settings(db, month, trend_manual=tm, step=st / 100 if st else None,
                                    step_bonus=sb, base_rate=br, max_rate=mr, note=note or None)
        return RedirectResponse(f"/month/{month}", status_code=303)

    def _back(month: str, msg: str) -> RedirectResponse:
        return RedirectResponse(f"/month/{month}?msg=" + quote(msg), status_code=303)

    @app.post("/month/{month}/plans", dependencies=[Depends(auth)])
    async def save_plans(month: str, request: Request):
        form = await request.form()
        bad = []
        for k, v in form.items():
            if not (k.startswith("p") and k[1:].isdigit()):
                continue
            try:
                plan = _num(v)
            except ValueError:
                plan = -1.0
            if plan is not None and plan < MIN_PLAN:  # аудит А1: «1e-9» или опечатка дали бы ставку 45 000
                s_ = db.seller(int(k[1:]))
                bad.append(s_["name"] if s_ else k)
                continue
            if service.save_plan_override(db, month, int(k[1:]), plan):
                service.apply_override(db, cfg, month, int(k[1:]))  # после фиксации — только этому продавцу
        if bad:
            return _back(month, f"Не сохранено для: {', '.join(bad)} — ручной план на полный месяц, от "
                                f"{_rub(MIN_PLAN)} ₽. Остальные ручные планы сохранены.")
        return RedirectResponse(f"/month/{month}", status_code=303)

    @app.post("/month/{month}/freeze", dependencies=[Depends(auth)])
    def freeze(month: str):
        service.freeze(db, cfg, month)
        return RedirectResponse(f"/month/{month}", status_code=303)

    # ------------------------------------------------------------------ итоги
    @app.get("/month/{month}/results", dependencies=[Depends(auth)])
    def results(month: str):
        rows = db.q("SELECT r.*, s.name, s.role FROM results r JOIN sellers s ON s.id=r.seller_id WHERE month=? "
                    "ORDER BY s.role DESC, s.name", month)
        miss = sorted({d for d, _ in db.month_missing(month, cfg.stores, upto=date.today() - timedelta(days=1))})
        mc = db.one("SELECT * FROM month_close WHERE month=?", month)
        state = ("Итоги ещё не подведены." if not rows else
                 f"Итоги окончательные (все дни выгружены), посчитаны {mc['closed_at'][:16].replace('T', ' ')}."
                 if mc and mc["final"] else
                 "Итоги ПРЕДВАРИТЕЛЬНЫЕ: месяц не закончился или выгружены не все дни. Бот пересчитает их сам, когда "
                 "придут все дни месяца, и только тогда разошлёт продавцам.")
        if cfg.trial(month):
            state += f" {month} — пробный месяц (первый расчётный — {cfg.first_month}): бот итоги не подводит и не рассылает."
        if miss:
            state += f" Нет выгрузки 1С за: {', '.join(d[8:] + '.' + d[5:7] for d in miss[:12])}" + \
                     (" …" if len(miss) > 12 else "") + " (если магазины не работали — «График работы» → «Закрытые дни»)."
        b = [f'<p class=note>{html.escape(state)}</p>',
             f'<form method=post action="/month/{month}/close"><button class=btn>Посчитать итоги {month}</button>'
             '<span class=note> — для просмотра: продавцам кнопка ничего не отправляет; пока выгружены не все дни, '
             'итоги предварительные.</span></form>',
             '<table><tr><th>Продавец</th><th>Продажи</th><th>Отработано</th><th>Выполнение</th><th>Шагов</th>'
             '<th>Ставка (полный мес.)</th></tr>']
        for r in rows:
            perf = "" if r["performance"] is None else pct(r["performance"])
            b.append(f'<tr><td>{html.escape(r["name"])}</td><td class=n>{_rub(r["sales"])}</td><td class=n>{r["fte"]:.2f}</td>'
                     f'<td class=n>{perf}</td>'
                     f'<td class=n>{r["steps"]}</td><td class=n>{_rub(r["rate"])}</td></tr>')
        b.append(f'</table><a href="/month/{month}/payroll.xlsx">Скачать ведомость (xlsx)</a>')
        return page(f"Итоги {month}", "".join(b))

    @app.post("/month/{month}/close", dependencies=[Depends(auth)])
    def close(month: str):
        service.close_month(db, cfg, month)  # окончательные — только при всех выгруженных днях (аудит Д2)
        return RedirectResponse(f"/month/{month}/results", status_code=303)

    @app.get("/month/{month}/payroll.xlsx", dependencies=[Depends(auth)])
    def payroll(month: str):
        p = service.month_params(db, cfg, month)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = month
        ws.append(["Продавец", "Роль", "Продажи, ₽", "Отработано (доля мес.)", "План (полный мес.)", "Выполнение",
                   "Ставка (полный мес.)", "Ставка к начислению", f"{p.sales_percent:.0%} от продаж", "Итого"])
        for r in db.q("SELECT r.*, s.name, s.role, pl.plan_full FROM results r JOIN sellers s ON s.id=r.seller_id "
                      "LEFT JOIN plans pl ON pl.month=r.month AND pl.seller_id=r.seller_id WHERE r.month=?", month):
            i = ws.max_row + 1
            ws.append([r["name"], r["role"], r["sales"], r["fte"], r["plan_full"], r["performance"], r["rate"],
                       f"=G{i}*MIN(1,D{i})", f"=C{i}*{p.sales_percent}", f"=H{i}+I{i}"])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f"attachment; filename=payroll_{month}.xlsx"})

    # ------------------------------------------------------------------ приём выгрузки из 1С по HTTP
    @app.post("/api/upload")
    async def upload(request: Request):
        token = os.environ.get("UPLOAD_TOKEN", "")
        if not token or not secrets.compare_digest(request.headers.get("X-Token", ""), token):
            raise HTTPException(401, "bad token")
        from urllib.parse import unquote

        from .importers.onec_sales import archive_name, import_file
        name = os.path.basename(unquote(request.headers.get("X-Filename", "upload.csv"))) or "upload.csv"
        inbox, archive = cfg.onec.inbox, cfg.onec.archive
        inbox.mkdir(parents=True, exist_ok=True)
        archive.mkdir(parents=True, exist_ok=True)
        path = inbox / f".upload-{secrets.token_hex(4)}-{name}"  # не под итоговым именем: import_inbox не заберёт
        path.write_bytes(await request.body())
        try:
            rep = import_file(path, db, cfg)
        except Exception as e:  # noqa: BLE001 — битый файл в archive/bad, ответ 1С с причиной (аудит В1)
            (archive / "bad").mkdir(exist_ok=True)
            path.replace(archive_name(archive / "bad", name))
            raise HTTPException(422, f"{name}: файл не распознан — {e}") from e
        path.replace(archive_name(archive, name))
        return {"lines": rep.lines, "days": rep.days, "unknown_sellers": sorted(rep.unknown_sellers)}

    # ------------------------------------------------------------------ проверка плана на прошлых месяцах
    def _check_range(frm: str, to: str) -> list[str]:
        cur = date.today().strftime("%Y-%m")
        ok = lambda m: len(m) == 7 and m[4] == "-" and m[:4].isdigit() and m[5:].isdigit() and "01" <= m[5:] <= "12"  # noqa: E731
        frm = frm.strip() if ok(frm.strip()) else month_add(cur, -2)
        to = to.strip() if ok(to.strip()) else cur
        return backtest.months_between(min(frm, to), max(frm, to))

    def _pct(x) -> str:
        return pct(x)

    @app.get("/check", dependencies=[Depends(auth)])
    def check_page(frm: str = "", to: str = ""):
        months = _check_range(frm, to)
        res = [backtest.run(db, cfg, m) for m in months]
        p = res[-1].params
        b = ['<p>Какой план бот поставил бы каждому продавцу на месяц — по данным <b>до начала</b> месяца, как 1-го '
             'числа, — и сколько продавец продал на самом деле. Ничего не сохраняется и никому не отправляется.</p>',
             f'<p class=note>По нынешним правилам: норма {cfg.norm_shifts} рабочих дней, каждые +{p.step:.0%} сверх '
             f'плана — +{_rub(p.step_bonus)} ₽ к ставке {_rub(p.base_rate)} ₽, максимум {_rub(p.max_rate)} ₽, старшим '
             f'+{_rub(p.admin_bonus)} ₽. Факт — из 1С, а за месяцы без выгрузки 1С — из ведомостей зарплаты. '
             'Выполнение считается на отработанное время: (продажи ÷ доля отработанного) ÷ план.</p>',
             f'<form method=get action="/check">С <input type=text name=frm value="{months[0]}" style="width:80px"> '
             f'по <input type=text name=to value="{months[-1]}" style="width:80px"> <button class=btn>Показать</button> '
             f'<a href="/check.xlsx?frm={months[0]}&to={months[-1]}">Скачать в Excel</a></form>',
             '<h2>Магазины вместе</h2><table><tr><th>Месяц</th><th>Данные</th><th>План магазинов</th><th>Факт</th>'
             '<th>Факт / план</th><th>Коэф. год к году</th><th>Выполнили план</th><th>Ставки по схеме</th></tr>']
        for r in res:
            plan = (r.store.plan or 0) * r.coverage
            cov = f' <span class=note>({r.coverage:.0%} дней)</span>' if r.coverage < 1 else ""
            b.append(f'<tr><td>{r.month}</td><td>{r.source}{cov}</td><td class=n>{_rub(plan or None)}</td>'
                     f'<td class=n>{_rub(r.store_fact or None)}</td>'
                     f'<td class=n>{_pct(r.store_fact / plan if plan and r.store_fact else None)}</td>'
                     f'<td class=n>{"—" if r.store.trend is None else f"{r.store.trend:.3f}".replace(".", ",")}</td>'
                     f'<td class=n>{r.done} из {len(r.rows)}</td><td class=n>{_rub(r.rates or None)}</td></tr>')
        b.append('</table><p class=note>План магазинов при неполном месяце — пропорционально дням с выгрузкой. '
                 'Ставки по схеме — сумма ставок продавцов за отработанное время, без процента от продаж.</p>')
        sids = sorted({x.seller_id: x for r in res for x in r.rows}.values(), key=lambda x: (x.role != "admin", x.name))
        if sids:
            b.append('<h2>Выполнение плана по продавцам</h2><table><tr><th>Продавец</th>' +
                     "".join(f'<th>{r.month}</th>' for r in res) + '</tr>')
            for x in sids:
                cells = []
                for r in res:
                    y = next((z for z in r.rows if z.seller_id == x.seller_id), None)
                    ok = y is not None and y.performance is not None and y.performance >= 1
                    cells.append(f'<td class="n{" ok" if ok else ""}">{_pct(y.performance) if y else ""}</td>')
                b.append(f'<tr><td>{html.escape(x.name)}</td>{"".join(cells)}</tr>')
            b.append('</table>')
        how = METHOD_NAMES
        for r in res:
            b.append(f'<h2>{r.month}</h2>')
            if not r.rows:
                b.append('<p class=note>Нет данных о продажах за этот месяц — загрузите выгрузку 1С.</p>')
                continue
            b.append('<table><tr><th>Продавец</th><th>Как считали план</th><th>План (полный мес.)</th>'
                     '<th>Отработано</th><th>План на отработанное</th><th>Факт</th><th>Факт − план</th>'
                     '<th>Выполнение</th><th>Ставка</th></tr>')
            for x in r.rows:
                worked = f"{x.days} раб. дн." if x.days is not None else f"{x.fte:.2f} ставки".replace(".", ",")
                diff = x.sales - x.plan_worked if x.plan_worked else None
                role = ' <span class=note>старший</span>' if x.role == "admin" else ""
                b.append(f'<tr><td>{html.escape(x.name)}{role}</td><td>{how.get(x.method, x.method)}</td>'
                         f'<td class=n>{_rub(x.plan)}</td><td class=n>{worked}</td><td class=n>{_rub(x.plan_worked)}</td>'
                         f'<td class=n>{_rub(x.sales)}</td>'
                         f'<td class=n>{"" if diff is None else ("+" if diff > 0 else "") + _rub(diff)}</td>'
                         f'<td class=n>{_pct(x.performance)}</td><td class=n>{_rub(x.rate)}</td></tr>')
            b.append('</table>')
        return page("Проверка плана на прошлых месяцах", "".join(b))

    RUB = '#,##0" ₽"'

    def _xlsx_formats(sheet, fmt: dict[int, str], width: int) -> None:
        """Числа — в рублях и процентах, как на странице; шапка жирная, колонки пошире."""
        for c in sheet[1]:
            c.font = Font(bold=True)
            sheet.column_dimensions[c.column_letter].width = width
        sheet.column_dimensions["A"].width = 26
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                if c.column in fmt:
                    c.number_format = fmt[c.column]

    @app.get("/check.xlsx", dependencies=[Depends(auth)])
    def check_xlsx(frm: str = "", to: str = ""):
        months = _check_range(frm, to)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Магазины"
        ws.append(["Месяц", "Данные", "Доля дней", "План магазинов", "Факт", "Коэф. год к году", "Выполнили план",
                   "Продавцов", "Ставки по схеме"])
        for m in months:
            r = backtest.run(db, cfg, m)
            ws.append([m, r.source, r.coverage, (r.store.plan or 0) * r.coverage or None, r.store_fact or None,
                       r.store.trend, r.done, len(r.rows), r.rates or None])
            d = wb.create_sheet(m)
            d.append(["Продавец", "Роль", "Как считали план", "План (полный мес.)", "Рабочих дней", "Отработано (доля)",
                      "План на отработанное", "Факт", "Факт − план", "Выполнение", "Ступеней", "Ставка (полный мес.)"])
            for x in r.rows:
                d.append([x.name, "старший" if x.role == "admin" else "продавец", METHOD_NAMES.get(x.method, x.method),
                          x.plan, x.days, x.fte, x.plan_worked, x.sales,
                          x.sales - x.plan_worked if x.plan_worked else None, x.performance, x.steps, x.rate])
            _xlsx_formats(d, {4: RUB, 6: "0.00", 7: RUB, 8: RUB, 9: RUB, 10: "0%", 12: RUB}, 16)
        _xlsx_formats(ws, {3: "0%", 4: RUB, 5: RUB, 6: "0.000", 9: RUB}, 14)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f"attachment; filename=check_{months[0]}_{months[-1]}.xlsx"})

    # ------------------------------------------------------------------ бот MAX: токены, ссылки, владелец
    def _client(token: str):
        if max_client:
            return max_client(token)
        from .max_api import MaxClient
        return MaxClient(token, cfg.max.base_url, cfg.max.ca_file, timeout=15, pause=0.5)

    def _check(bot: str) -> str:
        """GET /me: имя бота или причина ошибки — в kv для страницы; ссылку на бота подставляем сами."""
        from .max_api import MaxAuthError
        tok = botcfg.token(bot)
        if not tok:
            return ""
        try:
            me = _client(tok).me()
        except MaxAuthError:
            res = "❌ MAX не принял токен — скопируйте его из business.max.ru заново"
        except Exception as e:  # noqa: BLE001
            res = f"❌ нет связи с MAX ({type(e).__name__})"
        else:
            nick = me.get("username") or ""
            if nick and not botcfg.link(cfg, bot):
                botcfg.set_link(db, cfg, bot, f"https://max.ru/{nick}")
            res = f"✅ на связи: {me.get('first_name') or me.get('name') or nick}" + (f" (@{nick})" if nick else "")
        db.set_kv(f"bot_check:{bot}", f"{res} — проверено {date.today():%d.%m}")
        return res

    @app.get("/bot", dependencies=[Depends(auth)])
    def bot_page(msg: str = ""):
        names = {"plan": "«План» — план и показатели", "training": "«Обучение» — уроки, тесты, база знаний"}
        cell = lambda f: "".join(f"<td>{f(b)}</td>" for b in names)  # noqa: E731
        b = [f'<p class=note>{html.escape(msg)}</p>' if msg else "",
             '<p class=note>Токен и ссылку на бота MAX выдаёт после модерации (business.max.ru → бот). Токен — как '
             'пароль: он хранится только в файле .env на этом компьютере, здесь видны последние 4 знака. Второй бот '
             'можно подключить позже — до этого уроки и тесты приходят в «План».</p>',
             '<h2>1. Боты</h2><form method=post action="/bot"><table>',
             '<tr><th></th>' + "".join(f"<th>{t}</th>" for t in names.values()) + '</tr>',
             '<tr><td>Токен</td>' + cell(lambda bt: f'<input type=password name="token_{bt}" autocomplete=off '
                                        f'style="width:260px" placeholder="вставьте токен"><br><span class=note>'
                                        f'{botcfg.mask(botcfg.token(bt))}</span>') + '</tr>',
             '<tr><td>Ссылка на бота</td>' + cell(lambda bt: f'<input type=text name="link_{bt}" style="width:260px" '
                                                 f'value="{html.escape(botcfg.link(cfg, bt))}" '
                                                 'placeholder="заполнится сама после проверки">') + '</tr>',
             '<tr><td>Связь с MAX</td>' + cell(lambda bt: html.escape(db.get_kv(f"bot_check:{bt}") or "—")) + '</tr>',
             '</table><button class=btn name=action value=save>Сохранить и проверить связь</button>']
        if runner and runner.running:
            b.append('<button class=btn name=action value=restart title="Нужно после смены токена">'
                     'Перезапустить службу</button><span class=note> — админка будет недоступна ~15 секунд</span>')
        b.append('</form>')

        own = botcfg.owners(db, cfg)
        b.append('<h2>2. Владелец</h2><p class=note>Владельцу каждое утро — сводка по всем продавцам, «Кого похвалить», '
                 'предупреждения о графике и выгрузке 1С. Продавцы этого не видят.</p>')
        if own:
            b.append('<form method=post action="/bot"><table><tr><th>Имя в MAX</th><th>user_id</th><th>Подключён</th>'
                     '<th></th></tr>')
            for o in own:
                b.append(f'<tr><td>{html.escape(o.get("name") or "—")}</td><td class=n>{o["id"]}</td>'
                         f'<td class=n>{_date((o.get("at") or "")[:10] or None)}</td><td><button class=btn name=action '
                         f'value="owner_del:{o["id"]}">Отключить</button></td></tr>')
            b.append('</table></form>')
        else:
            b.append('<p class=off>Владелец ещё не подключён — сводку получать некому.</p>')
        code = botcfg.owner_code(db)
        if code:
            lk = botcfg.link(cfg, "plan")
            if lk:
                b += [f'<div class=qr>{qr_svg(f"{lk}?start={code}")}</div>',
                      '<p>Наведите камеру телефона на QR-код и нажмите «Начать» в боте «План».</p>',
                      f'<p class=note>Ссылка: {html.escape(f"{lk}?start={code}")}</p>']
            b.append(f'<p>Или отправьте боту «План» код <b class=code>{code}</b>. Код одноразовый, действует сутки.</p>')
        else:
            b.append('<form method=post action="/bot"><button class=btn name=action value=owner_code>'
                     'Подключить владельца</button></form>')

        rows = [s for s in db.sellers() if s["role"] != "system"]
        plan = [s for s in rows if s["max_user_id"]]
        wait = [s for s in rows if not s["max_user_id"]]
        b.append('<h2>3. Продавцы</h2><p class=note>Каждый продавец подключается своим QR-кодом из карточки на странице '
                 '«Продавцы». Бот запоминает его аккаунт MAX и присылает только его личный план и показатели — чужих '
                 'цифр продавец не видит. Рядом с подключённым видно имя из его профиля MAX: проверьте, что подключился '
                 'тот самый человек.</p>')
        line = f'<p>Подключены к «Плану»: <b>{len(plan)} из {len(rows)}</b>'
        if botcfg.token("training"):
            line += f' · к «Обучению»: <b>{len([s for s in rows if s["training_user_id"]])} из {len(rows)}</b>'
        b.append(line + ' · <a href="/sellers">Продавцы →</a></p>')
        if wait:
            b.append('<p class=note>Ещё не подключены: ' + ", ".join(html.escape(s["name"]) for s in wait) + '</p>')
        return page("Бот MAX", "".join(b))

    @app.post("/bot", dependencies=[Depends(auth)])
    async def bot_save(request: Request):
        form = await request.form()
        action = str(form.get("action") or "")
        msg = []
        if action == "owner_code":
            botcfg.new_owner_code(db)
        elif action.startswith("owner_del:"):
            botcfg.remove_owner(db, cfg, int(action.split(":", 1)[1]))
            msg.append("Владелец отключён: сводка ему больше не приходит.")
        elif action == "restart" and runner:
            runner.restart()
            return HTMLResponse('<!doctype html><meta charset=utf-8><meta http-equiv=refresh content="20;url=/bot">'
                                '<p style="font-family:Arial">Служба перезапускается… Страница обновится сама через '
                                '20 секунд.</p>')
        elif action == "save":
            changed = False
            for bot in ("plan", "training"):
                tok = str(form.get(f"token_{bot}") or "").strip()
                if tok and tok != botcfg.token(bot):
                    try:
                        botcfg.set_token(cfg.env_path, bot, tok)
                        changed = True
                    except ValueError as e:
                        msg.append(f"{'«План»' if bot == 'plan' else '«Обучение»'}: {e}.")
                lk = str(form.get(f"link_{bot}") or "").strip()
                if lk != botcfg.link(cfg, bot):
                    if lk and not lk.startswith("https://max.ru/"):
                        msg.append("Ссылка на бота должна начинаться с https://max.ru/ — оставьте поле пустым, "
                                   "она заполнится сама.")
                    else:
                        botcfg.set_link(db, cfg, bot, lk)
                _check(bot)
            if runner and changed:
                msg.append("Бот запущен." if runner.start() else
                           "Новый токен заработает после перезапуска службы — кнопка «Перезапустить службу».")
            elif changed:
                msg.append("Сохранено. Перезапустите службу AwerSalesBot, чтобы бот начал работать с новым токеном.")
            else:
                msg.append("Сохранено.")
        return RedirectResponse(f"/bot?msg={quote(' '.join(msg))}" if msg else "/bot", status_code=303)

    # ------------------------------------------------------------------ смены и график работы
    def _short(store: str) -> str:
        return store[:1].upper()

    @app.get("/crews", dependencies=[Depends(auth)])
    def crews_page(msg: str = ""):
        crews_mod.ensure_crews(db, cfg)
        today = date.today()
        wed = crews_mod.wednesday(today)
        cs = crews_mod.crews(db)
        b = [f'<p class=note>{html.escape(msg)}</p>' if msg else "",
             '<p class=note>Смена — состав продавцов одного магазина. Каждую среду смены меняются магазинами. '
             'Магазин на эту неделю можно поправить — дальше бот чередует сам. Заболел кто-то и его заменили — '
             'это правится в «Графике работы» в конкретном дне, а не здесь.</p>',
             f'<form method=post action="/crews"><h2>Неделя со среды {wed:%d.%m.%Y}</h2><table>'
             '<tr><th>Смена</th><th>Магазин на этой неделе</th><th>Со следующей среды</th></tr>']
        for c in cs:
            now = crews_mod.crew_store(c, today, cfg.stores)
            nxt = crews_mod.crew_store(c, wed + timedelta(days=7), cfg.stores)
            opts = "".join(f'<option {"selected" if st == now else ""}>{html.escape(st)}</option>' for st in cfg.stores)
            b.append(f'<tr><td><input type=text name="name{c["id"]}" value="{html.escape(c["name"])}"></td>'
                     f'<td><select name="store{c["id"]}">{opts}</select></td><td>{html.escape(nxt)}</td></tr>')
        b.append('</table><h2>Состав</h2><table><tr><th>Продавец</th><th>Смена</th>'
                 '<th>Сильный продавец — ставить на сб и вс</th></tr>')
        for sl in db.sellers():
            opts = '<option value="">—</option>' + "".join(
                f'<option value="{c["id"]}" {"selected" if sl["crew_id"] == c["id"] else ""}>{html.escape(c["name"])}</option>'
                for c in cs)
            b.append(f'<tr><td>{html.escape(sl["name"])}</td><td><select name="crew{sl["id"]}">{opts}</select></td>'
                     f'<td><input type=checkbox name="key{sl["id"]}" {"checked" if sl["weekend_key"] else ""}></td></tr>')
        b.append('</table><button class=btn>Сохранить смены</button></form>'
                 f'<p class=note>Минимум в каждом магазине: в будни {cfg.staffing.weekday_min}, в субботу и воскресенье '
                 f'{cfg.staffing.weekend_min} продавца. Проверка — в «Графике работы».</p>')
        return page("Смены", "".join(b))

    @app.post("/crews", dependencies=[Depends(auth)])
    async def crews_save(request: Request):
        form = await request.form()
        wed = crews_mod.wednesday(date.today()).isoformat()
        with db.tx() as c:
            for cr in crews_mod.crews(db):
                name = " ".join(str(form.get(f"name{cr['id']}", cr["name"])).split()) or cr["name"]
                store = form.get(f"store{cr['id']}")
                store = store if store in cfg.stores else crews_mod.crew_store(cr, date.today(), cfg.stores)
                c.execute("UPDATE crews SET name=?, store=?, anchor=? WHERE id=?", (name, store, wed, cr["id"]))
            for sl in db.sellers():
                crew = form.get(f"crew{sl['id']}")
                c.execute("UPDATE sellers SET crew_id=?, weekend_key=? WHERE id=?",
                          (int(crew) if crew and str(crew).isdigit() else None,
                           1 if form.get(f"key{sl['id']}") else 0, sl["id"]))
        return RedirectResponse("/crews?msg=" + quote("Сохранено"), status_code=303)

    @app.get("/schedule/{month}", dependencies=[Depends(auth)])
    def schedule(month: str, msg: str = "", suggest: int = 0):
        """График по магазинам: в каждой таблице — кто работает в этом магазине; смены меняются магазинами по средам.
        suggest=1 — показать предложение графика по правилам (crews.suggest); сохраняет владелец кнопкой."""
        crews_mod.ensure_crews(db, cfg)
        e = html.escape
        y, mo = map(int, month.split("-"))
        n = service.days_in(month)
        dates = [date(y, mo, d) for d in range(1, n + 1)]
        saved = db.schedule_days(month)
        short: list[str] = []
        if suggest:
            proposal, short = crews_mod.suggest(db, cfg, month)
            saved = {**saved, **proposal}  # продавцы без смены и тестовые — как были
        cs = crews_mod.crews(db)
        cnum = {c["id"]: i for i, c in enumerate(cs, 1)}
        sellers = sorted(db.sellers(), key=lambda s: (cnum.get(s["crew_id"], 99), s["name"]))
        own = {(s["id"], d.day): crews_mod.default_store(db, cfg, s["id"], d) for s in sellers for d in dates}
        where = {(sid, d): (st or own[(sid, d)] or cfg.stores[0]) for sid, ds in saved.items() for d, st in ds.items()}
        we = lambda d: " we" if d.weekday() >= 5 else ""  # noqa: E731
        # пустой месяц — не список «0 из 3» на каждый день: что делать, скажет рекомендация
        probs = short if suggest else crews_mod.problems(db, cfg, dates[0], dates[-1]) if saved else []
        b = [f'<p class=links><a href="/schedule/{month_add(month, -1)}">← {month_add(month, -1)}</a>'
             f'<a href="/schedule/{month_add(month, 1)}">{month_add(month, 1)} →</a><a href="/crews">Смены</a>'
             f'<a href="/closed">Закрытые дни</a><a class=act href="/schedule/{month}?suggest=1">Предложить график на '
             'месяц</a></p>',
             f'<p class=note>{e(msg)}</p>' if msg else ""]
        if suggest:
            b.append('<p class=off>Это предложение бота — ещё не сохранено. Проверьте, поправьте галочки и нажмите '
                     '«Сохранить график» внизу. Чтобы отказаться, просто уйдите со страницы.</p>')
        tips = crews_mod.advice(db, cfg, month)
        if tips and not suggest:
            b.append('<h2>Рекомендации по рабочим дням</h2><ul>' + "".join(f"<li>{e(t)}</li>" for t in tips) + "</ul>")
        b.append('<p class=note>Отдельная таблица на каждый магазин. Строка «смена» — какая смена в этом магазине: каждую '
                 'среду смены меняются магазинами. Белые клетки — дни, когда смена продавца в этом магазине; серые — '
                 'другой магазин: галочка там — выход на замену (выделяется цветом). Один продавец в один день — только '
                 f'в одном магазине. Норма — {cfg.norm_shifts} рабочих дней. Рабочие дни считаются по графику (и у '
                 'старших); у кого графика на месяц нет — по дням с продажами в 1С.</p>')
        pend = db.offschedule_pending()
        if pend:
            b.append('<h2>Продажи в дни вне графика — нужен ваш ответ</h2><p class=note>Продавец продал в день, которого нет '
                     'в его графике. Вышел на замену — «Работал(а)»: день добавится в график и станет рабочим. Не работал '
                     '(покупатель забрал товар, отложенный в его смену) — «Не работал(а)». Продажа в обоих случаях остаётся '
                     'у продавца. Пока ответа нет, день не считается рабочим.</p><table><tr><th>День</th><th>Продавец</th>'
                     '<th>Магазин</th><th>Продажи, ₽</th><th></th></tr>')
            for x in pend:
                b.append(f'<tr><td>{date.fromisoformat(x["day"]):%d.%m.%Y}</td><td>{e(x["name"])}</td><td>{e(x["store"])}</td>'
                         f'<td class=n>{_rub(x["amount"])}</td><td><form method=post action="/offschedule">'
                         f'<input type=hidden name=day value="{x["day"]}"><input type=hidden name=sid value="{x["seller_id"]}">'
                         f'<input type=hidden name=back value="{e(month)}">'
                         '<button class=btn name=worked value=1>Работал(а) — добавить в график</button> '
                         '<button class=btn name=worked value=0>Не работал(а)</button></form></td></tr>')
            b.append('</table>')
        if probs:
            head = ("В сменах не хватает людей — добавьте продавцов в смену («Смены») или замену из другой смены:"
                    if suggest else "Нарушены правила графика:")
            b.append(f'<p class=off>{head}</p><ul>' + "".join(f"<li>{e(x)}</li>" for x in probs) + "</ul>")
        b.append(f'<form method=post action="/schedule/{month}">')
        head = "".join(f'<th class="{we(d)}">{d.day}<br>{crews_mod.WD[d.weekday()]}</th>' for d in dates)
        for si, st in enumerate(cfg.stores):
            shifts = "".join(
                f'<td class="sh{we(d)}">' + ", ".join(f"С{cnum[c['id']]}" for c in cs
                                                      if crews_mod.crew_store(c, d, cfg.stores) == st) + "</td>"
                for d in dates)
            b.append(f'<h2>{e(st)}</h2><table class=grid><tr><th>Продавец</th>{head}</tr>'
                     f'<tr><td class=nm><span class=note>смена</span></td>{shifts}</tr>')
            for sl in sellers:
                sid = sl["id"]
                cells = []
                for d in dates:
                    mine = own[(sid, d.day)] == st
                    on = where.get((sid, d.day)) == st
                    cls = ("" if mine else "alien") + we(d) + (" swap" if on and not mine else "")
                    cells.append(f'<td class="{cls.strip()}"><input type=checkbox name="w{sid}_{d.day}_{si}" '
                                 f'data-s="{sid}" data-d="{d.day}" data-st="{si}" data-own="{int(mine)}" '
                                 f'data-we="{int(d.weekday() >= 5)}" {"checked" if on else ""}></td>')
                crew = f"С{cnum[sl['crew_id']]}" if sl["crew_id"] in cnum else "без смены"
                b.append(f'<tr><td class=nm>{e(sl["name"])} <span class=note>{crew}'
                         f'{" · сильный" if sl["weekend_key"] else ""}</span></td>{"".join(cells)}</tr>')
            cnt = "".join(f'<td id="n{si}_{d.day}" class="{we(d).strip()}" data-need="{crews_mod.minimum(cfg, d)}"></td>'
                          for d in dates)
            b.append(f'<tr><td class=nm><b>Продавцов</b> <span class=note>(нужно {cfg.staffing.weekday_min} в будни, '
                     f'{cfg.staffing.weekend_min} в сб и вс)</span></td>{cnt}</tr></table>')
        b.append('<h2>Итого по продавцам</h2><table><tr><th>Продавец</th><th>Смена</th><th>Рабочих дней</th>'
                 + "".join(f"<th>{e(st)}</th>" for st in cfg.stores) + '<th>Заполнить по схеме (с учётом ротации)</th></tr>')
        for sl in sellers:
            sid = sl["id"]
            crew = next((c["name"] for c in cs if c["id"] == sl["crew_id"]), "без смены")
            fill = (f'<select id="p{sid}"><option value="3,2">3/2</option><option value="2,2">2/2</option>'
                    f'<option value="5,2">5/2</option></select> с <input type=number id="k{sid}" value=1 min=1 max=7 '
                    f'style="width:42px">-го <button type=button class=btn onclick="fill({sid})">Заполнить</button>'
                    f'<button type=button class=btn onclick="weekends({sid})">+ сб и вс</button>'
                    f'<button type=button class=btn onclick="fill({sid},true)">Очистить</button>')
            b.append(f'<tr><td>{e(sl["name"])}{" · сильный" if sl["weekend_key"] else ""}</td><td>{e(crew)}</td>'
                     f'<td id="t{sid}" class=n></td>' + "".join(f'<td id="t{sid}_{si}" class=n></td>'
                                                              for si in range(len(cfg.stores))) + f'<td>{fill}</td></tr>')
        b.append(f'</table><button class=btn>Сохранить график</button></form>'
                 f'<script>const NORM={cfg.norm_shifts},NS={len(cfg.stores)};'
                 'const Q=s=>[...document.querySelectorAll(s)];'
                 'function count(){Q("[id^=n]").forEach(c=>{const [si,d]=c.id.slice(1).split("_");'
                 'const k=Q(`input[data-st="${si}"][data-d="${d}"]`).filter(x=>x.checked).length;'
                 'c.textContent=k;c.className=(c.className.replace(/ ?off/g,""))+(k<+c.dataset.need?" off":"");});'
                 'Q("[id^=t]").forEach(c=>{const p=c.id.slice(1).split("_");const sel=p.length>1?'
                 '`input[data-s="${p[0]}"][data-st="${p[1]}"]`:`input[data-s="${p[0]}"]`;'
                 'const k=Q(sel).filter(x=>x.checked).length;c.textContent=k;'
                 'if(p.length===1)c.className="n"+(k===NORM?"":" off");});}'
                 'function fill(s,clear){const [on,off]=document.getElementById("p"+s).value.split(",").map(Number);'
                 'const k=+document.getElementById("k"+s).value||1;'
                 'Q(`input[data-s="${s}"]`).forEach(x=>{const d=+x.dataset.d;'
                 'x.checked=!clear&&x.dataset.own==="1"&&d>=k&&((d-k)%(on+off))<on;mark(x);});count();}'
                 'function weekends(s){const busy=new Set(Q(`input[data-s="${s}"]`).filter(x=>x.checked).map(x=>x.dataset.d));'
                 'Q(`input[data-s="${s}"][data-we="1"][data-own="1"]`).forEach(x=>{if(!busy.has(x.dataset.d)){x.checked=true;mark(x);}});count();}'
                 'function mark(x){const td=x.parentElement;td.classList.toggle("swap",x.checked&&x.dataset.own!=="1");}'
                 'document.addEventListener("change",e=>{const x=e.target;if(!x.dataset||!x.dataset.s)return;'
                 'if(x.checked)Q(`input[data-s="${x.dataset.s}"][data-d="${x.dataset.d}"]`).forEach(y=>{if(y!==x){y.checked=false;mark(y);}});'
                 'mark(x);count();});count();</script>')
        return page(f"График работы {month}", "".join(b))

    @app.post("/schedule/{month}", dependencies=[Depends(auth)])
    async def schedule_save(month: str, request: Request):
        form = await request.form()
        n = service.days_in(month)
        y, mo = map(int, month.split("-"))
        for sl in db.sellers():
            days = {}
            for d in range(1, n + 1):
                picked = [st for si, st in enumerate(cfg.stores) if form.get(f"w{sl['id']}_{d}_{si}")]
                if picked:  # магазин храним, только если это не магазин его смены (замена)
                    dflt = crews_mod.default_store(db, cfg, sl["id"], date(y, mo, d))
                    st = dflt if dflt in picked else picked[0]
                    days[d] = None if st == dflt else st
            db.set_schedule(month, sl["id"], days)
        return RedirectResponse(f"/schedule/{month}?msg=" + quote("График сохранён"), status_code=303)

    @app.post("/offschedule", dependencies=[Depends(auth)])
    async def offschedule(request: Request):
        """Ответ на продажу в день вне графика. Итоги уже закрытого месяца пересчитываются (рабочих дней стало иначе)."""
        form = await request.form()
        try:
            day = date.fromisoformat(str(form.get("day")))
            sid = int(str(form.get("sid")))
        except ValueError:
            raise HTTPException(400, "Неверный день или продавец")
        if not db.seller(sid):
            raise HTTPException(404, "Продавец не найден")
        worked = form.get("worked") == "1"
        db.decide_offschedule(day.isoformat(), sid, worked)
        month = day.strftime("%Y-%m")
        msg = f"{day:%d.%m} — {'рабочий день, добавлен в график' if worked else 'не рабочий день'}"
        if service.refresh_results(db, cfg, month):
            msg += f". Итоги {month} пересчитаны — проверьте «Итоги и ведомость»"
        back = str(form.get("back") or month)
        if not re.fullmatch(r"\d{4}-\d{2}", back):
            back = month
        return RedirectResponse(f"/schedule/{back}?msg=" + quote(msg), status_code=303)

    # ------------------------------------------------------------------ дни, когда магазины не работали
    @app.get("/closed", dependencies=[Depends(auth)])
    def closed_days(msg: str = ""):
        fmt = lambda d: date.fromisoformat(d).strftime("%d.%m.%Y")  # noqa: E731
        b = [f'<p class=note>{html.escape(msg)}</p>' if msg else "",
             '<p class=note>Праздник, ремонт, отключение света — в такой день продаж нет не из-за 1С. Отмеченный день '
             'бот не считает пропуском данных (иначе досчитывает месяц, и план следующего года завышается), не ждёт за '
             'него выгрузку и не шлёт продавцам «продаж нет». Если продажи за день всё же придут из 1С, отметка '
             'снимется сама.</p>']
        gaps = db.data_gaps(date.today() - timedelta(days=2))  # вчера и позавчера выгрузка ещё может прийти
        if gaps:
            b.append('<h2>Дни без данных из 1С</h2><p class=note>Ни продаж, ни отметки. Если магазины работали — '
                     'выгрузите эти дни из 1С ещё раз.</p><form method=post action="/closed"><table>'
                     + "".join(f'<tr><td>{fmt(d)}</td><td><button class=btn name=add value="{d}">Магазины были '
                               f'закрыты</button></td></tr>' for d in gaps) + '</table></form>')
        marked = db.store_closed_days()
        b.append('<h2>Отмечены закрытыми</h2>' + ('<form method=post action="/closed"><table>' + "".join(
            f'<tr><td>{fmt(d)}</td><td><button class=btn name=remove value="{d}">Снять отметку</button></td></tr>'
            for d in marked) + '</table></form>' if marked else '<p class=note>Пока нет.</p>'))
        b.append('<h2>Отметить день</h2><form method=post action="/closed">Дата <input type=text name=day '
                 'placeholder="25.06.2026" style="width:110px"> <button class=btn>Магазины были закрыты</button></form>')
        return page("Дни, когда магазины были закрыты", "".join(b))

    @app.post("/closed", dependencies=[Depends(auth)])
    def closed_save(add: str = Form(""), remove: str = Form(""), day: str = Form("")):
        try:
            if remove:
                d = date.fromisoformat(remove)
                db.set_store_closed(d, cfg.stores, closed=False)
                msg = f"Отметка снята: {d:%d.%m.%Y}"
            else:
                d = date.fromisoformat(add) if add else datetime.strptime(day.strip(), "%d.%m.%Y").date()
                if d >= date.today():
                    raise ValueError("отмечается только прошедший день")
                db.set_store_closed(d, cfg.stores)
                msg = f"Отмечено: {d:%d.%m.%Y} магазины были закрыты"
        except ValueError as e:
            msg = f"Не отмечено: {e}. Дата — как 25.06.2026"
        return RedirectResponse("/closed?msg=" + quote(msg), status_code=303)

    # ------------------------------------------------------------------ обучение: правка контента
    def _learn_tabs(kind: str) -> str:
        return "<p class=links>" + "".join(f'<a href="/learn?kind={k}"{" class=act" if k == kind else ""}>{t}</a>'
                                           for k, t in content.KINDS.items()) + "</p>"

    @app.get("/learn", dependencies=[Depends(auth)])
    def learn(kind: str = "lesson", msg: str = ""):
        if kind not in content.KINDS:
            raise HTTPException(404)
        rows = content.items(db, cfg, kind, active_only=False)
        b = [_learn_tabs(kind), f'<p class=note>{html.escape(msg)}</p>' if msg else "",
             '<p class=note>Тексты можно менять здесь — исправленное больше не перезаписывается стартовыми файлами. '
             'Черновики новых уроков, статей и тестов может написать Claude (скилл «Продажи в магазине»): вы проверяете '
             'и включаете. Выключенное не рассылается.</p>',
             f'<p><a class=btn href="/learn/new?kind={kind}">+ Добавить</a></p>'
             '<form method=post action="/learn/sync" style="display:inline"><button class=btn>Подтянуть новые тексты '
             'из файлов</button></form>',
             '<table><tr><th>Название</th><th>Тема / раздел</th><th>Включено</th><th>Отправлено / ответы</th>'
             '<th>Кем изменено</th></tr>']
        for x in rows:
            sent = db.one("SELECT COUNT(*) n FROM lessons_sent WHERE lesson_id=?",
                          x["id"] if kind != "onboarding" else f"onb-{x['data'].get('shift')}")["n"]
            q = db.one("SELECT COUNT(*) n, COALESCE(SUM(correct),0) ok FROM quiz_answers WHERE lesson_id=?", x["id"])
            stat = f"{sent}" + (f" · квиз {q['ok']} из {q['n']} верно" if q["n"] else "")
            if kind == "article":
                v = db.one("SELECT COUNT(*) n, COUNT(DISTINCT seller_id) p FROM kb_views WHERE article_id=?", x["id"])
                stat = f"открыли {v['n']} раз ({v['p']} чел.)"
            topic = (content.LESSON_TOPICS.get(x["topic"]) or content.ARTICLE_SECTIONS.get(x["topic"]) or
                     (f"рабочий день {x['data'].get('shift')}" if kind == "onboarding" else ""))
            b.append(f'<tr><td><a href="/learn/{quote(x["id"])}">{html.escape(x["title"])}</a></td>'
                     f'<td>{html.escape(topic or "")}</td><td>{"да" if x["active"] else "<span class=off>нет</span>"}'
                     f'</td><td class=n>{stat}</td><td>{"вы (админка)" if x["origin"] == "admin" else "стартовый файл"}'
                     f'</td></tr>')
        b.append("</table>")
        return page(f"Обучение: {content.KINDS[kind]}", "".join(b))

    @app.post("/learn/sync", dependencies=[Depends(auth)])
    def learn_sync():
        n = content.sync_from_files(db, cfg, force=True)
        return RedirectResponse("/learn?msg=" + quote(f"Из файлов добавлено или обновлено: {n}"), status_code=303)

    def _learn_form(x: dict, err: str = "") -> HTMLResponse:
        kind = x["kind"]
        e = html.escape
        d = x.get("data") or {}
        f = [_learn_tabs(kind), f'<p class=off>{e(err)}</p>' if err else "",
             f'<form method=post action="/learn/{quote(x["id"])}"><input type=hidden name=kind value="{kind}"><table>',
             f'<tr><td>Название</td><td><input type=text name=title style="width:520px" value="{e(x.get("title") or "")}"></td></tr>']
        if kind == "lesson":
            opts = "".join(f'<option value="{k}" {"selected" if x.get("topic") == k else ""}>{t}</option>'
                           for k, t in content.LESSON_TOPICS.items())
            f.append(f'<tr><td>Тема (по цели недели)</td><td><select name=topic>{opts}</select></td></tr>')
        if kind == "article":
            opts = "".join(f'<option value="{k}" {"selected" if x.get("topic") == k else ""}>{t}</option>'
                           for k, t in content.ARTICLE_SECTIONS.items())
            f.append(f'<tr><td>Раздел</td><td><select name=topic>{opts}</select></td></tr>')
        if kind == "onboarding":
            f.append(f'<tr><td>Рабочий день новичка №</td><td><input type=number name=shift min=1 max=60 '
                     f'value="{d.get("shift") or 1}"></td></tr>')
        if kind != "checklist":
            f.append(f'<tr><td>Текст</td><td><textarea name=body rows=12 cols=90>{e(x.get("body") or "")}</textarea>'
                     '<br><span class=note>Без символов * и _ (разметка MAX). Фразы для покупателя — в «ёлочках».</span></td></tr>')
        if kind in ("lesson", "onboarding"):
            f.append(f'<tr><td>Попробуйте сегодня</td><td><input type=text name=try_today style="width:520px" '
                     f'value="{e(x.get("try_today") or "")}"></td></tr>')
        if kind == "article":
            f.append(f'<tr><td>Слова для поиска</td><td><input type=text name=tags style="width:520px" '
                     f'value="{e(", ".join(d.get("tags") or []))}"></td></tr>')
        if kind == "test":
            f.append(f'<tr><td>Статьи перед тестом (id через запятую)</td><td><input type=text name=article_ids '
                     f'style="width:520px" value="{e(", ".join(d.get("article_ids") or []))}"></td></tr>')
        if kind in ("lesson", "test", "survey"):
            qs = d.get("quiz") if kind == "lesson" else d.get("questions")
            hint = ("? вопрос · - вариант · ~ свой ответ (можно ответить своими словами)" if kind == "survey" else
                    "? вопрос · + верный вариант · - неверный · = объяснение")
            f.append(f'<tr><td>{"Квиз после урока" if kind == "lesson" else "Вопросы"}</td><td>'
                     f'<textarea name=questions rows=14 cols=90>'
                     f'{e(x.get("questions_raw") or content.format_questions(qs or [], survey=kind == "survey"))}'
                     f'</textarea><br><span class=note>{hint}</span></td></tr>')
        if kind == "checklist":
            f.append(f'<tr><td>Пункты (по одному в строке)</td><td><textarea name=items rows=12 cols=90>'
                     f'{e(chr(10).join(d.get("items") or []))}</textarea></td></tr>')
        f.append(f'<tr><td>Включено</td><td><input type=checkbox name=active {"checked" if x.get("active", 1) else ""}>'
                 '</td></tr></table><button class=btn>Сохранить</button></form>'
                 f'<p class=note>id: {e(x["id"])} — не меняется (по нему бот помнит, что уже отправлял).</p>')
        return page(f"{content.KINDS[kind]}: {e(x.get('title') or 'новый')}", "".join(f))

    @app.get("/learn/new", dependencies=[Depends(auth)])
    def learn_new(kind: str = "lesson"):
        if kind not in content.KINDS:
            raise HTTPException(404)
        return _learn_form({"id": content.new_id(db, kind), "kind": kind, "active": 1,
                            "topic": "general" if kind == "lesson" else ("tkani" if kind == "article" else None)})

    @app.get("/learn/{cid}", dependencies=[Depends(auth)])
    def learn_edit(cid: str):
        x = content.get(db, cfg, cid)
        if not x:
            raise HTTPException(404)
        return _learn_form(x)

    @app.post("/learn/{cid}", dependencies=[Depends(auth)])
    async def learn_save(cid: str, request: Request):
        form = await request.form()
        kind = str(form.get("kind") or "")
        if kind not in content.KINDS:
            raise HTTPException(400)
        old = content.get(db, cfg, cid) or {}
        title = " ".join(str(form.get("title") or "").split())
        body = str(form.get("body") or "").replace("\r\n", "\n").strip() or None
        data = dict(old.get("data") or {})
        x = {"id": cid, "kind": kind, "title": title, "body": body, "topic": form.get("topic") or None,
             "try_today": str(form.get("try_today") or "").strip() or None, "active": 1 if form.get("active") else 0,
             "data": data}
        try:
            if not title:
                raise ValueError("Нужно название")
            if any(ch in (title + (body or "") + (x["try_today"] or "")) for ch in "*_"):
                raise ValueError("Символы * и _ ломают разметку MAX — уберите их")
            if kind in ("lesson", "test", "survey"):
                raw = str(form.get("questions") or "")
                x["questions_raw"] = raw
                qs = content.parse_questions(raw, survey=kind == "survey")
                data["quiz" if kind == "lesson" else "questions"] = qs
                if kind in ("test", "survey") and not qs:
                    raise ValueError("Нужен хотя бы один вопрос")
            if kind == "onboarding":
                data["shift"] = int(form.get("shift") or 1)
            if kind == "article":
                data["tags"] = [t.strip().lower() for t in str(form.get("tags") or "").split(",") if t.strip()]
            if kind == "test":
                data["article_ids"] = [t.strip() for t in str(form.get("article_ids") or "").split(",") if t.strip()]
            if kind == "checklist":
                items_ = [i.strip() for i in str(form.get("items") or "").splitlines() if i.strip()]
                if not items_:
                    raise ValueError("Нужен хотя бы один пункт")
                if any(len(i) > 120 for i in items_):
                    raise ValueError("Пункт длиннее 120 символов — кнопка не поместится")
                data["items"] = items_
        except ValueError as err:
            return _learn_form(x, str(err))
        content.save(db, cid, kind, title, body, x["try_today"], x["topic"], data, bool(x["active"]))
        return RedirectResponse(f"/learn?kind={kind}&msg=" + quote(f"Сохранено: {title}"), status_code=303)

    # ------------------------------------------------------------------ задания: тесты, опросы, чек-листы
    @app.get("/tasks", dependencies=[Depends(auth)])
    def tasks_page(msg: str = ""):
        crews_mod.ensure_crews(db, cfg)
        e = html.escape
        cs = db.q("SELECT * FROM campaigns ORDER BY active DESC, id DESC")
        quiet = tasks_mod.quiet_hours(db, date.today())
        b = [f'<p class=note>{e(msg)}</p>' if msg else "",
             '<p class=note>Тесты, опросы и чек-листы приходят в бот «Обучение» в тестовые часы — только тем, у кого '
             'сегодня рабочий день по графику. Тексты — на странице «Обучение». Рейтинг по результатам — '
             '<a href="/rating">информативный</a>, на ставку не влияет.</p>',
             '<h2>Рассылки</h2><table><tr><th>Что</th><th>Кому</th><th>Когда</th><th>Часы</th><th>Получили</th>'
             '<th>Выполнили</th><th>Включена</th></tr>']
        for c in cs:
            x = content.get(db, cfg, c["content_id"])
            st = db.one("SELECT COUNT(*) n, SUM(done_at IS NOT NULL) d FROM deliveries WHERE campaign_id=?", c["id"])
            a, z = tasks_mod.hours(cfg, c)
            when = tasks_mod.REPEATS[c["repeat"]] + (
                f" ({['пн', 'вт', 'ср', 'чт', 'пт', 'сб', 'вс'][c['weekday']]})" if c["repeat"] == "weekly" and
                c["weekday"] is not None else "") + f", с {date.fromisoformat(c['start_day']):%d.%m}"
            b.append(f'<tr><td><a href="/tasks/{c["id"]}">{e(content.KINDS[x["kind"]][:-1] if x else "?")}: '
                     f'{e(x["title"] if x else c["content_id"])}</a></td><td>{e(tasks_mod.audience_label(db, c["audience"]))}'
                     f'</td><td>{e(when)}</td><td>{a:%H:%M}–{z:%H:%M}</td><td class=n>{st["n"]}</td>'
                     f'<td class=n>{st["d"] or 0}</td><td>{"да" if c["active"] else "нет"}</td></tr>')
        b.append("</table>")
        opts = "".join(f'<option value="{e(x["id"])}">{e(content.KINDS[x["kind"]])}: {e(x["title"])}</option>'
                       for k in ("test", "survey", "checklist") for x in content.items(db, cfg, k))
        aud = "".join(f'<option value="{k}">{v}</option>' for k, v in tasks_mod.AUDIENCES.items())
        aud += "".join(f'<option value="store:{e(st)}">магазин «{e(st)}» (кто сегодня там)</option>' for st in cfg.stores)
        aud += "".join(f'<option value="crew:{c["id"]}">смена «{e(c["name"])}»</option>' for c in crews_mod.crews(db))
        rep = "".join(f'<option value="{k}">{v}</option>' for k, v in tasks_mod.REPEATS.items())
        wd = "".join(f'<option value="{i}">{n}</option>' for i, n in enumerate(crews_mod.WD))
        a, z = tasks_mod.hours(cfg)
        hint = (f" Самые тихие часы по чекам 1С: {', '.join(f'{h}:00–{h + 1}:00' for h in quiet)}." if quiet else "")
        b += ['<h2>Новая рассылка</h2><form method=post action="/tasks"><table>',
              f'<tr><td>Что</td><td><select name=content_id>{opts}</select></td></tr>',
              f'<tr><td>Кому</td><td><select name=audience>{aud}</select></td></tr>',
              f'<tr><td>Когда</td><td><select name=repeat>{rep}</select> день недели (для «раз в неделю»): '
              f'<select name=weekday>{wd}</select></td></tr>',
              f'<tr><td>Начиная с</td><td><input type=date name=start_day value="{date.today().isoformat()}"></td></tr>',
              f'<tr><td>Тестовые часы</td><td>с <input type=time name=hour_from value="{a:%H:%M}"> до '
              f'<input type=time name=hour_to value="{z:%H:%M}"><span class=note>{hint}</span></td></tr>',
              '<tr><td>Только в рабочий день по графику</td><td><input type=checkbox name=only_working checked></td></tr>',
              '</table><button class=btn>Создать рассылку</button></form>']
        return page("Задания: тесты, опросы, чек-листы", "".join(b))

    @app.post("/tasks", dependencies=[Depends(auth)])
    def tasks_create(content_id: str = Form(...), audience: str = Form("all"), repeat: str = Form("once"),
                     weekday: int = Form(0), start_day: str = Form(""), hour_from: str = Form(""),
                     hour_to: str = Form(""), only_working: str = Form("")):
        x = content.get(db, cfg, content_id)
        if not x or x["kind"] not in ("test", "survey", "checklist") or repeat not in tasks_mod.REPEATS:
            raise HTTPException(400)
        with db.tx() as c:
            c.execute("INSERT INTO campaigns(content_id,audience,repeat,start_day,weekday,hour_from,hour_to,only_working,"
                      "active,created_at) VALUES(?,?,?,?,?,?,?,?,1,?)",
                      (content_id, audience, repeat, start_day or date.today().isoformat(), weekday,
                       hour_from or None, hour_to or None, 1 if only_working else 0, date.today().isoformat()))
        return RedirectResponse("/tasks?msg=" + quote(f"Рассылка создана: {x['title']}"), status_code=303)

    @app.get("/tasks/{cid}", dependencies=[Depends(auth)])
    def task_results(cid: int):
        e = html.escape
        c = db.one("SELECT * FROM campaigns WHERE id=?", cid)
        if not c:
            raise HTTPException(404)
        x = content.get(db, cfg, c["content_id"]) or {"title": c["content_id"], "kind": "test", "data": {}}
        qs = x["data"].get("questions") or [{"q": it} for it in x["data"].get("items") or []]
        b = [f'<p>{e(tasks_mod.audience_label(db, c["audience"]))} · {e(tasks_mod.REPEATS[c["repeat"]])} · '
             f'<a href="/learn/{quote(c["content_id"])}">текст</a></p>',
             f'<form method=post action="/tasks/{cid}/toggle"><button class=btn>'
             f'{"Выключить рассылку" if c["active"] else "Включить рассылку"}</button></form>',
             '<table><tr><th>Продавец</th><th>День</th><th>Выполнено</th><th>Результат</th>'
             + "".join(f'<th>{e(q["q"][:40])}</th>' for q in qs) + '</tr>']
        for d in db.q("SELECT d.*, s.name FROM deliveries d JOIN sellers s ON s.id=d.seller_id WHERE campaign_id=? "
                      "ORDER BY day DESC, s.name", cid):
            ans = {r["q"]: r for r in db.q("SELECT * FROM answers WHERE campaign_id=? AND seller_id=? AND day=?",
                                           cid, d["seller_id"], d["day"])}
            cells = []
            for i, q in enumerate(qs):
                r = ans.get(i)
                if r is None:
                    cells.append("<td></td>")
                elif x["kind"] == "test":
                    cells.append(f'<td>{"✅" if r["correct"] else "✗ " + e(q["options"][r["answer"]][:30])}</td>')
                elif x["kind"] == "checklist":
                    cells.append("<td>✅</td>")
                else:
                    cells.append(f'<td>{e((r["text"] or "")[:200])}</td>')
            score = "" if d["score"] is None else f'{d["score"]:.0%}'
            b.append(f'<tr><td>{e(d["name"])}</td><td>{date.fromisoformat(d["day"]):%d.%m}</td>'
                     f'<td>{"да" if d["done_at"] else "нет"}</td><td class=n>{score}</td>{"".join(cells)}</tr>')
        b.append("</table>")
        return page(f"{content.KINDS[x['kind']][:-1]}: {e(x['title'])}", "".join(b))

    @app.post("/tasks/{cid}/toggle", dependencies=[Depends(auth)])
    def task_toggle(cid: int):
        with db.tx() as c:
            c.execute("UPDATE campaigns SET active=1-active WHERE id=?", (cid,))
        return RedirectResponse(f"/tasks/{cid}", status_code=303)

    @app.get("/rating", dependencies=[Depends(auth)])
    def rating_page(month: str = ""):
        month = month or date.today().strftime("%Y-%m")
        pct_ = lambda v: "" if v is None else f"{v:.0%}"  # noqa: E731
        b = [f'<p class=links><a href="/rating?month={month_add(month, -1)}">← {month_add(month, -1)}</a>'
             f'<a href="/rating?month={month_add(month, 1)}">{month_add(month, 1)} →</a></p>',
             '<p class=note>Информативно, на ставку и ЗП не влияет (решение владельца). Итог — среднее из: тесты '
             '(доля верных), квизы после уроков, задания (выполнено из полученных). Продавцам не показывается.</p>',
             '<table><tr><th>#</th><th>Продавец</th><th>Итог</th><th>Тесты</th><th>Квизы</th><th>Задания</th></tr>']
        for i, r in enumerate(tasks_mod.rating(db, month), 1):
            b.append(f'<tr><td>{i}</td><td>{html.escape(r["name"])}</td><td class=n><b>{pct_(r["score"])}</b></td>'
                     f'<td class=n>{pct_(r["test"])}</td><td class=n>{pct_(r["quiz"])}</td>'
                     f'<td class=n>{r["done"]} из {r["sent"]}</td></tr>')
        b.append("</table>")
        return page(f"Рейтинг по обучению {month}", "".join(b))

    # ------------------------------------------------------------------ продавцы: подключение и увольнение
    def _date(d: str | None) -> str:
        return date.fromisoformat(d).strftime("%d.%m.%Y") if d else "—"

    def _role_select(s) -> str:
        return (f'<select name="role{s["id"]}"><option value=seller {"selected" if s["role"] == "seller" else ""}>'
                f'продавец</option><option value=test {"selected" if s["role"] == "test" else ""}>тестовый</option>'
                f'<option value=admin {"selected" if s["role"] == "admin" else ""}>'
                'админ (старший)</option></select>')

    def _max_name(s) -> str:
        """Имя из профиля MAX при подключении — владелец сверяет, что подключился тот самый человек."""
        return f' <span class=note>(в MAX: {html.escape(s["max_name"])})</span>' if s["max_name"] else ""

    def _training_mark(s) -> str:
        if not cfg.training_token:
            return ""
        if s["training_user_id"]:
            return " · обучение ✅"
        return f' · <a href="/sellers/{s["id"]}/card">обучение — подключить</a>'

    def _fire_btn(s) -> str:
        msg = (f'Уволить {s["name"]}? Бот перестанет присылать сообщения и отвечать ему. '
               'Продажи этого месяца останутся в итогах и ведомости.')
        return (f'<button class=btn name=action value="fire:{s["id"]}" '
                f'onclick="return confirm({html.escape(repr(msg))})">Уволить</button>')

    @app.get("/sellers", dependencies=[Depends(auth)])
    def sellers(msg: str = ""):
        rows = db.sellers(active_only=False)
        waiting = [s for s in rows if s["active"] and not s["max_user_id"]]
        working = [s for s in rows if s["active"] and s["max_user_id"]]
        fired = [s for s in rows if not s["active"]]
        b = [f'<p class=note>{html.escape(msg)}</p>' if msg else "",
             '<p class=note>Новые продавцы появляются здесь сами, когда впервые встречаются в выгрузке 1С. '
             'Можно добавить заранее — тогда первый урок программы новичка придёт сразу после подключения к боту.</p>',
             names_link(),
             '<h2>Добавить продавца</h2><form method=post action="/sellers/add"><table>',
             '<tr><td>Фамилия Имя — точно как в 1С</td><td><input type=text name=name style="width:260px" required></td></tr>',
             '<tr><td>Обращение в сообщениях</td><td><input type=text name=first_name style="width:260px" placeholder="по умолчанию — имя из 1С"></td></tr>',
             '<tr><td>Роль</td><td><select name=role><option value=seller>продавец</option>'
             '<option value=admin>админ (старший)</option><option value=test>тестовый — проверить бота: сообщения '
             'как у продавца, в ведомость и средние по команде не входит</option></select></td></tr>',
             '<tr><td>Если в 1С записан по-другому</td><td><input type=text name=alias style="width:260px"></td></tr>',
             '</table><button class=btn>Добавить и выдать код подключения</button></form>',
             '<form method=post action="/sellers">',
             f'<h2>Ждут подключения к боту ({len(waiting)})</h2>']
        if waiting:
            b.append('<table><tr><th>Продавец (как в 1С)</th><th>Обращение</th><th>Роль</th><th>Первый рабочий день</th>'
                     '<th>Код</th><th></th></tr>')
            for s in waiting:
                first, _ = db.shift_range(s["id"])
                code = (f'<a href="/sellers/{s["id"]}/card">{s["invite_code"]} — карточка с QR</a>' if s["invite_code"]
                        else f'<button class=btn name=action value="invite:{s["id"]}">Выдать код</button>')
                b.append(f'<tr><td>{html.escape(s["name"])}</td>'
                         f'<td><input type=text name="fn{s["id"]}" value="{html.escape(s["first_name"] or "")}"></td>'
                         f'<td>{_role_select(s)}</td><td class=n>{_date(first or s["hired_at"])}</td>'
                         f'<td>{code}</td><td>{_fire_btn(s)}</td></tr>')
            b.append('</table>')
        else:
            b.append('<p class=note>Все работающие продавцы подключены.</p>')
        b.append(f'<h2>Работают ({len(working)})</h2>')
        if working:
            b.append('<table><tr><th>Продавец (как в 1С)</th><th>Обращение</th><th>Роль</th><th>Последний рабочий день</th>'
                     '<th>MAX</th><th></th></tr>')
            for s in working:
                _, last = db.shift_range(s["id"])
                b.append(f'<tr><td>{html.escape(s["name"])}</td>'
                         f'<td><input type=text name="fn{s["id"]}" value="{html.escape(s["first_name"] or "")}"></td>'
                         f'<td>{_role_select(s)}</td><td class=n>{_date(last)}</td>'
                         f'<td>✅ подключён{_max_name(s)}{_training_mark(s)} <button class=btn name=action value="invite:{s["id"]}" '
                         'title="Если сменил телефон или аккаунт MAX">Новый код</button></td>'
                         f'<td>{_fire_btn(s)}</td></tr>')
            b.append('</table>')
        if waiting or working:
            b.append('<button class=btn>Сохранить обращение и роли</button>')
        b.append(f'<h2>Уволены ({len(fired)})</h2>')
        if fired:
            b.append('<table><tr><th>Продавец</th><th>Уволен</th><th>Последний рабочий день</th><th></th></tr>')
            for s in fired:
                _, last = db.shift_range(s["id"])
                b.append(f'<tr><td>{html.escape(s["name"])}</td><td class=n>{_date(s["fired_at"])}</td>'
                         f'<td class=n>{_date(last)}</td><td><button class=btn name=action '
                         f'value="restore:{s["id"]}">Вернуть на работу</button></td></tr>')
            b.append('</table>')
        b.append('</form><p class=note>Уволенный не получает сообщений, не участвует в планах и целях со следующего '
                 'дня; его продажи остаются в выторге магазина, а за текущий месяц — в итогах и ведомости.</p>')
        return page("Продавцы: подключение и увольнение", "".join(b))

    @app.post("/sellers", dependencies=[Depends(auth)])
    async def sellers_save(request: Request):
        form = await request.form()
        action = str(form.get("action") or "")
        if ":" in action:
            what, sid = action.split(":", 1)
            sid = int(sid)
            s = db.seller(sid)
            if not s:
                raise HTTPException(404)
            if what == "invite":
                db.new_invite(sid)
                return RedirectResponse(f"/sellers/{sid}/card", status_code=303)
            if what == "fire":
                db.fire(sid)
                return RedirectResponse(f"/sellers?msg={quote(s['name'] + ' уволен(а), бот отключён')}", status_code=303)
            if what == "restore":
                db.restore(sid)
                db.new_invite(sid)
                return RedirectResponse(f"/sellers/{sid}/card", status_code=303)
        with db.tx() as c:
            for s in db.sellers():
                sid = s["id"]
                role = form.get(f"role{sid}", s["role"])
                c.execute("UPDATE sellers SET first_name=?, role=? WHERE id=?",
                          (form.get(f"fn{sid}") or s["first_name"], role if role in ROLES else s["role"],
                           sid))
        return RedirectResponse("/sellers?msg=" + quote("Сохранено"), status_code=303)

    @app.post("/sellers/add", dependencies=[Depends(auth)])
    def sellers_add(name: str = Form(...), first_name: str = Form(""), role: str = Form("seller"),
                    alias: str = Form("")):
        try:
            sid = db.add_seller(name, first_name.strip() or None, role if role in ROLES else "seller")
        except ValueError as e:
            return RedirectResponse("/sellers?msg=" + quote(f"Не добавлен: {e}. Если он уволен — «Вернуть на работу»."),
                                    status_code=303)
        alias = " ".join(alias.split())
        if alias and not db.seller_id(alias, create=False):
            with db.tx() as c:
                c.execute("INSERT INTO seller_aliases(alias, seller_id) VALUES(?,?)", (alias, sid))
        db.new_invite(sid)
        return RedirectResponse(f"/sellers/{sid}/card", status_code=303)

    # ------------------------------------------------------------------ имена в 1С ↔ продавец ↔ MAX
    def names_link() -> str:
        n = len(names.suspects(db))
        warn = f' <b class=off>— похоже, есть повторы: {n}</b>' if n else ""
        return f'<p class=links><a href="/sellers/names">Имена в 1С и в боте — сопоставление →</a></p>' + \
            (f"<p>{warn}</p>" if warn else "")

    @app.get("/sellers/names", dependencies=[Depends(auth)])
    def names_page(msg: str = ""):
        rows = db.sellers(active_only=False)
        b = [f'<p class=note>{html.escape(msg)}</p>' if msg else "",
             '<p>Как бот понимает, чей это план: <b>имя в 1С</b> → продавец в этом списке → его личный QR-код → '
             '<b>аккаунт MAX</b>. Продажи из 1С и история из ведомостей зарплаты собираются по имени продавца, '
             'поэтому один человек должен быть здесь <b>одной строкой</b>.</p>',
             '<p class=note>Если в 1С или в ведомостях человек записан по-другому (другой порядок слов, отчество, '
             'опечатка, новая фамилия) — добавьте это написание в его строку. «Фамилия Имя МСЛ» в 1С — тот же '
             'продавец в магазине МСЛ, добавлять не нужно.</p>']
        sus = names.suspects(db)
        if sus:
            b.append('<h2>Похоже, это один человек</h2><form method=post action="/sellers/names"><table>'
                     '<tr><th>Останется (как в 1С)</th><th>Будет другим написанием</th><th></th></tr>')
            for keep, drop in sus:
                b.append(f'<tr><td>{html.escape(keep["name"])}</td><td>{html.escape(drop["name"])}</td><td>'
                         f'<button class=btn name=action value="merge:{keep["id"]}:{drop["id"]}" '
                         f'onclick="return confirm(\'Объединить? Продажи, история, план и подключение к боту '
                         f'перейдут к одному продавцу.\')">Объединить</button>'
                         f'<button class=btn name=action value="notdup:{keep["id"]}:{drop["id"]}">Это разные люди'
                         '</button></td></tr>')
            b.append('</table></form>')
        b.append('<h2>Все продавцы</h2><form method=post action="/sellers/names"><table>'
                 '<tr><th>Имя в 1С</th><th>Другие написания (1С, ведомости)</th><th>Продажи в 1С</th>'
                 '<th>Ведомости</th><th>Аккаунт MAX</th><th>Это тот же человек, что…</th></tr>')
        opts = "".join(f'<option value={s["id"]}>{html.escape(s["name"])}</option>' for s in rows)
        for s in rows:
            sid = s["id"]
            al = "".join(f'{html.escape(x)} <button class=btn name=action value="unalias:{html.escape(x)}" '
                         'title="Убрать это написание" style="padding:0 6px">×</button><br>' for x in db.aliases(sid))
            first, last = db.shift_range(sid)
            hist = db.one("SELECT COUNT(*) n, MIN(month) a, MAX(month) b FROM monthly_history WHERE seller_id=?", sid)
            hist_s = f'{hist["n"]} мес. ({hist["a"]} – {hist["b"]})' if hist["n"] else "—"
            mx = ("✅ " + html.escape(s["max_name"] or "подключён")) if s["max_user_id"] else "не подключён"
            fired = ' <span class=note>(уволен)</span>' if not s["active"] else ""
            b.append(f'<tr><td>{html.escape(s["name"])}{fired}</td>'
                     f'<td>{al}<input type=text name="alias{sid}" placeholder="другое написание" style="width:180px">'
                     f'<button class=btn name=action value="alias:{sid}">Добавить</button></td>'
                     f'<td class=n>{f"{_date(first)} – {_date(last)}" if first else "—"}</td><td class=n>{hist_s}</td><td>{mx}</td>'
                     f'<td><select name="same{sid}"><option value="">—</option>{opts}</select>'
                     f'<button class=btn name=action value="same:{sid}" onclick="return confirm(\'Объединить '
                     f'двух продавцов в одного?\')">Объединить</button></td></tr>')
        b.append('</table></form><p class=links><a href="/sellers">← к списку продавцов</a></p>')
        return page("Имена в 1С и в боте", "".join(b))

    @app.post("/sellers/names", dependencies=[Depends(auth)])
    async def names_save(request: Request):
        form = await request.form()
        action = str(form.get("action") or "")
        kind, _, arg = action.partition(":")
        msg = ""
        try:
            if kind == "alias":
                alias = " ".join(str(form.get(f"alias{arg}") or "").split())
                db.add_alias(int(arg), alias)
                msg = f"«{alias}» теперь считается тем же продавцом, что «{db.seller(int(arg))['name']}»."
            elif kind == "unalias":
                db.remove_alias(arg)
                msg = f"Написание «{arg}» убрано."
            elif kind in ("merge", "same", "notdup"):
                if kind == "same":
                    other = str(form.get(f"same{arg}") or "")
                    if not other or other == arg:
                        raise ValueError("выберите, с кем объединить")
                    keep, drop = names.main_of(db, db.seller(int(arg)), db.seller(int(other)))
                else:
                    a, bb = map(int, arg.split(":"))
                    keep, drop = db.seller(a), db.seller(bb)
                if kind == "notdup":
                    names.not_duplicate(db, keep["id"], drop["id"])
                    msg = f"Хорошо: «{keep['name']}» и «{drop['name']}» — разные люди."
                else:
                    both = db.merge_sellers(keep["id"], drop["id"])
                    msg = (f"Объединено: «{drop['name']}» теперь другое написание «{keep['name']}». "
                           "Продажи, история, план и подключение к боту — у одного продавца.")
                    if both:
                        msg += f" В месяцах {', '.join(both)} история была у обоих — взята большая сумма."
                    msg += " Если план этого месяца уже зафиксирован — проверьте его в «Настройки месяца»."
                    redone = service.refresh_results(db, cfg)  # итоги прошлых месяцев — по объединённым продажам (Д3)
                    if redone:
                        msg += f" Итоги {', '.join(redone)} пересчитаны — проверьте «Итоги и ведомость»."
        except ValueError as e:
            msg = f"Не сделано: {e}"
        return RedirectResponse(f"/sellers/names?msg={quote(msg)}", status_code=303)

    @app.get("/sellers/{sid}/card", dependencies=[Depends(auth)])
    def seller_card(sid: int):
        s = db.seller(sid)
        if not s:
            raise HTTPException(404)
        name = html.escape(s["first_name"] or s["name"])
        back = '<p class=links><a href="">Проверить, подключился ли</a><a href="/sellers">← к списку продавцов</a></p>'
        if s["max_user_id"] and cfg.training_token and not s["training_user_id"] and s["active"]:
            code = s["training_code"] or db.new_training_code(sid)
            b = [f'<p>✅ Бот «План» подключён. Осталось подключить бот «Обучение» — уроки, тесты и база знаний.</p>',
                 '<p class=note>Бот «План» уже прислал продавцу кнопку «Подключить обучение». Если кнопки нет — '
                 'покажите этот QR-код.</p>']
            if cfg.max_training.bot_link:
                link = f"{cfg.max_training.bot_link}?start={code}"
                b += [f'<div class=qr>{qr_svg(link)}</div>', f'<p class=note>Ссылка: {html.escape(link)}</p>']
            b += [f'<p>Или найдите бот «Обучение» в MAX и отправьте ему код <b class=code>{code}</b>.</p>', back]
            return page(f"Подключение к обучению: {html.escape(s['name'])}", "".join(b))
        if s["max_user_id"]:
            both = " и к боту «Обучение»" if cfg.training_token else ""
            body = (f'<p>✅ <b>{html.escape(s["name"])}</b> подключён к боту «План»{both}{_max_name(s)}. Утром после '
                    'следующего рабочего дня придёт первое сообщение.</p><p class=links><a href="/sellers">← к списку '
                    'продавцов</a></p>')
            return page(f"Подключение: {html.escape(s['name'])}", body)
        if not s["active"] or not s["invite_code"]:
            return RedirectResponse("/sellers", status_code=303)
        code = s["invite_code"]
        b = [f'<p>Покажите эту страницу продавцу <b>{name}</b> или распечатайте её.</p>']
        if cfg.max.bot_link:
            link = f"{cfg.max.bot_link}?start={code}"
            b += [f'<div class=qr>{qr_svg(link)}</div>',
                  '<ol><li>Наведите камеру телефона на QR-код (или откройте ссылку ниже).</li>',
                  '<li>Откроется бот в MAX — нажмите «Начать».</li>',
                  f'<li>Бот ответит «Готово, {name}!» — подключение закончено.</li></ol>',
                  f'<p class=note>Ссылка: {html.escape(link)}</p>',
                  f'<p>Если QR-код не открывается — найдите бота в MAX и отправьте ему код <b class=code>{code}</b>.</p>']
        else:
            b += ['<ol><li>Откройте в MAX бота магазина.</li>',
                  f'<li>Отправьте ему код <b class=code>{code}</b>.</li>',
                  f'<li>Бот ответит «Готово, {name}!» — подключение закончено.</li></ol>',
                  '<p class=note>Чтобы здесь был QR-код, подключите бота на странице <a href="/bot">«Бот MAX»</a>.</p>']
        if cfg.training_token:
            b.append('<p class=note>После подключения бот «План» пришлёт кнопку «Подключить обучение» — нажать её '
                     'и «Начать» во втором боте.</p>')
        b.append(back)
        return page(f"Подключение: {html.escape(s['name'])}", "".join(b))

    return app
