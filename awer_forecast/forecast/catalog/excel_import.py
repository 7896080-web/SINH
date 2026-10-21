"""Импорт справочника из Excel-карты мэппинга (раздел 3 ТЗ: начальная загрузка — из уже
существующего файла сопоставления, не с нуля через API).

Точный макет файла у владельца свой, поэтому колонки ищутся по названиям (синонимы
ниже, регистр/пробелы/«ё» не важны), а строка заголовков — среди первых строк листа.
Обязательные колонки: артикул, размер, баркод. Остальные — если есть.

Одна строка = один цветоразмерный SKU и его баркоды: в ячейке баркода их может быть
несколько (через запятую, точку с запятой, пробел или перенос), и один SKU может
встречаться в нескольких строках — все баркоды уходят в общий пул `barcode_pool`.

Баркод, уже привязанный к ДРУГОМУ размеру, не перепривязывается молча: это конфликт,
он попадает в отчёт. Повторный импорт того же файла ничего не меняет.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date

from openpyxl import load_workbook
from sqlalchemy import and_, select, update
from sqlalchemy.engine import Connection

from forecast.catalog.resolver import normalize_code
from forecast.db.schema import (
    t_barcode_pool, t_product_models, t_product_sizes, t_suppliers,
)

DEFAULT_CATEGORY = "Без категории"
HEADER_SCAN_ROWS = 15

# поле → варианты заголовка (после нормализации: нижний регистр, без пробелов и знаков)
HEADER_SYNONYMS: dict[str, tuple[str, ...]] = {
    "article": ("артикул", "артикулмодели", "article", "артикул1с"),
    "color": ("цвет", "color", "colour", "цветмодели"),
    "size": ("размер", "size", "размерпродавца", "размер1с"),
    "barcode": ("баркод", "баркоды", "штрихкод", "штрихкоды", "barcode", "barcodes", "ean", "gtin"),
    "seller_sku": ("sellersku", "артикулпродавца", "externalsku", "skuпродавца", "skuселлера"),
    "lamoda_sku": ("lamodasku", "skulamoda", "skuламода", "ламодаsku", "lamodaskuразмера", "размерныйskulamoda"),
    "lamoda_parent_sku": ("lamodaparentsku", "parentsku", "skuмоделиlamoda", "lamodaskuмодели"),
    "wb_sku": ("wbsku", "nmid", "артикулwb", "wb"),
    "ozon_sku": ("ozonsku", "артикулozon", "offerid", "ozon"),
    "category": ("категория", "category"),
    "subcategory": ("подкатегория", "subcategory"),
    "season": ("сезон", "season"),
    "supplier": ("поставщик", "supplier"),
    "production_days": ("срокпроизводства", "срокпроизводствадн", "срокпроизводствадней", "productiondays"),
    "delivery_days": ("срокдоставки", "срокдоставкидн", "срокдоставкидней", "deliverydays"),
    "moq": ("moq", "минпартия", "минимальнаяпартия"),
    "available_to_reorder": ("доступенкзаказу", "доступенкзаказуупоставщика",
                             "можнодозаказать", "availabletoreorder"),
    "status": ("статус", "status"),
    "collection_type": ("типколлекции", "коллекция", "collectiontype"),
}
REQUIRED = ("article", "size", "barcode")
STATUS_VALUES = {
    "новинка": "new", "new": "new", "активный": "active", "active": "active",
    "затухающий": "fading", "fading": "fading", "снят": "discontinued", "discontinued": "discontinued",
}
COLLECTION_VALUES = {"базовая": "base", "base": "base", "капсульная": "capsule", "capsule": "capsule"}


def normalize_header(value) -> str:
    text = str(value or "").lower().replace("ё", "е")
    return re.sub(r"[^0-9a-zа-я]", "", text)


def parse_bool_ru(value) -> bool | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    text = str(value).strip().lower()
    if text in ("да", "yes", "true", "1", "+", "истина", "д", "y"):
        return True
    if text in ("нет", "no", "false", "0", "-", "ложь", "н", "n"):
        return False
    return None


def split_barcodes(value) -> list[str]:
    if value is None:
        return []
    if isinstance(value, (int, float)):
        code = normalize_code(value)
        return [code] if code else []
    parts = re.split(r"[,;\s]+", str(value))
    return [normalize_code(p) for p in parts if normalize_code(p)]


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _int_or_none(value) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(float(str(value).replace(",", ".")))
    except ValueError:
        return None


@dataclass
class ImportReport:
    sheet: str = ""
    header_row: int = 0
    columns: dict[str, str] = field(default_factory=dict)
    rows_total: int = 0
    rows_imported: int = 0
    models_created: int = 0
    models_updated: int = 0
    sizes_created: int = 0
    sizes_updated: int = 0
    barcodes_added: int = 0
    barcodes_known: int = 0
    suppliers_created: list[str] = field(default_factory=list)
    conflicts: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        return {k: v for k, v in self.__dict__.items()}


class CatalogFileError(Exception):
    """Файл не подходит целиком (не .xlsx, нет обязательных колонок)."""


def _find_header(ws) -> tuple[int, dict[str, int]]:
    """Первая строка среди HEADER_SCAN_ROWS, где нашлись все обязательные колонки."""
    lookup = {syn: name for name, syns in HEADER_SYNONYMS.items() for syn in syns}
    best: tuple[int, dict[str, int]] = (0, {})
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS, values_only=True), start=1):
        found: dict[str, int] = {}
        for col_idx, value in enumerate(row):
            name = lookup.get(normalize_header(value))
            if name and name not in found:
                found[name] = col_idx
        if all(r in found for r in REQUIRED):
            return row_idx, found
        if len(found) > len(best[1]):
            best = (row_idx, found)
    return 0, best[1]


def read_catalog_rows(file_bytes: bytes) -> tuple[list[tuple[int, dict]], ImportReport]:
    if not file_bytes:
        raise CatalogFileError("Файл пустой.")
    try:
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:  # BadZipFile, InvalidFileException и т.п. — файл не .xlsx
        raise CatalogFileError(f"Файл не читается как .xlsx ({exc.__class__.__name__}). "
                               "Сохраните карту в формате Excel (.xlsx).") from exc
    report = ImportReport()
    try:
        chosen = None
        partial: dict[str, int] = {}
        for ws in wb.worksheets:
            header_row, columns = _find_header(ws)
            if header_row:
                chosen = (ws, header_row, columns)
                break
            if len(columns) > len(partial):
                partial = columns
        if chosen is None:
            missing = [r for r in REQUIRED if r not in partial]
            names = {"article": "Артикул", "size": "Размер", "barcode": "Баркод"}
            raise CatalogFileError(
                "Не нашёл обязательные колонки: " + ", ".join(names[m] for m in missing)
                + ". Заголовки ищутся в первых строках листа (синонимы: Штрихкод, EAN, GTIN…)."
            )
        ws, header_row, columns = chosen
        report.sheet = ws.title
        report.header_row = header_row
        headers = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
        report.columns = {name: str(headers[idx]) for name, idx in columns.items()}
        rows = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            if row is None or all(v is None or str(v).strip() == "" for v in row):
                continue
            rows.append((row_idx, {name: (row[idx] if idx < len(row) else None) for name, idx in columns.items()}))
        report.rows_total = len(rows)
        return rows, report
    finally:
        wb.close()


def _get_or_create_supplier(conn: Connection, name: str, values: dict, report: ImportReport,
                            default_lead_days: int) -> int:
    row = conn.execute(select(t_suppliers.c.supplier_id).where(t_suppliers.c.name == name)).first()
    production = _int_or_none(values.get("production_days"))
    delivery = _int_or_none(values.get("delivery_days"))
    moq = _int_or_none(values.get("moq"))
    if row:
        changes = {k: v for k, v in (("production_days", production), ("delivery_days", delivery), ("moq", moq))
                   if v is not None}
        if changes:
            conn.execute(update(t_suppliers).where(t_suppliers.c.supplier_id == row[0]).values(**changes))
        return row[0]
    if production is None and delivery is None:
        report.suppliers_created.append(f"{name} (сроки не указаны — поставлено {default_lead_days} дн., проверьте)")
        production, delivery = default_lead_days, 0
    else:
        report.suppliers_created.append(name)
    result = conn.execute(t_suppliers.insert().values(
        name=name, production_days=production or 0, delivery_days=delivery or 0, moq=moq or 1,
    ))
    return result.inserted_primary_key[0]


def import_catalog(conn: Connection, file_bytes: bytes, *, default_supplier_lead_days: int = 60) -> ImportReport:
    """Импортирует карту в текущей транзакции `conn`. Строки с ошибками пропускаются и
    попадают в отчёт; файл без обязательных колонок отклоняется целиком (CatalogFileError)."""
    rows, report = read_catalog_rows(file_bytes)

    models: dict[tuple[str, str], int] = {}
    for psid_row in conn.execute(select(t_product_models.c.product_model_id, t_product_models.c.article,
                                        t_product_models.c.color)):
        models[(psid_row.article, psid_row.color or "")] = psid_row.product_model_id
    sizes: dict[tuple[int, str], int] = {
        (r.product_model_id, r.size_label): r.product_size_id
        for r in conn.execute(select(t_product_sizes.c.product_size_id, t_product_sizes.c.product_model_id,
                                     t_product_sizes.c.size_label))
    }
    size_lamoda = {r.product_size_id: r.lamoda_sku for r in conn.execute(
        select(t_product_sizes.c.product_size_id, t_product_sizes.c.lamoda_sku))}
    pool: dict[str, tuple[int, str | None]] = {
        r.barcode: (r.product_size_id, r.seller_sku)
        for r in conn.execute(select(t_barcode_pool.c.barcode, t_barcode_pool.c.product_size_id,
                                     t_barcode_pool.c.seller_sku))
    }
    size_names = {r.product_size_id: f"{r.article} / {r.color or '—'} / {r.size_label}" for r in conn.execute(
        select(t_product_sizes.c.product_size_id, t_product_sizes.c.size_label,
               t_product_models.c.article, t_product_models.c.color)
        .join(t_product_models, t_product_models.c.product_model_id == t_product_sizes.c.product_model_id))}
    suppliers_cache: dict[str, int] = {}
    updated_models: set[int] = set()
    updated_sizes: set[int] = set()

    for row_idx, values in rows:
        article = _text(values.get("article"))
        color = _text(values.get("color"))
        size = _text(values.get("size"))
        barcodes = split_barcodes(values.get("barcode"))
        if not article or not size:
            report.errors.append(f"строка {row_idx}: пустой артикул или размер — пропущена")
            continue
        if not barcodes:
            report.errors.append(f"строка {row_idx}: нет баркода ({article} / {color or '—'} / {size}) — пропущена")
            continue

        model_values: dict = {}
        category = _text(values.get("category"))
        if category:
            model_values["category"] = category
        for field_name in ("subcategory", "season", "wb_sku", "ozon_sku"):
            text = _text(values.get(field_name))
            if text:
                model_values[field_name] = text
        parent_sku = normalize_code(values.get("lamoda_parent_sku"))
        if parent_sku:
            model_values["lamoda_sku"] = parent_sku
        available = parse_bool_ru(values.get("available_to_reorder"))
        if available is not None:
            model_values["available_to_reorder"] = available
        status = STATUS_VALUES.get(_text(values.get("status")).lower())
        if status:
            model_values["status"] = status
        collection = COLLECTION_VALUES.get(_text(values.get("collection_type")).lower())
        if collection:
            model_values["collection_type"] = collection
        supplier_name = _text(values.get("supplier"))
        if supplier_name:
            if supplier_name not in suppliers_cache:
                suppliers_cache[supplier_name] = _get_or_create_supplier(
                    conn, supplier_name, values, report, default_supplier_lead_days)
            model_values["supplier_id"] = suppliers_cache[supplier_name]

        model_key = (article, color)
        model_id = models.get(model_key)
        if model_id is None:
            result = conn.execute(t_product_models.insert().values(
                article=article, color=color or None,
                **{"category": DEFAULT_CATEGORY, **model_values}))
            model_id = result.inserted_primary_key[0]
            models[model_key] = model_id
            report.models_created += 1
        elif model_values:
            current = conn.execute(select(*[t_product_models.c[k] for k in model_values])
                                   .where(t_product_models.c.product_model_id == model_id)).first()
            changes = {k: v for k, v in model_values.items() if current._mapping[k] != v}
            if changes:
                conn.execute(update(t_product_models)
                             .where(t_product_models.c.product_model_id == model_id).values(**changes))
                if model_id not in updated_models:
                    updated_models.add(model_id)
                    report.models_updated += 1

        lamoda_sku = normalize_code(values.get("lamoda_sku"))
        size_key = (model_id, size)
        psid = sizes.get(size_key)
        if psid is None:
            result = conn.execute(t_product_sizes.insert().values(
                product_model_id=model_id, size_label=size, lamoda_sku=lamoda_sku or None))
            psid = result.inserted_primary_key[0]
            sizes[size_key] = psid
            size_lamoda[psid] = lamoda_sku or None
            size_names[psid] = f"{article} / {color or '—'} / {size}"
            report.sizes_created += 1
        elif lamoda_sku and size_lamoda.get(psid) != lamoda_sku:
            conn.execute(update(t_product_sizes).where(t_product_sizes.c.product_size_id == psid)
                         .values(lamoda_sku=lamoda_sku))
            size_lamoda[psid] = lamoda_sku
            if psid not in updated_sizes:
                updated_sizes.add(psid)
                report.sizes_updated += 1

        seller_sku = normalize_code(values.get("seller_sku")) or None
        for barcode in barcodes:
            known = pool.get(barcode)
            if known is None:
                conn.execute(t_barcode_pool.insert().values(
                    product_size_id=psid, barcode=barcode, seller_sku=seller_sku, valid_from=date.today()))
                pool[barcode] = (psid, seller_sku)
                report.barcodes_added += 1
            elif known[0] != psid:
                report.conflicts.append(
                    f"строка {row_idx}: баркод {barcode} уже привязан к «{size_names.get(known[0], known[0])}», "
                    f"в файле — «{article} / {color or '—'} / {size}». Оставлен прежний."
                )
            else:
                report.barcodes_known += 1
                if seller_sku and known[1] != seller_sku:
                    conn.execute(update(t_barcode_pool).where(and_(
                        t_barcode_pool.c.barcode == barcode)).values(seller_sku=seller_sku))
                    pool[barcode] = (psid, seller_sku)
        report.rows_imported += 1
    return report
