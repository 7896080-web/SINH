"""Excel-выгрузки страниц (openpyxl)."""
from __future__ import annotations

import io
from urllib.parse import quote

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="E8EEF4")
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def build_workbook(sheets: list[tuple[str, list[str], list[list]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        for idx, header in enumerate(headers, start=1):
            letter = ws.cell(row=1, column=idx).column_letter
            longest = max([len(str(header))] + [len(str(r[idx - 1])) for r in rows if r[idx - 1] is not None])
            ws.column_dimensions[letter].width = min(max(longest + 2, 8), 60)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def xlsx_response(content: bytes, filename: str) -> Response:
    # Content-Disposition — только latin-1: кириллическое имя по RFC 5987 + ASCII-запасное.
    disposition = f"attachment; filename=\"export.xlsx\"; filename*=UTF-8''{quote(filename)}"
    return Response(content, media_type=XLSX_MIME, headers={"Content-Disposition": disposition})
