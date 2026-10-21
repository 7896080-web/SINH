"""Страницы: вход, поставка на FBO (подтверждение, сборка, Excel), закупка, импорт карты."""
import io
from datetime import timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from forecast.config import PlanningSettings, Settings
from forecast.db.schema import t_purchase_recommendations, t_shipment_recommendations, t_stock_1c, t_stock_fbo
from forecast.planning.recalc import recalc
from forecast.security import hash_password, verify_password
from forecast.jobs import finish_job, start_job
from forecast.timeutil import today_msk, utcnow
from forecast.web.queries import FBO_COLUMNS
from factories import add_size, make_xlsx
from test_planning import _seed_sales

PASSWORD = "correct horse"


def settings_for(db_url: str) -> Settings:
    return Settings(database_url=db_url, session_secret="test-secret", login="manager",
                    password_hash=hash_password(PASSWORD, iterations=1000), lamoda_base_url="http://lamoda.test",
                    lamoda_client_id="", lamoda_client_secret="", lamoda_seller_id="", lamoda_country="RU",
                    orders_backfill_days=30, orders_overlap_hours=6, onec_exchange_dirs=(), onec_sales_dir=None,
                    planning=PlanningSettings())


@pytest.fixture
def world(engine):
    """Три размера: два к отгрузке, у третьего нет остатка на ЦС."""
    as_of = today_msk()
    with engine.begin() as conn:
        ids = []
        for size, cs in (("48", 20), ("50", 20), ("52", 0)):
            psid = add_size(conn, "AW-1", "чёрный", size, [f"46000000000{size}"], lamoda_sku=f"MPAW1IN{size}")
            _seed_sales(conn, psid, {i: 1 for i in range(60)}, start=as_of - timedelta(days=60))
            conn.execute(t_stock_1c.insert().values(product_size_id=psid, snapshot_date=as_of, qty=cs))
            conn.execute(t_stock_fbo.insert().values(product_size_id=psid, warehouse_code="BYKOVO",
                                                     snapshot_date=as_of, qty=1))
            ids.append(psid)
    run_id = start_job(engine, "recalc", "test")              # номер партии — как в жизни, из job_runs
    with engine.begin() as conn:
        recalc(conn, PlanningSettings(), as_of, run_id, utcnow())
    finish_job(engine, run_id, "ok", "test")
    return ids


@pytest.fixture
def client(engine, db_url):
    from forecast.web.app import create_app
    app = create_app(settings_for(db_url), engine)
    with TestClient(app) as c:
        yield c


def login(client):
    r = client.post("/login", data={"login": "manager", "password": PASSWORD}, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/"


def recs(engine):
    with engine.connect() as conn:
        return {r.recommendation_id: r for r in conn.execute(select(t_shipment_recommendations))}


def test_password_hash_roundtrip():
    h = hash_password("секрет", iterations=1000)
    assert verify_password("секрет", h) and not verify_password("другой", h) and not verify_password("x", "garbage")


def test_pages_require_login(client):
    for path in ("/", "/fbo", "/purchase", "/catalog", "/fbo/export.xlsx"):
        r = client.get(path, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == "/login", path
    assert client.get("/health").json() == {"status": "ok"}


def test_wrong_password_is_rejected(client, monkeypatch):
    monkeypatch.setattr("forecast.web.app.time.sleep", lambda s: None)
    client.post("/login", data={"login": "manager", "password": "nope"})
    assert client.get("/fbo", follow_redirects=False).status_code == 303


def test_fbo_page_shows_only_sizes_with_central_stock(client, world):
    login(client)
    page = client.get("/fbo").text
    assert "Рекомендации по поставке с ЦС на FBO Lamoda" in page
    assert "MPAW1IN48" in page and "MPAW1IN50" in page and "MPAW1IN52" not in page
    for column in FBO_COLUMNS[:9]:
        assert column in page


def test_confirm_assemble_and_excel_contains_only_confirmed(client, engine, world):
    login(client)
    rows = recs(engine)
    rec48, rec50 = sorted(rows)                        # обе строки последнего пересчёта
    assert all(r.confirmed_qty is None for r in rows.values())        # по умолчанию пусто
    r = client.post("/fbo/save", data={f"c_{rec48}": "5", f"c_{rec50}": ""}, follow_redirects=False)
    assert r.status_code == 303
    after = recs(engine)
    assert after[rec48].status == "confirmed" and float(after[rec48].confirmed_qty) == 5
    assert after[rec48].confirmed_by == "manager" and after[rec50].status == "pending"

    r = client.post("/fbo/assemble", data={f"c_{rec48}": "5"}, follow_redirects=False)
    assembly_id = int(r.headers["location"].split("assembled=")[1])
    after = recs(engine)
    assert after[rec48].status == "sent_to_edo" and after[rec48].assembly_id == assembly_id
    assert after[rec50].status == "pending" and after[rec50].assembly_id is None

    page_xlsx = load_workbook(io.BytesIO(client.get("/fbo/export.xlsx").content)).active
    assembly = client.get(f"/fbo/assembly/{assembly_id}.xlsx")
    assert assembly.status_code == 200
    ws = load_workbook(io.BytesIO(assembly.content)).active
    header = [c.value for c in ws[1]]
    assert header == [c.value for c in page_xlsx[1]] == FBO_COLUMNS      # тот же формат, что страница
    data = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(data) == 1 and data[0][2] == "48" and data[0][8] == 5
    assert data[0][9] == "MPAW1IN48" and data[0][10] == "4600000000048"


def test_assemble_without_confirmed_rows_does_nothing(client, engine, world):
    login(client)
    r = client.post("/fbo/assemble", data={}, follow_redirects=False)
    assert r.headers["location"] == "/fbo"
    assert all(r.status == "pending" for r in recs(engine).values())


def test_invalid_quantity_is_not_saved(client, engine, world):
    login(client)
    rec = sorted(recs(engine))[0]
    client.post("/fbo/save", data={f"c_{rec}": "-3"})
    client.post("/fbo/save", data={f"c_{rec}": "2,5"})
    assert recs(engine)[rec].confirmed_qty is None


def test_assembled_rows_count_as_in_transit_in_next_recalc(client, engine, world):
    login(client)
    rec48 = sorted(recs(engine))[0]
    before = float(recs(engine)[rec48].qty_to_ship)
    client.post("/fbo/assemble", data={f"c_{rec48}": str(int(before))})
    with engine.begin() as conn:
        recalc(conn, PlanningSettings(), today_msk(), 999, utcnow() + timedelta(minutes=1))
        new = conn.execute(select(t_shipment_recommendations).where(
            t_shipment_recommendations.c.batch_id == 999,
            t_shipment_recommendations.c.product_size_id == world[0])).one()
    # собранное ещё не создано в Lamoda, но уже распределено — повторно не предлагается
    assert float(new.qty_to_ship) == 0 and "собрано у нас" in new.confidence_explanation


def test_purchase_page_mark_ordered_and_export(client, engine, world):
    login(client)
    page = client.get("/purchase").text
    assert "Рекомендации по закупке у поставщика" in page
    with engine.connect() as conn:
        rec = conn.execute(select(t_purchase_recommendations).where(
            t_purchase_recommendations.c.status == "pending")).first()
    assert rec is not None
    client.post("/purchase/save", data={f"o_{rec.recommendation_id}": "30"})
    with engine.connect() as conn:
        row = conn.execute(select(t_purchase_recommendations).where(
            t_purchase_recommendations.c.recommendation_id == rec.recommendation_id)).one()
    assert row.status == "ordered" and float(row.ordered_qty) == 30
    wb = load_workbook(io.BytesIO(client.get("/purchase/export.xlsx").content))
    assert wb.sheetnames == ["К закупке", "Без пополнения", "Уже заказано"]
    assert wb["Уже заказано"].max_row == 2


def test_catalog_import_via_page(client, engine):
    login(client)
    content = make_xlsx([["AW-9", "синий", "M", "4609999999999", None, None, None, None, None]],
                        ["Артикул", "Цвет", "Размер", "Баркод", "x", "y", "z", "w", "v"])
    r = client.post("/catalog/import", files={"file": ("карта.xlsx", content)}, follow_redirects=True)
    assert "баркодов +1" in r.text
    r = client.post("/catalog/import", files={"file": ("bad.xlsx", b"nope")}, follow_redirects=True)
    assert "не читается как .xlsx" in r.text


def test_dashboard_renders_and_runs_recalc(client, engine, world):
    login(client)
    assert "Пересчёт рекомендаций" in client.get("/").text
    r = client.post("/jobs/recalc", follow_redirects=True)
    assert r.status_code == 200
    with engine.connect() as conn:
        batches = {b for (b,) in conn.execute(select(t_shipment_recommendations.c.batch_id).distinct())}
    assert len(batches) == 2
