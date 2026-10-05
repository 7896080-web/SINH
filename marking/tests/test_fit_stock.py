"""Правка количеств одной формой и «Уменьшить до остатка 1С» (05.10.2026):
поставка 12561 — 7 строк из 38 «не хватает», править по одной было долго,
а страница после каждой правки прыгала наверх."""
from datetime import date
from decimal import Decimal

import pytest

from markapp import onec, settings
from markapp import supplies as S
from markapp.models import OnecTask, Supply, SupplyRow


def _checked_short(db):
    org = settings.lamoda_org(db)
    s = Supply(number="12599", doc_number="12599", organization_id=org.id, status="draft",
               supply_date=date(2026, 10, 6))
    s.rows = [SupplyRow(position=1, supplier_sku="A", qty=5, ean="2000000000011", price=Decimal("100")),
              SupplyRow(position=2, supplier_sku="B", qty=3, ean="2000000000028", price=Decimal("100")),
              SupplyRow(position=3, supplier_sku="C", qty=2, ean="2000000000035", price=Decimal("100"))]
    db.add(s)
    db.flush()
    task = onec.enqueue_check(db, s)
    task.status = "sent"
    db.commit()
    # баркод|ID товара|артикул|наименование|размер|цвет|остаток ЦС|нужно|статус
    check = ("2000000000011|ID1|A|a|M|red|1|5|short\n"
             "2000000000028|ID2|B|b|M|red|-2|3|short\n"
             "2000000000035|ID3|C|c|M|red|9|2|ok\n")
    onec.apply_result_text(db, f"{task.order_id}|ERROR|не хватает|SUPPLY_CHECK", check)
    db.commit()
    return s


def test_fit_to_stock_reduces_short_rows_and_checks_without_new_request(db):
    s = _checked_short(db)
    assert s.status == "draft" and [r.onec_status for r in s.rows] == ["short", "short", "ok"]
    changes = onec.fit_to_stock(db, s)
    db.commit()
    assert {r.supplier_sku: r.qty for r in s.rows} == {"A": 1, "C": 2}       # B: остаток -2 — убрана
    assert len(changes) == 2 and "строка убрана" in changes[1]
    assert s.status == "checked"
    assert db.query(OnecTask).filter(OnecTask.command == "SUPPLY_CHECK").count() == 1   # в 1С не ходили


def test_fit_to_stock_refuses_after_composition_changed(db):
    s = _checked_short(db)
    S.update_quantities(s, {s.rows[2].id: 1})          # правка после проверки
    with pytest.raises(S.SupplyError, match="состав менялся"):
        onec.fit_to_stock(db, s)


def test_fit_to_stock_needs_an_answer(db):
    org = settings.lamoda_org(db)
    s = Supply(number="12598", doc_number="12598", organization_id=org.id, status="draft")
    s.rows = [SupplyRow(position=1, supplier_sku="A", qty=5, ean="2000000000011")]
    db.add(s)
    db.commit()
    with pytest.raises(S.SupplyError, match="нет ответа 1С"):
        onec.fit_to_stock(db, s)


def test_update_quantities_in_one_go(db):
    s = _checked_short(db)
    ids = [r.id for r in s.rows]
    assert S.update_quantities(s, {ids[0]: 1, ids[1]: 0, ids[2]: 2}) == 2   # C не менялась
    assert {r.supplier_sku: r.qty for r in s.rows} == {"A": 1, "C": 2}
    assert S.update_quantities(s, {ids[0]: 1}) == 0
    with pytest.raises(S.SupplyError):
        S.update_quantities(s, {ids[0]: -1})


def test_rows_form_and_fit_button_return_to_rows(client, db):
    s = _checked_short(db)
    page = client.get(f"/supplies/{s.id}").text
    assert 'id="rows"' in page and 'form="rowsForm"' in page and "Уменьшить до остатка 1С" in page
    r = client.post(f"/supplies/{s.id}/rows", data={f"qty_{s.rows[2].id}": "1"}, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].endswith("#rows")
    db.expire_all()
    assert db.get(Supply, s.id).rows[2].qty == 1


# --- Разрешительный документ из карточки НК (решение заказчика 05.10.2026) ----------

def _card(gtin, tn, doc_attr, doc_value):
    from markapp.models import NkCard
    return NkCard(gtin=gtin, status="ok", tn_ved=tn,
                  attrs=[{"attr_name": "Код ТНВЭД", "attr_value": tn},
                         {"attr_name": "Товарный знак", "attr_value": "AWER"},
                         {"attr_name": doc_attr, "attr_value": doc_value}])


def test_permit_is_taken_from_the_nk_card():
    from markapp import codes as C
    kind, number, day = C.nk_permit(_card("04620180403734", "6201400000", "Декларация о соответствии",
                                          "ЕАЭС N RU Д-RU.РА08.В.85411/26:::2026-09-25"))
    assert (kind, number, day) == ("CONFORMITY_DECLARATION", "ЕАЭС N RU Д-RU.РА08.В.85411/26", "2026-09-25")
    kind, number, _ = C.nk_permit(_card("04620180403734", "6109100000", "Сертификат соответствия",
                                        "ЕАЭС KG 417/043.RU.02.06303:::2024-12-26"))
    assert kind == "CONFORMITY_CERTIFICATE" and number.endswith("06303")
    # Два документа — самый новый.
    _, number, _ = C.nk_permit(_card("04620180403734", "6203429000", "Декларация о соответствии",
                                     "ЕАЭС N RU Д-TR.РА03.В.54968/21:::2021-12-15;"
                                     "ЕАЭС N RU Д-RU.РА08.В.73958/26:::2026-09-24"))
    assert number.endswith("73958/26")


def test_document_uses_nk_permit_and_flags_file_mismatch(db):
    from markapp import codes as C
    from markapp.models import GtinPair
    org = settings.lamoda_org(db)
    s = Supply(number="12597", doc_number="12597", organization_id=org.id, status="moved",
               extra_headers=["Номер разрешительного документа", "Дата начала действия"])
    s.rows = [SupplyRow(position=1, supplier_sku="A", qty=1, extras=["", ""]),
              SupplyRow(position=2, supplier_sku="B", qty=1, extras=["ЕАЭС N RU Д-RU.X.1", "01.01.2026"])]
    db.add(s)
    db.add_all([GtinPair(supplier_sku="A", gtin="04620180403734", source="manual"),
                GtinPair(supplier_sku="B", gtin="04630688318072", source="manual"),
                _card("04620180403734", "6201400000", "Декларация о соответствии",
                      "ЕАЭС N RU Д-RU.РА08.В.85411/26:::2026-09-25"),
                _card("04630688318072", "6110209100", "Декларация о соответствии",
                      "ЕАЭС N RU Д-RU.РА08.В.73918/26:::2026-09-24")])
    db.commit()
    attrs, problems = C.attrs_by_sku(db, s)
    assert attrs["A"]["cert_number"].endswith("85411/26") and attrs["A"]["cert_src"] == "Нацкаталог"
    assert attrs["A"]["cert_date"] == "2026-09-25" and attrs["A"]["tnved"] == "6201400000"
    assert any(p.startswith("B: документ в файле") for p in problems)


def test_rows_export_to_excel_like_the_screen(client, db):
    import io
    import openpyxl
    s = _checked_short(db)
    r = client.get(f"/supplies/{s.id}/rows.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert "Поставка 12599" in ws["A1"].value
    assert [c.value for c in ws[4]][:5] == ["#", "Размерный артикул", "Кол-во", "Цена", "EAN"]
    assert ws["B5"].value == "A" and ws["C5"].value == 5 and ws["E5"].value == "2000000000011"
    row_b = [c.value for c in ws[6]]
    assert -2 in row_b and "не хватает" in row_b                  # остаток и ответ 1С как на экране
    assert ws.cell(row=ws.max_row, column=2).value == "Итого" and ws.cell(row=ws.max_row, column=3).value == 10
