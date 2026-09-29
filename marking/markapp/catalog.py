"""Справочник «Одежда полный» — выгрузка каталога Lamoda Seller (ТЗ, 5.2).

Лист «Upload Template»: строка 1 служебная, строка 2 — описания, строка 3 —
заголовки вида `Название #код`, данные с 4-й. Колонки ищем по КОДУ после `#`:
русские названия Lamoda меняет, коды стабильнее, а позиции тем более не
держатся.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

import openpyxl
from sqlalchemy.orm import Session

from markapp.models import CatalogItem
from markapp.timeutils import now_utc

SHEET = "Upload Template"
# код колонки -> поле справочника
COLUMNS = {
    "supplier_sku": "supplier_sku",
    "product_identifier": "ean",
    "price": "price",
    "sku": "lamoda_sku",
    "supplier_parent_sku": "parent_sku",
    "size_value": "size",
    "color_family": "color",
    "title": "title",
    "tn_ved": "tn_ved",
    "tax_class": "tax_class",
}
REQUIRED = ("supplier_sku", "product_identifier", "price")
_CODE = re.compile(r"#\s*([A-Za-z0-9_]+)\s*$")


def norm_sku(value) -> str:
    """Обрезаем края; пробелы ВНУТРИ значимы (в каталоге бывают двойные)."""
    return str(value or "").strip()


def loose_sku(value) -> str:
    """Для подсказки «может быть, вы имели в виду»: двойные пробелы схлопнуты."""
    return re.sub(r"\s+", " ", norm_sku(value))


def cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def to_price(v) -> Decimal | None:
    if v in (None, ""):
        return None
    try:
        return Decimal(str(v).replace(",", ".").replace(" ", "")).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


@dataclass
class CatalogParse:
    rows: list[dict] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def parse_catalog(data: bytes) -> CatalogParse:
    out = CatalogParse()
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        out.errors.append(f"нет листа «{SHEET}» — это не выгрузка каталога Lamoda Seller")
        return out
    ws = wb[SHEET]
    rows = ws.iter_rows(values_only=True)
    header = None
    for i, row in enumerate(rows, start=1):
        if i == 3:
            header = row
            break
    if header is None:
        out.errors.append("в файле меньше трёх строк")
        return out
    idx = {}
    for j, h in enumerate(header):
        m = _CODE.search(str(h or ""))
        if m and m.group(1) in COLUMNS and m.group(1) not in idx:
            idx[m.group(1)] = j
    missing = [c for c in REQUIRED if c not in idx]
    if missing:
        out.errors.append("нет колонок с кодами: " + ", ".join(missing))
        return out
    seen: dict[str, int] = {}
    dups: set[str] = set()
    for excel_row, row in enumerate(rows, start=4):
        get = lambda code: row[idx[code]] if code in idx and idx[code] < len(row) else None
        sku = norm_sku(get("supplier_sku"))
        if not sku:
            continue
        if sku in seen:
            dups.add(sku)
            continue
        seen[sku] = excel_row
        rec = {field_: cell_text(get(code)) for code, field_ in COLUMNS.items() if field_ != "price"}
        rec["supplier_sku"] = sku
        rec["price"] = to_price(get("price"))
        rec["_row"] = excel_row
        out.rows.append(rec)
    if dups:
        # Дубль в файле — отказ ЭТИХ строк: какую из двух считать верной, мы не знаем.
        out.rows = [r for r in out.rows if r["supplier_sku"] not in dups]
        out.errors.append(f"артикулы повторяются в файле и не загружены ({len(dups)}): "
                          + "; ".join(sorted(dups)[:5]))
    return out


@dataclass
class ImportResult:
    added: int = 0
    changed: int = 0
    unchanged: int = 0
    changes: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def import_catalog(db: Session, data: bytes) -> ImportResult:
    """Обновляет справочник по `supplier_sku`. Отсутствующие не удаляет.

    Не коммитит — это делает вызывающий вместе с записью в журнал.
    """
    parsed = parse_catalog(data)
    res = ImportResult(errors=list(parsed.errors))
    now = now_utc()
    existing = {c.supplier_sku: c for c in db.query(CatalogItem).all()}
    for rec in parsed.rows:
        item = existing.get(rec["supplier_sku"])
        if item is None:
            item = CatalogItem(supplier_sku=rec["supplier_sku"], first_seen_at=now)
            db.add(item)
            existing[item.supplier_sku] = item
            res.added += 1
            changed_fields = []
        else:
            changed_fields = []
            if item.ean != rec["ean"]:
                changed_fields.append(f"штрихкод {item.ean} → {rec['ean']}")
            if item.price != rec["price"]:
                changed_fields.append(f"цена {item.price} → {rec['price']}")
        for k, v in rec.items():
            if not k.startswith("_"):
                setattr(item, k, v)
        item.last_seen_at = now
        if changed_fields:
            item.updated_at = now
            res.changed += 1
            res.changes.append(f"{rec['supplier_sku']}: " + ", ".join(changed_fields))
        elif item.id is not None:
            res.unchanged += 1
    return res


def find(db: Session, sku: str) -> tuple[CatalogItem | None, str]:
    """Артикул в справочнике. Второе значение — подсказка, если точного нет."""
    sku = norm_sku(sku)
    item = db.query(CatalogItem).filter(CatalogItem.supplier_sku == sku).first()
    if item is not None:
        return item, ""
    loose = loose_sku(sku)
    for cand in db.query(CatalogItem.supplier_sku).all():
        if loose_sku(cand[0]) == loose:
            return None, f"похоже на «{cand[0]}» (отличаются пробелы) — проверьте артикул"
    return None, ""
