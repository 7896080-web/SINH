from datetime import date
from decimal import Decimal
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, Request, UploadFile
from fastapi.responses import RedirectResponse, Response
from sqlalchemy.orm import Session

from markapp import audit, gtin as G, nk, onec, settings, stickers, supplies as S, upd_service as U
from markapp.database import get_db
from markapp.deps import get_current_user
from markapp.flash import flash
from markapp.models import (EDITABLE_STATUSES, FboUpload, OnecTask, Organization, Supply,
                            SupplyStatus, UpdDocument, User)
from markapp.pages import render
from markapp.timeutils import parse_ru, ru, today_local

router = APIRouter()
MAX_UPLOAD = 20 * 1024 * 1024


def _back(supply_id: int) -> RedirectResponse:
    return RedirectResponse(f"/supplies/{supply_id}", status_code=303)


def _get(db: Session, supply_id: int) -> Supply:
    supply = db.get(Supply, supply_id)
    if supply is None:
        raise S.SupplyError("поставка не найдена")
    return supply


def _date_or_none(text: str) -> date | None:
    text = (text or "").strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text) if "-" in text else parse_ru(text)
    except ValueError:
        raise S.SupplyError(f"дата «{text}» не разобрана — ДД.ММ.ГГГГ")


async def _read(upload: UploadFile) -> bytes:
    data = await upload.read(MAX_UPLOAD + 1)
    if len(data) > MAX_UPLOAD:
        raise S.SupplyError("файл больше 20 МБ")
    return data


@router.get("/supplies")
def supply_list(request: Request, db: Session = Depends(get_db),
                user: User = Depends(get_current_user)):
    items = db.query(Supply).order_by(Supply.id.desc()).limit(200).all()
    return render(request, "supplies.html", user, "supplies", supplies=items,
                  totals={s.id: S.totals(s) for s in items})


@router.get("/supplies/new")
def supply_new(request: Request, db: Session = Depends(get_db),
               user: User = Depends(get_current_user)):
    return render(request, "supply_new.html", user, "supplies",
                  next_number=S.next_number(db), org=settings.lamoda_org(db))


@router.post("/supplies/new")
async def supply_create(request: Request, file: UploadFile = File(...),
                        number: str = Form(""), doc_number: str = Form(""),
                        supply_date: str = Form(""), db: Session = Depends(get_db),
                        user: User = Depends(get_current_user)):
    org = settings.lamoda_org(db)
    if org is None:
        flash(request, "Не выбран ИП для Lamoda — укажите его на странице «Организации».", "error")
        return RedirectResponse("/supplies/new", status_code=303)
    try:
        data = await _read(file)
        parsed = S.parse_input(data)
        # Номер: поле формы → B2 файла → следующий по нумерации. Форма больше не
        # подставляет номер заранее: заполненное поле молча перебивало B2.
        number = (number.strip() or parsed.number or S.next_number(db)).strip()
        # Правило Lamoda: номер поставки = номер УПД (= B3, = имена файлов).
        # Поле формы оставлено только ради старых закладок и не читается.
        doc_number = number
        d = _date_or_none(supply_date) or parsed.supply_date
        supply = S.create_supply(db, parsed, organization_id=org.id, number=number,
                                 doc_number=doc_number, supply_date=d,
                                 filename=file.filename or "", username=user.username,
                                 source_file=data)
        audit.log(db, user.username, "supply_created", f"поставка {number}",
                  f"{len(supply.rows)} строк, файл {file.filename}")
        db.commit()
    except S.SupplyError as e:
        db.rollback()
        flash(request, f"Поставка не создана: {e}", "error")
        return RedirectResponse("/supplies/new", status_code=303)
    flash(request, f"Поставка {supply.number} создана: {len(supply.rows)} строк.", "ok")
    return _back(supply.id)


@router.get("/supplies/{supply_id}")
def supply_detail(supply_id: int, request: Request, db: Session = Depends(get_db),
                  user: User = Depends(get_current_user)):
    supply = db.get(Supply, supply_id)
    if supply is None:
        return RedirectResponse("/supplies", status_code=303)
    tasks = (db.query(OnecTask).filter(OnecTask.supply_id == supply.id)
             .order_by(OnecTask.id.desc()).all())
    fbo = (db.query(FboUpload).filter(FboUpload.supply_id == supply.id)
           .order_by(FboUpload.id.desc()).first())
    upd = (db.query(UpdDocument).filter(UpdDocument.supply_id == supply.id)
           .order_by(UpdDocument.id.desc()).first())
    scheme, scheme_warn = U.effective_scheme(db, supply, upd.doc_date if upd else None)
    sticker_warn = ""
    if upd is not None:
        planned, _ = U.effective_scheme(db, supply, None)
        if planned != upd.scheme:
            sticker_warn = (f"УПД выпущен по схеме «{U.SCHEME_LABELS[upd.scheme]}», а по плановой дате "
                            f"стикеры печатались бы как «{U.SCHEME_LABELS[planned]}». Перепечатайте стикеры.")
    pending = [t for t in tasks if t.status in ("pending", "sent")]
    gmap = G.gtin_map(db)
    gtins = {r.id: gmap.get(r.supplier_sku, "") for r in supply.rows}
    from markapp.models import NkCard
    cards = {c.gtin: c for c in db.query(NkCard).filter(NkCard.gtin.in_([g for g in gtins.values() if g])).all()}
    editable = (supply.status in [s.value for s in EDITABLE_STATUSES]
                and not S.movement_sent(supply))
    problems = S.blocking_problems(supply) + (S.catalog_drift(db, supply) if editable else [])
    intro = None
    if fbo is not None and fbo.ok:
        try:
            intro = U.introduction_check(db, supply, [r.kiz for r in U.read_fbo(fbo.content).rows])
        except Exception:  # noqa: BLE001 — страница не падает из-за файла; выпуск проверит сам
            intro = None
    return render(request, "supply.html", user, "supplies", supply=supply, tasks=tasks,
                  totals=S.totals(supply), problems=problems,
                  move_problems=S.movement_problems(supply),
                  editable=editable,
                  epf_ready=onec.epf_ready(db), fbo=fbo, upd=upd, scheme=scheme, intro=intro,
                  scheme_warn=scheme_warn, sticker_warn=sticker_warn,
                  scheme_labels=U.SCHEME_LABELS, pending=pending, today=today_local(),
                  gtins=gtins, cards=cards, onec_notes=S.shared_onec_items(supply),
                  no_gtin=len({r.supplier_sku for r in supply.rows if not gtins[r.id]}),
                  organizations=db.query(Organization).filter(Organization.is_active.is_(True)).all())


ONEC_LABELS = {"ok": "хватает", "short": "не хватает", "not_found": "не найден",
               "ambiguous": "штрихкод у нескольких товаров"}


def rows_workbook(db: Session, supply) -> bytes:
    """Таблица «Строки» поставки как на экране — для склада (05.10.2026)."""
    import io

    import openpyxl
    from openpyxl.styles import Alignment, Font
    from markapp.models import STATUS_LABELS, NkCard
    gmap = G.gtin_map(db)
    gtins = {r.id: gmap.get(r.supplier_sku, "") for r in supply.rows}
    cards = {c.gtin: c for c in db.query(NkCard).filter(NkCard.gtin.in_([g for g in gtins.values() if g])).all()}
    notes = S.shared_onec_items(supply)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Поставка {supply.number}"
    t = S.totals(supply)
    ws["A1"] = (f"Поставка {supply.number} от {ru(supply.supply_date) if supply.supply_date else '—'} · "
                f"{STATUS_LABELS.get(SupplyStatus(supply.status), supply.status)}")
    ws["A1"].font = Font(bold=True, size=13)
    from markapp.templating import money
    ws["A2"] = f"Артикулов {t['articles']}, штук {t['units']}, сумма {money(t['money'])}"
    head = (["#", "Размерный артикул", "Кол-во", "Цена", "EAN"] + list(supply.extra_headers or [])
            + ["GTIN", "НК: цвет", "НК: размер", "1С: товар", "Остаток ЦС", "1С", "Замечания"])
    ws.append([])
    ws.append(head)
    for c in ws[4]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in supply.rows:
        g = gtins[r.id]
        card = cards.get(g)
        ok = card is not None and card.status == "ok"
        onec = " · ".join(x for x in (r.onec_name, r.onec_size, r.onec_color) if x)
        ws.append([r.position, r.supplier_sku, r.qty, float(r.price) if r.price is not None else None, r.ean or ""]
                  + [str(v or "") for v in (r.extras or [])]
                  + [g, card.color if ok else "", card.size if ok else "", onec, r.onec_stock,
                     ONEC_LABELS.get(r.onec_status, r.onec_status or ""),
                     "; ".join(x for x in (r.warnings, notes.get(r.id, "")) if x)])
        row = ws.max_row
        ws.cell(row=row, column=4).number_format = "# ##0.00"
        for col in (5, 6 + len(supply.extra_headers or [])):        # EAN и GTIN — текстом
            ws.cell(row=row, column=col).number_format = "@"
    ws.append([None, "Итого", t["units"], None])
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=3).font = Font(bold=True)
    widths = [5, 42, 8, 11, 16] + [14] * len(supply.extra_headers or []) + [16, 14, 10, 40, 11, 14, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "C5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/supplies/{supply_id}/rows.xlsx")
def supply_rows_xlsx(supply_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    from urllib.parse import quote

    from fastapi.responses import Response
    supply = db.get(Supply, supply_id)
    if supply is None:
        return RedirectResponse("/supplies", status_code=303)
    name = f"Поставка_{supply.number}_строки.xlsx"
    return Response(rows_workbook(db, supply),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"})


@router.get("/supplies/{supply_id}/source")
def supply_source(supply_id: int, db: Session = Depends(get_db),
                  user: User = Depends(get_current_user)):
    """Исходный файл, из которого собран черновик, — как был загружен."""
    supply = db.get(Supply, supply_id)
    if supply is None or not supply.source_file:
        return RedirectResponse(f"/supplies/{supply_id}", status_code=303)
    return _attachment(supply.source_file, supply.source_filename or f"поставка_{supply.number}.xlsx",
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@router.post("/supplies/{supply_id}/header")
def supply_header(supply_id: int, request: Request, number: str = Form(...),
                  doc_number: str = Form(""), supply_date: str = Form(""),
                  planned_upd_date: str = Form(""), scheme_choice: str = Form("auto"),
                  scheme_reason: str = Form(""), db: Session = Depends(get_db),
                  user: User = Depends(get_current_user)):
    try:
        supply = _get(db, supply_id)
        # Номер УПД не вводится отдельно — он равен номеру поставки (правило Lamoda).
        number = number.strip()
        doc_number = number
        editable = (supply.status in [s.value for s in EDITABLE_STATUSES]
                    and not S.movement_sent(supply))
        if not editable and (number != supply.number or doc_number != supply.doc_number
                             or _date_or_none(supply_date) != supply.supply_date):
            raise S.SupplyError("номер и дата поставки после отправки перемещения в 1С не меняются")
        if db.query(UpdDocument).filter(UpdDocument.supply_id == supply.id).first() \
                and doc_number != supply.doc_number:
            raise S.SupplyError("УПД уже выпущен — номер поставки (он же номер УПД) не меняется")
        errs = S.validate_numbers(db, number, doc_number, supply.id)
        if errs:
            raise S.SupplyError("; ".join(errs))
        if scheme_choice not in U.SCHEME_LABELS:
            raise S.SupplyError("неизвестная схема")
        if scheme_choice != "auto" and scheme_choice != supply.scheme_choice and not scheme_reason.strip():
            raise S.SupplyError("ручной выбор схемы договора требует причины")
        if editable and number != supply.number:
            S._remember_number(db, number)
        supply.number, supply.doc_number = number, doc_number
        supply.supply_date = _date_or_none(supply_date)
        supply.planned_upd_date = _date_or_none(planned_upd_date) or supply.supply_date
        if scheme_choice != supply.scheme_choice:
            audit.log(db, user.username, "scheme_changed", f"поставка {supply.number}",
                      f"{supply.scheme_choice} → {scheme_choice}; причина: {scheme_reason.strip()}")
            supply.scheme_choice = scheme_choice
            supply.scheme_reason = scheme_reason.strip()
        db.commit()
        flash(request, "Сохранено.", "ok")
    except S.SupplyError as e:
        db.rollback()
        flash(request, str(e), "error")
    return _back(supply_id)


def _back_rows(supply_id: int) -> RedirectResponse:
    # К таблице строк, а не в начало страницы: правят строки — туда и вернуться.
    return RedirectResponse(f"/supplies/{supply_id}#rows", status_code=303)


@router.post("/supplies/{supply_id}/rows")
async def supply_rows_qty(supply_id: int, request: Request, db: Session = Depends(get_db),
                          user: User = Depends(get_current_user)):
    """Количества всех строк одной формой: поля `qty_<id строки>`."""
    form = await request.form()
    try:
        changes = {int(k[4:]): int(str(v).strip() or "0") for k, v in form.items() if k.startswith("qty_")}
    except ValueError:
        flash(request, "Количество — целое число.", "error")
        return _back_rows(supply_id)
    try:
        supply = _get(db, supply_id)
        n = S.update_quantities(supply, changes)
        if n:
            audit.log(db, user.username, "supply_rows_edited", f"поставка {supply.number}", f"строк изменено: {n}")
        db.commit()
        flash(request, f"Изменено строк: {n}. Поставка — черновик: проверьте остаток в 1С заново." if n
              else "Количества не изменились.", "ok")
    except S.SupplyError as e:
        db.rollback()
        flash(request, str(e), "error")
    return _back_rows(supply_id)


@router.post("/supplies/{supply_id}/fit-stock")
def supply_fit_stock(supply_id: int, request: Request, db: Session = Depends(get_db),
                     user: User = Depends(get_current_user)):
    try:
        supply = _get(db, supply_id)
        changes = onec.fit_to_stock(db, supply)
        audit.log(db, user.username, "supply_fit_stock", f"поставка {supply.number}", "; ".join(changes)[:2000])
        db.commit()
        done = supply.status == SupplyStatus.checked.value
        flash(request, f"Уменьшено до остатка 1С: {len(changes)} строк. "
              + ("Поставка — «проверено в 1С», можно перемещать." if done
                 else "Поставка — черновик: есть другие препятствия, см. замечания."), "ok")
    except S.SupplyError as e:
        db.rollback()
        flash(request, str(e), "error")
    return _back_rows(supply_id)


@router.post("/supplies/{supply_id}/rows/{row_id}")
def supply_row_qty(supply_id: int, row_id: int, request: Request, qty: int = Form(...),
                   db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    try:
        supply = _get(db, supply_id)
        S.update_row_qty(supply, row_id, qty)
        db.commit()
        flash(request, "Строка изменена; поставка вернулась в черновик — проверьте остаток в 1С заново.", "ok")
    except S.SupplyError as e:
        db.rollback()
        flash(request, str(e), "error")
    return _back_rows(supply_id)


@router.post("/supplies/{supply_id}/refresh")
def supply_refresh(supply_id: int, request: Request, db: Session = Depends(get_db),
                   user: User = Depends(get_current_user)):
    try:
        supply = _get(db, supply_id)
        S.refresh_from_catalog(db, supply)
        db.commit()
        flash(request, "Штрихкоды и цены перечитаны из справочника «Одежда полный».", "ok")
    except S.SupplyError as e:
        db.rollback()
        flash(request, str(e), "error")
    return _back(supply_id)


@router.post("/supplies/{supply_id}/delete")
def supply_delete(supply_id: int, request: Request, db: Session = Depends(get_db),
                  user: User = Depends(get_current_user)):
    try:
        supply = _get(db, supply_id)
        S.ensure_editable(supply)
        if db.query(OnecTask).filter(OnecTask.supply_id == supply.id,
                                     OnecTask.command == "SUPPLY_MOVEMENT").first():
            raise S.SupplyError("по поставке уже отправлялось перемещение — удалять нельзя")
        db.query(OnecTask).filter(OnecTask.supply_id == supply.id).delete()
        audit.log(db, user.username, "supply_deleted", f"поставка {supply.number}")
        db.delete(supply)
        db.commit()
        flash(request, f"Черновик поставки {supply.number} удалён. Номер больше не выдаётся.", "ok")
        return RedirectResponse("/supplies", status_code=303)
    except S.SupplyError as e:
        db.rollback()
        flash(request, str(e), "error")
        return _back(supply_id)


@router.post("/supplies/{supply_id}/check")
def supply_check(supply_id: int, request: Request, db: Session = Depends(get_db),
                 user: User = Depends(get_current_user)):
    try:
        supply = _get(db, supply_id)
        S.ensure_editable(supply)
        problems = S.blocking_problems(supply) + S.catalog_drift(db, supply)
        if problems:
            raise S.SupplyError("; ".join(problems))
        if not onec.epf_ready(db):
            raise S.SupplyError("обработка 1С ещё не обновлена под маркировку (нет ответа на PING) — "
                                "см. «Диагностика»")
        task = onec.enqueue_check(db, supply)
        audit.log(db, user.username, "onec_check", f"поставка {supply.number}", task.order_id)
        db.commit()
        flash(request, "Проверка остатка отправлена в 1С. Ответ придёт через минуту-две — обновите страницу.", "ok")
    except (S.SupplyError, onec.OnecError) as e:
        db.rollback()
        flash(request, str(e), "error")
    return _back(supply_id)


@router.post("/supplies/{supply_id}/move")
def supply_move(supply_id: int, request: Request, db: Session = Depends(get_db),
                user: User = Depends(get_current_user)):
    try:
        supply = _get(db, supply_id)
        if supply.status != SupplyStatus.checked.value:
            raise S.SupplyError("переместить можно только после успешной проверки остатка в 1С")
        problems = (S.blocking_problems(supply) + S.catalog_drift(db, supply)
                    + S.movement_problems(supply))
        if problems:
            raise S.SupplyError("; ".join(problems))
        if not onec.epf_ready(db):
            raise S.SupplyError("обработка 1С не обновлена под маркировку")
        busy = (db.query(OnecTask).filter(OnecTask.supply_id == supply.id,
                                          OnecTask.command == "SUPPLY_MOVEMENT",
                                          OnecTask.status.in_(("pending", "sent", "timeout"))).first())
        if busy and busy.status == "timeout":
            raise S.SupplyError("перемещение уже отправлялось и зависло — сначала посмотрите документ в 1С, "
                                "затем «Повторить» у задания внизу страницы")
        if busy:
            raise S.SupplyError("перемещение уже отправлено, ждём ответа 1С")
        task = onec.enqueue_movement(db, supply)
        audit.log(db, user.username, "onec_movement", f"поставка {supply.number}", task.order_id)
        db.commit()
        flash(request, "Перемещение отправлено в 1С одним документом. Ответ — через минуту-две.", "ok")
    except (S.SupplyError, onec.OnecError) as e:
        db.rollback()
        flash(request, str(e), "error")
    return _back(supply_id)


@router.post("/supplies/{supply_id}/retry/{task_id}")
def supply_retry(supply_id: int, task_id: int, request: Request, db: Session = Depends(get_db),
                 user: User = Depends(get_current_user)):
    """Повтор зависшего перемещения. Безопасен: 1С находит уже проведённый
    документ по order_id и возвращает его номер, второго не создаёт."""
    try:
        supply = _get(db, supply_id)
        task = db.get(OnecTask, task_id)
        if task is None or task.supply_id != supply.id or task.command != "SUPPLY_MOVEMENT" \
                or task.status != "timeout":
            raise S.SupplyError("повторить можно только зависшее перемещение")
        if supply.status != SupplyStatus.checked.value:
            # Перемещена (ответ лёг на другое задание) или вернулась в черновик —
            # повтор либо лишний, либо обходит проверку остатка.
            raise S.SupplyError("повтор не нужен: поставка уже не в статусе «проверено в 1С»")
        if not onec.epf_ready(db):
            raise S.SupplyError("обработка 1С не обновлена под маркировку")
        new = onec.enqueue_movement(db, supply)
        audit.log(db, user.username, "onec_movement_retry", f"поставка {supply.number}",
                  f"после задания {task.id}: {new.order_id}")
        db.commit()
        flash(request, "Перемещение отправлено повторно. Если документ уже есть, 1С вернёт его номер.", "ok")
    except (S.SupplyError, onec.OnecError) as e:
        db.rollback()
        flash(request, str(e), "error")
    return _back(supply_id)


@router.post("/supplies/{supply_id}/fbo")
async def supply_fbo(supply_id: int, request: Request, file: UploadFile = File(...),
                     db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    try:
        supply = _get(db, supply_id)
        if supply.status not in (SupplyStatus.moved.value, SupplyStatus.upd_issued.value):
            raise S.SupplyError("выгрузку «Поставки FBO» загружают после перемещения в 1С")
        data = await _read(file)
        try:
            res = U.check_fbo_against_supply(data, supply, db)
        except Exception as e:
            raise S.SupplyError(f"выгрузка не разобрана: {e}")
        report = "\n".join(res.problems + getattr(res, "_diffs", []) + res.notes)
        db.add(FboUpload(supply_id=supply.id, filename=file.filename or "", content=data,
                         ok=res.ok, report=report, username=user.username))
        pairs_note = ""
        if res.ok:
            # Выгрузка — полный источник связки артикул ↔ GTIN (GTIN из кода строки).
            added = G.import_fbo_pairs(db, U.read_fbo(data).rows, file.filename or "", user.username)
            nk.queue_missing(db)
            pairs_note = f" Справочник GTIN: новых пар {added.added}"
            if added.conflicts:
                pairs_note += f", КОНФЛИКТОВ {len(added.conflicts)}: " + "; ".join(added.conflicts[:2])
            pairs_note += "."
        db.commit()
        if res.ok:
            flash(request, f"Выгрузка сверена с поставкой: {res.rows} строк, {res.units} шт. Можно выпускать УПД."
                  + pairs_note, "warn" if "КОНФЛИКТ" in pairs_note else "ok")
        else:
            flash(request, "Выгрузка НЕ сходится с поставкой: " + "; ".join(res.problems), "error")
    except S.SupplyError as e:
        db.rollback()
        flash(request, str(e), "error")
    return _back(supply_id)


@router.post("/supplies/{supply_id}/upd")
def supply_upd(supply_id: int, request: Request, doc_date: str = Form(...),
               totals_mode: str = Form("rows"), force: str = Form(""),
               replace_id: str = Form(""), replace_reason: str = Form(""),
               db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    try:
        supply = _get(db, supply_id)
        if supply.status not in (SupplyStatus.moved.value, SupplyStatus.upd_issued.value):
            raise S.SupplyError("УПД выпускается после перемещения в 1С")
        fbo = (db.query(FboUpload).filter(FboUpload.supply_id == supply.id)
               .order_by(FboUpload.id.desc()).first())
        if fbo is None or not fbo.ok:
            raise S.SupplyError("нет сверенной выгрузки «Поставки FBO»")
        if totals_mode not in ("rows", "reference"):
            raise S.SupplyError("режим итогов: rows или reference")
        d = _date_or_none(doc_date)
        if d is None:
            raise S.SupplyError("укажите дату УПД")
        # Ввод в оборот — ДО сборки и без «force»: «выпустить с ошибками проверки»
        # снимает замечания к XML, а это не замечание — ЧЗ откажет в передаче кодов
        # уже после того, как Lamoda подпишет УПД.
        intro = U.introduction_check(db, supply, [r.kiz for r in U.read_fbo(fbo.content).rows])
        if intro.blocking:
            raise S.SupplyError("УПД не выпущен: " + intro.blocking)
        built = U.build_for_supply(db, supply, fbo.content, d, totals_mode)
        errors = U.has_errors(built.findings)
        if errors and not force:
            raise S.SupplyError("проверка УПД нашла ошибки, документ не выпущен: "
                                + U.findings_text(built.findings)[:300])
        old = db.query(UpdDocument).filter(UpdDocument.supply_id == supply.id).all()
        # Перевыпуск — это новый документ под ТЕМ ЖЕ номером. Если прежний уже
        # ушёл в Lamoda по ЭДО, это корректировка, и решение принимает человек:
        # сервер требует причину и то, что человек видел именно этот УПД.
        # Подтверждения в браузере мало — повторная отправка формы (F5, двойной
        # клик, открытая вчера вкладка) заменила бы документ молча.
        # Метка «видел этот УПД» — ИдФайл (в нём uuid), а не id строки: SQLite
        # отдаёт id удалённой строки новой, и вчерашняя вкладка с id=1 совпала
        # бы с перевыпущенным документом, тоже получившим id=1.
        if old:
            current = max(old, key=lambda o: o.id).id_file
            if replace_id != current:
                raise S.SupplyError("УПД уже выпущен (или перевыпущен в другой вкладке) — "
                                    "обновите страницу и укажите причину перевыпуска")
            if not replace_reason.strip():
                raise S.SupplyError("перевыпуск УПД требует причины: прежний мог уже уйти в Lamoda")
        for o in old:
            audit.log(db, user.username, "upd_replaced", f"УПД {o.doc_number}",
                      f"прежний от {ru(o.doc_date)} заменён; причина: {replace_reason.strip()}")
            db.delete(o)
        db.flush()
        doc = UpdDocument(supply_id=supply.id, fbo_upload_id=fbo.id, doc_number=supply.doc_number,
                          doc_date=d, scheme=built.scheme,
                          scheme_manual=supply.scheme_choice != "auto",
                          scheme_reason=supply.scheme_reason, id_file=built.id_file, xml=built.xml,
                          positions=built.positions, total_with_vat=built.total_with_vat,
                          check_report=U.findings_text(built.findings), has_errors=errors,
                          username=user.username)
        db.add(doc)
        supply.status = SupplyStatus.upd_issued.value
        audit.log(db, user.username, "upd_issued", f"УПД {supply.doc_number}",
                  f"от {ru(d)}, схема {built.scheme}, {built.positions} поз., {built.total_with_vat}"
                  + (" — ВЫПУЩЕН С ОШИБКАМИ ПРОВЕРКИ" if errors else ""))
        db.commit()
        msg = f"УПД {supply.doc_number} выпущен: {built.positions} поз., с НДС {built.total_with_vat}."
        if built.scheme_warning:
            msg += " Внимание: " + built.scheme_warning
        if intro.warning:
            msg += " Внимание: " + intro.warning + "."
        flash(request, msg, "warn" if (built.scheme_warning or errors or intro.warning) else "ok")
    except S.SupplyError as e:
        db.rollback()
        flash(request, str(e), "error")
    return _back(supply_id)


def _attachment(data: bytes, name: str, media: str) -> Response:
    return Response(data, media_type=media, headers={
        "Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"})


@router.get("/supplies/{supply_id}/upd.xml")
def supply_upd_download(supply_id: int, db: Session = Depends(get_db),
                        user: User = Depends(get_current_user)):
    doc = (db.query(UpdDocument).filter(UpdDocument.supply_id == supply_id)
           .order_by(UpdDocument.id.desc()).first())
    if doc is None:
        return RedirectResponse(f"/supplies/{supply_id}", status_code=303)
    return _attachment(doc.xml, f"{doc.doc_number}.xml", "application/xml")


@router.post("/supplies/{supply_id}/stickers")
def supply_stickers(supply_id: int, request: Request, boxes: int = Form(...),
                    db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    try:
        supply = _get(db, supply_id)
        upd = (db.query(UpdDocument).filter(UpdDocument.supply_id == supply.id)
               .order_by(UpdDocument.id.desc()).first())
        data, scheme = stickers.build(db, supply, boxes, upd.doc_date if upd else None)
        audit.log(db, user.username, "stickers", f"поставка {supply.number}", f"{boxes} коробов, {scheme}")
        db.commit()
        return _attachment(data, stickers.filename(supply, boxes),
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except (S.SupplyError, ValueError) as e:
        db.rollback()
        flash(request, str(e), "error")
        return _back(supply_id)
