"""Справочник «размерный артикул ↔ GTIN» (ТЗ, 5.3).

Три источника, и во всех трёх конфликт (тот же артикул с другим GTIN или тот
же GTIN у другого артикула) НЕ перезаписывается молча, а показывается:
перезапись переводила бы заказ кодов на чужой товар.
1. файлы `product_gtin`, которые заказчик уже загружал в Lamoda Seller;
2. выгрузка «Поставки FBO» — GTIN берётся из КИЗ каждой строки;
3. ручное сопоставление.

Выгрузка `product_gtin` для Lamoda — только пары, которых Lamoda ещё не
получала, по правилам шаблона: GTIN ровно 14 цифр и начинается с 046, не
больше 1000 строк в файле, лист инструкций сохраняется.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from sqlalchemy.orm import Session

from markapp.catalog import cell_text, norm_sku
from markapp.models import GtinPair
from markapp.timeutils import now_utc

TEMPLATE = Path(__file__).resolve().parent / "xlsx" / "product_gtin_template.xlsx"
MAX_ROWS_PER_FILE = 1000
_GTIN14 = re.compile(r"^\d{14}$")


class GtinError(ValueError):
    pass


def check_digit_ok(gtin: str) -> bool:
    """Контрольная цифра GS1: веса 3 и 1 справа налево, начиная с 3."""
    if not _GTIN14.match(gtin):
        return False
    digits = [int(c) for c in gtin]
    total = sum(d * (3 if i % 2 == 0 else 1) for i, d in enumerate(reversed(digits[:-1])))
    return (10 - total % 10) % 10 == digits[-1]


def normalize_gtin(value) -> str:
    """GTIN из ячейки: число теряет ведущие нули — дополняем до 14 (GTIN-8/12/13
    и GTIN-14 с несколькими нулями впереди, записанный числом)."""
    text = cell_text(value)
    if text.isdigit() and 8 <= len(text) < 14:
        text = text.zfill(14)
    return text


def validate_gtin(gtin: str) -> str:
    """Пустая строка — всё в порядке, иначе причина."""
    if not _GTIN14.match(gtin):
        return f"GTIN «{gtin}»: нужно ровно 14 цифр"
    if not check_digit_ok(gtin):
        return f"GTIN «{gtin}»: неверная контрольная цифра"
    return ""


def gtin_of_code(code: str) -> str:
    """GTIN из кода маркировки (короткого или полного): `01` + 14 цифр."""
    code = (code or "").strip()
    if code.startswith("01") and len(code) >= 16 and code[2:16].isdigit():
        return code[2:16]
    return ""


@dataclass
class AddResult:
    added: int = 0
    same: int = 0
    conflicts: list[str] = field(default_factory=list)
    invalid: list[str] = field(default_factory=list)

    def merge(self, other: "AddResult") -> None:
        self.added += other.added
        self.same += other.same
        self.conflicts += other.conflicts
        self.invalid += other.invalid


class _Index:
    """Пары по артикулу и по GTIN — с учётом добавленных в этой же сессии
    (autoflush выключен: запрос их не увидел бы, и повтор внутри одного файла
    ушёл бы вторым INSERT под уникальный индекс)."""

    def __init__(self, db: Session):
        self.by_sku = {p.supplier_sku: p for p in db.query(GtinPair).all()}
        self.by_gtin = {p.gtin: p for p in self.by_sku.values()}


def add_pairs(db: Session, pairs: list[tuple[str, str]], *, source: str, source_name: str,
              exported: bool, username: str = "") -> AddResult:
    res = AddResult()
    idx = _Index(db)
    now = now_utc()
    for raw_sku, raw_gtin in pairs:
        sku, gtin = norm_sku(raw_sku), normalize_gtin(raw_gtin)
        if not sku and not gtin:
            continue
        problem = "нет артикула" if not sku else validate_gtin(gtin)
        if problem:
            res.invalid.append(f"{sku or '—'}: {problem}")
            continue
        by_sku, by_gtin = idx.by_sku.get(sku), idx.by_gtin.get(gtin)
        if by_sku is not None and by_sku.gtin == gtin:
            res.same += 1
            if exported and by_sku.exported_at is None:
                by_sku.exported_at = now     # Lamoda её уже получила
            continue
        if by_sku is not None:
            res.conflicts.append(f"{sku}: в справочнике GTIN {by_sku.gtin}, в {source_name} — {gtin}")
            continue
        if by_gtin is not None:
            res.conflicts.append(f"GTIN {gtin}: в справочнике у «{by_gtin.supplier_sku}», "
                                 f"в {source_name} — у «{sku}»")
            continue
        pair = GtinPair(supplier_sku=sku, gtin=gtin, source=source, source_name=source_name,
                        exported_at=now if exported else None, created_by=username)
        db.add(pair)
        idx.by_sku[sku] = pair
        idx.by_gtin[gtin] = pair
        res.added += 1
    return res


def read_product_gtin(data: bytes) -> list[tuple[str, str]]:
    """Файл `product_gtin`: лист «Лист1», колонки `Supplier SKU` и `Gtin`."""
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb["Лист1"] if "Лист1" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip().lower() for h in next(rows, [])]
    if "supplier sku" not in header or "gtin" not in header:
        raise GtinError("это не файл product_gtin: нет колонок «Supplier SKU» и «Gtin»")
    i_sku, i_gtin = header.index("supplier sku"), header.index("gtin")
    out = []
    for row in rows:
        if row is None:
            continue
        sku = row[i_sku] if i_sku < len(row) else None
        gtin = row[i_gtin] if i_gtin < len(row) else None
        if sku in (None, "") and gtin in (None, ""):
            continue
        out.append((cell_text(sku), gtin))
    return out


def import_product_gtin(db: Session, data: bytes, filename: str, username: str = "") -> AddResult:
    return add_pairs(db, read_product_gtin(data), source="product_gtin", source_name=filename,
                     exported=True, username=username)


def pairs_from_fbo(rows) -> list[tuple[str, str]]:
    """Строки выгрузки «Поставки FBO» (`upd_constructor.lamoda_xlsx.RawRow`):
    артикул — колонка SKU, GTIN — из кода маркировки строки."""
    seen = {}
    for r in rows:
        g = gtin_of_code(r.kiz)
        if g and r.name.strip():
            seen.setdefault((r.name.strip(), g), None)
    return list(seen)


def import_fbo_pairs(db: Session, fbo_rows, filename: str, username: str = "") -> AddResult:
    # Lamoda приняла поставку с этими кодами, но это не значит, что у неё есть
    # пара «артикул ↔ GTIN» в справочнике: по 12550 так пришли 54 новые пары.
    return add_pairs(db, pairs_from_fbo(fbo_rows), source="fbo", source_name=filename,
                     exported=False, username=username)


def pending_export(db: Session) -> list[GtinPair]:
    """Пары, которых Lamoda ещё не получала и которые пройдут правила шаблона."""
    pairs = (db.query(GtinPair).filter(GtinPair.exported_at.is_(None))
             .order_by(GtinPair.supplier_sku).all())
    return [p for p in pairs if p.gtin.startswith("046")]


def build_product_gtin_files(pairs: list[GtinPair]) -> list[bytes]:
    """Файлы по шаблону Lamoda, не больше 1000 строк в каждом."""
    files = []
    for start in range(0, len(pairs), MAX_ROWS_PER_FILE):
        wb = openpyxl.load_workbook(TEMPLATE)
        ws = wb["Лист1"]
        for i, p in enumerate(pairs[start:start + MAX_ROWS_PER_FILE], start=2):
            ws.cell(i, 1, p.supplier_sku)
            c = ws.cell(i, 2, p.gtin)
            c.number_format = "@"     # текст: число потеряло бы ведущий ноль
        buf = io.BytesIO()
        wb.save(buf)
        files.append(buf.getvalue())
    return files


def mark_exported(pairs: list[GtinPair]) -> None:
    now = now_utc()
    for p in pairs:
        p.exported_at = now


def gtin_map(db: Session) -> dict[str, str]:
    return {p.supplier_sku: p.gtin for p in db.query(GtinPair).all()}
