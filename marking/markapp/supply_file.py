"""Файл поставки для Lamoda с кодами (ТЗ, 6.3).

Заполняется ОФИЦИАЛЬНЫЙ шаблон `xlsx/lamoda_fulfilment_shipment_template.xlsx`
(лист инструкций сохраняется), а не рисуется свой. Развёртка: 1 строка =
1 артикул = 1 КИЗ; повторы строки поставки идут подряд, в её порядке; все
колонки как во входящей строке, кроме «Количество» = 1 и своего кода.

Запись — как в файле kiz-tool по 12550, который Lamoda приняла
(`tests/fixtures/reference_kiz_tool_supply_12550.xlsx`): дата B1 — число
Excel, номер — текст, количество — число, цена/EAN/код — текст; GS внутри
кода — `_x001D_` (сам символ openpyxl писать отказывается).
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
from sqlalchemy.orm import Session

from markapp.crypto import decrypt_value
from markapp.labels import FULL_RE, GS, normalize
from markapp.models import GtinPair, MarkCode, Supply, SupplyStatus

TEMPLATE = Path(__file__).resolve().parent / "xlsx" / "lamoda_fulfilment_shipment_template.xlsx"
SHEET = "Лист1"
FIRST_ROW = 5
READY_STATUSES = (SupplyStatus.moved.value, SupplyStatus.upd_issued.value, SupplyStatus.accepted.value)
EXCEL_EPOCH = date(1899, 12, 30)


class SupplyFileError(Exception):
    pass


def _excel_serial(d: date) -> int:
    return (d - EXCEL_EPOCH).days


def _price_text(p: Decimal | None) -> str:
    if p is None:
        return ""
    p = Decimal(p)
    return str(int(p)) if p == p.to_integral_value() else format(p.normalize(), "f")


def codes_by_row(db: Session, supply: Supply) -> list[tuple]:
    """[(строка поставки, [полные коды])] с проверками ТЗ 6.3. Коды артикула,
    встречающегося в нескольких строках, раздаются строкам по порядку."""
    if supply.status not in READY_STATUSES:
        raise SupplyFileError("файл поставки выдаётся после перемещения в 1С")
    codes = db.query(MarkCode).filter(MarkCode.supply_id == supply.id).order_by(MarkCode.id).all()
    if not codes:
        raise SupplyFileError("у поставки нет кодов")
    not_in = [c for c in codes if c.status != "INTRODUCED"]
    if not_in:
        raise SupplyFileError(f"не в обороте кодов: {len(not_in)} — файл для Lamoda только с кодами «в обороте»")
    pairs = {p.supplier_sku: p.gtin for p in db.query(GtinPair).all()}
    pool: dict[str, list[MarkCode]] = defaultdict(list)
    for c in codes:
        pool[c.supplier_sku].append(c)
    need: dict[str, int] = defaultdict(int)
    for r in supply.rows:
        need[r.supplier_sku] += r.qty
    problems = []
    for sku, n in need.items():
        if len(pool.get(sku, [])) != n:
            problems.append(f"{sku}: штук {n}, кодов {len(pool.get(sku, []))}")
        gtin = pairs.get(sku)
        bad = [c for c in pool.get(sku, []) if c.gtin != gtin or c.cis[2:16] != gtin]
        if bad:
            problems.append(f"{sku}: {len(bad)} кодов не с GTIN артикула ({gtin or 'нет GTIN'})")
    extra = sorted(set(pool) - set(need))
    if extra:
        problems.append(f"коды артикулов не из поставки: {', '.join(extra[:3])}")
    if problems:
        raise SupplyFileError("; ".join(problems[:5]) + (f" и ещё {len(problems) - 5}" if len(problems) > 5 else ""))
    out, seen = [], set()
    for r in supply.rows:
        mine = [pool[r.supplier_sku].pop(0) for _ in range(r.qty)]
        full = []
        for c in mine:
            code = normalize(decrypt_value(c.full_enc))
            if not FULL_RE.match(code):
                raise SupplyFileError(f"{r.supplier_sku}: код {c.cis} не полный")
            if code in seen:
                raise SupplyFileError(f"код {c.cis} встречается дважды")
            seen.add(code)
            full.append(code)
        out.append((r, full))
    return out


def build(db: Session, supply: Supply) -> bytes:
    rows = codes_by_row(db, supply)
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb[SHEET]
    if supply.supply_date is not None:
        ws["B1"] = _excel_serial(supply.supply_date)
    ws["B2"] = str(supply.number)
    ws["B3"] = str(supply.doc_number)
    for i, h in enumerate(supply.extra_headers or []):
        ws.cell(row=FIRST_ROW - 1, column=7 + i, value=str(h))
    n = FIRST_ROW
    for r, full in rows:
        for code in full:
            ws.cell(row=n, column=1, value=r.supplier_sku)
            ws.cell(row=n, column=2, value=1)
            ws.cell(row=n, column=3, value=_price_text(r.price))
            ws.cell(row=n, column=4, value=r.ean or None)
            ws.cell(row=n, column=6, value=code.replace(GS, "_x001D_"))
            for i, v in enumerate(r.extras or []):
                if v not in (None, ""):
                    ws.cell(row=n, column=7 + i, value=str(v))
            n += 1
    total = sum(len(f) for _, f in rows)
    if n - FIRST_ROW != total or total != sum(r.qty for r in supply.rows):
        raise SupplyFileError("число строк файла не равно сумме количеств")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename(supply: Supply) -> str:
    return f"{supply.doc_number}.xlsx"
