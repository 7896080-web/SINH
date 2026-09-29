"""Поставка Lamoda: входящий файл, нумерация, правка (ТЗ, разд. 4 и 6.1).

Во входящем файле нужны только размерный артикул и количество. Штрихкод и
цена берутся из справочника «Одежда полный»: цена там та, что потом уйдёт и
в файл Lamoda, и в УПД, — одна цена на строку, расхождение между документами
исключено конструктивно.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date

import openpyxl
from sqlalchemy.orm import Session

from markapp import settings
from markapp.catalog import cell_text, loose_sku, norm_sku, to_price
from markapp.models import CatalogItem, EDITABLE_STATUSES, Supply, SupplyRow, SupplyStatus
from markapp.timeutils import parse_ru

HEADER_ROW = 4
BASE_HEADERS = ("seller sku", "количество", "стоимость", "ean", "вес юи", "data matrix")
MAX_ARTICLES = 5000
MAX_UNITS = 10000
# Номер поставки и номер документа по правилам шаблона Lamoda.
NUMBER_RE = re.compile(r"^[A-Za-z0-9_-]{1,20}$")


class SupplyError(ValueError):
    pass


@dataclass
class InputRow:
    excel_row: int
    sku: str
    qty: int
    file_price: str
    file_ean: str
    extras: list


@dataclass
class InputFile:
    rows: list[InputRow] = field(default_factory=list)
    extra_headers: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    supply_date: date | None = None
    number: str = ""


def _norm_header(h) -> str:
    return str(h or "").strip().lower()


def parse_input(data: bytes) -> InputFile:
    """Шаблон `fulfilment_shipment`: A1:B3 шапка, строка 4 — заголовки."""
    out = InputFile()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:
        out.errors.append(f"файл не читается как Excel: {e}")
        return out
    ws = wb["Лист1"] if "Лист1" in wb.sheetnames else wb[wb.sheetnames[0]]
    b1, b2 = ws.cell(1, 2).value, ws.cell(2, 2).value
    if isinstance(b1, date):
        out.supply_date = b1.date() if hasattr(b1, "date") else b1
    elif isinstance(b1, str) and b1.strip():
        try:
            out.supply_date = parse_ru(b1)
        except ValueError:
            out.errors.append(f"B1: дата поставки «{b1}» — ждём ДД.ММ.ГГГГ")
    if b2 not in (None, ""):
        out.number = cell_text(b2)
    header = [c.value for c in ws[HEADER_ROW]]
    names = [_norm_header(h) for h in header]
    if not names or names[0] != "seller sku" or "количество" not in names:
        out.errors.append("строка 4 — не заголовки шаблона Lamoda (ждём «Seller SKU», «Количество» …)")
        return out
    col = {n: i for i, n in enumerate(names) if n}
    extra_idx = [i for i, n in enumerate(names) if n and n not in BASE_HEADERS]
    out.extra_headers = [str(header[i]).strip() for i in extra_idx]
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        vals = [ws.cell(r, c + 1).value for c in range(len(header))]
        if not any(v not in (None, "") for v in vals):
            continue
        sku = norm_sku(vals[col["seller sku"]])
        raw_qty = vals[col["количество"]]
        if not sku:
            out.errors.append(f"строка {r}: нет артикула")
            continue
        try:
            as_float = float(str(raw_qty).strip().replace(",", "."))
            if not as_float.is_integer():
                raise ValueError
            qty = int(as_float)
        except (ValueError, TypeError):
            out.errors.append(f"строка {r}: количество «{raw_qty}» — не целое число")
            continue
        if qty <= 0:
            out.errors.append(f"строка {r}: количество должно быть больше нуля")
            continue
        out.rows.append(InputRow(
            excel_row=r, sku=sku, qty=qty,
            file_price=cell_text(vals[col["стоимость"]]) if "стоимость" in col else "",
            file_ean=cell_text(vals[col["ean"]]) if "ean" in col else "",
            extras=[cell_text(vals[i]) for i in extra_idx],
        ))
    if not out.rows and not out.errors:
        out.errors.append("в файле нет ни одной строки товара")
    if len({r.sku for r in out.rows}) > MAX_ARTICLES:
        out.errors.append(f"артикулов больше {MAX_ARTICLES} — разбейте на несколько поставок")
    if sum(r.qty for r in out.rows) > MAX_UNITS:
        out.errors.append(f"единиц больше {MAX_UNITS} — разбейте на несколько поставок")
    return out


# --- Нумерация ----------------------------------------------------------------

def next_number(db: Session) -> str:
    """Последний номер + шаг. Номер не повторяется даже у удалённой поставки:
    последний выданный хранится в настройке, а не вычисляется по существующим."""
    last = int(settings.get(db, settings.SUPPLY_LAST_NUMBER) or 0)
    step = int(settings.get(db, settings.SUPPLY_STEP) or 10)
    taken = {n for (n,) in db.query(Supply.number).all()}
    candidate = last + step
    while str(candidate) in taken:
        candidate += step
    return str(candidate)


def _remember_number(db: Session, number: str) -> None:
    if number.isdigit() and int(number) > int(settings.get(db, settings.SUPPLY_LAST_NUMBER) or 0):
        settings.put(db, settings.SUPPLY_LAST_NUMBER, number)


def validate_numbers(db: Session, number: str, doc_number: str, supply_id: int | None = None) -> list[str]:
    errs = []
    for label, value in (("номер поставки", number), ("номер документа", doc_number)):
        if not NUMBER_RE.match(value or ""):
            errs.append(f"{label} «{value}»: до 20 символов — латиница, цифры, «_», «-»")
    q = db.query(Supply)
    if supply_id:
        q = q.filter(Supply.id != supply_id)
    if q.filter(Supply.number == number).first():
        errs.append(f"поставка с номером {number} уже есть")
    if q.filter(Supply.doc_number == doc_number).first():
        errs.append(f"номер документа {doc_number} уже занят другой поставкой")
    return errs


# --- Создание и правка --------------------------------------------------------

def _catalog_maps(db: Session):
    items = db.query(CatalogItem).all()
    exact = {c.supplier_sku: c for c in items}
    loose = {}
    for c in items:
        loose.setdefault(loose_sku(c.supplier_sku), c.supplier_sku)
    return exact, loose


def fill_row_from_catalog(row: SupplyRow, exact, loose) -> None:
    """EAN и цена — из справочника; расхождение с файлом — предупреждение."""
    warn = []
    item = exact.get(row.supplier_sku)
    if item is None:
        row.ean = ""
        row.price = None
        hint = loose.get(loose_sku(row.supplier_sku))
        warn.append("нет в справочнике «Одежда полный»"
                    + (f"; похоже на «{hint}» (отличаются пробелы)" if hint else ""))
    else:
        row.ean = item.ean
        row.price = item.price
        if not item.ean:
            warn.append("в справочнике нет штрихкода")
        if item.price is None or item.price <= 0:
            warn.append("в справочнике нет цены")
        if row.file_ean and row.file_ean != item.ean:
            warn.append(f"EAN в файле {row.file_ean}, в справочнике {item.ean}")
        fp = to_price(row.file_price)
        if fp is not None and item.price is not None and fp != item.price:
            warn.append(f"цена в файле {fp}, в справочнике {item.price}")
    row.warnings = "; ".join(warn)


def create_supply(db: Session, parsed: InputFile, *, organization_id: int, number: str,
                  doc_number: str, supply_date: date | None, filename: str, username: str) -> Supply:
    if parsed.errors:
        raise SupplyError("; ".join(parsed.errors))
    errs = validate_numbers(db, number, doc_number)
    if errs:
        raise SupplyError("; ".join(errs))
    supply = Supply(number=number, doc_number=doc_number, supply_date=supply_date,
                    planned_upd_date=supply_date, organization_id=organization_id,
                    status=SupplyStatus.draft.value, source_filename=filename,
                    extra_headers=parsed.extra_headers, created_by=username)
    db.add(supply)
    exact, loose = _catalog_maps(db)
    counts: dict[str, int] = {}
    for r in parsed.rows:
        counts[r.sku] = counts.get(r.sku, 0) + 1
    for pos, r in enumerate(parsed.rows, start=1):
        row = SupplyRow(position=pos, supplier_sku=r.sku, qty=r.qty, file_price=r.file_price,
                        file_ean=r.file_ean, extras=r.extras)
        fill_row_from_catalog(row, exact, loose)
        if counts[r.sku] > 1:
            row.warnings = "; ".join(filter(None, [row.warnings, "артикул стоит в файле не одной строкой"]))
        supply.rows.append(row)
    _remember_number(db, number)
    db.flush()
    return supply


def ensure_editable(supply: Supply) -> None:
    if supply.status not in [s.value for s in EDITABLE_STATUSES]:
        raise SupplyError("поставка уже перемещена в 1С — состав зафиксирован. "
                          "Исправление перемещения делается вручную в 1С")


def mark_edited(supply: Supply) -> None:
    """Любая правка состава возвращает поставку в черновик: прежняя проверка
    остатка относилась к другому составу."""
    supply.status = SupplyStatus.draft.value
    for row in supply.rows:
        row.onec_status = ""


def refresh_from_catalog(db: Session, supply: Supply) -> None:
    ensure_editable(supply)
    exact, loose = _catalog_maps(db)
    for row in supply.rows:
        fill_row_from_catalog(row, exact, loose)
    mark_edited(supply)


def update_row_qty(supply: Supply, row_id: int, qty: int) -> None:
    ensure_editable(supply)
    if qty < 0:
        raise SupplyError("количество не может быть отрицательным")
    for row in list(supply.rows):
        if row.id == row_id:
            if qty == 0:
                supply.rows.remove(row)
            else:
                row.qty = qty
            mark_edited(supply)
            return
    raise SupplyError("строка не найдена")


def blocking_problems(supply: Supply) -> list[str]:
    """Что мешает отправить поставку в 1С."""
    out = []
    if not supply.rows:
        out.append("в поставке нет строк")
    if supply.supply_date is None:
        out.append("не указана дата поставки")
    no_cat = [r for r in supply.rows if not r.ean or r.price is None or r.price <= 0]
    if no_cat:
        out.append(f"строк без штрихкода или цены из справочника: {len(no_cat)}")
    return out


def totals(supply: Supply) -> dict:
    units = sum(r.qty for r in supply.rows)
    money = sum((r.price or 0) * r.qty for r in supply.rows)
    return {"articles": len({r.supplier_sku for r in supply.rows}), "rows": len(supply.rows),
            "units": units, "money": money}
