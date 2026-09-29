#!/usr/bin/env python3
"""Генератор стикеров коробов Lamoda («Короб marketplace») для AWER.

Использование:
    python gen_stickers.py --boxes 27 --invoice 12350 --date 28.09.2026
    python gen_stickers.py -n 10 -i 12250 -d 23.09.2026 -o output/my.xlsx

Результат: Excel, книжная A4, 8 стикеров на странице (2x4),
стикеры в оригинальном размере образца, пустой отступ между рядами.
"""
import argparse
import copy
from datetime import date as _date
from pathlib import Path

import openpyxl
from openpyxl.styles import Border, PatternFill
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break

# ---- Постоянные реквизиты (менять здесь) -------------------------------
SENDER = "Отправитель: ИП Яворская Т.Н"
# (бренд и адрес берутся из шаблона)

# С 01.10.2026 Lamoda переводит FBO с комиссии на агентскую модель: слово «комиссия» из
# сопровождающих поставку документов убирается. Схема — по ДАТЕ ПОСТАВКИ (-d), как в УПД-конструкторе
# (там — по дате УПД). Поставки по 30.09.2026 включительно печатаются по-старому.
AGENCY_FROM = _date(2026, 10, 1)
SCHEME_MARK = {"commission": "КОМИССИЯ", "agency": "АГЕНТ"}               # строка под брендом
RECIPIENT = {"commission": 'Получатель: ООО "КУПИШУЗ" Комиссия',
             "agency": 'Получатель: ООО "КУПИШУЗ" Агент'}
MARK_ROW = 6        # строка «КОМИССИЯ»/«АГЕНТ» внутри блока стикера (B8 шаблона)


def scheme_for(date: str) -> str:
    """'commission' или 'agency' по дате поставки ДД.ММ.ГГГГ."""
    dd, mm, yy = date.strip().split(".")
    return "agency" if _date(int(yy), int(mm), int(dd)) >= AGENCY_FROM else "commission"

# ---- Вёрстка ------------------------------------------------------------
TEMPLATE = Path(__file__).parent / "template" / "sample_template.xlsx"
SPACER_HEIGHT = 9          # pt, отступ между рядами стикеров
BLOCK_ROWS = 12            # строк на один ряд стикеров (строки 2..13 шаблона)
COLS = list(range(2, 15))  # B..N
STICKERS_PER_PAGE = 8      # 2 в ряд x 4 ряда
# Явные ширины колонок (см. CLAUDE.md, п.1)
WIDTHS = {"A": 0.6, "C": 13, "D": 2, "E": 13, "F": 2,
          "G": 1.4, "H": 1.4, "J": 13, "K": 2, "L": 13, "M": 2, "N": 0.6}


def build(total: int, invoice: str, date: str, out: Path, scheme: str | None = None,
          sender: str = SENDER) -> None:
    """sender — строка «Отправитель: …»; в программе маркировки берётся из организации."""
    scheme = scheme or scheme_for(date)
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb["Лист1"]

    row_heights = {r: ws.row_dimensions[r].height for r in range(2, 14)}
    cells = {}
    for r in range(2, 14):
        for c in COLS:
            cell = ws.cell(row=r, column=c)
            cells[(r - 2, c)] = dict(
                value=cell.value, font=copy.copy(cell.font),
                border=copy.copy(cell.border), fill=copy.copy(cell.fill),
                alignment=copy.copy(cell.alignment),
                number_format=cell.number_format,
                protection=copy.copy(cell.protection))
    merges = [(m.min_row - 2, m.min_col, m.max_row - 2, m.max_col)
              for m in ws.merged_cells.ranges]
    # ВАЖНО: обращаться к column_dimensions только для B и I
    b_w = ws.column_dimensions["B"].width
    i_w = ws.column_dimensions["I"].width

    for m in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(m))
    ws.delete_rows(1, ws.max_row + 1)

    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["B"].width = b_w
    ws.column_dimensions["I"].width = i_w
    if "O" in ws.column_dimensions:
        del ws.column_dimensions["O"]

    n_blocks = (total + 1) // 2
    blocks_per_page = STICKERS_PER_PAGE // 2
    cur = 2
    for block in range(n_blocks):
        base = cur
        for i, r in enumerate(range(2, 14)):
            ws.row_dimensions[base + i].height = row_heights[r]
        for (rel, c), d in cells.items():
            cell = ws.cell(row=base + rel, column=c)
            cell.font, cell.border, cell.fill = d["font"], d["border"], d["fill"]
            cell.alignment, cell.number_format = d["alignment"], d["number_format"]
            cell.protection, cell.value = d["protection"], d["value"]

        last_num = None
        for slot in (0, 1):
            num = block * 2 + slot + 1
            off = 0 if slot == 0 else 7
            if num > total:  # пустой слот последнего ряда
                for i in range(BLOCK_ROWS):
                    for c in range(2 + off, 8 + off):
                        cell = ws.cell(row=base + i, column=c)
                        cell.value, cell.border, cell.fill = None, Border(), PatternFill()
                continue
            last_num = num
            b, c_, e = 2 + off, 3 + off, 5 + off
            ws.cell(row=base + 2, column=b).value = sender
            ws.cell(row=base + 3, column=b).value = f"Номер накладной: {invoice}"
            ws.cell(row=base + 8, column=b).value = f"Дата:{date}"
            ws.cell(row=base + 8, column=c_).value = f"Номер короба:{num}"
            ws.cell(row=base + 8, column=e).value = f"Всего коробов:{total}"
            ws.cell(row=base + MARK_ROW, column=b).value = SCHEME_MARK[scheme]
            ws.cell(row=base + 9, column=b).value = RECIPIENT[scheme]

        for r0, c0, r1, c1 in merges:
            slot = 0 if c0 < 8 else 1
            if block * 2 + slot + 1 > total:
                continue
            ws.merge_cells(start_row=base + r0, start_column=c0,
                           end_row=base + r1, end_column=c1)

        ws.row_dimensions[base + BLOCK_ROWS].height = SPACER_HEIGHT
        cur = base + BLOCK_ROWS + 1
        if (block + 1) % blocks_per_page == 0 and last_num is not None and last_num < total:
            ws.row_breaks.append(Break(id=cur - 1))

    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.scale = 100
    ws.page_setup.fitToWidth = False
    ws.page_setup.fitToHeight = False
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    ws.page_margins = PageMargins(left=0.15, right=0.15, top=0.3, bottom=0.3,
                                  header=0.1, footer=0.1)
    ws.print_area = f"A1:M{cur - 1}"
    ws.sheet_view.showGridLines = False

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("-n", "--boxes", type=int, required=True, help="количество коробов")
    p.add_argument("-i", "--invoice", required=True, help="номер накладной/поставки")
    p.add_argument("-d", "--date", required=True, help="дата поставки, ДД.ММ.ГГГГ")
    p.add_argument("-o", "--out", type=Path, help="путь к выходному xlsx")
    p.add_argument("--scheme", choices=["auto", "commission", "agency"], default="auto",
                   help="auto — по дате: по 30.09.2026 «КОМИССИЯ», с 01.10.2026 «АГЕНТ»")
    a = p.parse_args()
    out = a.out or Path("output") / f"Маркировка_короба_Лемода_{a.invoice}_{a.boxes}шт.xlsx"
    scheme = scheme_for(a.date) if a.scheme == "auto" else a.scheme
    build(a.boxes, a.invoice, a.date, out, scheme)
    pages = -(-a.boxes // STICKERS_PER_PAGE)
    print(f"OK: {out}  ({a.boxes} коробов, ~{pages} стр. A4, «{SCHEME_MARK[scheme]}»)")


if __name__ == "__main__":
    main()
