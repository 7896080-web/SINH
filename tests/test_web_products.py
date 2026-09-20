"""Единая страница «Товары и остатки» (/products) — заменила «Синхронизируемые
товары» и «Управление остатками». Здесь же проверяется главное свойство новой
страницы: по каждому кабинету показано число, которое РЕАЛЬНО уйдёт, и причина нуля.
"""
from datetime import date

from app.models import (Product, Barcode, PlatformAccount, Platform, SyncSetting,
                        DispatchQueueItem)


def _accounts(web_db, *specs):
    made = []
    for platform, name in specs:
        a = PlatformAccount(platform=platform, name=name, warehouse_id="wh")
        web_db.add(a)
        made.append(a)
    web_db.commit()
    for a in made:
        web_db.refresh(a)
    return made


def _product(web_db, uid="u1", stock=10, reserve=0, offset=None, override=None,
             broadcast=True, barcode="111"):
    p = Product(uid_1c=uid, article="A-" + uid, name="Товар " + uid, stock_on_hand=stock,
                reserve=reserve, broadcast_offset=offset, transmit_override=override,
                broadcast_enabled=broadcast)
    web_db.add(p)
    if barcode:
        web_db.add(Barcode(barcode=barcode + uid, uid_1c=uid))
    web_db.commit()
    return p


def _sync(web_db, uid, account, enabled=True, threshold=0):
    s = SyncSetting(uid_1c=uid, account_id=account.id, enabled=enabled, min_threshold=threshold)
    web_db.add(s)
    web_db.commit()
    return s


# --------------------------------------------------------------- страница и объяснения

def test_page_renders_with_product_and_ladder(logged_in_client, web_db):
    _product(web_db)
    r = logged_in_client.get("/products")
    assert r.status_code == 200
    assert "Товары и остатки" in r.text
    assert "Товар u1" in r.text
    assert "Как считается количество для площадки" in r.text


def test_old_page_urls_redirect_to_merged_page(logged_in_client, web_db):
    for old in ("/sync-products", "/stock-control"):
        r = logged_in_client.get(old, follow_redirects=False)
        assert r.status_code == 301, old
        assert r.headers["location"] == "/products"


def test_each_cabinet_shows_real_quantity(logged_in_client, web_db):
    a1, a2 = _accounts(web_db, (Platform.wb, "WB-1"), (Platform.ozon, "OZ-1"))
    _product(web_db, stock=29, offset=11)
    _sync(web_db, "u1", a1, enabled=True)
    _sync(web_db, "u1", a2, enabled=False)

    r = logged_in_client.get("/products")
    assert "→ 18" in r.text          # 29 − 11 в отмеченный кабинет
    assert "не передаётся" in r.text  # неотмеченный кабинет


def test_disabled_broadcast_shows_zero_and_reason(logged_in_client, web_db):
    a1, = _accounts(web_db, (Platform.wb, "WB-1"))
    _product(web_db, stock=29, offset=11, broadcast=False)
    _sync(web_db, "u1", a1, enabled=True)

    r = logged_in_client.get("/products")
    assert "трансляция товара выключена" in r.text


def test_cabinet_threshold_reason_is_shown(logged_in_client, web_db):
    a1, = _accounts(web_db, (Platform.wb, "WB-1"))
    _product(web_db, stock=10)                    # авторежим: 10 − 0 = 10
    _sync(web_db, "u1", a1, enabled=True, threshold=32)

    r = logged_in_client.get("/products")
    assert "порог кабинета 32" in r.text


def test_paused_platform_reason_is_shown(logged_in_client, web_db):
    a1, = _accounts(web_db, (Platform.wb, "WB-1"))
    a1.dispatch_enabled = False
    _product(web_db, stock=29, offset=11)
    _sync(web_db, "u1", a1, enabled=True)

    r = logged_in_client.get("/products")
    assert "на паузе" in r.text


def test_only_blocked_filter_keeps_problem_rows(logged_in_client, web_db):
    a1, = _accounts(web_db, (Platform.wb, "WB-1"))
    _product(web_db, uid="ok", stock=29, offset=11)
    _product(web_db, uid="bad", stock=29, offset=11, broadcast=False)
    _sync(web_db, "ok", a1, enabled=True)
    _sync(web_db, "bad", a1, enabled=True)

    r = logged_in_client.get("/products/rows?only_blocked=true")
    assert "Товар bad" in r.text
    assert "Товар ok" not in r.text


# --------------------------------------------------------------- правки строки

def test_reserve_update_and_resend(logged_in_client, web_db):
    a1, = _accounts(web_db, (Platform.wb, "WB-1"))
    _product(web_db, stock=10)
    _sync(web_db, "u1", a1, enabled=True)

    logged_in_client.post("/products/u1/reserve", data={"reserve": "2"})
    web_db.expire_all()
    assert web_db.query(Product).first().reserve == 2
    assert web_db.query(DispatchQueueItem).filter(DispatchQueueItem.uid_1c == "u1").count() >= 1


def test_reserve_negative_clamped(logged_in_client, web_db):
    _product(web_db, stock=10, reserve=3)
    logged_in_client.post("/products/u1/reserve", data={"reserve": "-5"})
    web_db.expire_all()
    assert web_db.query(Product).first().reserve == 0


def test_offset_direct_value(logged_in_client, web_db):
    _product(web_db, stock=29, override=7)
    logged_in_client.post("/products/u1/offset", data={"value": "11"})
    web_db.expire_all()
    p = web_db.query(Product).first()
    assert p.broadcast_offset == 11
    assert p.transmit_override is None       # порог гасит устаревшую ручную цифру


def test_offset_computed_from_recount(logged_in_client, web_db):
    _product(web_db, stock=29)
    logged_in_client.post("/products/u1/offset", data={"stock_at_date": "43", "available": "32"})
    web_db.expire_all()
    assert web_db.query(Product).first().broadcast_offset == 11


def test_offset_can_be_negative(logged_in_client, web_db):
    _product(web_db, stock=10)
    logged_in_client.post("/products/u1/offset", data={"stock_at_date": "10", "available": "15"})
    web_db.expire_all()
    assert web_db.query(Product).first().broadcast_offset == -5


def test_offset_cleared_by_empty_value(logged_in_client, web_db):
    _product(web_db, stock=10, offset=3)
    logged_in_client.post("/products/u1/offset", data={"value": ""})
    web_db.expire_all()
    assert web_db.query(Product).first().broadcast_offset is None


def test_offset_rejects_non_number(logged_in_client, web_db):
    _product(web_db, stock=10, offset=3)
    logged_in_client.post("/products/u1/offset", data={"value": "abc"})
    web_db.expire_all()
    assert web_db.query(Product).first().broadcast_offset == 3


def test_clear_legacy_override(logged_in_client, web_db):
    _product(web_db, stock=10, override=44)
    logged_in_client.post("/products/u1/clear-override")
    web_db.expire_all()
    assert web_db.query(Product).first().transmit_override is None


def _ready_for_broadcast(web_db, uid="u1", account=None):
    """Строка, доведённая до «актуализирован»: кабинет отмечен, дата задана, 1С
    ответила, факт подтверждён, расчёт прошёл и покрыл этот кабинет.

    Трансляция включается только с этого состояния — правило «сначала расчёт,
    потом трансляция». Раньше включить можно было что угодно и когда угодно, и
    строка ZJYM269002 XL показала, чем это кончается: «ждём 1С», а рядом
    «→ 21 после включения» от порога, оставшегося с другой даты.
    """
    from datetime import date as _date

    from app.timeutils import now_utc

    account = account or _accounts(web_db, (Platform.wb, "WB-1"))[0]
    _sync(web_db, uid, account)
    product = web_db.query(Product).filter(Product.uid_1c == uid).one()
    product.offset_base_date = _date(2026, 8, 7)
    product.offset_base_stock = 10
    product.fact_at_date = 10
    product.recalc_done_at = now_utc()
    product.recalc_account_ids = str(account.id)
    web_db.commit()
    return account


def test_broadcast_toggle(logged_in_client, web_db):
    _product(web_db, broadcast=False)
    _ready_for_broadcast(web_db)

    logged_in_client.post("/products/u1/broadcast", data={"enabled": "true"})
    web_db.expire_all()
    assert web_db.query(Product).first().broadcast_enabled is True

    logged_in_client.post("/products/u1/broadcast", data={"enabled": "false"})
    web_db.expire_all()
    assert web_db.query(Product).first().broadcast_enabled is False


def test_broadcast_cannot_be_switched_on_before_the_catch_up(logged_in_client, web_db):
    """Ровно случай ZJYM269002 XL: расчёта не было, 1С на новую дату не ответила,
    а порог остался с прошлой — включение отправило бы на площадки остаток,
    не сверенный ни с чем."""
    _product(web_db, broadcast=False, offset=0)
    account = _accounts(web_db, (Platform.wb, "WB-1"))[0]
    _sync(web_db, "u1", account)

    r = logged_in_client.post("/products/u1/broadcast", data={"enabled": "true"})

    web_db.expire_all()
    assert web_db.query(Product).first().broadcast_enabled is False
    assert "Включать трансляцию рано" in r.text


def test_switching_broadcast_off_is_never_blocked(logged_in_client, web_db):
    """Снять с трансляции должно быть можно всегда: товар мог начать
    транслироваться до того, как появилось это правило."""
    _product(web_db, broadcast=True)
    account = _accounts(web_db, (Platform.wb, "WB-1"))[0]
    _sync(web_db, "u1", account)

    logged_in_client.post("/products/u1/broadcast", data={"enabled": "false"})

    web_db.expire_all()
    assert web_db.query(Product).first().broadcast_enabled is False


def test_active_since_set_and_clear(logged_in_client, web_db):
    _product(web_db)
    logged_in_client.post("/products/u1/active-since", data={"value": "2026-09-01"})
    web_db.expire_all()
    assert web_db.query(Product).first().broadcast_active_since == date(2026, 9, 1)

    logged_in_client.post("/products/u1/active-since", data={"value": ""})
    web_db.expire_all()
    assert web_db.query(Product).first().broadcast_active_since is None


def test_toggle_cabinet_enables_and_enqueues_resend(logged_in_client, web_db):
    a1, = _accounts(web_db, (Platform.wb, "WB-1"))
    _product(web_db, stock=10)

    logged_in_client.post(f"/products/u1/{a1.id}/toggle", data={"enabled": "true"})
    web_db.expire_all()
    setting = web_db.query(SyncSetting).filter(SyncSetting.account_id == a1.id).first()
    assert setting.enabled is True
    items = web_db.query(DispatchQueueItem).filter(DispatchQueueItem.reason == "manual_enable").all()
    assert [i.quantity for i in items] == [10]


def test_toggle_off_does_not_enqueue(logged_in_client, web_db):
    a1, = _accounts(web_db, (Platform.wb, "WB-1"))
    _product(web_db, stock=10)
    _sync(web_db, "u1", a1, enabled=True)

    logged_in_client.post(f"/products/u1/{a1.id}/toggle", data={"enabled": "false"})
    web_db.expire_all()
    assert web_db.query(SyncSetting).first().enabled is False
    assert web_db.query(DispatchQueueItem).filter(DispatchQueueItem.reason == "manual_enable").count() == 0


def test_cabinet_threshold_update_and_clamp(logged_in_client, web_db):
    a1, = _accounts(web_db, (Platform.wb, "WB-1"))
    _product(web_db, stock=10)

    logged_in_client.post(f"/products/u1/{a1.id}/threshold", data={"min_threshold": "3"})
    web_db.expire_all()
    assert web_db.query(SyncSetting).first().min_threshold == 3

    logged_in_client.post(f"/products/u1/{a1.id}/threshold", data={"min_threshold": "-5"})
    web_db.expire_all()
    assert web_db.query(SyncSetting).first().min_threshold == 0


def test_three_wb_cabinets_are_separate_columns(logged_in_client, web_db):
    _accounts(web_db, (Platform.wb, "ИП Первый"), (Platform.wb, "ИП Второй"), (Platform.wb, "ИП Третий"))
    _product(web_db)
    r = logged_in_client.get("/products")
    for name in ("ИП Первый", "ИП Второй", "ИП Третий"):
        assert name in r.text


# --------------------------------------------------------------- массовые действия

def test_bulk_reserve_offset_broadcast(logged_in_client, web_db):
    _product(web_db, uid="a", stock=10)
    _product(web_db, uid="b", stock=10)

    logged_in_client.post("/products/bulk", data={
        "action": "set_reserve", "uids": ["a", "b"], "int_value": "4", "q": ""})
    web_db.expire_all()
    assert [p.reserve for p in web_db.query(Product).order_by(Product.uid_1c)] == [4, 4]

    logged_in_client.post("/products/bulk", data={
        "action": "set_offset", "uids": ["a"], "int_value": "-3", "q": ""})
    web_db.expire_all()
    assert web_db.query(Product).filter(Product.uid_1c == "a").first().broadcast_offset == -3

    logged_in_client.post("/products/bulk", data={
        "action": "clear_offset", "uids": ["a"], "q": ""})
    web_db.expire_all()
    assert web_db.query(Product).filter(Product.uid_1c == "a").first().broadcast_offset is None

    logged_in_client.post("/products/bulk", data={
        "action": "broadcast_off", "uids": ["a", "b"], "q": ""})
    web_db.expire_all()
    assert not any(p.broadcast_enabled for p in web_db.query(Product))


def test_bulk_without_selection_is_noop(logged_in_client, web_db):
    _product(web_db, stock=10, reserve=1)
    logged_in_client.post("/products/bulk", data={"action": "set_reserve", "int_value": "5", "q": ""})
    web_db.expire_all()
    assert web_db.query(Product).first().reserve == 1


def test_dispatch_toggle_per_platform_and_all(logged_in_client, web_db):
    wb, oz = _accounts(web_db, (Platform.wb, "WB-1"), (Platform.ozon, "OZ-1"))

    logged_in_client.post("/products/dispatch-toggle", data={"scope": "wb", "enabled": "false", "q": ""})
    web_db.expire_all()
    assert web_db.query(PlatformAccount).filter(PlatformAccount.id == wb.id).first().dispatch_enabled is False
    assert web_db.query(PlatformAccount).filter(PlatformAccount.id == oz.id).first().dispatch_enabled is True

    logged_in_client.post("/products/dispatch-toggle", data={"scope": "all", "enabled": "true", "q": ""})
    web_db.expire_all()
    assert all(a.dispatch_enabled for a in web_db.query(PlatformAccount))


# --------------------------------------------------------------- Excel

def test_export_contains_new_columns(logged_in_client, web_db):
    a1, = _accounts(web_db, (Platform.wb, "WB-1"))
    _product(web_db, stock=29, offset=11)
    _sync(web_db, "u1", a1, enabled=True)

    r = logged_in_client.get("/products/export")
    assert r.status_code == 200
    assert "spreadsheet" in r.headers["content-type"]

    import io
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(r.content)).active
    headers = [c.value for c in ws[1]]
    assert "Порог трансляции" in headers
    assert "Трансляция" in headers
    assert "Уходит на площадки" in headers
    assert "WB-1 (WB) — Синхронизировать" in headers
    values = dict(zip(headers, [c.value for c in ws[2]]))
    assert values["Порог трансляции"] == 11
    assert values["Уходит на площадки"] == 18


def test_import_sets_the_manual_offset_and_the_cabinet(logged_in_client, web_db):
    """Товар без даты расчёта: порог у него живёт как введённое руками число, и
    файл его задаёт. Трансляцию такому товару не включить ни отсюда, ни со
    страницы — расчёт не начат, и остаток ничем не сверен."""
    a1, = _accounts(web_db, (Platform.wb, "WB-1"))
    _product(web_db, stock=29, broadcast=False)

    import io
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["ID_1С", "Резерв", "Порог трансляции", "WB-1 (WB) — Синхронизировать", "WB-1 (WB) — Порог"])
    ws.append(["u1", 0, 11, "Да", 0])
    buf = io.BytesIO()
    wb.save(buf)

    logged_in_client.post("/products/import",
                          files={"file": ("p.xlsx", buf.getvalue(),
                                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    web_db.expire_all()
    p = web_db.query(Product).first()
    assert p.broadcast_offset == 11
    assert web_db.query(SyncSetting).first().enabled is True


def test_import_switches_broadcast_on_for_a_calculated_product(logged_in_client, web_db):
    """Полный сценарий массовой настройки: файлом отмечают кабинет И включают
    трансляцию. Кабинеты поэтому разбираются РАНЬШЕ «Трансляции» — гейт
    включения спрашивает, есть ли отмеченный кабинет, покрытый расчётом, и при
    обратном порядке такой файл всегда упирался бы в отказ."""
    from datetime import date

    from app.timeutils import now_utc

    a1, = _accounts(web_db, (Platform.wb, "WB-1"))
    product = _product(web_db, stock=29, broadcast=False)
    product.offset_base_date = date(2026, 9, 1)
    product.offset_base_stock = 30
    product.fact_at_date = 30
    product.recalc_done_at = now_utc()
    product.recalc_account_ids = str(a1.id)
    web_db.commit()

    import io
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["ID_1С", "Трансляция", "WB-1 (WB) — Синхронизировать", "WB-1 (WB) — Порог"])
    ws.append(["u1", "Да", "Да", 0])
    buf = io.BytesIO()
    wb.save(buf)

    logged_in_client.post("/products/import",
                          files={"file": ("p.xlsx", buf.getvalue(),
                                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    web_db.expire_all()
    assert web_db.query(Product).first().broadcast_enabled is True
    assert web_db.query(SyncSetting).first().enabled is True


# --------------------------------------------------------------- фильтры до ограничения строк

def _many(web_db, n, with_proposal_from=None, blocked_from=None):
    """n товаров по алфавиту; предложения и «уходит 0» — только у ХВОСТА списка,
    чтобы поймать старую ошибку: ограничение в 300 строк применялось до фильтра."""
    from app.routers.products import PAGE_LIMIT
    a1, = _accounts(web_db, (Platform.wb, "WB-1"))
    for i in range(n):
        uid = f"p{i:05d}"
        web_db.add(Product(uid_1c=uid, article=f"A{i:05d}", name=f"Товар {i:05d}",
                           stock_on_hand=10,
                           broadcast_enabled=not (blocked_from is not None and i >= blocked_from)))
        web_db.add(Barcode(barcode=f"bc{i:05d}", uid_1c=uid))
        web_db.add(SyncSetting(uid_1c=uid, account_id=a1.id, enabled=True,
                               has_proposal=(with_proposal_from is not None and i >= with_proposal_from)))
    web_db.commit()
    return a1, PAGE_LIMIT


def test_proposal_filter_finds_rows_beyond_the_page_limit(logged_in_client, web_db):
    """Товары с предложениями лежат за пределами первых 300 по алфавиту."""
    _many(web_db, 350, with_proposal_from=340)

    r = logged_in_client.get("/products/rows?only_proposals=true")

    assert r.text.count('<tr id="row-') == 10
    assert "Товар 00345" in r.text


def test_blocked_filter_finds_rows_beyond_the_page_limit(logged_in_client, web_db):
    """То же для фильтра «уходит 0»: раньше он показывал пусто на большом каталоге."""
    _many(web_db, 350, blocked_from=340)

    r = logged_in_client.get("/products/rows?only_blocked=true")

    assert r.text.count('<tr id="row-') == 10
    assert "Товар 00345" in r.text


def test_page_says_how_many_rows_were_hidden(logged_in_client, web_db):
    _many(web_db, 400)

    r = logged_in_client.get("/products/rows")

    assert "Показано 300 из 400" in r.text


def test_export_is_not_silently_truncated(logged_in_client, web_db):
    """Обрезанная выгрузка помечается прямо в файле: оператор правит её в Excel и
    импортирует обратно, считая, что охватил весь каталог."""
    import io
    from openpyxl import load_workbook
    from app.routers import products as products_router

    _many(web_db, 50)
    products_router.EXPORT_LIMIT = 10          # искусственно занижаем потолок
    try:
        r = logged_in_client.get("/products/export")
    finally:
        products_router.EXPORT_LIMIT = 50000

    ws = load_workbook(io.BytesIO(r.content)).active
    last = [c.value for c in ws[ws.max_row]]
    assert "показаны первые 10 строк из 50" in str(last[0])


# --------------------------------------------------------------- фильтр размера U

def test_hide_size_u_filter(logged_in_client, web_db):
    """«Скрыть размер U» убирает безразмерные позиции и не трогает остальные."""
    for uid, size in [("a", "XL"), ("b", "U"), ("c", "u"), ("d", " U "), ("e", None)]:
        web_db.add(Product(uid_1c=uid, article=f"ART-{uid}", name=f"Товар {uid}",
                           size=size, stock_on_hand=5, broadcast_enabled=True))
        web_db.add(Barcode(barcode=f"bc-{uid}", uid_1c=uid))
    web_db.commit()

    full = logged_in_client.get("/products/rows")
    assert full.text.count('<tr id="row-') == 5

    filtered = logged_in_client.get("/products/rows?hide_size_u=true")
    assert filtered.text.count('<tr id="row-') == 2      # XL и товар без размера
    assert "Товар a" in filtered.text
    assert "Товар e" in filtered.text                     # размера нет — не прячем
    for hidden in ("Товар b", "Товар c", "Товар d"):      # U, u и « U » — регистр и пробелы
        assert hidden not in filtered.text


def test_hide_size_u_applies_to_export(logged_in_client, web_db):
    import io
    from openpyxl import load_workbook
    for uid, size in [("a", "XL"), ("b", "U")]:
        web_db.add(Product(uid_1c=uid, article=f"ART-{uid}", name=f"Товар {uid}",
                           size=size, stock_on_hand=5, broadcast_enabled=True))
        web_db.add(Barcode(barcode=f"bc-{uid}", uid_1c=uid))
    web_db.commit()

    r = logged_in_client.get("/products/export?hide_size_u=true")
    ws = load_workbook(io.BytesIO(r.content)).active
    sizes = [row[2] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert sizes == ["XL"]


def _cab_block(html: str, account_name: str) -> str:
    """Кусок строки товара, относящийся к одному кабинету (одна строка внутри
    единственной колонки «Кабинеты»)."""
    blocks = html.split('<div class="pr-cab">')[1:]
    matching = [b for b in blocks if account_name in b]
    assert matching, f"в строке товара нет блока кабинета «{account_name}»"
    return matching[0]


def test_proposal_marker_is_visible_in_the_product_row(logged_in_client, web_db):
    """Кабинеты собраны в одну колонку, каждый своей строкой: молния должна стоять
    в строке того кабинета, где нашлась карточка, и только там — иначе фильтр
    «только с предложениями» показывает строки без единой молнии."""
    a1, a2 = _accounts(web_db, (Platform.wb, "ИП ЯВОРСКАЯ"), (Platform.kit, "КИТ"))
    _product(web_db, stock=5)
    _sync(web_db, "u1", a1, enabled=False)
    _sync(web_db, "u1", a2, enabled=False)
    web_db.query(SyncSetting).filter(SyncSetting.account_id == a1.id).first().has_proposal = True
    web_db.commit()

    r = logged_in_client.get("/products/rows?only_proposals=true")

    with_proposal = _cab_block(r.text, "ИП ЯВОРСКАЯ")
    assert "⚡" in with_proposal
    assert f"/products/u1/{a1.id}/toggle" in with_proposal
    assert "Карточка найдена в кабинете «ИП ЯВОРСКАЯ»" in with_proposal

    assert "⚡" not in _cab_block(r.text, "КИТ")   # у этого кабинета предложения нет
    assert "⚡ КИТ" not in r.text


def test_proposal_marker_shows_inactive_cabinet_too(logged_in_client, web_db):
    """Предложение по кабинету, который сейчас неактивен: колонки у него нет,
    и без этой пометки молния была бы не видна нигде."""
    a1, a2 = _accounts(web_db, (Platform.wb, "ИП ЯВОРСКАЯ"), (Platform.kit, "КИТ"))
    _product(web_db, stock=5)
    _sync(web_db, "u1", a2, enabled=False)
    web_db.query(SyncSetting).filter(SyncSetting.account_id == a2.id).first().has_proposal = True
    a2.is_active = False
    web_db.commit()

    r = logged_in_client.get("/products/rows?only_proposals=true")

    assert "⚡ КИТ" in r.text
    assert "кабинет сейчас неактивен" in r.text


def test_enabled_cabinet_has_no_proposal_marker(logged_in_client, web_db):
    """Товар уже передаётся в кабинет — предлагать нечего."""
    a1, = _accounts(web_db, (Platform.wb, "ИП ЯВОРСКАЯ"))
    _product(web_db, stock=5)
    _sync(web_db, "u1", a1, enabled=True)
    web_db.query(SyncSetting).first().has_proposal = True
    web_db.commit()

    r = logged_in_client.get("/products/rows")

    assert "⚡" not in _cab_block(r.text, "ИП ЯВОРСКАЯ")
    assert "⚡ ИП ЯВОРСКАЯ" not in r.text
