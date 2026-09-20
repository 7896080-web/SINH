"""Файл Excel как рабочий сценарий: дата, факт, кабинеты, трансляция — разом.

Оператор думает о работе одной мыслью: «вот эти триста позиций считаем от
первого сентября и включаем». В файле он так и пишет. Дальше расчёт идёт
минутами и заканчивается уже без него, поэтому включать трансляцию приходилось
вторым заходом — назавтра, разыскав те же строки в каталоге на сто пятьдесят
тысяч позиций. На практике это означало, что половина строк оставалась
невключённой и молчала об этом.

Теперь файл ставит ПРОСЬБУ, а включает её тот, кто имеет на это право: расчёт,
поставивший отметку «актуализирован». Гейт не ослаблен ни на грамм — здесь это
и закреплено: включается ровно то, что в ту же минуту включила бы страница.
"""

import io
from datetime import date

from openpyxl import Workbook

from app.models import (Barcode, DispatchQueueItem, Platform, Product, RecalcItem,
                        RecalcJob, SyncSetting)
from app.recalc import catch_up_product
from app.workers.scheduler import PENDING_WAREHOUSE_NAME
from tests.factories import make_account

DAY = date(2026, 8, 7)


class FakeClient:
    last_unresolved = 0

    def get_orders_since(self, since):
        return []


def _wh(platform):
    return PENDING_WAREHOUSE_NAME.get(platform, "Ожидает")


def _file(headers, *rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _import(client, content):
    return client.post("/products/import", files={"file": ("t.xlsx", content)},
                       follow_redirects=False)


def _seed(db, uid="u1", **kwargs):
    fields = dict(article="A-1", name="Товар", stock_on_hand=20, reserve=0)
    fields.update(kwargs)
    product = Product(uid_1c=uid, **fields)
    db.add(product)
    db.add(Barcode(barcode=f"bc-{uid}", uid_1c=uid))
    db.commit()
    return product


def _label(account):
    return f"{account.name} ({account.platform.value.upper()})"


# ------------------------------------------------------- просьба, а не ошибка

def test_a_row_awaiting_the_calculation_records_the_request(logged_in_client, web_db):
    account = make_account(web_db, name="ИП ЯВОРСКАЯ")
    _seed(web_db, offset_base_date=DAY, offset_base_stock=20, fact_at_date=20)

    r = _import(logged_in_client, _file(
        ["ID_1С", "Трансляция", f"{_label(account)} — Синхронизировать"],
        ["u1", "Да", "Да"]))

    assert r.status_code == 303
    product = web_db.query(Product).filter(Product.uid_1c == "u1").one()
    assert product.broadcast_enabled is False, "включать до расчёта нельзя"
    assert product.broadcast_requested_at is not None, "просьба должна быть записана"


def test_the_request_does_not_open_the_gate_by_itself(logged_in_client, web_db):
    """Самое важное: пока расчёта не было, наружу не уходит ничего. Просьба — это
    запись о намерении, а не включённая трансляция."""
    account = make_account(web_db, name="ИП ЯВОРСКАЯ")
    _seed(web_db, offset_base_date=DAY, offset_base_stock=20, fact_at_date=20)

    _import(logged_in_client, _file(
        ["ID_1С", "Трансляция", f"{_label(account)} — Синхронизировать"],
        ["u1", "Да", "Да"]))

    assert web_db.query(DispatchQueueItem).count() == 0


def test_a_row_that_will_never_get_there_is_still_an_error(logged_in_client, web_db):
    """Без отмеченного кабинета заказы спрашивать негде, и расчёт эту строку не
    вылечит никогда. Запомнить просьбу значило бы дать обещание, которое никто
    не выполнит, — а оператор ушёл бы, считая, что всё настроено."""
    _seed(web_db)

    _import(logged_in_client, _file(["ID_1С", "Трансляция"], ["u1", "Да"]))

    product = web_db.query(Product).filter(Product.uid_1c == "u1").one()
    assert product.broadcast_requested_at is None


def test_a_ready_row_is_switched_on_at_once(logged_in_client, web_db):
    """Строке, доведённой до «актуализирован», ждать нечего."""
    account = make_account(web_db, name="ИП ЯВОРСКАЯ")
    from app.timeutils import now_utc
    product = _seed(web_db, offset_base_date=DAY, offset_base_stock=20, fact_at_date=20,
                    recalc_done_at=now_utc())
    web_db.add(SyncSetting(uid_1c="u1", account_id=account.id, enabled=True))
    product.recalc_account_ids = str(account.id)
    web_db.commit()

    _import(logged_in_client, _file(["ID_1С", "Трансляция"], ["u1", "Да"]))

    product = web_db.query(Product).filter(Product.uid_1c == "u1").one()
    assert product.broadcast_enabled is True
    assert product.broadcast_requested_at is None


# ------------------------------------------------------------- кто включает

def test_the_calculation_switches_the_broadcast_on(db):
    account = make_account(db, name="ИП ЯВОРСКАЯ")
    from app.timeutils import now_utc
    product = _seed(db, offset_base_date=DAY, offset_base_stock=20, fact_at_date=20,
                    broadcast_requested_at=now_utc())
    db.add(SyncSetting(uid_1c="u1", account_id=account.id, enabled=True))
    db.commit()

    catch_up_product(db, product, lambda d, aid: FakeClient(), _wh)
    db.commit()

    db.refresh(product)
    assert product.broadcast_enabled is True
    assert product.broadcast_requested_at is None, "просьба выполнена и снята"


def test_the_stock_leaves_right_after_that(db):
    """Включить и не отправить — худший исход: всё зелено, а на площадке чужое
    число. Кабинет покрыт расчётом не впервые, значит доотправку по нему надо
    поставить отдельно: цикл «покрыт впервые» его не касается."""
    account = make_account(db, name="ИП ЯВОРСКАЯ")
    from app.timeutils import now_utc
    product = _seed(db, offset_base_date=DAY, offset_base_stock=20, fact_at_date=20,
                    broadcast_requested_at=now_utc(), recalc_done_at=now_utc())
    product.recalc_account_ids = str(account.id)
    db.add(SyncSetting(uid_1c="u1", account_id=account.id, enabled=True))
    db.commit()

    catch_up_product(db, product, lambda d, aid: FakeClient(), _wh)
    db.commit()

    assert db.query(DispatchQueueItem).filter(
        DispatchQueueItem.account_id == account.id).count() == 1


def test_a_calculation_with_problems_switches_nothing_on(db):
    """Отметку «актуализирован» расчёт с проблемами не ставит — значит и ворота
    не открывает. Иначе просьба из файла обходила бы ровно ту проверку, ради
    которой отметка и существует."""
    account = make_account(db, Platform.kit, name="КИТ", warehouse_id="wh-2")
    from app.timeutils import now_utc
    product = _seed(db, offset_base_date=DAY, offset_base_stock=20, fact_at_date=20,
                    broadcast_requested_at=now_utc())
    db.add(SyncSetting(uid_1c="u1", account_id=account.id, enabled=True))
    db.commit()

    class Losing(FakeClient):
        last_unresolved = 3

    catch_up_product(db, product, lambda d, aid: Losing(), _wh)
    db.commit()

    db.refresh(product)
    assert product.broadcast_enabled is False
    assert product.broadcast_requested_at is not None, "просьба остаётся до следующего расчёта"


def test_the_answer_from_1c_can_open_the_gate_too(db):
    """Порядок прихода не определён: расчёт мог закончиться раньше, чем 1С
    ответила на заявку о дате. Спрашивай мы только в расчёте — такая строка
    ждала бы вечно."""
    from app.models import StockDateRow, StockDateSnapshot, StockDateStatus
    from app.offset_base import fill_waiting_products
    from app.timeutils import now_utc

    account = make_account(db, name="ИП ЯВОРСКАЯ")
    product = _seed(db, offset_base_date=DAY, fact_at_date=20,
                    broadcast_requested_at=now_utc(), recalc_done_at=now_utc())
    product.recalc_account_ids = str(account.id)
    db.add(SyncSetting(uid_1c="u1", account_id=account.id, enabled=True))
    snapshot = StockDateSnapshot(snapshot_date=DAY, status=StockDateStatus.done,
                                 requested_by="admin")
    db.add(snapshot)
    db.commit()
    db.add(StockDateRow(snapshot_id=snapshot.id, uid_1c="u1", quantity=20))
    db.commit()

    fill_waiting_products(db, snapshot)

    db.refresh(product)
    assert product.broadcast_enabled is True


# ------------------------------------------------------- просьбу можно снять

def test_no_in_the_file_cancels_the_request(logged_in_client, web_db):
    """Иначе трансляция вернулась бы сама после ближайшего расчёта — молча и
    вопреки тому, что оператор только что написал в файле."""
    from app.timeutils import now_utc
    account = make_account(web_db, name="ИП ЯВОРСКАЯ")
    _seed(web_db, offset_base_date=DAY, offset_base_stock=20, fact_at_date=20,
          broadcast_requested_at=now_utc())
    web_db.add(SyncSetting(uid_1c="u1", account_id=account.id, enabled=True))
    web_db.commit()

    _import(logged_in_client, _file(["ID_1С", "Трансляция"], ["u1", "Нет"]))

    assert web_db.query(Product).filter(
        Product.uid_1c == "u1").one().broadcast_requested_at is None


def test_switching_off_by_hand_cancels_the_request(logged_in_client, web_db):
    from app.timeutils import now_utc
    _seed(web_db, broadcast_enabled=True, broadcast_requested_at=now_utc())

    logged_in_client.post("/products/u1/broadcast", data={"enabled": "false"})

    assert web_db.query(Product).filter(
        Product.uid_1c == "u1").one().broadcast_requested_at is None


def test_the_row_says_it_is_waiting(logged_in_client, web_db):
    """Ожидание обязано быть видно: иначе строка выглядит просто выключенной, и
    оператор включает её вторым заходом руками."""
    from app.timeutils import now_utc
    _seed(web_db, broadcast_requested_at=now_utc())

    page = logged_in_client.get("/products?q=A-1")

    assert "включится сама после расчёта" in page.text


# ----------------------------------------------------- расчёт запускает файл

def test_the_import_starts_the_calculation(logged_in_client, web_db):
    """«Расчёт от заданного числа» — то, ради чего файл и заполняли. Запускать
    его руками означало искать в каталоге ровно те же строки, которых в отборе
    на странице уже нет."""
    account = make_account(web_db, name="ИП ЯВОРСКАЯ")
    _seed(web_db)

    _import(logged_in_client, _file(
        ["ID_1С", "Дата расчёта", "Факт на дату", f"{_label(account)} — Синхронизировать"],
        ["u1", "2026-08-07", 20, "Да"]))

    job = web_db.query(RecalcJob).one()
    assert job.total == 1
    assert web_db.query(RecalcItem).filter(RecalcItem.job_id == job.id).one().uid_1c == "u1"


def test_a_row_without_a_cabinet_is_not_sent_to_the_calculation(logged_in_client, web_db):
    """Расчёт по такой строке — пустой проход: заказы спрашивать негде."""
    _seed(web_db)

    _import(logged_in_client, _file(
        ["ID_1С", "Дата расчёта", "Факт на дату"], ["u1", "2026-08-07", 20]))

    assert web_db.query(RecalcJob).count() == 0


def test_a_ready_row_is_not_recalculated_again(logged_in_client, web_db):
    """Файл правят и ради мелочей — брони, порога кабинета. Гонять по такому
    файлу расчёт заново значило бы часами опрашивать площадки ни за чем."""
    from app.timeutils import now_utc
    account = make_account(web_db, name="ИП ЯВОРСКАЯ")
    product = _seed(web_db, offset_base_date=DAY, offset_base_stock=20, fact_at_date=20,
                    recalc_done_at=now_utc())
    product.recalc_account_ids = str(account.id)
    web_db.add(SyncSetting(uid_1c="u1", account_id=account.id, enabled=True))
    web_db.commit()

    _import(logged_in_client, _file(["ID_1С", "Резерв"], ["u1", 3]))

    assert web_db.query(RecalcJob).count() == 0


def test_a_running_job_is_not_joined_by_a_second_one(logged_in_client, web_db):
    """Два задания шли бы по одним товарам и дублировали обращения к площадкам —
    то же правило, что и у кнопки на странице."""
    from app.models import RecalcStatus
    account = make_account(web_db, name="ИП ЯВОРСКАЯ")
    _seed(web_db)
    web_db.add(RecalcJob(created_by="admin", total=1, status=RecalcStatus.running))
    web_db.commit()

    r = _import(logged_in_client, _file(
        ["ID_1С", "Дата расчёта", f"{_label(account)} — Синхронизировать"],
        ["u1", "2026-08-07", "Да"]))

    assert r.status_code == 303
    assert web_db.query(RecalcJob).count() == 1


# --------------------------------------------------------- Да/Нет — выбором

def _export(client):
    import asyncio
    from openpyxl import load_workbook
    r = client.get("/products/export")
    return load_workbook(io.BytesIO(r.content)).active


def test_the_broadcast_column_is_a_dropdown(logged_in_client, web_db):
    """Опечатку в этой колонке импорт читает как «Нет» и молча выключает то, что
    оператор включал."""
    _seed(web_db)

    ws = _export(logged_in_client)
    ranges = [str(rng) for dv in ws.data_validations.dataValidation for rng in dv.sqref.ranges]
    headers = [c.value for c in ws[1]]
    letter = ws.cell(row=1, column=headers.index("Трансляция") + 1).column_letter

    assert any(r.startswith(f"{letter}2") for r in ranges)


def test_the_cabinet_column_is_a_dropdown_too(logged_in_client, web_db):
    account = make_account(web_db, name="ИП ЯВОРСКАЯ")
    _seed(web_db)

    ws = _export(logged_in_client)
    headers = [c.value for c in ws[1]]
    letter = ws.cell(row=1,
                     column=headers.index(f"{_label(account)} — Синхронизировать") + 1).column_letter
    ranges = [str(rng) for dv in ws.data_validations.dataValidation for rng in dv.sqref.ranges]

    assert any(r.startswith(f"{letter}2") for r in ranges)


def test_the_dropdown_offers_exactly_yes_and_no(logged_in_client, web_db):
    _seed(web_db)

    ws = _export(logged_in_client)

    assert all(dv.formula1 == '"Да,Нет"' for dv in ws.data_validations.dataValidation)


def test_numbers_are_left_alone(logged_in_client, web_db):
    """Список на числовой колонке сделал бы файл неправимым: порог и факт —
    произвольные числа."""
    _seed(web_db)

    ws = _export(logged_in_client)
    headers = [c.value for c in ws[1]]
    letter = ws.cell(row=1, column=headers.index("Факт на дату") + 1).column_letter
    ranges = [str(rng) for dv in ws.data_validations.dataValidation for rng in dv.sqref.ranges]

    assert not any(r.startswith(letter) for r in ranges)
