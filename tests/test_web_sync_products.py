import io

from openpyxl import Workbook, load_workbook


def _seed_account(web_db, platform="wb", name="ИП Яворская", warehouse_id="wh-1"):
    from app.models import PlatformAccount

    account = PlatformAccount(platform=platform, name=name, warehouse_id=warehouse_id)
    web_db.add(account)
    web_db.commit()
    web_db.refresh(account)
    return account


def _seed(web_db):
    from app.models import Product, SyncSetting

    a1 = _seed_account(web_db, name="ИП Яворская")
    a2 = _seed_account(web_db, name="Ozon-Основной", platform="ozon")

    web_db.add(Product(uid_1c="u1", article="ART-001", name="Кроссовки", stock_on_hand=5))
    web_db.add(SyncSetting(uid_1c="u1", account_id=a1.id, enabled=True))
    web_db.add(SyncSetting(uid_1c="u1", account_id=a2.id, enabled=False))
    web_db.commit()
    return a1, a2


def test_sync_products_page_renders(logged_in_client, web_db):
    _seed(web_db)
    r = logged_in_client.get("/sync-products")
    assert r.status_code == 200
    assert "Кроссовки" in r.text
    assert "ИП Яворская" in r.text


def test_sync_products_shows_three_wb_cabinets_as_separate_columns(logged_in_client, web_db):
    from app.models import Product

    _seed_account(web_db, name="ИП Яворская")
    _seed_account(web_db, name="ИП Ребрик")
    _seed_account(web_db, name="ИП Караман")
    web_db.add(Product(uid_1c="u1", article="A1", name="Товар", stock_on_hand=5))
    web_db.commit()

    r = logged_in_client.get("/sync-products")
    for name in ("ИП Яворская", "ИП Ребрик", "ИП Караман"):
        assert name in r.text


def test_toggle_checkbox_enables_account(logged_in_client, web_db):
    a1, a2 = _seed(web_db)
    r = logged_in_client.post(
        f"/sync-products/u1/{a2.id}/toggle", data={"enabled": "true"},
    )
    assert r.status_code == 200
    assert "checked" in r.text

    from app.models import SyncSetting
    setting = web_db.query(SyncSetting).filter(
        SyncSetting.uid_1c == "u1", SyncSetting.account_id == a2.id,
    ).first()
    assert setting.enabled is True

    from app.models import DispatchQueueItem
    queued = web_db.query(DispatchQueueItem).filter(DispatchQueueItem.uid_1c == "u1").all()
    assert any(q.account_id == a2.id for q in queued)


def test_toggle_off_does_not_enqueue_resend(logged_in_client, web_db):
    a1, a2 = _seed(web_db)
    logged_in_client.post(f"/sync-products/u1/{a1.id}/toggle", data={"enabled": "false"})

    from app.models import SyncSetting, DispatchQueueItem
    setting = web_db.query(SyncSetting).filter(
        SyncSetting.uid_1c == "u1", SyncSetting.account_id == a1.id,
    ).first()
    assert setting.enabled is False
    assert web_db.query(DispatchQueueItem).count() == 0


def test_update_threshold(logged_in_client, web_db):
    a1, a2 = _seed(web_db)
    r = logged_in_client.post(
        f"/sync-products/u1/{a1.id}/threshold", data={"min_threshold": "3"},
    )
    assert r.status_code == 200

    from app.models import SyncSetting
    setting = web_db.query(SyncSetting).filter(
        SyncSetting.uid_1c == "u1", SyncSetting.account_id == a1.id,
    ).first()
    assert setting.min_threshold == 3


def test_update_threshold_negative_clamped_to_zero(logged_in_client, web_db):
    a1, a2 = _seed(web_db)
    logged_in_client.post(f"/sync-products/u1/{a1.id}/threshold", data={"min_threshold": "-5"})

    from app.models import SyncSetting
    setting = web_db.query(SyncSetting).filter(
        SyncSetting.uid_1c == "u1", SyncSetting.account_id == a1.id,
    ).first()
    assert setting.min_threshold == 0


def test_sync_products_export(logged_in_client, web_db):
    a1, a2 = _seed(web_db)
    r = logged_in_client.get("/sync-products/export")
    assert r.status_code == 200

    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert headers[0:4] == ["ID_1С", "Артикул", "Наименование", "Остаток"]
    assert any("Синхронизировать" in h for h in headers)
    assert any("Порог" in h for h in headers)

    row = next(ws.iter_rows(min_row=2, values_only=True))
    row_dict = dict(zip(headers, row))
    assert row_dict["ИП Яворская (WB) — Синхронизировать"] == "Да"
    assert row_dict["Ozon-Основной (OZON) — Синхронизировать"] == "Нет"


def test_sync_products_import_bulk_toggle_and_threshold(logged_in_client, web_db):
    a1, a2 = _seed(web_db)

    wb = Workbook()
    ws = wb.active
    headers = ["ID_1С", "Артикул", "Наименование", "Остаток",
               "ИП Яворская (WB) — Синхронизировать", "ИП Яворская (WB) — Порог",
               "Ozon-Основной (OZON) — Синхронизировать", "Ozon-Основной (OZON) — Порог"]
    ws.append(headers)
    ws.append(["u1", "ART-001", "Кроссовки", 5, "Нет", 2, "Да", 0])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    r = logged_in_client.post(
        "/sync-products/import",
        files={"file": ("upd.xlsx", buf.getvalue(),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )
    assert r.status_code == 303

    r2 = logged_in_client.get("/sync-products")
    assert "Изменено строк: 2" in r2.text

    from app.models import SyncSetting
    s1 = web_db.query(SyncSetting).filter(SyncSetting.uid_1c == "u1", SyncSetting.account_id == a1.id).first()
    s2 = web_db.query(SyncSetting).filter(SyncSetting.uid_1c == "u1", SyncSetting.account_id == a2.id).first()
    assert s1.enabled is False
    assert s1.min_threshold == 2
    assert s2.enabled is True


def test_update_reserve_sets_value_and_enqueues_for_enabled(logged_in_client, web_db):
    from app.models import Product, DispatchQueueItem
    a1, a2 = _seed(web_db)  # a1 включён, a2 выключен

    r = logged_in_client.post("/sync-products/u1/reserve", data={"reserve": 2})
    assert r.status_code == 200

    web_db.expire_all()
    p = web_db.query(Product).filter(Product.uid_1c == "u1").first()
    assert p.reserve == 2

    queued = web_db.query(DispatchQueueItem).all()
    assert any(i.account_id == a1.id for i in queued)      # переотправка на включённый
    assert all(i.account_id != a2.id for i in queued)      # не на выключенный


def test_update_reserve_negative_clamped_to_zero(logged_in_client, web_db):
    from app.models import Product
    _seed(web_db)
    logged_in_client.post("/sync-products/u1/reserve", data={"reserve": -5})
    web_db.expire_all()
    assert web_db.query(Product).filter(Product.uid_1c == "u1").first().reserve == 0


def test_export_has_reserve_column(logged_in_client, web_db):
    from app.models import Product
    _seed_account(web_db)
    web_db.add(Product(uid_1c="u1", article="A", name="Т", stock_on_hand=5, reserve=3))
    web_db.commit()

    r = logged_in_client.get("/sync-products/export")
    ws = load_workbook(io.BytesIO(r.content)).active
    headers = [c.value for c in ws[1]]
    assert "Резерв" in headers
    assert ws[2][headers.index("Резерв")].value == 3


def test_import_updates_reserve(logged_in_client, web_db):
    from app.models import Product
    _seed_account(web_db)
    web_db.add(Product(uid_1c="u1", article="A", name="Т", stock_on_hand=5, reserve=0))
    web_db.commit()

    wb = Workbook()
    ws = wb.active
    ws.append(["ID_1С", "Артикул", "Наименование", "Остаток", "Резерв"])
    ws.append(["u1", "A", "Т", 5, 4])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    logged_in_client.post(
        "/sync-products/import",
        files={"file": ("f.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    web_db.expire_all()
    assert web_db.query(Product).filter(Product.uid_1c == "u1").first().reserve == 4
