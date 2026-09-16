import io

from openpyxl import Workbook, load_workbook


def _seed_account(web_db, platform="wb", name="Кабинет", warehouse_id="wh-1"):
    from app.models import PlatformAccount

    account = PlatformAccount(platform=platform, name=name, warehouse_id=warehouse_id)
    web_db.add(account)
    web_db.commit()
    web_db.refresh(account)
    return account


def _seed(web_db):
    from app.models import Product, SyncAnomaly, AnomalyReason, AnomalyStatus

    a1 = _seed_account(web_db, name="Кабинет WB")
    a2 = _seed_account(web_db, platform="ozon", name="Кабинет Ozon")

    web_db.add(Product(uid_1c="u1", article="ART-100", name="Куртка", stock_on_hand=3))
    web_db.add(Product(uid_1c="u2", article="ART-200", name="Ботинки", stock_on_hand=7))
    web_db.add(SyncAnomaly(uid_1c="u1", account_id=a1.id, reason=AnomalyReason.order_on_disabled,
                            order_id="o1", status=AnomalyStatus.new))
    web_db.add(SyncAnomaly(uid_1c="u2", account_id=a2.id, reason=AnomalyReason.missing_barcode,
                            status=AnomalyStatus.new))
    web_db.commit()
    return a1, a2


def test_anomalies_page_groups_and_sorts_by_urgency(logged_in_client, web_db):
    _seed(web_db)
    r = logged_in_client.get("/anomalies")
    assert "Куртка" in r.text and "Ботинки" in r.text
    assert r.text.index("Куртка") < r.text.index("Ботинки")
    assert "Кабинет WB" in r.text


def test_anomalies_filter_by_reason(logged_in_client, web_db):
    _seed(web_db)
    r = logged_in_client.get("/anomalies/rows?status=new&reason=missing_barcode")
    assert "Ботинки" in r.text and "Куртка" not in r.text


def test_anomalies_filter_by_account(logged_in_client, web_db):
    a1, a2 = _seed(web_db)
    r = logged_in_client.get(f"/anomalies/rows?status=new&account_id={a1.id}")
    assert "Куртка" in r.text and "Ботинки" not in r.text


def test_anomaly_resolve_button_enables_sync_and_closes(logged_in_client, web_db):
    a1, a2 = _seed(web_db)
    r = logged_in_client.post(f"/anomalies/u2/{a2.id}/resolve")
    assert r.status_code == 200

    from app.models import SyncSetting, SyncAnomaly, AnomalyStatus
    setting = web_db.query(SyncSetting).filter(
        SyncSetting.uid_1c == "u2", SyncSetting.account_id == a2.id,
    ).first()
    assert setting.enabled is True

    anomaly = web_db.query(SyncAnomaly).filter(SyncAnomaly.uid_1c == "u2").first()
    assert anomaly.status == AnomalyStatus.resolved

    r2 = logged_in_client.get("/anomalies?status=new")
    assert "Ботинки" not in r2.text


def test_anomalies_bulk_resolve_via_excel(logged_in_client, web_db):
    _seed(web_db)

    r = logged_in_client.get("/anomalies/export")
    wb = load_workbook(io.BytesIO(r.content))
    headers = [c.value for c in wb.active[1]]
    assert "Кабинет" in headers
    rows = list(wb.active.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2

    wb2 = Workbook()
    ws2 = wb2.active
    ws2.append(headers)
    for row in rows:
        row = list(row)
        if row[0] == "u1":  # помечаем только Куртку
            row[-1] = "Да"
        ws2.append(row)
    buf = io.BytesIO()
    wb2.save(buf)
    buf.seek(0)

    logged_in_client.post(
        "/anomalies/import",
        files={"file": ("imp.xlsx", buf.getvalue(),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    r2 = logged_in_client.get("/anomalies?status=new")
    assert "Куртка" not in r2.text
    assert "Ботинки" in r2.text
