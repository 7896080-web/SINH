"""Страница «Расхождения учёта со складом».

Число `Product.stock_discrepancy` решает, сколько уходит на площадки
(`порог = расхождение + бронь`), а увидеть его можно было только построчно в
каталоге на 152 тысячи позиций либо скриптом с консоли сервера. Происхождение
числа система записывает (`StockDiscrepancyLog`), но читателя у этой записи не
было вовсе — писатель без читателя.
"""
import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from app.models import DiscrepancySource, Product, StockDiscrepancyLog
from app.timeutils import now_utc


@pytest.fixture()
def catalog(web_db):
    rows = [
        # uid,      расхожд, бронь, трансляция
        ("u-плюс",       12,  2, True),
        ("u-минус",      -7,  0, True),
        ("u-большой",   200,  0, False),
        # Большой МИНУС обязателен в наборе: без него сортировка по модулю и
        # по знаку дают один порядок, и подмена одной другой не ловится.
        ("u-край",     -213,  0, True),     # реальный край на бою
        ("u-ноль",        0,  0, True),    # измеряли, склад сошёлся — НЕ расхождение
        ("u-не-мерян", None,  0, True),    # не измеряли вовсе
    ]
    for uid, gap, reserve, on in rows:
        web_db.add(Product(uid_1c=uid, article=uid, name=f"Свитшот {uid}",
                           size="5XL", color="LACIVERT", stock_on_hand=50,
                           reserve=reserve, stock_discrepancy=gap,
                           broadcast_offset=(gap + reserve) if gap is not None else None,
                           broadcast_enabled=on))
    # Происхождение есть только у одного — у остальных число старше истории.
    web_db.add(StockDiscrepancyLog(
        uid_1c="u-минус", old_value=None, new_value=-7,
        source=DiscrepancySource.fact, username="operator",
        base_date=date(2026, 8, 7), base_stock=43, fact=50,
        created_at=now_utc() - timedelta(days=3)))
    web_db.commit()


def _page(client, query=""):
    answer = client.get(f"/discrepancies{query}")
    assert answer.status_code == 200, answer.text[:400]
    return answer.text


def test_only_measured_non_zero_rows_are_listed(logged_in_client, catalog):
    """Ноль — это «измеряли, склад сошёлся», то есть ОТСУТСТВИЕ расхождения.

    Держать такие строки в списке «с расхождениями» значит утопить в них то,
    ради чего список открыли. NULL («не измеряли») — тем более."""
    page = _page(logged_in_client)
    assert "u-плюс" in page and "u-минус" in page and "u-большой" in page
    assert "u-ноль" not in page, "ноль — не расхождение"
    assert "u-не-мерян" not in page, "не измеряли — не расхождение"


def test_the_biggest_disagreement_comes_first(logged_in_client, catalog):
    """Сортировка ПО МОДУЛЮ, а не по знаку.

    Разбирают список сверху, и величина расхождения — мера того, насколько учёт
    разошёлся со складом. Сортируй мы по знаку, первым шёл бы самый
    отрицательный, а плюс на 200 штук стоит разбора не меньше минуса на 213."""
    page = _page(logged_in_client)
    order = [page.index(uid) for uid in ("u-край", "u-большой", "u-плюс", "u-минус")]
    assert order == sorted(order), (
        "порядок не по модулю расхождения: −213, 200, 12, −7")


def test_the_negative_filter_keeps_only_the_overselling_side(logged_in_client, catalog):
    """По отрицательным наружу осознанно уходит БОЛЬШЕ учётного остатка —
    устаревшее измерение здесь уже оверселл, и смотрят их отдельно."""
    page = _page(logged_in_client, "?only_negative=true")
    assert "u-минус" in page
    assert "u-край" in page
    assert "u-плюс" not in page and "u-большой" not in page


def test_the_broadcasting_filter_keeps_only_what_goes_out_now(logged_in_client, catalog):
    """По невещающему товару неверное расхождение ничем пока не грозит."""
    page = _page(logged_in_client, "?only_broadcasting=true")
    assert "u-плюс" in page and "u-минус" in page
    assert "u-большой" not in page       # трансляция выключена


def test_the_row_says_where_the_number_came_from(logged_in_client, catalog):
    """Без происхождения строка называет число и молчит о том, на каком
    основании оно принято, — а пересчитывать склад идут именно с этим вопросом."""
    page = _page(logged_in_client, "?only_negative=true")
    assert "введён факт" in page
    assert "operator" in page
    assert "07.08.2026" in page and "43" in page      # контекст измерения


def test_a_number_older_than_the_history_says_so(logged_in_client, catalog):
    """История заведена 23.09. Пустая ячейка читалась бы как «никто не ставил»."""
    page = _page(logged_in_client)
    assert "число старше истории" in page


def test_the_export_carries_the_same_filter(logged_in_client, catalog):
    """Ссылка выгрузки передаёт отбор — эндпоинт обязан его принять.

    `only_unfinished` на «Товарах» на этом уже спотыкался: ссылка передавала,
    эндпоинт не принимал, FastAPI молча отбрасывал, оператор получал весь
    каталог вместо отобранного."""
    dump = logged_in_client.get("/discrepancies/export?only_negative=true")
    assert dump.status_code == 200
    sheet = load_workbook(io.BytesIO(dump.content)).active
    ids = {row[0] for row in sheet.iter_rows(min_row=2, values_only=True) if row[0]}
    assert ids == {"u-минус", "u-край"}, ids


def test_the_export_carries_the_origin_of_every_number(logged_in_client, catalog):
    """Разбирают такие строки пачкой в файле, а не глазами по одной."""
    dump = logged_in_client.get("/discrepancies/export")
    sheet = load_workbook(io.BytesIO(dump.content)).active
    headers = [c.value for c in next(sheet.iter_rows(max_row=1))]
    for column in ("Расхождение", "Чем поставлено", "Кто", "Когда",
                   "Дата расчёта", "Учёт 1С на дату", "Факт на дату"):
        assert column in headers, f"в выгрузке нет колонки «{column}»: {headers}"


def test_a_clean_catalogue_says_so_plainly(logged_in_client, web_db):
    """Молчание на исправной системе — обязательное свойство: пустая таблица без
    объяснения читается как поломка страницы."""
    web_db.add(Product(uid_1c="u1", article="a", name="n", stock_on_hand=5,
                       reserve=0, stock_discrepancy=None))
    web_db.commit()
    page = _page(logged_in_client)
    assert "Расхождений нет" in page


def test_the_page_is_in_the_menu(logged_in_client, catalog):
    """Страница, до которой нет ссылки, — это страница, которой нет."""
    page = _page(logged_in_client)
    assert 'href="/discrepancies"' in page
    assert 'href="#i-discrepancies"' in page, "пункт меню без иконки — пустое место"
    assert 'id="i-discrepancies"' in page, "иконки нет в спрайте"
