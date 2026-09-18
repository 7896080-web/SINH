import io

from openpyxl import Workbook, load_workbook


def _seed_basic_product(web_db):
    from app.models import Product, Barcode

    web_db.add(Product(uid_1c="u1", article="ART-001", name="Кроссовки", stock_on_hand=5))
    web_db.add(Barcode(barcode="4600000000001", uid_1c="u1", source_platform="wb"))
    web_db.commit()
    return "u1"


def _seed_account(web_db, platform="ozon", name="Основной"):
    from app.models import PlatformAccount

    account = PlatformAccount(platform=platform, name=name, warehouse_id="wh-1")
    web_db.add(account)
    web_db.commit()
    web_db.refresh(account)
    return account


def test_mapping_page_shows_mapped_barcode(logged_in_client, web_db):
    _seed_basic_product(web_db)
    r = logged_in_client.get("/mapping")
    assert r.status_code == 200
    assert "4600000000001" in r.text
    assert "ART-001" in r.text


def test_mapping_rows_htmx_fragment_filters_by_query(logged_in_client, web_db):
    _seed_basic_product(web_db)
    r = logged_in_client.get("/mapping/rows?q=4600000000001")
    assert "4600000000001" in r.text
    r2 = logged_in_client.get("/mapping/rows?q=nothing-matches")
    assert "4600000000001" not in r2.text


def test_mapping_conflicts_tab_empty_by_default(logged_in_client, web_db):
    r = logged_in_client.get("/mapping?view=conflicts")
    assert "Конфликтов сопоставления не найдено" in r.text


def test_mapping_export_mapped_view(logged_in_client, web_db):
    _seed_basic_product(web_db)
    r = logged_in_client.get("/mapping/export")
    assert r.status_code == 200

    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert [c.value for c in ws[1]] == ["ID_1С", "Баркод", "Артикул", "Наименование",
                                        "Размер", "Цвет", "Источник", "Добавлен"]
    assert ws.cell(row=2, column=2).value == "4600000000001"


def test_mapping_import_adds_new_barcode(logged_in_client, web_db):
    uid = _seed_basic_product(web_db)

    wb = Workbook()
    ws = wb.active
    ws.append(["ID_1С", "Баркод", "Артикул", "Наименование", "Источник", "Добавлен"])
    ws.append([uid, "4600000000099", "ART-001", "Кроссовки", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    r = logged_in_client.post(
        "/mapping/import",
        files={"file": ("import.xlsx", buf.getvalue(),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )
    assert r.status_code == 303

    r2 = logged_in_client.get("/mapping")
    assert "Добавлено баркодов: 1" in r2.text
    assert "4600000000099" in r2.text


def test_mapping_import_reports_error_for_unknown_product(logged_in_client, web_db):
    wb = Workbook()
    ws = wb.active
    ws.append(["ID_1С", "Баркод", "Артикул", "Наименование", "Источник", "Добавлен"])
    ws.append(["does-not-exist", "111", "", "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    logged_in_client.post(
        "/mapping/import",
        files={"file": ("import.xlsx", buf.getvalue(),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    r = logged_in_client.get("/mapping")
    assert "не найден" in r.text


def test_mapping_conflicts_page_shows_account_name(logged_in_client, web_db):
    from app.models import MappingConflict

    account = _seed_account(web_db, platform="ozon", name="ИП Яворская")
    web_db.add(MappingConflict(barcode="777", account_id=account.id, attempts=1))
    web_db.commit()

    r = logged_in_client.get("/mapping?view=conflicts")
    assert "777" in r.text
    assert "ИП Яворская" in r.text


def test_mapping_conflicts_filter_by_account(logged_in_client, web_db):
    from app.models import MappingConflict

    a1 = _seed_account(web_db, platform="wb", name="Кабинет 1")
    a2 = _seed_account(web_db, platform="wb", name="Кабинет 2")
    web_db.add(MappingConflict(barcode="111", account_id=a1.id, attempts=1))
    web_db.add(MappingConflict(barcode="222", account_id=a2.id, attempts=1))
    web_db.commit()

    r = logged_in_client.get(f"/mapping/rows?view=conflicts&account_id={a1.id}")
    assert "111" in r.text and "222" not in r.text


def test_mapping_load_catalog_without_credentials_shows_friendly_error(logged_in_client, web_db):
    account = _seed_account(web_db, platform="ozon", name="Основной")

    r = logged_in_client.post(f"/mapping/load-catalog/{account.id}", follow_redirects=False)
    assert r.status_code == 303

    r2 = logged_in_client.get("/mapping?view=conflicts")
    assert "Не удалось загрузить" in r2.text


def test_mapping_load_catalog_enriches_conflict_with_platform_name(logged_in_client, web_db, monkeypatch):
    from app.models import MappingConflict
    from app.workers.platform_clients.base import CatalogItem
    import app.routers.mapping as mapping_router

    account = _seed_account(web_db, platform="ozon", name="Основной")
    web_db.add(MappingConflict(barcode="777", account_id=account.id, attempts=1))
    web_db.commit()

    class FakeClient:
        def get_catalog_items(self):
            return [CatalogItem(external_id="e1", barcode="777", article="ART-777", name="Товар с Ozon")]

    monkeypatch.setattr(mapping_router, "build_client", lambda db, account_id: FakeClient())

    r = logged_in_client.post(f"/mapping/load-catalog/{account.id}", follow_redirects=False)
    assert r.status_code == 303

    r2 = logged_in_client.get("/mapping?view=conflicts")
    assert "загружено карточек 1" in r2.text
    assert "Товар с Ozon" in r2.text
    assert "ART-777" in r2.text


def test_mapping_conflict_resolution_full_cycle_via_excel(logged_in_client, web_db):
    """Экспорт конфликта с пустым ID_1С -> заполнить -> загрузить обратно ->
    баркод переезжает в «Сопоставленные», конфликт исчезает."""
    from app.models import Product, MappingConflict

    account = _seed_account(web_db, platform="wb", name="ИП Яворская")
    web_db.add(Product(uid_1c="u9", article="A9", name="Шапка", stock_on_hand=1))
    web_db.add(MappingConflict(barcode="555", account_id=account.id, attempts=2))
    web_db.commit()

    r = logged_in_client.get("/mapping/export?view=conflicts")
    wb = load_workbook(io.BytesIO(r.content))
    headers = [c.value for c in wb.active[1]]
    assert headers[0] == "ID_1С"

    wb2 = Workbook()
    ws2 = wb2.active
    ws2.append(headers)
    ws2.append(["u9", "555", "ИП Яворская", "wb", "", "", 2, "", ""])
    buf = io.BytesIO()
    wb2.save(buf)
    buf.seek(0)

    logged_in_client.post(
        "/mapping/import",
        files={"file": ("resolve.xlsx", buf.getvalue(),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    r2 = logged_in_client.get("/mapping?view=conflicts")
    assert "555" not in r2.text

    r3 = logged_in_client.get("/mapping")
    assert "555" in r3.text


# ------------------------------------------------ переподвязка перепутанных баркодов

def _seed_two_sizes(web_db):
    """Ровно случай с боя: баркоды размеров перепутаны местами.

    В карточках 1С баркод размера M принадлежал L и наоборот. В 1С это
    исправили, а к нам исправление не доезжает: приём справочника заводит
    только отсутствующие баркоды, а импорт из Excel до этой правки отвечал
    «уже привязан к другому товару» и строку пропускал. Пока привязка неверна,
    заказ на M списывает L.
    """
    from app.models import Barcode, Product
    from app.timeutils import now_utc

    for uid, size in (("u-l", "L"), ("u-m", "M")):
        web_db.add(Product(uid_1c=uid, article="ZJYM269002", name="Джемпер",
                           size=size, color="011/коричневый", stock_on_hand=10,
                           recalc_done_at=now_utc(), recalc_account_ids="1"))
    web_db.add(Barcode(barcode="bc-L", uid_1c="u-m"))     # перепутано
    web_db.add(Barcode(barcode="bc-M", uid_1c="u-l"))     # перепутано
    web_db.commit()


def _repoint_file(rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["ID_1С", "Баркод", "Артикул", "Наименование", "Размер", "Цвет",
               "Источник", "Добавлен"])
    for uid, barcode in rows:
        ws.append([uid, barcode, "ZJYM269002", "Джемпер", "", "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _post_import(client, payload, repoint=False):
    return client.post(
        "/mapping/import",
        files={"file": ("import.xlsx", payload,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"repoint": "true"} if repoint else {},
        follow_redirects=False,
    )


def test_import_refuses_to_repoint_without_the_checkbox(logged_in_client, web_db):
    """Переподвязка меняет, на какой товар спишется заказ. Случайно загруженный
    старый файл не должен молча переразнести каталог."""
    from app.models import Barcode

    _seed_two_sizes(web_db)

    _post_import(logged_in_client, _repoint_file([("u-l", "bc-L"), ("u-m", "bc-M")]))

    web_db.expire_all()
    assert web_db.query(Barcode).filter(Barcode.barcode == "bc-L").one().uid_1c == "u-m"
    assert "разрешить переподвязку" in logged_in_client.get("/mapping").text


def test_swapped_barcodes_are_put_back_in_one_import(logged_in_client, web_db):
    """Обмен местами проходит одним файлом: баркод уникален сам по себе, мы его
    не пересоздаём, а переставляем ссылку — двум строкам столкнуться не на чем."""
    from app.models import Barcode

    _seed_two_sizes(web_db)

    _post_import(logged_in_client, _repoint_file([("u-l", "bc-L"), ("u-m", "bc-M")]),
                 repoint=True)

    web_db.expire_all()
    assert web_db.query(Barcode).filter(Barcode.barcode == "bc-L").one().uid_1c == "u-l"
    assert web_db.query(Barcode).filter(Barcode.barcode == "bc-M").one().uid_1c == "u-m"
    assert "Переподвязано: 2" in logged_in_client.get("/mapping").text


def test_repointing_drops_the_done_mark_on_both_products(logged_in_client, web_db):
    """Расчёт собирал заказы по прежнему набору баркодов — к новому его вывод
    не относится. Оставить «актуализирован» значило бы разрешить трансляцию
    остатка, сверенного не по тем продажам."""
    from app.models import Product

    _seed_two_sizes(web_db)

    _post_import(logged_in_client, _repoint_file([("u-l", "bc-L")]), repoint=True)

    web_db.expire_all()
    for uid in ("u-l", "u-m"):
        product = web_db.query(Product).filter(Product.uid_1c == uid).one()
        assert product.recalc_done_at is None, uid
        assert product.recalc_account_ids is None, uid


def test_an_unchanged_row_is_not_counted_as_repointed(logged_in_client, web_db):
    """Файл выгружают целиком, а правят одну-две строки. Остальные должны
    проходить как «уже сопоставлены», а не тревожить сбросом расчёта."""
    from app.models import Product

    _seed_two_sizes(web_db)

    _post_import(logged_in_client, _repoint_file([("u-m", "bc-L")]), repoint=True)

    web_db.expire_all()
    body = logged_in_client.get("/mapping").text
    assert "Уже были сопоставлены: 1" in body
    assert "Переподвязано" not in body
    assert web_db.query(Product).filter(Product.uid_1c == "u-m").one().recalc_done_at is not None


def test_export_carries_size_and_colour(logged_in_client, web_db):
    """Без них строки файла различаются только двумя непрозрачными
    идентификаторами, и перепутать размеры в нём проще, чем исправить."""
    _seed_two_sizes(web_db)

    r = logged_in_client.get("/mapping/export?view=mapped")

    ws = load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    assert "Размер" in header and "Цвет" in header
    rows = {r[header.index("Баркод")].value: r[header.index("Размер")].value
            for r in ws.iter_rows(min_row=2)}
    assert rows["bc-L"] == "M"        # пока перепутано — файл это и показывает
    assert rows["bc-M"] == "L"
