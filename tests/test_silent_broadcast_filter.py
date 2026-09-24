"""Фильтр «только с выключенной трансляцией».

Обратная сторона `only_broadcasting`. Спрашивают им не «что сейчас уходит
наружу», а «что молчит»: что ещё предстоит подключить и что могло замолчать не
по чьему-то решению. На каталоге в 152 тысячи SKU иначе этого не увидеть.
"""
import io

import pytest
from openpyxl import load_workbook

from app.models import Product, PlatformAccount, SyncSetting
from app.timeutils import now_utc


@pytest.fixture()
def catalog(web_db):
    account = PlatformAccount(platform="wb", name="ИП А", warehouse_id="1",
                              is_active=True, dispatch_enabled=True)
    web_db.add(account)
    web_db.flush()
    for uid, on in (("u-вкл", True), ("u-выкл", False), ("u-выкл-2", False)):
        web_db.add(Product(uid_1c=uid, article=uid, name=uid, stock_on_hand=10,
                           reserve=0, broadcast_enabled=on))
        web_db.add(SyncSetting(uid_1c=uid, account_id=account.id, enabled=True))
    # Строка с непогашенной просьбой включить трансляцию: сама она выключена,
    # но включится, когда закончится расчёт.
    web_db.add(Product(uid_1c="u-просит", article="u-просит", name="u-просит",
                       stock_on_hand=10, reserve=0, broadcast_enabled=False,
                       broadcast_requested_at=now_utc()))
    web_db.commit()
    return account


def test_the_filter_shows_only_the_silent_rows(logged_in_client, catalog):
    page = logged_in_client.get("/products?only_silent=true")
    assert page.status_code == 200
    assert "u-выкл" in page.text and "u-выкл-2" in page.text
    assert "u-вкл<" not in page.text      # включённая строка сюда не попадает


def test_the_two_filters_add_up_to_the_whole_catalogue(logged_in_client, catalog):
    """Прямой и обратный фильтр в сумме дают каталог — без дыры и без нахлёста.

    Это и есть причина, по которой обратный фильтр спрашивает СТРОГО обратное, а
    не «выключена и не просила включения». Исключи мы отсюда ждущие расчёта
    строки — они не показывались бы НИ ПОД ОДНИМ фильтром, и увидеть их можно
    было бы, только специально их разыскивая.
    """
    def uids(query):
        text = logged_in_client.get(f"/products?{query}").text
        return {u for u in ("u-вкл", "u-выкл", "u-выкл-2", "u-просит")
                if f">{u}<" in text or f"value=\"{u}\"" in text}

    on = uids("only_broadcasting=true")
    off = uids("only_silent=true")
    assert on & off == set(), f"строка попала под оба фильтра: {on & off}"
    assert on | off == {"u-вкл", "u-выкл", "u-выкл-2", "u-просит"}, (
        f"не под одним фильтром: {{'u-вкл','u-выкл','u-выкл-2','u-просит'}} - {on | off}")


def test_a_row_waiting_for_a_calc_counts_as_silent(logged_in_client, catalog):
    """Просьба включить трансляцию — это ещё не трансляция.

    Пока расчёт не закончен, товар молчит, и оператор, отбирающий молчащие
    строки, обязан его видеть."""
    page = logged_in_client.get("/products?only_silent=true").text
    assert "u-просит" in page
    assert "u-вкл<" not in page      # а включённая строка — нет


def test_both_filters_at_once_say_why_the_list_is_empty(logged_in_client, catalog):
    """Пустая таблица без объяснения читается как «таких товаров нет».

    А это ответ не про товары, а про сами фильтры: они спрашивают об одном поле
    противоположное. Молчаливо отдать пустой список значило бы соврать про
    каталог."""
    page = logged_in_client.get("/products?only_broadcasting=true&only_silent=true")
    assert page.status_code == 200
    assert "Снимите одну из галочек" in page.text
    assert "u-вкл" not in page.text and "u-выкл" not in page.text


def test_the_export_gets_the_filter_too(logged_in_client, catalog):
    """Ссылка выгрузки передаёт фильтр — эндпоинт обязан его принять.

    `only_unfinished` на этом уже спотыкался: ссылка его передавала, эндпоинт не
    принимал, FastAPI молча отбрасывал, и оператор получал весь каталог вместо
    отобранного."""
    dump = logged_in_client.get("/products/export?only_silent=true")
    assert dump.status_code == 200
    sheet = load_workbook(io.BytesIO(dump.content)).active
    ids = {row[0] for row in sheet.iter_rows(min_row=2, values_only=True) if row[0]}
    assert ids == {"u-выкл", "u-выкл-2", "u-просит"}, ids


def test_bulk_edit_by_the_whole_selection_respects_the_filter(logged_in_client,
                                                              web_db, catalog):
    """Правка «по всему отбору» обязана видеть тот же отбор, что и страница.

    Фильтр уезжает в форму скрытым полем из ФРАГМЕНТА таблицы: он
    перерисовывается при каждом изменении фильтров, а шапка страницы — нет.
    Потеряйся он по дороге, правка ушла бы по всему каталогу — включая строки,
    которые оператор отбором как раз исключил."""
    answer = logged_in_client.post("/products/bulk", data={
        "action": "set_reserve", "int_value": "7",
        "only_silent": "true", "all_filtered": "true",
    }, follow_redirects=False)
    assert answer.status_code == 303
    assert "only_silent=true" in answer.headers["location"], (
        "возврат потерял фильтр — оператор попадёт на полный список")

    web_db.expire_all()
    by_uid = {p.uid_1c: p.reserve
              for p in web_db.query(Product).all()}
    assert by_uid == {"u-вкл": 0, "u-выкл": 7, "u-выкл-2": 7, "u-просит": 7}, by_uid


def test_the_table_fragment_carries_the_filter_into_the_bulk_form(logged_in_client,
                                                                  catalog):
    """Скрытое поле фильтра живёт во ФРАГМЕНТЕ таблицы, и это надо проверять там.

    Предыдущий тест шлёт фильтр формой напрямую — то есть проверяет обработчик, а
    не страницу. Потеряйся поле в шаблоне, он всё равно был бы зелёным, а
    оператор, нажавший «изменить весь отбор», правил бы ВЕСЬ КАТАЛОГ: сервер
    получил бы форму без фильтра и понял бы её буквально. Ровно этот класс уже
    ловили на странице «Мэппинга»: запросы были верными, ошибка — в том, как их
    зовёт страница."""
    fragment = logged_in_client.get("/products/rows?only_silent=true").text
    assert 'name="only_silent"' in fragment and 'form="bulk-form"' in fragment
    # И обратное: без фильтра поля быть не должно — иначе форма отправляла бы
    # отбор, которого оператор не просил.
    assert 'name="only_silent"' not in logged_in_client.get("/products/rows").text
