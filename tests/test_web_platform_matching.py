import io

from openpyxl import load_workbook


def _account(web_db, platform, name):
    from app.models import PlatformAccount
    a = PlatformAccount(platform=platform, name=name, warehouse_id="wh")
    web_db.add(a)
    web_db.commit()
    web_db.refresh(a)
    return a


def _catalog(web_db, account, external_id, barcode, article="ART", name="Товар"):
    from app.models import PlatformCatalogItem
    web_db.add(PlatformCatalogItem(
        account_id=account.id, external_id=external_id,
        barcode=barcode, article=article, name=name,
    ))
    web_db.commit()


def _seed_pool_overlap(web_db):
    """WB-карточка = пул {SHARED, WBONLY}; Ozon = {SHARED} (по размеру).
    Пулы пересекаются по SHARED, значит WBONLY тоже в этом товаре, хотя на
    Ozon его нет — это и проверяем."""
    wb = _account(web_db, "wb", "ИП Яворская")
    oz = _account(web_db, "ozon", "ОЗОН")
    _catalog(web_db, wb, "wbcard", "SHARED", article="A-SH", name="Куртка")
    _catalog(web_db, wb, "wbcard", "WBONLY", article="A-SH", name="Куртка")
    _catalog(web_db, oz, "oz-1", "SHARED", article="A-SH-OZ", name="Куртка Ozon")
    return wb, oz


def test_page_renders(logged_in_client, web_db):
    _seed_pool_overlap(web_db)
    r = logged_in_client.get("/platform-matching")
    assert r.status_code == 200
    assert "Куртка" in r.text


def test_pool_clusters_across_platforms_via_shared_barcode(logged_in_client, web_db):
    """WBONLY нет на Ozon, но он в одном пуле с SHARED (общий баркод) —
    кластер должен охватывать обе площадки и содержать 2 баркода."""
    _seed_pool_overlap(web_db)
    from app.routers.platform_matching import _build_clusters
    clusters = _build_clusters(web_db)
    assert len(clusters) == 1
    c = clusters[0]
    assert set(c["platforms"]) == {"wb", "ozon"}
    assert set(c["barcodes"]) == {"SHARED", "WBONLY"}
    assert c["barcode_count"] == 2
    assert c["status"] == "unmapped"


def test_coverage_multi(logged_in_client, web_db):
    wb, oz = _seed_pool_overlap(web_db)
    # добавим одиночный Ozon-товар, не связанный ни с чем
    _catalog(web_db, oz, "oz-solo", "SOLO", article="A-SOLO", name="Одиночка")
    r = logged_in_client.get("/platform-matching/rows?coverage=multi")
    assert "Куртка" in r.text
    assert "Одиночка" not in r.text  # одноплощадочный отфильтрован


def test_partial_status_when_one_size_mapped(logged_in_client, web_db):
    """Привязали ОДИН баркод пула к 1С: по правилу SKU сопоставлен (partial),
    но не все размеры привязаны для резолвинга заказов."""
    from app.models import Product, Barcode
    _seed_pool_overlap(web_db)
    web_db.add(Product(uid_1c="u-sh", article="A-SH", name="Куртка", stock_on_hand=0))
    web_db.add(Barcode(barcode="SHARED", uid_1c="u-sh"))
    web_db.commit()

    from app.routers.platform_matching import _build_clusters
    c = _build_clusters(web_db)[0]
    assert c["status"] == "partial"
    assert c["uid_1cs"] == ["u-sh"]

    # оба баркода привязаны -> mapped
    web_db.add(Barcode(barcode="WBONLY", uid_1c="u-sh"))
    web_db.commit()
    c2 = _build_clusters(web_db)[0]
    assert c2["status"] == "mapped"


def test_mapped_unmapped_filter(logged_in_client, web_db):
    from app.models import Product, Barcode
    _seed_pool_overlap(web_db)
    web_db.add(Product(uid_1c="u-sh", article="A-SH", name="Куртка", stock_on_hand=0))
    web_db.add(Barcode(barcode="SHARED", uid_1c="u-sh"))
    web_db.commit()

    r_mapped = logged_in_client.get("/platform-matching/rows?mapped=mapped")
    assert "Куртка" in r_mapped.text
    r_unmapped = logged_in_client.get("/platform-matching/rows?mapped=unmapped")
    assert "Куртка" not in r_unmapped.text


def test_search_by_barcode_in_pool(logged_in_client, web_db):
    _seed_pool_overlap(web_db)
    r = logged_in_client.get("/platform-matching/rows?q=WBONLY")
    assert "Куртка" in r.text  # ищется по баркоду внутри пула
    r2 = logged_in_client.get("/platform-matching/rows?q=нет-такого")
    assert "Куртка" not in r2.text


def test_export_one_row_per_barcode(logged_in_client, web_db):
    _seed_pool_overlap(web_db)
    r = logged_in_client.get("/platform-matching/export")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert "ID_1С" in headers and "Баркод" in headers
    barcodes = {row[headers.index("Баркод")].value for row in ws.iter_rows(min_row=2)}
    assert barcodes == {"SHARED", "WBONLY"}  # по строке на каждый баркод пула


def test_page_in_nav(logged_in_client, web_db):
    r = logged_in_client.get("/mapping")
    assert "/platform-matching" in r.text
    assert "Сопоставление площадок" in r.text
