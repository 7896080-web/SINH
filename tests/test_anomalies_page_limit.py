"""Страница аномалий обязана открываться и через год.

Строки в статусе `new` не чистит НИКТО намеренно: аномалия в `new` — это
незаконченное дело, и `retention` её не трогает, сколько бы ей ни было. Пока
идёт переход со старой системы, «заказ на невключённый товар» приходит
десятками в сутки и сам не закрывается — пару товар+кабинет никто не включит,
потому что этим каталогом мы ещё не управляем. То есть список растёт
монотонно, а страница обновляет себя каждые две минуты.

Отсюда предел показа и, главное, три вещи вокруг него: счётчик обязан называть
ЧИСЛО ОТБОРА, а не длину показанного, страница — говорить, что список обрезан, а
выгрузка — брать весь отбор. Тот же тройной дефект уже находили на конфликтах
сопоставления: чинить надо все три, иначе оператор всё равно не узнает.
"""

from datetime import timedelta

import pytest

from app.models import (AnomalyReason, AnomalyStatus, Platform, PlatformAccount,
                        Product, SyncAnomaly)
from app.routers import anomalies as R
from app.timeutils import now_utc


@pytest.fixture()
def account(web_db):
    row = PlatformAccount(platform=Platform.wb, name="ИП ЯВОРСКАЯ",
                          warehouse_id="wh-1")
    web_db.add(row)
    web_db.commit()
    web_db.refresh(row)
    return row


def _anomaly(web_db, account, uid, minutes_ago=0,
             reason=AnomalyReason.order_on_disabled):
    web_db.add(Product(uid_1c=uid, article=f"АРТ-{uid}", name=f"Товар {uid}",
                       stock_on_hand=1))
    web_db.add(SyncAnomaly(uid_1c=uid, account_id=account.id, reason=reason,
                           order_id=f"o-{uid}", status=AnomalyStatus.new,
                           detected_at=now_utc() - timedelta(minutes=minutes_ago)))


def test_the_page_shows_the_limit_and_says_how_many_there_are(
        monkeypatch, logged_in_client, web_db, account):
    monkeypatch.setattr(R, "PAGE_LIMIT", 2)
    for i in range(5):
        _anomaly(web_db, account, f"u{i}", minutes_ago=i)
    web_db.commit()

    page = logged_in_client.get("/anomalies")

    assert page.text.count("<tr id=\"anomaly-") == 2, "предел показа не действует"
    assert "из 5" in page.text, (
        "«показаны первые 2» без «из 5» читается как «столько их и есть»")


def test_the_notice_does_not_show_when_everything_fits(
        monkeypatch, logged_in_client, web_db, account):
    """Оговорка на полном списке — лишний шум, от которого перестают читать
    оговорки вообще."""
    monkeypatch.setattr(R, "PAGE_LIMIT", 10)
    _anomaly(web_db, account, "u1")
    web_db.commit()

    page = logged_in_client.get("/anomalies")
    assert "Показаны первые" not in page.text


def test_the_total_is_not_recounted_when_everything_fits(
        monkeypatch, logged_in_client, web_db, account):
    """Счёт отбора — лишний `GROUP BY`, и на полном списке он не нужен вовсе.

    Страница обновляет себя каждые две минуты, то есть этот запрос пошёл бы на
    каждое обновление — ровно та цена, от которой предел показа и спасает.
    Пришло меньше предела: это и есть всё, считать нечего."""
    counted: list[int] = []
    real = R._count_grouped_anomalies
    monkeypatch.setattr(R, "_count_grouped_anomalies",
                        lambda *a, **kw: counted.append(1) or real(*a, **kw))
    monkeypatch.setattr(R, "PAGE_LIMIT", 10)
    _anomaly(web_db, account, "u1")
    web_db.commit()

    page = logged_in_client.get("/anomalies")

    assert "Показаны первые" not in page.text
    assert counted == [], "лишний пересчёт отбора на полном списке"


def test_the_total_is_counted_when_the_list_is_cut(
        monkeypatch, logged_in_client, web_db, account):
    """Обратная сторона: обрезали — обязаны назвать число отбора.

    Без этой проверки оптимизация выше однажды выключила бы счёт совсем, и
    «показаны первые 300» снова читалось бы как «столько их и есть»."""
    counted: list[int] = []
    real = R._count_grouped_anomalies
    monkeypatch.setattr(R, "_count_grouped_anomalies",
                        lambda *a, **kw: counted.append(1) or real(*a, **kw))
    monkeypatch.setattr(R, "PAGE_LIMIT", 2)
    for i in range(5):
        _anomaly(web_db, account, f"u{i}", minutes_ago=i)
    web_db.commit()

    page = logged_in_client.get("/anomalies")

    assert counted == [1]
    assert "из 5" in page.text


def test_the_limit_cuts_the_oldest_rows_not_random_ones(
        monkeypatch, logged_in_client, web_db, account):
    """Сортировка обязана быть В SQL.

    Отсортируй мы после `LIMIT`, страница показывала бы произвольные строки и
    разложила бы по порядку уже их — то есть «первые две» были бы не самыми
    свежими, а случайными, и сказать об этом было бы нечем."""
    monkeypatch.setattr(R, "PAGE_LIMIT", 2)
    _anomaly(web_db, account, "свежая", minutes_ago=1)
    _anomaly(web_db, account, "средняя", minutes_ago=100)
    _anomaly(web_db, account, "древняя", minutes_ago=10000)
    web_db.commit()

    page = logged_in_client.get("/anomalies")

    assert "АРТ-свежая" in page.text
    assert "АРТ-средняя" in page.text
    assert "АРТ-древняя" not in page.text


def test_the_reason_order_survives_the_limit(
        monkeypatch, logged_in_client, web_db, account):
    """Порядок причин берётся из `REASON_PRIORITY`, и он один.

    По алфавиту значения встают наоборот («missing_barcode» раньше
    «order_on_disabled»), то есть предел показа выкидывал бы со страницы ровно
    те строки, которые разбирают первыми."""
    monkeypatch.setattr(R, "PAGE_LIMIT", 1)
    _anomaly(web_db, account, "свежийбаркод", minutes_ago=1,
             reason=AnomalyReason.missing_barcode)
    _anomaly(web_db, account, "старыйзаказ", minutes_ago=5000,
             reason=AnomalyReason.order_on_disabled)
    web_db.commit()

    page = logged_in_client.get("/anomalies")

    assert "АРТ-старыйзаказ" in page.text, (
        "причина с высшим приоритетом не попала на страницу")
    assert "АРТ-свежийбаркод" not in page.text


def test_the_export_takes_the_whole_selection_not_the_page(
        monkeypatch, logged_in_client, web_db, account):
    """Файл правят руками и загружают обратно: «моей строки в файле нет»
    оператор читает как «её нет в системе»."""
    import io

    from openpyxl import load_workbook

    monkeypatch.setattr(R, "PAGE_LIMIT", 2)
    for i in range(7):
        _anomaly(web_db, account, f"u{i}", minutes_ago=i)
    web_db.commit()

    answer = logged_in_client.get("/anomalies/export")

    book = load_workbook(io.BytesIO(answer.content))
    # Минус шапка.
    assert book.active.max_row - 1 == 7, "в файл ушла страница, а не отбор"


def test_the_refreshed_fragment_carries_the_notice_too(
        monkeypatch, logged_in_client, web_db, account):
    """Страница обновляет себя сама, и подменяется именно фрагмент.

    Живи оговорка только в полной странице, она исчезала бы через две минуты
    после открытия — то есть ровно тогда, когда человек начинает работать."""
    monkeypatch.setattr(R, "PAGE_LIMIT", 1)
    for i in range(4):
        _anomaly(web_db, account, f"u{i}", minutes_ago=i)
    web_db.commit()

    fragment = logged_in_client.get("/anomalies/rows")

    assert "из 4" in fragment.text
