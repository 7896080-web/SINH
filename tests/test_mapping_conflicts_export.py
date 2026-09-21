"""Конфликты сопоставления: выгрузка берёт весь отбор, счётчик говорит правду.

Конфликт заводится в момент, когда по баркоду пришёл РЕАЛЬНЫЙ заказ и разнести
его не удалось. То есть каждая строка — уже случившаяся непроведённая продажа, а
пока баркод не сопоставлен, каждый следующий заказ повторит то же самое: продажи
идут мимо нас, остаток завышен ровно на них.

Выгрузка для того и нужна, чтобы разобрать их пачкой. Отдавать в файле первые
триста и молчать об этом — значит оставить остальные неразобранными. А страница
вдобавок показывала «300» вместо «300 из 2412»: число читалось как «столько их и
есть», и искать остальные было незачем.
"""
from app.models import MappingConflict, Platform
from app.routers.mapping import EXPORT_LIMIT, PAGE_LIMIT, _count_conflicts, _query_conflicts
from tests.factories import make_account


def _conflicts(db, count, barcode_prefix="b"):
    account = make_account(db, Platform.wb)
    for i in range(count):
        db.add(MappingConflict(barcode=f"{barcode_prefix}{i:05d}",
                               account_id=account.id, attempts=1))
    db.commit()
    return account


def test_the_page_still_shows_a_page(db):
    """Потолок страницы не трогаем: строка весит немало, а разбирают конфликты
    не глазами по списку, а файлом."""
    _conflicts(db, PAGE_LIMIT + 25)

    assert len(_query_conflicts(db, "", "")) == PAGE_LIMIT


def test_the_export_takes_the_whole_selection(db):
    _conflicts(db, PAGE_LIMIT + 25)

    rows = _query_conflicts(db, "", "", limit=EXPORT_LIMIT)

    assert len(rows) == PAGE_LIMIT + 25


def test_the_counter_tells_the_truth(db):
    """Раньше здесь стояло `len(rows)`, то есть всегда не больше трёхсот."""
    _conflicts(db, PAGE_LIMIT + 25)

    assert _count_conflicts(db, "", "") == PAGE_LIMIT + 25


def test_the_counter_respects_the_search(db):
    """Счётчик обязан считать ТО ЖЕ, что показывает страница: разойдись они, «300
    из 2412» врало бы ровно там, где оператор сузил отбор."""
    _conflicts(db, 5, barcode_prefix="aaa")
    _conflicts(db, 3, barcode_prefix="bbb")

    assert _count_conflicts(db, "aaa", "") == 5
    assert len(_query_conflicts(db, "aaa", "")) == 5


def test_the_counter_respects_the_account_filter(db):
    first = _conflicts(db, 4, barcode_prefix="x")
    second = make_account(db, Platform.ozon, name="Второй")
    db.add(MappingConflict(barcode="y0", account_id=second.id, attempts=1))
    db.commit()

    assert _count_conflicts(db, "", str(first.id)) == 4
    assert _count_conflicts(db, "", str(second.id)) == 1


# ------------------------------------------------ через страницу, а не мимо неё

def _seed_web(web_db, count):
    from app.models import MappingConflict, PlatformAccount

    account = PlatformAccount(platform="wb", name="Кабинет", warehouse_id="wh")
    web_db.add(account)
    web_db.commit()
    web_db.refresh(account)
    for i in range(count):
        web_db.add(MappingConflict(barcode=f"b{i:05d}", account_id=account.id, attempts=1))
    web_db.commit()
    return account


def test_the_exported_file_holds_the_whole_selection(logged_in_client, web_db):
    """Проверка через HTTP обязательна: дефект был не в запросах, а в том, КАК
    их зовёт страница. Тесты на сами функции его не ловили — проверено откатом."""
    import io as _io

    from openpyxl import load_workbook

    _seed_web(web_db, PAGE_LIMIT + 17)

    r = logged_in_client.get("/mapping/export?view=conflicts")

    wb = load_workbook(_io.BytesIO(r.content), read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    assert len(rows) - 1 == PAGE_LIMIT + 17, "в файле должен быть весь отбор"


def test_the_page_says_the_list_is_cut(logged_in_client, web_db):
    """Без этой строки «300» читается как «столько их и есть», и остальные
    непроведённые продажи никто не ищет."""
    _seed_web(web_db, PAGE_LIMIT + 17)

    r = logged_in_client.get("/mapping?view=conflicts")

    assert f"из {PAGE_LIMIT + 17}" in r.text
