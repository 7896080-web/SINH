"""Сопоставление по артикулам (app/article_matching.py, страница /article-matching).

Закреплено главное: правила выводятся из уже сопоставленных по баркоду пар;
кандидат однозначен только тогда, когда правила дают ровно один товар 1С с
учётом размера; связь создаётся лишь подтверждением и никогда не перепривязывает
уже сопоставленный баркод.
"""
from app.article_matching import analyze, candidates, confirm, get_rule, normalize
from app.models import (ArticleMatchRule, AuditLog, Barcode, MappingConflict, Platform, PlatformAccount,
                        PlatformCatalogItem, Product)
from app.workers.platform_clients.wb import _parse_wb_cards
from tests.factories import make_account


def _p(db, uid, article, size=None, color=None, barcode=None):
    db.add(Product(uid_1c=uid, article=article, name=f"Товар {uid}", size=size, color=color))
    if barcode:
        db.add(Barcode(barcode=barcode, uid_1c=uid))


def _item(db, account, barcode, article, size=None, ext=None):
    db.add(PlatformCatalogItem(account_id=account.id, external_id=ext or f"x-{barcode}", barcode=barcode,
                               article=article, size=size, name="на площадке"))


# ------------------------------------------------------------- нормализация

def test_normalize_ignores_case_separators_and_cyrillic_lookalikes():
    assert normalize("jn-100 / 48") == normalize("JN100_48") == "JN10048"
    assert normalize("А100") == normalize("A100")           # кириллическая А
    assert normalize("Ёж-1") == normalize("еж 1") == "EЖ1"   # Ё→Е, затем Е → латинская E
    assert normalize("НС") == "HC"
    assert normalize(None) == ""


# ------------------------------------------------------------- разбор

def test_analysis_learns_rules_from_barcode_pairs(db):
    a = make_account(db, Platform.ozon)
    _p(db, "u1", "JN-100", "48", "синий", "b1")
    _p(db, "u2", "JN-100", "50", "синий", "b2")
    _p(db, "u3", "TS-7", "M", "белый", "b3")
    _p(db, "u4", "KS-1", None, None, "b4")
    _item(db, a, "b1", "jn100_48")         # артикул + размер
    _item(db, a, "b2", "JN-100-50")        # артикул + размер
    _item(db, a, "b3", "OZ-TS7-M")         # приставка + артикул + размер
    _item(db, a, "b4", "KS1")              # совпадает
    _item(db, a, "b9", "ZZ-9")             # без связи
    db.commit()

    an = analyze(db, a.id)

    assert an.total == 4 and an.unmatched == 1
    assert an.counts["size"] == 2 and an.counts["exact"] == 1 and an.counts["affix"] == 1
    assert an.prefixes["OZ"] == 1
    assert "size" in an.suggested_kinds and "exact" in an.suggested_kinds
    assert an.suggested_prefix == "OZ"
    assert an.examples["size"][0]["platform"] == "jn100_48"


def test_analysis_counts_sku_once_for_alternative_barcodes(db):
    a = make_account(db)
    _p(db, "u1", "A1", None, None, "b1")
    db.add(Barcode(barcode="b1alt", uid_1c="u1"))
    _item(db, a, "b1", "A1", ext="nm:1")
    _item(db, a, "b1alt", "A1", ext="nm:1")
    db.commit()
    assert analyze(db, a.id).total == 1


# ------------------------------------------------------------- кандидаты

def test_size_suffix_rule_finds_unique(db):
    a = make_account(db, Platform.ozon)
    _p(db, "u1", "JN-100", "48")
    _p(db, "u2", "JN-100", "50")
    _item(db, a, "n1", "JN100-50")
    db.commit()
    get_rule(db, a.id).kinds = "size"

    [c] = candidates(db, a.id)
    assert c.status == "unique" and c.products[0].uid_1c == "u2" and c.kinds == ["size"]


def test_wb_size_disambiguates_card_article(db):
    a = make_account(db, Platform.wb)
    _p(db, "u1", "JN-100", "48")
    _p(db, "u2", "JN-100", "50")
    _item(db, a, "n1", "jn-100", size="50")
    _item(db, a, "n2", "jn-100")                # размер неизвестен — неоднозначно
    _item(db, a, "n3", "jn-100", size="52")     # такого размера в 1С нет
    db.commit()
    get_rule(db, a.id).kinds = "exact"

    by_bc = {c.item.barcode: c for c in candidates(db, a.id)}
    assert by_bc["n1"].status == "unique" and by_bc["n1"].products[0].uid_1c == "u2"
    assert by_bc["n2"].status == "ambiguous" and len(by_bc["n2"].products) == 2
    assert by_bc["n3"].status == "size_mismatch" and by_bc["n3"].products == []


def test_prefix_is_stripped(db):
    a = make_account(db)
    _p(db, "u1", "TS-7")
    _item(db, a, "n1", "WB-TS-7")
    db.commit()
    rule = get_rule(db, a.id)
    rule.kinds = "exact"
    assert candidates(db, a.id)[0].status == "none"
    rule.strip_prefix = "wb-"
    assert candidates(db, a.id)[0].status == "unique"


def test_mapped_barcodes_are_not_candidates(db):
    a = make_account(db)
    _p(db, "u1", "A1", barcode="b1")
    _item(db, a, "b1", "A1")
    db.commit()
    assert candidates(db, a.id) == []


def test_disabled_kinds_give_nothing(db):
    a = make_account(db)
    _p(db, "u1", "A1")
    _item(db, a, "n1", "A1")
    db.commit()
    get_rule(db, a.id).kinds = ""
    assert candidates(db, a.id)[0].status == "none"


# ------------------------------------------------------------- подтверждение

def test_confirm_creates_barcode_and_closes_conflict(db):
    a = make_account(db)
    _p(db, "u1", "A1")
    _item(db, a, "n1", "A1")
    db.add(MappingConflict(barcode="n1", account_id=a.id))
    db.commit()
    get_rule(db, a.id).kinds = "exact"

    created, refused = confirm(db, a.id, [("n1", "u1"), ("n1", "u1")])
    db.commit()

    assert created == 1 and sum(refused.values()) == 1
    bc = db.query(Barcode).filter(Barcode.barcode == "n1").one()
    assert bc.uid_1c == "u1" and bc.source_platform == "article_match"
    assert db.query(MappingConflict).count() == 0


def test_confirm_refuses_wrong_or_ambiguous_uid(db):
    a = make_account(db)
    _p(db, "u1", "A1", "48")
    _p(db, "u2", "A1", "50")
    _item(db, a, "n1", "A1")                    # без размера — неоднозначно
    db.commit()
    get_rule(db, a.id).kinds = "exact"

    created, refused = confirm(db, a.id, [("n1", "u1"), ("zz", "u1")])
    assert created == 0 and sum(refused.values()) == 2
    assert db.query(Barcode).count() == 0


# ------------------------------------------------------------- WB размер

def test_wb_catalog_parses_tech_size():
    data = {"cards": [{"nmID": 5, "vendorCode": "JN-100", "title": "Джинсы",
                       "sizes": [{"chrtID": 1, "techSize": "48", "skus": ["b1"]},
                                 {"chrtID": 2, "wbSize": "50", "skus": ["b2"]}]}]}
    items = _parse_wb_cards(data)
    assert [(i.barcode, i.size) for i in items] == [("b1", "48"), ("b2", "50")]


# ------------------------------------------------------------- страница

def _web_setup(web_db):
    a = PlatformAccount(platform=Platform.ozon, name="Озон", warehouse_id="w")
    web_db.add(a)
    web_db.commit()
    web_db.refresh(a)
    web_db.add(Product(uid_1c="u1", article="JN-100", size="48", name="Джинсы"))
    web_db.add(Product(uid_1c="u2", article="JN-100", size="50", name="Джинсы"))
    web_db.add(Barcode(barcode="b1", uid_1c="u1"))
    web_db.add(PlatformCatalogItem(account_id=a.id, external_id="1", barcode="b1", article="JN100-48"))
    web_db.add(PlatformCatalogItem(account_id=a.id, external_id="2", barcode="n2", article="JN100-50"))
    web_db.commit()
    return a


def test_page_tabs_render(logged_in_client, web_db):
    a = _web_setup(web_db)
    r = logged_in_client.get("/article-matching")
    assert r.status_code == 200 and "артикул + размер" in r.text and "предлагается" in r.text
    r = logged_in_client.get(f"/article-matching?view=candidates&account_id={a.id}")
    assert r.status_code == 200 and "JN100-50" in r.text
    assert logged_in_client.get(f"/article-matching/rows?account_id={a.id}&status=unique").status_code == 200


def test_page_save_rule_and_confirm(logged_in_client, web_db):
    a = _web_setup(web_db)
    logged_in_client.post(f"/article-matching/rules/{a.id}", data={"kinds": ["size"], "strip_prefix": " OZ- "})
    rule = web_db.query(ArticleMatchRule).one()
    assert rule.kinds == "size" and rule.strip_prefix == "OZ-"

    r = logged_in_client.post("/article-matching/confirm", data={"account_id": str(a.id), "pick": ["n2|u2"]})
    assert "Сопоставлено баркодов: 1" in r.text
    assert web_db.query(Barcode).filter(Barcode.barcode == "n2").one().uid_1c == "u2"
    assert web_db.query(AuditLog).filter(AuditLog.action == "article_match_confirmed").count() == 1


def test_export_is_mapping_import_compatible(logged_in_client, web_db):
    import io
    from openpyxl import load_workbook
    a = _web_setup(web_db)
    logged_in_client.post(f"/article-matching/rules/{a.id}", data={"kinds": ["size"]})
    r = logged_in_client.get(f"/article-matching/export?account_id={a.id}")
    rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    assert rows[0][:2] == ("ID_1С", "Баркод")
    assert rows[1][:2] == ("u2", "n2")
