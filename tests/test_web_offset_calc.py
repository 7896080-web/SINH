"""Страница «Товары и остатки»: расчёт порога от даты.

Порог больше не вводится руками — оператор задаёт дату, видит остаток ЦС на неё
из выгрузки 1С, вписывает факт, и порог считается сам. Здесь проверяется, что
страница действительно так себя ведёт: и по одной строке, и массово.
"""
from datetime import date, timedelta

from app.models import (Platform, PlatformAccount, Product, StockDateRow, StockDateSnapshot,
                        StockDateStatus)
from app.timeutils import now_utc, today_local

DAY = date(2026, 8, 7)


def _product(web_db, uid="u1", stock=11, reserve=0) -> Product:
    p = Product(uid_1c=uid, article="A-1", name="Товар", stock_on_hand=stock,
                reserve=reserve, broadcast_enabled=True)
    web_db.add(p)
    web_db.commit()
    return p


def _account(web_db, name="WB-1") -> PlatformAccount:
    """Кабинет нужен не сам по себе: без единого активного кабинета страница
    показывает своё сообщение «добавьте кабинет» и перекрывает им результат
    импорта. На боевом кабинеты есть всегда."""
    a = PlatformAccount(platform=Platform.wb, name=name, warehouse_id="wh")
    web_db.add(a)
    web_db.commit()
    web_db.refresh(a)
    return a


def _snapshot(web_db, rows, day=DAY):
    snap = StockDateSnapshot(snapshot_date=day, status=StockDateStatus.done, rows_count=len(rows))
    web_db.add(snap)
    web_db.commit()
    web_db.refresh(snap)
    for uid, qty in rows:
        web_db.add(StockDateRow(snapshot_id=snap.id, uid_1c=uid, quantity=qty))
    web_db.commit()
    return snap


# ------------------------------------------------ по одной строке

def test_setting_a_date_pulls_the_stock_and_computes(logged_in_client, web_db):
    _snapshot(web_db, [("u1", 10)])
    product = _product(web_db, reserve=2)

    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})
    logged_in_client.post("/products/u1/fact", data={"value": "8"})

    web_db.refresh(product)
    assert product.offset_base_stock == 10
    assert product.broadcast_offset == 4


def test_the_row_shows_the_arithmetic(logged_in_client, web_db):
    """Оператор должен видеть не только число, но и из чего оно вышло — иначе
    порог выглядит как магия и ему нечем верить.

    Слагаемыми, а не формулой от даты: только так видно, что изменит правка
    расхождения и что — правка брони. 23.09 строка показывала «43 − (32 − бронь
    0)», и это читалось как формула, а не как утверждение о складе: оператор
    «поправил» подставленный факт на учётный и схлопнул порог."""
    _snapshot(web_db, [("u1", 10)])
    _product(web_db, reserve=2)
    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})

    r = logged_in_client.post("/products/u1/fact", data={"value": "8"})

    assert "расхождение" in r.text
    assert "бронь 2" in r.text
    assert "порог" in r.text

def test_waiting_for_1c_is_visible_in_the_row(logged_in_client, web_db):
    """Снимка на дату ещё нет. Строка обязана объяснить, что происходит, а не
    молча показывать пустой порог."""
    _product(web_db)

    r = logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})

    assert "ждём выгрузку 1С" in r.text


def test_future_date_is_refused(logged_in_client, web_db):
    product = _product(web_db)
    tomorrow = (today_local() + timedelta(days=1)).isoformat()

    r = logged_in_client.post("/products/u1/base-date", data={"value": tomorrow})

    assert "на будущую дату" in r.text
    web_db.refresh(product)
    assert product.offset_base_date is None


def test_changing_the_reserve_moves_the_threshold(logged_in_client, web_db):
    """Ради этого три величины и хранятся."""
    _snapshot(web_db, [("u1", 10)])
    product = _product(web_db, reserve=2)
    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})
    logged_in_client.post("/products/u1/fact", data={"value": "8"})

    logged_in_client.post("/products/u1/reserve", data={"reserve": "5"})

    web_db.refresh(product)
    assert product.broadcast_offset == 7


def test_clearing_the_fact_does_not_erase_the_discrepancy(logged_in_client, web_db):
    """Снять факт — не то же самое, что сказать «склад сошёлся с учётом».

    Факт был способом ИЗМЕРИТЬ расхождение; измеренное хранится на товаре и
    смену даты переживает, значит и очистку поля переживает тоже. Обнулить
    расхождение можно только вслух — поставив в его собственном поле ноль."""
    _snapshot(web_db, [("u1", 10)])
    product = _product(web_db, reserve=2)
    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})
    logged_in_client.post("/products/u1/fact", data={"value": "8"})

    logged_in_client.post("/products/u1/fact", data={"value": ""})

    web_db.refresh(product)
    assert product.fact_at_date is None
    assert product.stock_discrepancy == 2
    assert product.broadcast_offset == 4

    logged_in_client.post("/products/u1/discrepancy", data={"value": "0"})
    web_db.refresh(product)
    assert product.broadcast_offset == 2, "сказано вслух — порог свёлся к брони"

def test_clearing_the_threshold_removes_the_whole_calculation(logged_in_client, web_db):
    """«Сбросить порог» должно действительно сбрасывать: если оставить дату,
    первая же правка брони вернёт порог, и оператор решит, что кнопка не работает."""
    _snapshot(web_db, [("u1", 10)])
    product = _product(web_db, reserve=2)
    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})

    logged_in_client.post("/products/u1/offset", data={"clear": "1"})
    logged_in_client.post("/products/u1/reserve", data={"reserve": "4"})

    web_db.refresh(product)
    assert product.offset_base_date is None
    assert product.broadcast_offset is None


# ------------------------------------------------ массовые действия

def test_bulk_date_applies_to_the_selected_rows(logged_in_client, web_db):
    _snapshot(web_db, [("u1", 10), ("u2", 4)])
    a, b = _product(web_db, uid="u1", reserve=2), _product(web_db, uid="u2", reserve=1)

    logged_in_client.post("/products/bulk", data={
        "action": "set_base_date", "uids": ["u1", "u2"], "date_value": "2026-08-07"})

    web_db.refresh(a), web_db.refresh(b)
    assert (a.offset_base_stock, a.broadcast_offset) == (10, 2)
    assert (b.offset_base_stock, b.broadcast_offset) == (4, 1)


def test_bulk_fact_from_stock_does_not_overwrite_typed_values(logged_in_client, web_db):
    """Кнопка «Факт = остаток ЦС» заполняет только пустые: затирать введённое
    руками массовым действием нельзя."""
    _snapshot(web_db, [("u1", 10), ("u2", 4)])
    typed, empty = _product(web_db, uid="u1"), _product(web_db, uid="u2")
    logged_in_client.post("/products/bulk", data={
        "action": "set_base_date", "uids": ["u1", "u2"], "date_value": "2026-08-07"})
    logged_in_client.post("/products/u1/fact", data={"value": "3"})

    logged_in_client.post("/products/bulk", data={
        "action": "fact_from_stock", "uids": ["u1", "u2"]})

    web_db.refresh(typed), web_db.refresh(empty)
    assert typed.fact_at_date == 3              # не тронут
    assert empty.fact_at_date == 4              # подставлен остаток ЦС


def test_bulk_future_date_is_refused(logged_in_client, web_db):
    product = _product(web_db)
    tomorrow = (today_local() + timedelta(days=1)).isoformat()

    logged_in_client.post("/products/bulk", data={
        "action": "set_base_date", "uids": ["u1"], "date_value": tomorrow})

    web_db.refresh(product)
    assert product.offset_base_date is None


# ------------------------------------------------ фильтр незавершённых

def test_unfinished_filter_finds_rows_waiting_for_1c(logged_in_client, web_db):
    """На каталоге в 152 тысячи SKU недоделанные строки иначе не найти."""
    _product(web_db, uid="u1")
    _product(web_db, uid="u2")
    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})

    r = logged_in_client.get("/products/rows?only_unfinished=true")

    assert "row-u1" in r.text
    assert "row-u2" not in r.text


def test_unfinished_filter_finds_rows_without_a_fact(logged_in_client, web_db):
    _snapshot(web_db, [("u1", 10), ("u2", 4)])
    _product(web_db, uid="u1")
    _product(web_db, uid="u2")
    logged_in_client.post("/products/bulk", data={
        "action": "set_base_date", "uids": ["u1", "u2"], "date_value": "2026-08-07"})
    logged_in_client.post("/products/u1/fact", data={"value": "8"})

    r = logged_in_client.get("/products/rows?only_unfinished=true")

    assert "row-u2" in r.text
    assert "row-u1" not in r.text


def test_untouched_rows_are_not_called_unfinished(logged_in_client, web_db):
    """Товар, которому дату не задавали вовсе, в незавершённые не попадает —
    иначе фильтр вернул бы весь каталог и был бы бесполезен."""
    _product(web_db, uid="u1")

    r = logged_in_client.get("/products/rows?only_unfinished=true")

    assert "row-u1" not in r.text


# ------------------------------------------------ Excel

def _xlsx(headers, rows) -> bytes:
    import io
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client, payload):
    return client.post("/products/import", files={"file": (
        "p.xlsx", payload,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def _export_row(client):
    import io
    from openpyxl import load_workbook

    r = client.get("/products/export")
    ws = load_workbook(io.BytesIO(r.content)).active
    headers = [c.value for c in ws[1]]
    return headers, dict(zip(headers, [c.value for c in ws[2]]))


def test_export_carries_the_whole_calculation(logged_in_client, web_db):
    """В файле должно быть видно всё, из чего вышел порог, — иначе оператор не
    сможет ни проверить цифру, ни осмысленно её поправить."""
    _snapshot(web_db, [("u1", 10)])
    _product(web_db, reserve=2)
    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})
    logged_in_client.post("/products/u1/fact", data={"value": "8"})

    headers, values = _export_row(logged_in_client)

    assert "Дата расчёта" in headers and "Остаток ЦС на дату" in headers
    assert "Факт на дату" in headers
    assert values["Дата расчёта"] == "2026-08-07"
    assert values["Остаток ЦС на дату"] == 10
    assert values["Факт на дату"] == 8
    assert values["Порог трансляции"] == 4


def test_import_sets_the_date_and_the_fact(logged_in_client, web_db):
    """Главный сценарий массового ввода: выгрузили, проставили факт тысяче строк
    в Excel, залили обратно."""
    _snapshot(web_db, [("u1", 10)])
    product = _product(web_db, reserve=2)

    _upload(logged_in_client, _xlsx(
        ["ID_1С", "Дата расчёта", "Резерв", "Факт на дату"],
        [["u1", "2026-08-07", 2, 8]]))

    web_db.expire_all()
    product = web_db.query(Product).first()
    assert product.offset_base_date == DAY
    assert product.offset_base_stock == 10
    assert product.broadcast_offset == 4


def test_import_reads_a_date_cell_as_a_date(logged_in_client, web_db):
    """Excel сам превращает «2026-08-07» в дату при правке — файл вернётся с
    датой в ячейке, а не со строкой, и это не должно ломать импорт."""
    from datetime import datetime

    _snapshot(web_db, [("u1", 10)])
    _product(web_db, reserve=0)

    _upload(logged_in_client, _xlsx(
        ["ID_1С", "Дата расчёта", "Факт на дату"],
        [["u1", datetime(2026, 8, 7), 6]]))

    web_db.expire_all()
    assert web_db.query(Product).first().offset_base_date == DAY


def test_import_sets_the_threshold_and_keeps_the_basis(logged_in_client, web_db):
    """Порог задаётся файлом и у строки С ДАТОЙ — но связка остаётся согласованной.

    Раньше это был отказ, и причина была верной: записанный напрямую порог
    держался бы до первой правки брони, а потом формула молча вернула бы
    прежний. Теперь файл пишет РАСХОЖДЕНИЕ (порог − бронь), и круг «выгрузил →
    поправил → залил» по этой колонке замкнулся, не заводя второго источника
    правды."""
    _account(web_db)
    _snapshot(web_db, [("u1", 10)])
    product = _product(web_db, reserve=2)
    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})
    logged_in_client.post("/products/u1/fact", data={"value": "8"})
    web_db.expire_all()
    assert web_db.query(Product).first().broadcast_offset == 4

    _upload(logged_in_client, _xlsx(["ID_1С", "Порог трансляции"], [["u1", 6]]))

    web_db.expire_all()
    product = web_db.query(Product).first()
    assert product.broadcast_offset == 6
    assert product.stock_discrepancy == 4        # 6 − бронь 2
    # Факт — измерение склада, и правка порога его не переписывает: подставленное
    # под порог число оператор принял бы за собственное измерение.
    assert product.fact_at_date == 8

def test_a_threshold_above_the_1c_number_is_allowed(logged_in_client, web_db):
    """Порог больше учётного остатка — законное состояние, и отказывать в нём нельзя.

    Раньше здесь стоял отказ, и был он следствием механики: под порог
    подбирался ФАКТ, а склад в минусе не бывает. Подбирать больше нечего —
    расхождение хранится само, — и запрет вместе с механикой ушёл. На бою такие
    строки есть: учёт врёт сильнее, чем в 1С вообще числится товара, и наружу по
    ним не уходит ничего, что и требуется."""
    _account(web_db)
    _snapshot(web_db, [("u1", 10)])
    _product(web_db, reserve=2)
    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})
    logged_in_client.post("/products/u1/fact", data={"value": "8"})

    _upload(logged_in_client, _xlsx(["ID_1С", "Порог трансляции"], [["u1", 99]]))

    web_db.expire_all()
    product = web_db.query(Product).first()
    assert product.broadcast_offset == 99
    assert product.stock_discrepancy == 97

def test_a_plain_round_trip_does_not_complain(logged_in_client, web_db):
    """Обычный случай: выгрузили и залили обратно, ничего не меняя. Порог в файле
    совпадает с расчётным — ругаться не на что, иначе каждый круг давал бы
    ложную ошибку и оператор перестал бы читать сообщения."""
    _account(web_db)
    _snapshot(web_db, [("u1", 10)])
    _product(web_db, reserve=2)
    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})
    logged_in_client.post("/products/u1/fact", data={"value": "8"})

    r = _upload(logged_in_client, _xlsx(
        ["ID_1С", "Дата расчёта", "Остаток ЦС на дату", "Резерв", "Факт на дату",
         "Порог трансляции"],
        [["u1", "2026-08-07", 10, 2, 8, 4]]))

    assert "Ошибок" not in r.text
    web_db.expire_all()
    assert web_db.query(Product).first().broadcast_offset == 4


def test_import_still_honours_a_hand_typed_threshold_without_a_date(logged_in_client, web_db):
    """Товары, настроенные до этой правки, продолжают править через Excel
    по-старому: там порог живёт как введённое руками число."""
    _product(web_db)

    _upload(logged_in_client, _xlsx(["ID_1С", "Порог трансляции"], [["u1", 11]]))

    web_db.expire_all()
    assert web_db.query(Product).first().broadcast_offset == 11


def test_import_refuses_a_future_date(logged_in_client, web_db):
    product = _product(web_db)
    tomorrow = (today_local() + timedelta(days=1)).isoformat()

    _upload(logged_in_client, _xlsx(["ID_1С", "Дата расчёта"], [["u1", tomorrow]]))

    web_db.expire_all()
    assert web_db.query(Product).first().offset_base_date is None


def test_import_recomputes_once_after_all_three_values(logged_in_client, web_db):
    """Дата, бронь и факт приезжают одной строкой. Пересчитать порог надо после
    всех трёх — иначе он посчитается по половине данных."""
    _snapshot(web_db, [("u1", 10)])
    _product(web_db, reserve=0)

    _upload(logged_in_client, _xlsx(
        ["ID_1С", "Дата расчёта", "Резерв", "Факт на дату"],
        [["u1", "2026-08-07", 5, 8]]))

    web_db.expire_all()
    product = web_db.query(Product).first()
    assert product.reserve == 5
    assert product.broadcast_offset == 7        # 10 − (8 − 5), а не 10 − 8


def test_moving_the_date_clears_the_fact_on_the_page(logged_in_client, web_db):
    """Со стороны страницы то же правило: факт всегда «на дату». Перенесли дату —
    старое пересчитанное количество к новому числу отношения не имеет.

    А РАСХОЖДЕНИЕ переносится: оно свойство товара, и порог держится на нём.
    Новый пересчёт склада его перебьёт, а пока его нет — держим последнее
    измеренное: порог выше, наружу уходит меньше."""
    _snapshot(web_db, [("u1", 10)])
    _snapshot(web_db, [("u1", 12)], day=date(2026, 9, 1))
    product = _product(web_db, reserve=2)
    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})
    logged_in_client.post("/products/u1/fact", data={"value": "8"})

    logged_in_client.post("/products/u1/base-date", data={"value": "2026-09-01"})

    web_db.refresh(product)
    assert product.fact_at_date is None
    assert product.offset_base_stock == 12
    assert product.stock_discrepancy == 2
    assert product.broadcast_offset == 4

def test_bulk_date_reads_the_snapshot_once(logged_in_client, web_db):
    """Массовая простановка даты не должна ходить в базу за каждым товаром: на
    каталоге в 152 тысячи SKU это минуты ожидания и лишняя нагрузка на боевую
    базу. Считаем запросы — они не должны расти вместе с числом строк."""
    from sqlalchemy import event
    from app.database import engine

    rows = [(f"u{i:03d}", 10 + i) for i in range(20)]
    _snapshot(web_db, rows)
    for uid, _ in rows:
        _product(web_db, uid=uid, reserve=1)

    seen = []
    listener = lambda conn, cur, stmt, *a: seen.append(stmt)
    event.listen(engine, "before_cursor_execute", listener)
    try:
        logged_in_client.post("/products/bulk", data={
            "action": "set_base_date", "uids": [uid for uid, _ in rows],
            "date_value": "2026-08-07"})
    finally:
        event.remove(engine, "before_cursor_execute", listener)

    selects = [q for q in seen if q.lstrip().upper().startswith("SELECT")
               and "stock_date" in q.lower()]
    assert len(selects) <= 4, f"чтений снимка {len(selects)} на 20 товаров — растёт со строками"
    web_db.expire_all()
    p = web_db.query(Product).filter(Product.uid_1c == "u005").first()
    assert p.offset_base_stock == 15 and p.broadcast_offset == 1


# ---------------------------------------- расхождение: круг «выгрузил → залил»

def test_the_export_carries_the_stored_discrepancy(logged_in_client, web_db):
    """В файл едет ХРАНИМОЕ расхождение, а не «учёт минус факт».

    Раньше колонка была справкой и считалась из соседних ячеек. У строки, чью
    дату сдвинули назад, факт пуст — и справка выходила пустой при живом пороге,
    то есть файл показывал «расхождения нет» там, где оно есть.
    """
    _account(web_db)
    _snapshot(web_db, [("u1", 43)])
    _product(web_db, reserve=2)
    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})
    logged_in_client.post("/products/u1/fact", data={"value": "32"})
    logged_in_client.post("/products/u1/base-date", data={"value": "2026-09-01"})

    headers, row = _export_row(logged_in_client)

    assert "Расхождение" in headers
    assert not row["Факт на дату"], "факт всегда «на дату», новую дату он не описывает"
    assert row["Расхождение"] == 11
    assert row["Порог трансляции"] == 13          # 11 + бронь 2


def test_the_discrepancy_column_is_writable(logged_in_client, web_db):
    """Правка колонки применяется — это и есть способ сказать файлом «склад
    сошёлся» (ноль) или назвать новое число."""
    _account(web_db)
    _snapshot(web_db, [("u1", 43)])
    product = _product(web_db, reserve=2)
    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})
    logged_in_client.post("/products/u1/fact", data={"value": "32"})

    _upload(logged_in_client, _xlsx(["ID_1С", "Расхождение"], [["u1", 6]]))

    web_db.expire_all()
    product = web_db.query(Product).first()
    assert product.stock_discrepancy == 6
    assert product.broadcast_offset == 8          # 6 + бронь 2


def test_a_negative_discrepancy_survives_the_file(logged_in_client, web_db):
    """Минус — «на складе больше, чем знает 1С». На бою таких 53 товара, до −213.

    `max(0, ...)` здесь был бы дефектом: порог стал бы нулём, и наружу уехало бы
    ровно учётное число вместо того, что физически лежит."""
    _account(web_db)
    _snapshot(web_db, [("u1", 43)])
    _product(web_db)

    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})
    _upload(logged_in_client, _xlsx(["ID_1С", "Расхождение"], [["u1", -7]]))

    web_db.expire_all()
    product = web_db.query(Product).first()
    assert product.stock_discrepancy == -7
    assert product.broadcast_offset == -7


def test_an_untouched_file_changes_nothing(logged_in_client, web_db):
    """Круг «выгрузил → поправил ОДНУ колонку → залил» не должен трогать остального.

    Колонок, двигающих порог, теперь две — «Расхождение» и «Порог трансляции», —
    и обе выгрузка заполняет. Без снимка «как было в файле» вторая откатывала бы
    правку первой: оператор поправил расхождение, а нетронутая ячейка порога
    вернула бы старое число. Молча и по всему файлу.
    """
    _account(web_db)
    _snapshot(web_db, [("u1", 43)])
    _product(web_db, reserve=2)
    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})
    logged_in_client.post("/products/u1/fact", data={"value": "32"})

    _, row = _export_row(logged_in_client)
    # Правим ОДНУ ячейку, остальные возвращаем как есть — так и выглядит файл,
    # пришедший от оператора.
    _upload(logged_in_client, _xlsx(
        ["ID_1С", "Расхождение", "Порог трансляции", "Факт на дату"],
        [["u1", 6, row["Порог трансляции"], row["Факт на дату"]]]))

    web_db.expire_all()
    product = web_db.query(Product).first()
    assert product.stock_discrepancy == 6, "правку расхождения не откатила колонка порога"
    assert product.broadcast_offset == 8


def test_an_empty_discrepancy_cell_changes_nothing(logged_in_client, web_db):
    """Пустая ячейка — «не трогать», как и во всех остальных колонках файла.

    Иначе файл, собранный не из нашей выгрузки, снял бы измерения по всему
    каталогу разом и молча вернул на площадки полный остаток."""
    _account(web_db)
    _snapshot(web_db, [("u1", 43)])
    _product(web_db)
    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})
    logged_in_client.post("/products/u1/fact", data={"value": "32"})

    _upload(logged_in_client, _xlsx(["ID_1С", "Расхождение"], [["u1", ""]]))

    web_db.expire_all()
    assert web_db.query(Product).first().stock_discrepancy == 11


def test_the_row_edits_the_discrepancy_directly(logged_in_client, web_db):
    """То же и построчно: у поля на странице тот же смысл, что у колонки в файле.

    Массовый путь и построчный обязаны делать одно и то же — на расхождении
    этого правила больше всего и держится: разойдись они, оператор получил бы
    разный порог в зависимости от того, каким путём вводил число."""
    _snapshot(web_db, [("u1", 43)])
    product = _product(web_db, reserve=2)
    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})
    logged_in_client.post("/products/u1/fact", data={"value": "32"})

    logged_in_client.post("/products/u1/discrepancy", data={"value": "-7"})

    web_db.refresh(product)
    assert product.stock_discrepancy == -7
    assert product.broadcast_offset == -5         # −7 + бронь 2

    logged_in_client.post("/products/u1/discrepancy", data={"value": ""})
    web_db.refresh(product)
    assert product.stock_discrepancy is None, "пусто в поле строки — «не измеряли»"


def test_an_untouched_discrepancy_cell_does_not_undo_a_new_fact(logged_in_client, web_db):
    """Зеркало предыдущего: оператор поправил ФАКТ, а ячейку расхождения не трогал.

    Выгрузка заполняет обе, и обе двигают порог. Факт с новым числом записывает
    новое расхождение — а следом нетронутая ячейка расхождения вернула бы
    прежнее, то есть отменила бы физический пересчёт склада. Молча и по всему
    файлу. Поэтому каждая из двух колонок спрашивает про СЕБЯ: «человек изменил
    эту ячейку?», сравнивая со снимком до правок строки."""
    _account(web_db)
    _snapshot(web_db, [("u1", 43)])
    _product(web_db)
    logged_in_client.post("/products/u1/base-date", data={"value": "2026-08-07"})
    logged_in_client.post("/products/u1/fact", data={"value": "32"})

    _, row = _export_row(logged_in_client)
    assert row["Расхождение"] == 11

    _upload(logged_in_client, _xlsx(
        ["ID_1С", "Факт на дату", "Расхождение"],
        [["u1", 30, row["Расхождение"]]]))

    web_db.expire_all()
    product = web_db.query(Product).first()
    assert product.fact_at_date == 30
    assert product.stock_discrepancy == 13, "новый пересчёт склада не отменён"
    assert product.broadcast_offset == 13
