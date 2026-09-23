"""Находки аудита 22.09, блок «веб и оператор».

Общее у всех четырёх: оператор сделал ровно то, что велит интерфейс, и получил
не то, что ожидал, — причём молча.
"""
import io
from datetime import date, timedelta

from openpyxl import Workbook, load_workbook

from app.models import (DispatchQueueItem, Platform, PlatformAccount, Product,
                        SyncSetting)
from app.timeutils import now_utc
from tests.factories import make_account

DAY = date(2026, 8, 7)


def _account(web_db, name="ИП ЯВОРСКАЯ"):
    a = PlatformAccount(platform=Platform.wb, name=name, warehouse_id="wh")
    web_db.add(a)
    web_db.commit()
    return a


def _product(web_db, **kw):
    kw.setdefault("uid_1c", "u1")
    kw.setdefault("article", "A-1")
    kw.setdefault("name", "Товар")
    kw.setdefault("stock_on_hand", 20)
    kw.setdefault("reserve", 0)
    p = Product(**kw)
    web_db.add(p)
    web_db.commit()
    return p


def _file(headers, *rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _import(client, data):
    return client.post(
        "/products/import",
        files={"file": ("f.xlsx", data,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=True).text


def _cell(resp, column, uid="u1"):
    ws = load_workbook(io.BytesIO(resp.content)).active
    rows = list(ws.iter_rows(values_only=True))
    col = rows[0].index(column)
    for row in rows[1:]:
        if row[0] == uid:
            return row[col]
    raise AssertionError(f"строки {uid} нет в выгрузке")


# ---------------------------------------------------------------------------
# 1. Выгрузка убивала просьбу включить трансляцию
# ---------------------------------------------------------------------------

def _waiting_for_recalc(web_db):
    """Строка, попросившая трансляцию файлом и ждущая расчёта."""
    account = _account(web_db)
    _product(web_db, offset_base_date=DAY, offset_base_stock=20, fact_at_date=20,
             broadcast_requested_at=now_utc())
    web_db.add(SyncSetting(uid_1c="u1", account_id=account.id, enabled=True))
    web_db.commit()
    return account


def test_the_export_says_yes_for_a_row_waiting_for_the_recalc(logged_in_client, web_db):
    """Выгрузка писала «Нет» — по `broadcast_enabled`, — и это правда про
    сегодняшнее состояние, но неправда про решение оператора. Он попросил «Да»,
    гейт согласился и отложил.

    Цена — не в формулировке. Круг «выгрузил → поправил другую колонку → залил»
    возвращал это «Нет», импорт читал его как осознанное выключение и снимал
    просьбу: ту самую, про которую минуту назад пообещал «ждут расчёта и
    включатся сами». Трансляция не включалась никогда, товар молча не
    продавался, а человек считал, что настроил всё файлом.
    """
    _waiting_for_recalc(web_db)

    assert _cell(logged_in_client.get("/products/export"), "Трансляция") == "Да"


def test_a_round_trip_keeps_the_request(logged_in_client, web_db):
    """Главная проверка: файл, залитый обратно без единой правки в колонке
    «Трансляция», не должен менять ничего."""
    _waiting_for_recalc(web_db)
    resp = logged_in_client.get("/products/export")
    ws = load_workbook(io.BytesIO(resp.content)).active
    rows = list(ws.iter_rows(values_only=True))
    col = rows[0].index("Трансляция")

    _import(logged_in_client, _file(["ID_1С", "Резерв", "Трансляция"],
                                    ["u1", 3, rows[1][col]]))

    web_db.expire_all()
    product = web_db.query(Product).filter(Product.uid_1c == "u1").first()
    assert product.broadcast_requested_at is not None, "просьба снята нашим же эхом"
    assert product.reserve == 3, "правка соседней колонки не применилась"


def test_a_human_no_still_cancels_the_request(logged_in_client, web_db):
    """Правило импорта не ослаблено: «Нет» по-прежнему снимает и просьбу тоже.
    Просто теперь «Нет» в этой ячейке означает правку человека, а не наше эхо."""
    _waiting_for_recalc(web_db)

    _import(logged_in_client, _file(["ID_1С", "Трансляция"], ["u1", "Нет"]))

    web_db.expire_all()
    assert web_db.query(Product).filter(
        Product.uid_1c == "u1").first().broadcast_requested_at is None


def test_a_plain_off_row_still_exports_no(logged_in_client, web_db):
    """Обратная сторона: у строки без просьбы в колонке по-прежнему «Нет»."""
    _account(web_db)
    _product(web_db)

    assert _cell(logged_in_client.get("/products/export"), "Трансляция") == "Нет"


# ---------------------------------------------------------------------------
# 2. Ложная ошибка по «Порогу трансляции» при правке факта
# ---------------------------------------------------------------------------

def test_changing_the_fact_does_not_fault_the_untouched_threshold(logged_in_client, web_db):
    """Подсказка на странице велит порог не трогать и менять его через «Факт на
    дату». Оператор так и делает — и получает ошибку на КАЖДОЙ такой строке:
    порог из файла сравнивался с ПЕРЕСЧИТАННЫМ по только что применённому факту,
    а не с тем, что стояло в выгрузке.

    Ошибка при этом советует сделать ровно то, что человек и сделал. А место в
    сообщении не бесконечно: показываются первые пять, и всё режется по 440
    символам — настоящие ошибки вытесняются.
    """
    _account(web_db)
    _product(web_db, offset_base_date=DAY, offset_base_stock=10, reserve=2,
             fact_at_date=8, broadcast_offset=4)

    page = _import(logged_in_client, _file(
        ["ID_1С", "Дата расчёта", "Резерв", "Факт на дату", "Порог трансляции"],
        ["u1", DAY.isoformat(), 2, 6, 4]))

    assert "порог считается из даты и факта" not in page
    web_db.expire_all()
    assert web_db.query(Product).filter(Product.uid_1c == "u1").first().fact_at_date == 6


def test_editing_the_threshold_now_applies_and_keeps_the_basis(logged_in_client, web_db):
    """Обратная сторона: изменённая колонка порога ПРИМЕНЯЕТСЯ.

    Раньше это был отказ, и причина была верной: записанный мимо расхождения
    порог держался бы до первой правки брони, а потом формула молча вернула бы
    прежний — второй источник правды. Теперь файл пишет РАСХОЖДЕНИЕ (порог −
    бронь), связка остаётся согласованной, и отказ стал не нужен.

    Проверяем не запись в колонку, а СВОЙСТВО, ради которого всё делалось: после
    правки порога правка брони двигает порог ровно на изменение брони, а не
    возвращает прежний.
    """
    _account(web_db)
    _product(web_db, offset_base_date=DAY, offset_base_stock=10, reserve=2,
             fact_at_date=8, broadcast_offset=4)

    _import(logged_in_client, _file(["ID_1С", "Порог трансляции"], ["u1", 6]))

    web_db.expire_all()
    product = web_db.query(Product).filter(Product.uid_1c == "u1").first()
    assert product.broadcast_offset == 6
    assert product.stock_discrepancy == 4        # 6 − бронь 2
    # Факт — измерение склада, и правка порога его НЕ переписывает: раньше под
    # порог подбирался факт, и оператор видел в поле число, которого не вводил,
    # — ровно то, из-за чего 23.09 схлопнулись пороги у 62 товаров.
    assert product.fact_at_date == 8

    _import(logged_in_client, _file(["ID_1С", "Резерв"], ["u1", 5]))
    web_db.expire_all()
    product = web_db.query(Product).filter(Product.uid_1c == "u1").first()
    assert product.broadcast_offset == 9, "порог сдвинулся ровно на изменение брони"


# ---------------------------------------------------------------------------
# 3. Включение кабинета не сбрасывало предохранитель
# ---------------------------------------------------------------------------

def test_activating_an_account_resets_the_breaker(logged_in_client, web_db):
    """Кабинет гасится на пятом сбое подряд. Счётчик оставался на пяти, и первая
    же ошибка давала шестой — предохранитель гасил кабинет снова. Пяти попыток,
    ради которых он и сделан, у включённого кабинета не было вовсе.

    А включают его ровно тогда, когда первая попытка чаще всего и не проходит:
    оператор поправил ключи и даёт кабинету новый шанс.
    """
    account = PlatformAccount(platform=Platform.wb, name="WB", warehouse_id="wh",
                              is_active=False, consecutive_failures=5,
                              last_error="401 Unauthorized")
    web_db.add(account)
    web_db.commit()

    logged_in_client.post(f"/api-keys/accounts/{account.id}/activate",
                          follow_redirects=True)

    web_db.expire_all()
    account = web_db.query(PlatformAccount).first()
    assert account.is_active is True
    assert account.consecutive_failures == 0
    assert account.last_error is None


def test_the_reactivated_account_survives_one_failure(db):
    """Смысл сброса именно в этом: после включения кабинет обязан пережить
    ошибку, а не умереть от первой же."""
    from app.workers.circuit_breaker import record_failure

    account = make_account(db, Platform.wb)
    account.consecutive_failures = 0
    db.commit()

    disabled = record_failure(db, account, "таймаут")

    assert disabled is False and account.is_active is True


# ---------------------------------------------------------------------------
# 4. Расчёт порога со страницы «Тестирование» не доезжал до площадки
# ---------------------------------------------------------------------------

def test_the_testing_page_recalc_enqueues_the_new_number(logged_in_client, web_db):
    """Рассылка событийная и сама к порогу не вернётся: следующая отправка
    будет, только когда изменится остаток, а у медленного размера это месяцы.
    Близнец на странице «Товары» ставит запись в очередь, а этот — нет.

    Оператор видел новый порог в строке и считал, что карточка обновлена, — при
    том что на площадке лежало прежнее число.
    """
    account = _account(web_db)
    # Дата ТА ЖЕ, меняется только факт: смена даты снимала бы «актуализирован»,
    # и гейт справедливо не пустил бы наружу ничего — это не тот случай.
    _product(web_db, broadcast_enabled=True, offset_base_date=DAY,
             offset_base_stock=20, fact_at_date=20, recalc_done_at=now_utc())
    web_db.add(SyncSetting(uid_1c="u1", account_id=account.id, enabled=True))
    web_db.commit()
    product = web_db.query(Product).filter(Product.uid_1c == "u1").first()
    product.recalc_account_ids = str(account.id)
    web_db.commit()

    logged_in_client.post("/testing/offset-calc", data={
        "uid_1c": "u1", "account_id": "", "base_date": DAY.isoformat(), "fact": "5"})

    web_db.expire_all()
    queued = web_db.query(DispatchQueueItem).filter(
        DispatchQueueItem.uid_1c == "u1").all()
    assert queued, "новый порог не поставлен в очередь рассылки"


def test_the_testing_page_recalc_respects_the_gate(logged_in_client, web_db):
    """Через `enqueue_full_resend`, а не прямым `db.add`: страницу открывают
    ровно в том состоянии, где лестница даёт ноль, и мимо гейтов такая запись
    обнулила бы живую карточку."""
    account = _account(web_db)
    _product(web_db, broadcast_enabled=False)          # трансляция выключена
    web_db.add(SyncSetting(uid_1c="u1", account_id=account.id, enabled=True))
    web_db.commit()

    logged_in_client.post("/testing/offset-calc", data={
        "uid_1c": "u1", "account_id": "", "base_date": DAY.isoformat(), "fact": "5"})

    web_db.expire_all()
    assert web_db.query(DispatchQueueItem).count() == 0
