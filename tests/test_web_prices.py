"""Страница «Цены» (/prices): правила по кабинетам, расчёт, подтверждение, ручные
цены через Excel. Главное: со страницы ничего не уходит на площадку без явного
подтверждения, а пол минимальной наценки не обходится никаким путём."""
import io
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from app.models import (AuditLog, Barcode, Platform, PlatformAccount, PriceChange, PriceChangeStatus,
                        PriceRule, Product, ProductPrice, SyncSetting)


def _account(web_db, platform=Platform.wb, name="ИП Тест"):
    a = PlatformAccount(platform=platform, name=name, warehouse_id="wh")
    web_db.add(a)
    web_db.commit()
    web_db.refresh(a)
    return a


def _product(web_db, account, uid="u1", cost="500"):
    web_db.add(Product(uid_1c=uid, article="A-" + uid, name="Джинсы " + uid, size="46",
                       cost_price=Decimal(cost) if cost else None))
    web_db.add(Barcode(barcode="bc" + uid, uid_1c=uid))
    web_db.add(SyncSetting(uid_1c=uid, account_id=account.id, enabled=True))
    web_db.commit()


RULE = {"markup_percent": "100", "fixed_add": "0", "round_step": "100", "round_minus": "1",
        "min_margin_percent": "30", "max_change_percent": "20"}


def test_requires_login(client):
    r = client.get("/prices", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/login"


def test_all_tabs_render(logged_in_client, web_db):
    a = _account(web_db)
    _product(web_db, a)
    for view in ("rules", "products", "proposals", "log"):
        r = logged_in_client.get(f"/prices?view={view}")
        assert r.status_code == 200, view
    assert "Джинсы u1" in logged_in_client.get("/prices?view=products").text
    assert logged_in_client.get("/prices/rows?view=products&q=u1").status_code == 200


def test_save_rule_per_account(logged_in_client, web_db):
    wb = _account(web_db, Platform.wb, "WB")
    oz = _account(web_db, Platform.ozon, "Ozon")
    logged_in_client.post(f"/prices/rules/{wb.id}", data=RULE)
    logged_in_client.post(f"/prices/rules/{oz.id}", data={**RULE, "markup_percent": "140,5"})

    rules = {r.account_id: r.markup_percent for r in web_db.query(PriceRule).all()}
    assert rules == {wb.id: Decimal("100"), oz.id: Decimal("140.5")}
    assert web_db.query(AuditLog).filter(AuditLog.action == "price_rule_saved").count() == 2


def test_bad_rule_is_not_saved(logged_in_client, web_db):
    a = _account(web_db)
    r = logged_in_client.post(f"/prices/rules/{a.id}", data={**RULE, "round_minus": "100"})
    assert "не сохранено" in r.text
    r = logged_in_client.post(f"/prices/rules/{a.id}", data={**RULE, "markup_percent": "abc"})
    assert "не сохранено" in r.text
    r = logged_in_client.post(f"/prices/rules/{a.id}", data={**RULE, "min_margin_percent": "200"})
    assert "не сохранено" in r.text
    assert all(rule.markup_percent == 0 for rule in web_db.query(PriceRule).all())


def test_recalculate_approve_flow(logged_in_client, web_db):
    a = _account(web_db)
    _product(web_db, a)
    logged_in_client.post(f"/prices/rules/{a.id}", data=RULE)

    r = logged_in_client.post("/prices/recalculate", data={"account_id": ""})
    assert "предложено 1" in r.text
    change = web_db.query(PriceChange).one()
    assert change.new_price == 1099 and change.status == PriceChangeStatus.proposed

    logged_in_client.post("/prices/approve", data={"ids": [str(change.id)]})
    web_db.refresh(change)
    assert change.status == PriceChangeStatus.approved and change.decided_by == "admin"


def test_floor_cannot_be_approved_from_page(logged_in_client, web_db):
    a = _account(web_db)
    _product(web_db, a)
    logged_in_client.post(f"/prices/rules/{a.id}", data=RULE)
    logged_in_client.post(f"/prices/u1/{a.id}/manual", data={"value": "100"})
    logged_in_client.post("/prices/recalculate", data={"account_id": str(a.id)})
    change = web_db.query(PriceChange).one()
    assert change.status == PriceChangeStatus.blocked and change.block_reason == "floor"

    r = logged_in_client.post("/prices/approve", data={"ids": [str(change.id)], "confirm_large": "true"})
    assert "Не подтверждено" in r.text
    web_db.refresh(change)
    assert change.status == PriceChangeStatus.blocked


def test_large_change_needs_checkbox(logged_in_client, web_db):
    a = _account(web_db)
    _product(web_db, a)
    web_db.add(ProductPrice(uid_1c="u1", account_id=a.id, last_sent_price=500))
    web_db.commit()
    logged_in_client.post(f"/prices/rules/{a.id}", data=RULE)
    logged_in_client.post("/prices/recalculate", data={})
    change = web_db.query(PriceChange).one()
    assert change.block_reason == "max_change"

    logged_in_client.post("/prices/approve", data={"ids": [str(change.id)]})
    web_db.refresh(change)
    assert change.status == PriceChangeStatus.blocked
    logged_in_client.post("/prices/approve", data={"ids": [str(change.id)], "confirm_large": "true"})
    web_db.refresh(change)
    assert change.status == PriceChangeStatus.approved


def test_reject(logged_in_client, web_db):
    a = _account(web_db)
    _product(web_db, a)
    logged_in_client.post(f"/prices/rules/{a.id}", data=RULE)
    logged_in_client.post("/prices/recalculate", data={})
    change = web_db.query(PriceChange).one()
    logged_in_client.post("/prices/reject", data={"ids": [str(change.id)]})
    web_db.refresh(change)
    assert change.status == PriceChangeStatus.rejected


def test_manual_price_validation(logged_in_client, web_db):
    a = _account(web_db)
    _product(web_db, a)
    for bad in ("-5", "0", "12.5", "abc"):
        r = logged_in_client.post(f"/prices/u1/{a.id}/manual", data={"value": bad})
        assert "не принята" in r.text
    assert web_db.query(ProductPrice).count() == 0
    logged_in_client.post(f"/prices/u1/{a.id}/manual", data={"value": "1 500"})
    assert web_db.query(ProductPrice).one().manual_price == 1500
    logged_in_client.post(f"/prices/u1/{a.id}/manual", data={"value": ""})
    web_db.expire_all()
    assert web_db.query(ProductPrice).one().manual_price is None


def test_excel_roundtrip_sets_manual_prices(logged_in_client, web_db):
    a = _account(web_db)
    _product(web_db, a)
    r = logged_in_client.get("/prices/export?view=products")
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    col = headers.index(f"Ручная цена: {a.name} [{a.id}]") + 1
    ws.cell(row=2, column=col, value=1999)
    buf = io.BytesIO()
    wb.save(buf)

    r = logged_in_client.post("/prices/import", files={"file": ("p.xlsx", buf.getvalue())})
    assert "изменено: 1" in r.text
    assert web_db.query(ProductPrice).one().manual_price == 1999
    assert web_db.query(PriceChange).count() == 0          # импорт ничего не отправляет


def test_import_without_manual_columns_is_refused(logged_in_client, web_db):
    _account(web_db)
    wb = Workbook()
    wb.active.append(["ID_1С", "Цена"])
    wb.active.append(["u1", 10])
    buf = io.BytesIO()
    wb.save(buf)
    r = logged_in_client.post("/prices/import", files={"file": ("p.xlsx", buf.getvalue())})
    assert "нет колонок" in r.text


def test_proposals_export(logged_in_client, web_db):
    a = _account(web_db)
    _product(web_db, a)
    logged_in_client.post(f"/prices/rules/{a.id}", data=RULE)
    logged_in_client.post("/prices/recalculate", data={})
    r = logged_in_client.get("/prices/export?view=proposals")
    assert r.status_code == 200
    rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    assert len(rows) == 2 and rows[1][0] == "u1"
