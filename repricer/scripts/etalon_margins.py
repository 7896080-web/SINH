"""Наценки по категориям от эталона WB — ТОЛЬКО ЧТЕНИЕ, ничего не меняет и не отправляет.

Логика (задана владельцем 12.10):
  * эталон по артикулу 1С = средняя по всем кабинетам WB ЦЕНА ПРОДАВЦА СО СКИДКОЙ
    (то, что программа читает как цену продажи WB) минус 31% — оценка цены для покупателя;
  * Kit: цена для покупателя = эталон − 10%;
  * Ламода: цена для покупателя = эталон.

Наценка «маржинальность» m даёт цену для покупателя (цена продавца после его скидки)
= себестоимость, ₽ × m / (1 − комиссия) — так её считает `pricing.decide`. Отсюда для
каждого товара целевая m = цена для покупателя × (1 − комиссия) / себестоимость, ₽,
с комиссией ЭТОГО товара на ЭТОЙ площадке (тариф или правило + надбавка).

Наценка задаётся на КАТЕГОРИЮ площадки, а эталон у артикулов разный — поэтому по
категории берётся МЕДИАНА целевых m её товаров, и рядом видно, насколько товары от
неё расходятся (если сильно — категории мало одной наценки).

    C:\\repricer\\.venv\\Scripts\\python.exe C:\\repricer\\scripts\\etalon_margins.py --xlsx C:\\repricer\\эталон.xlsx
    ... --wb-discount 31 --kit -10 --lamoda 0       # проценты логики (по умолчанию — эти)

Лист «Импорт» — ровно в формате импорта «Цены товаров» по категориям (вид «по
категориям» → Excel → «Импорт из Excel»). Загружать его — решение человека, после
просмотра листа «Категории».
"""
from __future__ import annotations

import argparse
import os
import statistics
import sys
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.stdout.reconfigure(errors="replace")      # консоль Windows — cp1251

import priceapp  # noqa: E402,F401  (.env читается на импорте)
from priceapp import mapping, rates  # noqa: E402
from priceapp.database import SessionLocal  # noqa: E402
from priceapp.models import Account, OnecCost  # noqa: E402
from priceapp.platforms import PLATFORMS  # noqa: E402
from priceapp.pricing import article_map, article_of, base_price, get_rule, item_facts  # noqa: E402

NO_CATEGORY = "(без категории)"
WITHIN = 0.05            # «товар близко к медиане категории» — в пределах ±5%


def _accounts(db, platform):
    return (db.query(Account).filter(Account.is_active.is_(True), Account.platform == platform)
            .order_by(Account.name).all())


def etalon(db, articles, wb_discount: float) -> dict[str, float]:
    """Артикул 1С -> эталонная цена для покупателя: средняя цена продавца со скидкой по
    всем карточкам всех кабинетов WB, минус `wb_discount` %."""
    rule = get_rule(db, "wb")
    vals: dict[str, list[float]] = {}
    for a in _accounts(db, "wb"):
        for item_id, rows in mapping.account_items(db, a.id).items():
            sale = item_facts(rule, rows).sale
            if sale:
                vals.setdefault(article_of(articles, item_id), []).append(float(sale))
    return {art: sum(v) / len(v) * (1 - wb_discount / 100) for art, v in vals.items()}


def targets(db, platform: str, shift: float, et: dict[str, float], articles, costs, usd) -> list[dict]:
    """По каждому товару площадки: целевая маржинальность от эталона и текущая."""
    rule = get_rule(db, platform)
    out = []
    for a in _accounts(db, platform):
        for item_id, rows in mapping.account_items(db, a.id).items():
            art = article_of(articles, item_id)
            facts = item_facts(rule, rows)
            cost = costs.get(item_id)
            base = base_price(cost, usd)
            r = {"account": a.name, "item_id": item_id, "article": art,
                 "category": facts.category or NO_CATEGORY, "commission": facts.commission,
                 "base": base, "current_sale": facts.sale, "etalon": et.get(art), "target_price": None,
                 "target_m": None, "current_m": None, "why": ""}
            if base is None or base <= 0:
                r["why"] = "нет себестоимости или курса"
            elif facts.commission is None:
                r["why"] = "не задана комиссия площадки"
            else:
                keep = 1 - float(facts.commission) / 100
                if facts.sale:
                    r["current_m"] = facts.sale * keep / float(base)
                if r["etalon"] is None:
                    r["why"] = "нет цены на WB — эталона нет"
                else:
                    r["target_price"] = r["etalon"] * (1 + shift / 100)
                    r["target_m"] = r["target_price"] * keep / float(base)
            out.append(r)
    return out


def by_category(rows: list[dict]) -> list[dict]:
    groups: dict[str, list[dict]] = {}
    for r in rows:
        groups.setdefault(r["category"], []).append(r)
    out = []
    for cat, rs in sorted(groups.items()):
        ms = [r["target_m"] for r in rs if r["target_m"] is not None]
        cur = [r["current_m"] for r in rs if r["current_m"] is not None]
        g = {"category": cat, "items": len(rs), "articles": len({r["article"] for r in rs}),
             "with_target": len(ms), "no_target": len(rs) - len(ms), "median": None, "p25": None, "p75": None,
             "close": None, "current_median": statistics.median(cur) if cur else None}
        if ms:
            med = statistics.median(ms)
            q = statistics.quantiles(ms, n=4) if len(ms) > 1 else [ms[0], ms[0], ms[0]]
            g.update(median=med, p25=q[0], p75=q[2],
                     close=sum(1 for m in ms if abs(m / med - 1) <= WITHIN) / len(ms))
        out.append(g)
    return out


def _r(v, nd=3):
    return round(v, nd) if isinstance(v, float) else v


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--wb-discount", type=float, default=31.0, help="эталон = цена продавца WB со скидкой минус N %%")
    ap.add_argument("--kit", type=float, default=-10.0, help="Kit к эталону, %% (минус — дешевле)")
    ap.add_argument("--lamoda", type=float, default=0.0, help="Ламода к эталону, %%")
    ap.add_argument("--xlsx", default="", help="сохранить расчёт в .xlsx (листы «Категории», «Товары», «Импорт»)")
    args = ap.parse_args()
    db = SessionLocal()
    try:
        rate = rates.current(db)
        if rate is None:
            print("Курса доллара нет — себестоимость в рублях не посчитать. Страница «Курс $».")
            return 1
        articles = article_map(db)
        costs = {c.item_id: c.cost_usd for c in db.query(OnecCost)}
        et = etalon(db, articles, args.wb_discount)
        print(f"Курс {rate.usd_rub}. Эталон: средняя цена продавца со скидкой на WB минус {args.wb_discount:g}% — "
              f"посчитан по {len(et)} артикулам.")
        plan = [("kit", args.kit), ("lamoda", args.lamoda)]
        all_items, all_cats = [], []
        for platform, shift in plan:
            if not _accounts(db, platform):
                print(f"\n{PLATFORMS[platform]}: активных кабинетов нет — пропускаю.")
                continue
            items = targets(db, platform, shift, et, articles, costs, rate.usd_rub)
            cats = by_category(items)
            all_items += [dict(r, platform=platform) for r in items]
            all_cats += [dict(g, platform=platform, shift=shift) for g in cats]
            done = sum(g["with_target"] for g in cats)
            print(f"\n=== {PLATFORMS[platform]}: цена для покупателя = эталон {shift:+g}% — товаров {len(items)}, "
                  f"с эталоном {done}, категорий {len(cats)}")
            print(f"{'категория':40} {'товаров':>7} {'наценка':>8} {'разброс (25–75%)':>18} {'близко':>7} {'сейчас':>7}")
            for g in cats:
                if g["median"] is None:
                    print(f"{g['category'][:40]:40} {g['items']:>7} {'-':>8}   эталона нет ни у одного товара")
                    continue
                print(f"{g['category'][:40]:40} {g['items']:>7} {g['median']:8.2f} "
                      f"{g['p25']:8.2f} - {g['p75']:<7.2f} {g['close'] * 100:6.0f}% "
                      f"{'-' if g['current_median'] is None else format(g['current_median'], '.2f'):>7}")
            why: dict[str, int] = {}
            for r in items:
                if r["why"]:
                    why[r["why"]] = why.get(r["why"], 0) + 1
            for w, n in why.items():
                print(f"  без наценки: {w} — {n}")
        if args.xlsx:
            from priceapp.excel import xlsx_bytes
            from openpyxl import load_workbook
            import io
            cat_h = ["Площадка", "Категория", "Товаров", "Артикулов", "С эталоном", "Без эталона",
                     "Наценка (медиана)", "25%", "75%", "Близко к медиане (±5%), доля", "Сейчас (медиана)",
                     "К эталону, %"]
            cat_d = [[PLATFORMS[g["platform"]], g["category"], g["items"], g["articles"], g["with_target"],
                      g["no_target"], _r(g["median"]), _r(g["p25"]), _r(g["p75"]), _r(g["close"], 2),
                      _r(g["current_median"]), g["shift"]] for g in all_cats]
            it_h = ["Площадка", "Кабинет", "ID_1С", "Артикул", "Категория", "Комиссия, %", "Себестоимость, ₽",
                    "Эталон (покупатель WB), ₽", "Цель для покупателя, ₽", "Цена для покупателя сейчас, ₽",
                    "Наценка цель", "Наценка сейчас", "Почему нет"]
            it_d = [[PLATFORMS[r["platform"]], r["account"], r["item_id"], r["article"], r["category"],
                     float(r["commission"]) if r["commission"] is not None else None,
                     float(r["base"]) if r["base"] is not None else None, _r(r["etalon"], 0),
                     _r(r["target_price"], 0), r["current_sale"], _r(r["target_m"]), _r(r["current_m"]), r["why"]]
                    for r in all_items]
            imp_h = ["Где (код)", "Где", "Категория", "Вид наценки", "Наценка"]
            imp_d = [[g["platform"], PLATFORMS[g["platform"]], g["category"], "маржинальность", round(g["median"], 3)]
                     for g in all_cats if g["median"] is not None and g["category"] != NO_CATEGORY]
            wb = load_workbook(io.BytesIO(xlsx_bytes(cat_h, cat_d)))
            wb.active.title = "Категории"
            for title, (h, d) in (("Товары", (it_h, it_d)), ("Импорт", (imp_h, imp_d))):
                ws = wb.create_sheet(title)
                ws.append(h)
                for row in d:
                    ws.append(row)
            wb.move_sheet("Импорт", offset=-2)    # импорт читает первый лист
            wb.active = 0
            wb.save(args.xlsx)
            print(f"\nФайл: {args.xlsx} (первый лист «Импорт» — для «Цены товаров» → по категориям → Импорт)")
    finally:
        db.rollback()
        db.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
