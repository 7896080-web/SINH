"""Файл Ламоды «цены конкурентов» по мэппингу: карточка Ламоды -> та же вещь на WB.

Ламода просит шаблон (`competitors_prices.xlsx`): «Артикул Ламода» (parentSku), «Площадка
конкурента», «Артикул или ссылка на площадке конкурента», «Уникальный товар». Заполняем
по сопоставлению программы, а не по названиям: размеры карточки Ламоды -> SKU 1С -> строки
каталога кабинетов WB с теми же SKU -> карточка WB (nmID).

  * у вещи бывает несколько карточек WB (три кабинета, иногда две карточки в одном) — берём
    ту, где совпало больше всего размеров, при равенстве — с загруженной ценой, затем первый
    кабинет по имени; остальные перечислены на листе «Подробно»;
  * нет на WB, но есть на Ozon — строки в шаблоне нет (ссылку на Ozon по нашему каталогу не
    собрать: у нас id товара, а ссылка — по SKU витрины), карточка названа на «Подробно»;
  * нет ни на WB, ни на Ozon — «Уникальный товар» = «Да».

    .venv\\Scripts\\python.exe scripts\\lamoda_competitors.py --template C:\\temp\\competitors_prices.xlsx --out C:\\temp\\competitors_prices_filled.xlsx

Ничего не отправляет и ничего не меняет в базе.
"""
from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
try:
    sys.stdout.reconfigure(errors="replace")      # консоль Windows — cp1251
except Exception:
    pass

import priceapp  # noqa: E402,F401  (.env читается на импорте)
from priceapp import mapping  # noqa: E402
from priceapp.database import SessionLocal  # noqa: E402
from priceapp.models import Account, OnecBarcode  # noqa: E402
from priceapp.platforms import card_key, lamoda_parent  # noqa: E402
from priceapp.pricing import article_map, article_of  # noqa: E402

WB_URL = "https://www.wildberries.ru/catalog/{}/detail.aspx"
HEAD = ["Артикул Ламода", "Площадка конкурента", "Артикул или ссылка на площадке конкурента", "Уникальный товар"]


def _accounts(db, platform):
    return (db.query(Account).filter(Account.platform == platform, Account.is_active.is_(True))
            .order_by(Account.name).all())


def build(db) -> list[dict]:
    articles = article_map(db)
    info = {}
    for b in db.query(OnecBarcode):
        info.setdefault(b.item_id, b)
    # SKU 1С -> карточки WB: (кабинет, nmID) -> строки
    wb: dict[str, dict[tuple[str, str], list]] = {}
    for a in _accounts(db, "wb"):
        for item_id, rows in mapping.account_items(db, a.id).items():
            for r in rows:
                nm = card_key("wb", r)
                if nm.isdigit():
                    wb.setdefault(item_id, {}).setdefault((a.name, nm), []).append(r)
    ozon = set()
    for a in _accounts(db, "ozon"):
        ozon.update(mapping.account_items(db, a.id))
    out = []
    for a in _accounts(db, "lamoda"):
        cards: dict[str, dict] = {}
        for item_id, rows in mapping.account_items(db, a.id).items():
            for r in rows:
                p = lamoda_parent(r.external_id)
                if not p:
                    continue
                c = cards.setdefault(p, {"parent": p, "account": a.name, "items": set(), "lamoda_price": None,
                                         "lamoda_sale": None})
                c["items"].add(item_id)
                if r.current_price and (c["lamoda_price"] is None or r.current_price > c["lamoda_price"]):
                    c["lamoda_price"], c["lamoda_sale"] = r.current_price, r.current_sale_price or r.current_price
        for c in cards.values():
            votes: dict[tuple[str, str], set] = {}
            priced: dict[tuple[str, str], int] = {}
            for item_id in c["items"]:
                for key, rows in (wb.get(item_id) or {}).items():
                    votes.setdefault(key, set()).add(item_id)
                    for r in rows:
                        if r.current_price:
                            priced[key] = max(priced.get(key, 0), r.current_sale_price or r.current_price)
            ranked = sorted(votes, key=lambda k: (-len(votes[k]), 0 if k in priced else 1, k[0], k[1]))
            sku = [info.get(i) for i in sorted(c["items"])]
            c.update({
                "articles": ", ".join(sorted({article_of(articles, i) for i in c["items"]})),
                "colors": ", ".join(sorted({s.color for s in sku if s and s.color})),
                "sizes": len(c["items"]),
                "wb": ranked[0] if ranked else None,
                "wb_matched": len(votes[ranked[0]]) if ranked else 0,
                "wb_sale": priced.get(ranked[0]) if ranked else None,
                "wb_other": ", ".join(f"{k[0]} {k[1]}" for k in ranked[1:]),
                "on_ozon": bool(c["items"] & ozon),
            })
            out.append(c)
    return sorted(out, key=lambda c: (c["articles"], c["parent"]))


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--template", required=True, help="шаблон Ламоды competitors_prices.xlsx")
    ap.add_argument("--out", required=True)
    args = ap.parse_args(argv)
    from openpyxl import load_workbook
    if not Path(args.template).exists():
        print(f"Нет файла {args.template}")
        return 1
    db = SessionLocal()
    try:
        cards = build(db)
    finally:
        db.rollback()
        db.close()
    book = load_workbook(args.template)
    ws = book.worksheets[0]
    if [ws.cell(1, i + 1).value for i in range(4)] != HEAD:
        print(f"Шаблон не тот: в первой строке ожидаются колонки {HEAD}")
        return 1
    for row in range(2, ws.max_row + 1):          # прежние строки данных (колонки A-D) — заново
        for col in range(1, 5):
            ws.cell(row, col).value = None
    n = 2
    wb_n = uniq = ozon_only = 0
    for c in cards:
        if c["wb"]:
            ws.cell(n, 1).value, ws.cell(n, 2).value = c["parent"], "Wildberries"
            ws.cell(n, 3).value = WB_URL.format(c["wb"][1])
            wb_n += 1
        elif not c["on_ozon"]:
            ws.cell(n, 1).value, ws.cell(n, 4).value = c["parent"], "Да"
            uniq += 1
        else:
            ozon_only += 1
            continue
        n += 1
    det = book.create_sheet("Подробно")
    det.append(["Артикул Ламода", "Артикулы 1С", "Цвет", "Размеров на Ламоде", "Ламода: цена до скидки, ₽",
                "Ламода: со скидкой, ₽", "WB: кабинет", "WB: артикул (nmID)", "Ссылка WB",
                "Совпало размеров", "WB: со скидкой продавца, ₽", "Другие карточки WB", "Есть на Ozon", "В шаблоне"])
    for c in cards:
        w = c["wb"]
        det.append([c["parent"], c["articles"], c["colors"], c["sizes"], c["lamoda_price"], c["lamoda_sale"],
                    w[0] if w else None, int(w[1]) if w else None, WB_URL.format(w[1]) if w else None,
                    c["wb_matched"] or None, c["wb_sale"], c["wb_other"] or None, "да" if c["on_ozon"] else "",
                    "Wildberries" if w else ("уникальный" if not c["on_ozon"] else "нет — только Ozon")])
    book.save(args.out)
    print(f"Карточек Ламоды: {len(cards)}; со ссылкой на WB {wb_n}; уникальных {uniq}; "
          f"только на Ozon (в шаблон не вошли) {ozon_only}.")
    part = [c for c in cards if c["wb"] and c["wb_matched"] < c["sizes"]]
    if part:
        print(f"На WB совпали не все размеры: {len(part)} — см. «Подробно», колонка «Совпало размеров».")
    print(f"Файл: {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
