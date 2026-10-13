"""Цены Kit и Ламоды от эталона WB через скидку продавца — ТОЛЬКО ЧТЕНИЕ, ничего не отправляет.

Логика (задана владельцем 12.10, уточнена 13.10):
  * эталон по артикулу 1С = средняя по всем кабинетам WB ЭТАЛОННАЯ цена после скидки продавца
    (записанная `etalon_apply.py --adopt`; где её нет — нынешняя) минус 31%;
  * Ламода — эталон; за WB сама НЕ следует: план и файл, отправка — отдельно;
  * цель — цена ПОСЛЕ скидки продавца: Kit = эталон − 10%, Ламода = эталон;
  * цель достигается СКИДКОЙ продавца на площадке: при нынешней цене до скидки
    нужная скидка = 1 − цель / цена до скидки; коридор скидки — 35–50%;
  * скидка вне коридора — цену до скидки тоже менять: такую, чтобы скидка встала
    на ближнюю границу коридора (35% или 50%).

    C:\\repricer\\.venv\\Scripts\\python.exe C:\\repricer\\scripts\\etalon_prices.py --xlsx C:\\repricer\\эталон_цены.xlsx
    ... --only lamoda --xlsx C:\\repricer\\план_ламода.xlsx      только Ламода: лист «Сводка» по категориям и «Ламода»
    ... --wb-discount 31 --kit -10 --lamoda 0 --disc-min 35 --disc-max 50

WB — источник эталона: цена СО СКИДКОЙ продавца на нём не меняется; меняются цена до
скидки и скидка — так, чтобы скидка встала в коридор (по умолчанию тот же 35–50%,
`--wb-disc-min/--wb-disc-max`). Лист «WB» и сводка по категориям: цена со скидкой и
скидка в каждом кабинете, эталон, кабинет к средней по WB, что сделать.

Цель округляется тем же правилом, что и цены программы (шаг и окончание — «Правила
площадок»). Маржинальность — по цели: цель × (1 − комиссия) / себестоимость, ₽; ниже
пола площадки — пометка.
"""
from __future__ import annotations

import argparse
import os
import statistics
import sys
from decimal import ROUND_FLOOR, ROUND_HALF_UP, Decimal

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.stdout.reconfigure(errors="replace")      # консоль Windows — cp1251

import priceapp  # noqa: E402,F401  (.env читается на импорте)
from priceapp import mapping, platforms, rates  # noqa: E402
from priceapp.database import SessionLocal  # noqa: E402
from priceapp.models import Account, OnecBarcode, OnecCost  # noqa: E402
from priceapp.platforms import PLATFORMS  # noqa: E402
from priceapp.pricing import (article_map, article_of, base_price, get_rule, item_facts,  # noqa: E402
                              round_price)

NO_CATEGORY = "(без категории)"

# Что делать с товаром — одно слово на строку, по нему же сводка.
ONLY_DISCOUNT = "сменить скидку"
LIST_UP = "поднять цену до скидки"
LIST_DOWN = "снизить цену до скидки"
SAME = "уже так"
NO_ETALON = "нет эталона (нет на WB)"
NO_CURRENT = "нет текущей цены"
NO_COST = "нет себестоимости"


def _accounts(db, platform):
    return (db.query(Account).filter(Account.is_active.is_(True), Account.platform == platform)
            .order_by(Account.name).all())


def etalon(db, articles, wb_discount: Decimal, stats: dict | None = None) -> dict[str, Decimal]:
    """Эталон по артикулу 1С = средняя по всем кабинетам WB ЭТАЛОННАЯ цена после скидки
    продавца (записанная `etalon_apply.py --adopt` / `reference_tool.py adopt-current`) минус
    `wb_discount` %. Где эталона WB ещё нет — нынешняя цена после скидки продавца."""
    from priceapp import reference
    from priceapp.models import ReferencePrice
    rule = get_rule(db, "wb")
    vals: dict[str, list[int]] = {}
    st = stats if stats is not None else {}
    st.setdefault("ref", 0)
    st.setdefault("current", 0)
    for a in _accounts(db, "wb"):
        refs = {r.item_id: r for r in db.query(ReferencePrice).filter(ReferencePrice.account_id == a.id)}
        for item_id, rows in mapping.account_items(db, a.id).items():
            ref = refs.get(item_id)
            sale = reference.ref_sale(ref) if ref is not None else None
            if sale:
                st["ref"] += 1
            else:
                sale = item_facts(rule, rows).sale
                if sale:
                    st["current"] += 1
            if sale:
                vals.setdefault(article_of(articles, item_id), []).append(sale)
    return {art: Decimal(sum(v)) / len(v) * (1 - wb_discount / 100) for art, v in vals.items()}


def plan_item(target: int, current_list: int | None, current_sale: int | None,
              lo: Decimal, hi: Decimal) -> dict:
    """Как получить цену со скидкой `target`: какая скидка и какая цена до скидки."""
    if not current_list:
        return {"action": NO_CURRENT, "disc": None, "list": None}
    need = Decimal(1) - Decimal(target) / Decimal(current_list)
    if current_sale == target:
        return {"action": SAME, "disc": need, "list": current_list}
    if lo / 100 <= need <= hi / 100:
        return {"action": ONLY_DISCOUNT, "disc": need, "list": current_list}
    bound = lo / 100 if need < lo / 100 else hi / 100
    new_list = int((Decimal(target) / (1 - bound)).quantize(Decimal(1), ROUND_HALF_UP))
    return {"action": LIST_UP if new_list > current_list else LIST_DOWN, "disc": bound, "list": new_list,
            "need_at_current": need}


def compute(db, platform: str, shift: Decimal, et: dict, articles, costs, info, usd, lo, hi) -> list[dict]:
    rule = get_rule(db, platform)
    out = []
    for a in _accounts(db, platform):
        for item_id, rows in mapping.account_items(db, a.id).items():
            art = article_of(articles, item_id)
            facts = item_facts(rule, rows)
            sku = info.get(item_id)
            r = {"platform": platform, "account": a.name, "item_id": item_id, "article": art,
                 "size": sku.size if sku else "", "color": sku.color if sku else "",
                 "category": facts.category or NO_CATEGORY, "commission": facts.commission,
                 "current_list": facts.current, "current_sale": facts.sale,
                 "current_disc": facts.discount if facts.current else None,
                 "etalon": et.get(art), "target": None, "action": "", "disc": None, "list": None,
                 "change": None, "margin": None, "below_floor": False}
            if r["etalon"] is None:
                r["action"] = NO_ETALON
                out.append(r)
                continue
            r["target"] = round_price(r["etalon"] * (1 + shift / 100), rule.round_step, rule.round_minus)
            r.update(plan_item(r["target"], facts.current, facts.sale, lo, hi))
            if facts.sale:
                r["change"] = (r["target"] - facts.sale) * 100.0 / facts.sale
            base = base_price(costs.get(item_id), usd)
            if base and facts.commission is not None:
                r["margin"] = Decimal(r["target"]) * (1 - facts.commission / 100) / base
                r["below_floor"] = r["margin"] < Decimal(str(rule.min_markup_coef))
            elif not base:
                r["action"] = r["action"] or NO_COST
            out.append(r)
    return out


def wb_rows(db, et: dict, wb_discount: Decimal, articles, costs, info, usd,
            lo: Decimal = Decimal(35), hi: Decimal = Decimal(50)) -> list[dict]:
    """WB — источник эталона, цель на нём не ставится. По каждому товару каждого кабинета:
    цена до скидки, со скидкой, скидка продавца, средняя по кабинетам (от неё эталон) и
    насколько этот кабинет от неё отстоит, эталон, маржинальность по текущей."""
    rule = get_rule(db, "wb")
    keep = 1 - wb_discount / 100
    out = []
    for a in _accounts(db, "wb"):
        for item_id, rows in mapping.account_items(db, a.id).items():
            art = article_of(articles, item_id)
            facts = item_facts(rule, rows)
            sku = info.get(item_id)
            e = et.get(art)
            avg = e / keep if e is not None and keep else None
            base = base_price(costs.get(item_id), usd)
            margin = (Decimal(facts.sale) * (1 - facts.commission / 100) / base
                      if facts.sale and base and facts.commission is not None else None)
            out.append({"platform": "wb", "account": a.name, "item_id": item_id, "article": art,
                        "size": sku.size if sku else "", "color": sku.color if sku else "",
                        "category": facts.category or NO_CATEGORY, "current_list": facts.current,
                        "current_sale": facts.sale, "current_disc": facts.discount if facts.current else None,
                        "avg": avg, "etalon": e,
                        "vs_avg": (float(facts.sale) / float(avg) - 1) * 100 if facts.sale and avg else None,
                        "margin": margin, "commission": facts.commission,
                        "card": _wb_card(rows),
                        **wb_plan(facts.current, facts.sale, lo, hi)})
    mark_mixed_cards(out)
    return out


MIXED = "цены со скидкой по размерам разные — отправлять ТОЛЬКО по размерам"


def _wb_card(rows) -> str:
    for r in rows or []:
        if r is not None and getattr(r, "external_id", ""):
            return platforms.card_key("wb", r)
    return ""


def mark_mixed_cards(out: list[dict]) -> int:
    """Цена на WB уходит карточкой (nmID) — одна на все размеры, наибольшая. Если у размеров
    одной карточки цена со скидкой РАЗНАЯ, такая отправка подняла бы цену дешёвым размерам,
    а цена со скидкой не должна вырасти ни у одного. Помечаем такие карточки: им годится
    только цена по размерам. Возвращает число помеченных карточек."""
    sales: dict[tuple, set] = {}
    for r in out:
        if r["card"] and r["current_sale"]:
            sales.setdefault((r["account"], r["card"]), set()).add(r["current_sale"])
    mixed = {k for k, v in sales.items() if len(v) > 1}
    for r in out:
        r["mixed"] = (r["account"], r["card"]) in mixed
    return len(mixed)


IN_CORRIDOR = "скидка уже в коридоре"


def wb_plan(current_list: int | None, current_sale: int | None, lo: Decimal, hi: Decimal) -> dict:
    """WB: цену со скидкой продавца НЕ трогаем; цену до скидки и скидку — так, чтобы скидка
    была в коридоре, а цена со скидкой осталась прежней."""
    if not current_list or not current_sale:
        return {"action": NO_CURRENT, "disc": None, "list": None}
    need = Decimal(1) - Decimal(current_sale) / Decimal(current_list)
    if lo / 100 <= need <= hi / 100:
        return {"action": IN_CORRIDOR, "disc": need, "list": current_list}
    bound = lo / 100 if need < lo / 100 else hi / 100
    # вниз, а не к ближайшему: цена со скидкой = цена до скидки × (1 − скидка) не должна
    # стать больше нынешней даже на рубль
    new_list = int((Decimal(current_sale) / (1 - bound)).quantize(Decimal(1), ROUND_FLOOR))
    return {"action": LIST_UP if new_list > current_list else LIST_DOWN, "disc": bound, "list": new_list}


def _pct(d):
    return None if d is None else round(float(d) * 100, 1)


def _f(v, nd=2):
    return None if v is None else round(float(v), nd)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--wb-discount", type=Decimal, default=Decimal(31))
    ap.add_argument("--kit", type=Decimal, default=Decimal(-10))
    ap.add_argument("--lamoda", type=Decimal, default=Decimal(0))
    ap.add_argument("--disc-min", type=Decimal, default=Decimal(35))
    ap.add_argument("--disc-max", type=Decimal, default=Decimal(50))
    ap.add_argument("--wb-disc-min", type=Decimal, default=None, help="коридор скидки на WB (по умолчанию как у остальных)")
    ap.add_argument("--wb-disc-max", type=Decimal, default=None)
    ap.add_argument("--only", choices=("kit", "lamoda"), default="",
                    help="только эта площадка (без листа WB)")
    ap.add_argument("--show", type=int, default=15, help="примеров на площадку в консоли")
    ap.add_argument("--xlsx", default="", help="сохранить расчёт в .xlsx")
    args = ap.parse_args()
    lo, hi = args.disc_min, args.disc_max
    db = SessionLocal()
    try:
        rate = rates.current(db)
        if rate is None:
            print("Курса доллара нет — себестоимость в рублях не посчитать. Страница «Курс $».")
            return 1
        articles = article_map(db)
        costs = {c.item_id: c.cost_usd for c in db.query(OnecCost)}
        info = {}
        for b in db.query(OnecBarcode):
            info.setdefault(b.item_id, b)
        est: dict = {}
        et = etalon(db, articles, args.wb_discount, est)
        print(f"Эталон: средняя цена продавца со скидкой на WB минус {args.wb_discount}% — {len(et)} артикулов. "
              f"Коридор скидки продавца {lo}–{hi}%. НИЧЕГО НЕ ОТПРАВЛЯЕТСЯ.")
        wlo = args.wb_disc_min if args.wb_disc_min is not None else lo
        whi = args.wb_disc_max if args.wb_disc_max is not None else hi
        print(f"Эталон WB (цена после скидки продавца): из записанного эталона — {est['ref']} SKU, "
              f"из нынешней цены (эталона нет) — {est['current']} SKU.")
        wb = [] if args.only else wb_rows(db, et, args.wb_discount, articles, costs, info, rate.usd_rub, wlo, whi)
        if wb:
            print(f"\n=== Wildberries — цена со скидкой продавца НЕ меняется; цена до скидки и скидка — "
                  f"под коридор {wlo}–{whi}%. Товаров {len(wb)}")
            cnt: dict[str, int] = {}
            for r in wb:
                cnt[r["action"]] = cnt.get(r["action"], 0) + 1
            for act in (IN_CORRIDOR, LIST_UP, LIST_DOWN, NO_CURRENT):
                if cnt.get(act):
                    print(f"  {cnt[act]:6d}  {act}")
            mixed = {(r["account"], r["card"]) for r in wb if r["mixed"]}
            if mixed:
                print(f"  {len(mixed):6d}  карточек: {MIXED} (карточкой ушла бы наибольшая цена, "
                      f"и дешёвые размеры подорожали бы)")
            cats: dict[str, list[dict]] = {}
            for r in wb:
                cats.setdefault(r["category"], []).append(r)
            print(f"  {'категория':34} {'арт.':>4} {'со скидкой':>10} {'скидка':>7} {'эталон':>7}  "
                  + "  ".join(f"{a.name[:12]:>12}" for a in _accounts(db, "wb")))
            for cat, rs in sorted(cats.items(), key=lambda kv: -len(kv[1])):
                sale = [r["current_sale"] for r in rs if r["current_sale"]]
                disc = [float(r["current_disc"]) * 100 for r in rs if r["current_disc"] is not None]
                ets = [float(r["etalon"]) for r in rs if r["etalon"] is not None]
                per = []
                for a in _accounts(db, "wb"):
                    v = [r["vs_avg"] for r in rs if r["account"] == a.name and r["vs_avg"] is not None]
                    per.append(f"{statistics.median(v):+11.1f}%" if v else f"{'-':>12}")
                print(f"  {cat[:34]:34} {len({r['article'] for r in rs}):>4} "
                      f"{statistics.median(sale) if sale else '-':>10} "
                      f"{(format(statistics.median(disc), '.0f') + '%') if disc else '-':>7} "
                      f"{format(statistics.median(ets), '.0f') if ets else '-':>7}  " + "  ".join(per))
            print("  (по кабинетам — насколько цена со скидкой в кабинете выше/ниже средней по WB, от которой эталон)")
            print("  примеры:")
            for r in [r for r in wb if r["action"] in (LIST_UP, LIST_DOWN)][:args.show]:
                print(f"    {r['account']}: {r['article']} {r['size']} {r['color']}: сейчас {r['current_list']} / "
                      f"{r['current_sale']} ({_pct(r['current_disc'])}%) -> {r['action']}: цена до скидки "
                      f"{r['list']}, скидка {_pct(r['disc'])}%, со скидкой {r['current_sale']} (без изменений)")
        rows = []
        for platform, shift in (("kit", args.kit), ("lamoda", args.lamoda)):
            if args.only and platform != args.only:
                continue
            if not _accounts(db, platform):
                print(f"\n{PLATFORMS[platform]}: активных кабинетов нет.")
                continue
            items = compute(db, platform, shift, et, articles, costs, info, rate.usd_rub, lo, hi)
            rows += items
            cnt: dict[str, int] = {}
            for r in items:
                cnt[r["action"]] = cnt.get(r["action"], 0) + 1
            ch = [r["change"] for r in items if r["change"] is not None]
            print(f"\n=== {PLATFORMS[platform]}: цель со скидкой = эталон {shift:+}% — товаров {len(items)}")
            for act in (SAME, ONLY_DISCOUNT, LIST_UP, LIST_DOWN, NO_ETALON, NO_CURRENT, NO_COST):
                if cnt.get(act):
                    print(f"  {cnt[act]:6d}  {act}")
            if ch:
                print(f"  цена после скидки продавца изменится: медиана {statistics.median(ch):+.1f}%, "
                      f"от {min(ch):+.1f}% до {max(ch):+.1f}%")
            bf = sum(1 for r in items if r["below_floor"])
            if bf:
                print(f"  ниже пола по цели: {bf} — такая цена не ушла бы")
            print(f"  примеры:")
            shown = [r for r in items if r["target"]][:args.show]
            for r in shown:
                disc = f"скидка {_pct(r['disc'])}%" if r["disc"] is not None else ""
                lst = f"цена до скидки {r['list']}" if r["list"] else ""
                print(f"    {r['article']} {r['size']} {r['color']}: сейчас {r['current_list'] or '-'} / "
                      f"{r['current_sale'] or '-'} ({_pct(r['current_disc']) if r['current_disc'] is not None else '-'}%)"
                      f" -> цель {r['target']}: {r['action']}; {lst} {disc}".rstrip())
        if args.xlsx:
            from priceapp.excel import xlsx_bytes
            head = ["Площадка", "Кабинет", "Категория", "Артикул", "Размер", "Цвет", "ID_1С",
                    "Эталон WB после скидки продавца (средняя), ₽", f"Эталон (WB −{args.wb_discount}%), ₽",
                    "Цель со скидкой, ₽", "Сейчас: цена до скидки, ₽",
                    "Сейчас: цена со скидкой, ₽", "Сейчас: скидка, %", "Что сделать",
                    "Новая цена до скидки, ₽", "Новая скидка, %", "Изменение цены после скидки продавца, %",
                    "Маржинальность по цели", "Ниже пола", "Комиссия, %"]
            data = [[PLATFORMS[r["platform"]], r["account"], r["category"], r["article"], r["size"], r["color"],
                     r["item_id"], _f(r["etalon"] / (1 - args.wb_discount / 100), 0) if r["etalon"] is not None else None,
                     _f(r["etalon"], 0), r["target"], r["current_list"], r["current_sale"],
                     _pct(r["current_disc"]), r["action"], r["list"], _pct(r["disc"]),
                     _f(r["change"], 1), _f(r["margin"]), "да" if r["below_floor"] else "",
                     _f(r["commission"])]
                    for r in sorted(rows, key=lambda r: (r["platform"], r["category"], r["article"], r["size"]))]
            from openpyxl import load_workbook
            import io
            book = load_workbook(io.BytesIO(xlsx_bytes(head, data)))
            book.active.title = {"lamoda": "Ламода", "kit": "Kit"}.get(args.only, "Kit и Ламода")
            # Сводка по категориям: как изменится цена после скидки продавца
            sv = book.create_sheet("Сводка", 0)
            sv.append(["Площадка", "Категория", "Товаров", "С эталоном", "Дешевле", "Дороже", "Без изменений",
                       "Изменение: медиана, %", "мин, %", "макс, %", "Ниже пола"])
            groups: dict[tuple, list[dict]] = {}
            for r in rows:
                groups.setdefault((PLATFORMS[r["platform"]], r["category"]), []).append(r)
            for (plat, cat), rs in sorted(groups.items(), key=lambda kv: (kv[0][0], -len(kv[1]))):
                ch = [r["change"] for r in rs if r["change"] is not None]
                sv.append([plat, cat, len(rs), sum(1 for r in rs if r["target"]),
                           sum(1 for x in ch if x < -0.5), sum(1 for x in ch if x > 0.5),
                           sum(1 for x in ch if -0.5 <= x <= 0.5),
                           _f(statistics.median(ch), 1) if ch else None, _f(min(ch), 1) if ch else None,
                           _f(max(ch), 1) if ch else None, sum(1 for r in rs if r["below_floor"])])
            if args.only:
                book.save(args.xlsx)
                print(f"\nФайл: {args.xlsx}")
                return 0
            ws = book.create_sheet("WB")
            ws.append(["Кабинет", "Категория", "Артикул", "Размер", "Цвет", "ID_1С", "Сейчас: цена до скидки, ₽",
                       "Сейчас: цена со скидкой, ₽", "Сейчас: скидка, %", "Средняя со скидкой по WB, ₽",
                       "Кабинет к средней, %", f"Эталон (−{args.wb_discount}%), ₽", "Что сделать",
                       "Новая цена до скидки, ₽", "Новая скидка, %", "Цена со скидкой (не меняется), ₽",
                       "Маржинальность по текущей", "Комиссия, %", "Карточка"])
            for r in sorted(wb, key=lambda r: (r["category"], r["article"], r["size"], r["account"])):
                ws.append([r["account"], r["category"], r["article"], r["size"], r["color"], r["item_id"],
                           r["current_list"], r["current_sale"], _pct(r["current_disc"]), _f(r["avg"], 0),
                           _f(r["vs_avg"], 1), _f(r["etalon"], 0), r["action"], r["list"], _pct(r["disc"]),
                           r["current_sale"], _f(r["margin"]), _f(r["commission"]),
                           MIXED if r["mixed"] else ""])
            book.save(args.xlsx)
            print(f"\nФайл: {args.xlsx}")
    finally:
        db.rollback()
        db.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
