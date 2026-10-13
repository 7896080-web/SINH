r"""Новый эталон WB по заполненной форме маржинальности (scripts/etalon_cost.py).

  .venv\Scripts\python.exe scripts\etalon_apply.py --form маржинальность_wb.xlsx --xlsx план_эталон.xlsx
      ПЛАН: по каждой карточке WB — цена до скидки и скидка; ничего не отправляет
  ... --send      отправить план на WB (спросит подтверждение)
  ... --verify    дошёл ли план (по ценам, которые программа загрузила с WB после отправки)
  после сверки:   .venv\Scripts\python.exe scripts\reference_tool.py adopt-current — записать эталон

Правило (задано владельцем 13.10):
  * маржинальность кабинета для артикула — самая частная заполненная: «<кабинет>:
    маржинальность (если иная)», затем «Маржинальность артикула», затем маржинальность
    категории (лист «Категории»). Пусто везде — карточку не трогаем;
  * цель — цена СО СКИДКОЙ: себестоимость артикула, руб (из формы — та, что видел владелец)
    × маржинальность / (1 − комиссия), вверх до рубля; комиссия — нынешняя комиссия товара
    в кабинете (тариф WB по категории), как и в форме;
  * цель достигается ценой до скидки при нынешней скидке карточки, если та в коридоре
    35–50%; иначе скидка — на ближайшую границу. Цена до скидки — вверх до рубля, чтобы
    цена со скидкой была не ниже цели;
  * цена на WB — одна на карточку (nmID). Если у размеров карточки цели разные (разные
    артикулы 1С на одной карточке) — карточку не трогаем и называем.
"""
from __future__ import annotations

import argparse
import math
import os
import sys
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
try:
    sys.stdout.reconfigure(errors="replace")      # консоль Windows — cp1251
except Exception:
    pass

import priceapp  # noqa: E402,F401  (.env читается на импорте)
import wb_restore_discount as wrd  # noqa: E402
from etalon_cost import model_of  # noqa: E402
from priceapp import mapping  # noqa: E402
from priceapp.database import SessionLocal  # noqa: E402
from priceapp.models import Account  # noqa: E402
from priceapp.pricing import article_map, article_of, get_rule, item_facts  # noqa: E402


def _num(v):
    if v is None or v == "":
        return None
    try:
        x = float(str(v).replace(",", ".").strip())
    except ValueError:
        return None
    return x if x > 0 else None


def read_form(path: Path) -> tuple[dict, dict, list[str]]:
    """(артикул -> {cost_rub, art, over{кабинет: m}, category}, категория -> m, кабинеты)."""
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    for need in ("Артикулы", "Категории"):
        if need not in book.sheetnames:
            raise SystemExit(f"В файле нет листа «{need}» — нужна форма scripts/etalon_cost.py")
    rows = book["Артикулы"].iter_rows(values_only=True)
    head = [str(h or "").strip() for h in next(rows)]
    key = "Артикул" if "Артикул" in head else "Изделие"
    for need in (key, "Категория", "Себестоимость макс, руб", "Маржинальность артикула"):
        if need not in head:
            raise SystemExit(f"На листе «Артикулы» нет колонки «{need}»")
    ix = {h: i for i, h in enumerate(head) if h}
    cabs = [h[:-len(": маржинальность (если иная)")] for h in head if h.endswith(": маржинальность (если иная)")]
    arts = {}
    for r in rows:
        a = r[ix[key]]
        if a is None or str(a).strip() == "":
            continue
        arts[str(a).strip()] = {
            "category": str(r[ix["Категория"]] or ""),
            "cost_rub": _num(r[ix["Себестоимость макс, руб"]]),
            "art": _num(r[ix["Маржинальность артикула"]]),
            "over": {c: _num(r[ix[f"{c}: маржинальность (если иная)"]]) for c in cabs}}
    cats = {}
    for r in list(book["Категории"].iter_rows(values_only=True))[1:]:
        if r and r[0]:
            cats[str(r[0])] = _num(r[3] if len(r) > 3 else None)
    return arts, cats, cabs


def margin_for(arts: dict, cats: dict, art: str, cab: str):
    d = arts.get(art)
    if d is None:
        return None
    return d["over"].get(cab) or d["art"] or cats.get(d["category"])


def target_sale(cost_rub: float, m: float, commission) -> int:
    return math.ceil(cost_rub * m / (1 - float(commission) / 100) - 1e-9)


def list_for(target: int, cur_disc: int | None, lo: int, hi: int) -> tuple[int, int]:
    """Цена до скидки и целая скидка: скидка карточки, если в коридоре, иначе ближняя граница;
    цена до скидки вверх — цена со скидкой не ниже цели."""
    d = cur_disc if cur_disc is not None and lo <= cur_disc <= hi else (lo if (cur_disc or 0) < lo else hi)
    return math.ceil(target * 100 / (100 - d) - 1e-9), d


def build_plan(db, arts: dict, cats: dict, live: dict, lo: int, hi: int) -> list:
    """PlanCard на каждую карточку WB, где форма задаёт маржинальность."""
    rule = get_rule(db, "wb")
    articles = article_map(db)
    out = []
    for a in db.query(Account).filter(Account.is_active.is_(True), Account.platform == "wb").order_by(Account.name):
        cards: dict[str, dict] = {}
        for item_id, rows in mapping.account_items(db, a.id).items():
            art = model_of(article_of(articles, item_id))
            m = margin_for(arts, cats, art, a.name)
            cost = (arts.get(art) or {}).get("cost_rub")
            facts = item_facts(rule, rows)
            for r in rows:
                nm = (r.external_id or "").split(":")[0]
                if not nm.isdigit():
                    continue
                c = cards.setdefault(nm, {"targets": set(), "arts": set(), "why": ""})
                c["arts"].add(art)
                if m is None:
                    continue
                if not cost:
                    c["why"] = "нет себестоимости в форме"
                elif facts.commission is None:
                    c["why"] = "нет комиссии"
                else:
                    c["targets"].add(target_sale(cost, m, facts.commission))
        for nm, c in sorted(cards.items(), key=lambda kv: (sorted(kv[1]["arts"]), kv[0])):
            cur = live.get((a.id, nm))
            pc = wrd.PlanCard(a.id, a.name, nm, ", ".join(sorted(c["arts"]))[:120],
                              cur.price if cur else 0, cur.sale if cur else 0)
            if not c["targets"]:
                if c["why"]:
                    pc.skip = c["why"]
                    out.append(pc)
                continue                                # маржинальность не задана — не трогаем
            if len(c["targets"]) > 1:
                pc.skip = f"у размеров карточки разные цели ({', '.join(map(str, sorted(c['targets'])))})"
            elif cur is None:
                pc.skip = "карточки нет в загруженных ценах WB"
            else:
                t = next(iter(c["targets"]))
                pc.new_list, pc.disc = list_for(t, cur.discount, lo, hi)
                pc.target = t
            out.append(pc)
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Новый эталон WB по форме маржинальности")
    ap.add_argument("--form", required=True)
    ap.add_argument("--xlsx", default="")
    ap.add_argument("--disc-min", type=int, default=35)
    ap.add_argument("--disc-max", type=int, default=50)
    ap.add_argument("--send", action="store_true")
    ap.add_argument("--verify", action="store_true")
    ap.add_argument("--yes", action="store_true", help="не спрашивать подтверждение (для проверок)")
    ap.add_argument("--after", default="", help="база «сейчас» (по умолчанию рабочая)")
    args = ap.parse_args(argv)
    if not Path(args.form).exists():
        print(f"Нет файла {args.form}")
        return 1
    arts, cats, cabs = read_form(Path(args.form))
    live, accs = wrd.load(Path(args.after) if args.after else wrd._live_db())
    db = SessionLocal()
    try:
        cards = build_plan(db, arts, cats, live, args.disc_min, args.disc_max)
    finally:
        db.rollback()
        db.close()
    if args.verify:
        return wrd.verify_plan(cards, live, accs)
    go = [c for c in cards if not c.skip]
    print(f"Форма {args.form}: артикулов {len(arts)}; кабинеты {', '.join(cabs)}. "
          f"Скидка в коридоре {args.disc_min}-{args.disc_max}%.")
    for name, t in accs.values():
        print(f"  {name}: цены загружены с WB {t or '-'}")
    print(f"\nКарточек к отправке: {len(go)}")
    for name in sorted({c.account for c in go}):
        cs = [c for c in go if c.account == name]
        ch = [(c.target / c.old_sale - 1) * 100 for c in cs if c.old_sale]
        up = sum(1 for x in ch if x > 0.5)
        dn = sum(1 for x in ch if x < -0.5)
        med = sorted(ch)[len(ch) // 2] if ch else 0
        print(f"  {name}: {len(cs)}; цена со скидкой: дешевле {dn}, дороже {up}, без изменений {len(ch) - up - dn}; "
              f"медиана {med:+.1f}%")
    big = sorted((c for c in go if c.old_sale), key=lambda c: abs(c.target / c.old_sale - 1), reverse=True)[:15]
    if big:
        print("\nСамые большие изменения цены со скидкой:")
        for c in big:
            print(f"    {c.account}: {c.nm} {c.article}: {c.old_list} / {c.old_sale} -> {c.new_list}, скидка {c.disc}%,"
                  f" со скидкой {c.target} ({(c.target / c.old_sale - 1) * 100:+.1f}%)")
    skipped = [c for c in cards if c.skip]
    if skipped:
        print(f"\nНЕ отправляются ({len(skipped)}):")
        for c in skipped[:20]:
            print(f"    {c.account}: {c.nm} {c.article}: {c.skip}")
    if args.xlsx:
        from openpyxl import Workbook
        book = Workbook()
        book.remove(book.active)
        for name in sorted({c.account for c in cards}):
            ws = book.create_sheet(name[:31])
            ws.append(["Артикул WB", "Артикул", "Сейчас: цена до скидки", "Сейчас: со скидкой",
                       "Новая цена до скидки", "Новая скидка, %", "Цель со скидкой", "Изменение, %",
                       "Не отправляется"])
            for c in cards:
                if c.account == name:
                    t = getattr(c, "target", None)
                    ws.append([int(c.nm), c.article, c.old_list or None, c.old_sale or None, c.new_list, c.disc, t,
                               round((t / c.old_sale - 1) * 100, 1) if t and c.old_sale else None, c.skip])
        book.save(args.xlsx)
        print(f"\nФайл: {args.xlsx}")
    if not args.send:
        print("\nНичего не отправлено. Отправить: добавьте --send")
        return 0
    if not go:
        print("\nОтправлять нечего.")
        return 0
    if not args.yes:
        ans = input(f"\nОтправить на WB цену до скидки и скидку по {len(go)} карточкам? Введите ДА: ")
        if ans.strip().upper() != "ДА":
            print("Отменено, ничего не отправлено.")
            return 0
    db = SessionLocal()
    try:
        for name, res in wrd.send_plan(db, cards).items():
            print(f"  {name}: {res}")
    finally:
        db.close()
    print("\nДальше: перезапустите программу (загрузит цены WB), затем\n"
          f"  .venv\\Scripts\\python.exe scripts\\etalon_apply.py --form {args.form} --verify\n"
          "и, если «НЕ совпадает: 0», запишите эталон:\n"
          "  .venv\\Scripts\\python.exe scripts\\reference_tool.py adopt-current")
    return 0


if __name__ == "__main__":
    sys.exit(main())
