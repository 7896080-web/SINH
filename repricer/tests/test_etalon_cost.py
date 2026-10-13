"""Себестоимость по артикулу (верхняя планка, без цвета) и маржинальность от текущей цены WB."""
import os
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

from priceapp.models import OnecCost, PlatformItem, PriceChange
from tests import factories as F

ROOT = Path(__file__).resolve().parent.parent


def test_model_drops_three_digit_color_code_only():
    sys.path.insert(0, str(ROOT / "scripts"))
    import etalon_cost
    assert etalon_cost.model_of("256250-015") == "256250"
    assert etalon_cost.model_of("TC26-2775-1") == "TC26-2775-1"
    assert etalon_cost.model_of("048-SRT SAD-0032") == "048-SRT SAD-0032"


def test_cost_is_the_top_of_all_colors_and_margin_is_from_current_wb_price(db, tmp_path):
    w = F.account(db, "wb", "ИП 1", commission=25)
    F.manual_rate(db, "100")
    F.sku(db, "k1", "256250-015", "M", "Бежевый", barcodes=["b1"], cost_usd=20)
    F.sku(db, "k2", "256250-030", "M", "Черный", barcodes=["b2"], cost_usd=25)    # верхняя планка
    F.item(db, w, "b1", "256250-015", external_id="1:1"); F.item(db, w, "b2", "256250-030", external_id="2:1")
    for it in db.query(PlatformItem):
        it.current_price, it.current_sale_price = (13559, 7500) if it.barcode == "b1" else (12089, 6000)
    db.commit()
    out = tmp_path / "m.xlsx"
    r = subprocess.run([sys.executable, str(ROOT / "scripts" / "etalon_cost.py"), "--xlsx", str(out),
                        "--article", "256250"], capture_output=True, cwd=ROOT,
                       env={**os.environ, "PYTHONIOENCODING": "cp1251"})
    assert r.returncode == 0, r.stderr.decode("utf-8", "replace")
    text = r.stdout.decode("cp1251")
    assert "НИЧЕГО НЕ ОТПРАВЛЯЕТСЯ" in text and "Артикул 256250: себестоимость макс 25" in text
    arts = list(load_workbook(out)["Артикулы"].iter_rows(values_only=True))
    h = arts[0]
    row = dict(zip(h, arts[1]))
    assert row["Артикул"] == "256250" and row["Себестоимость макс, $"] == 25 and row["Себестоимость мин, $"] == 20
    goods = list(load_workbook(out)["Товары WB"].iter_rows(values_only=True))
    g = {x[goods[0].index("ID_1С")]: dict(zip(goods[0], x)) for x in goods[1:]}
    # 7500 × 0,75 = 5625 к получению; себестоимость артикула 25 $ × 100 = 2500 руб -> 2,25
    assert g["k1"]["Маржинальность (от макс)"] == 2.25 and g["k1"]["Маржинальность (от своей)"] == 2.81
    assert g["k2"]["Маржинальность (от макс)"] == 1.8
    assert db.query(PriceChange).count() == 0


def test_form_has_fill_columns_and_price_formulas(db, tmp_path):
    w = F.account(db, "wb", "ИП 1", commission=25)
    F.manual_rate(db, "100")
    F.sku(db, "k1", "256250-015", "M", "Бежевый", barcodes=["b1"], cost_usd=20)
    F.item(db, w, "b1", "256250-015", external_id="1:1")
    for it in db.query(PlatformItem):
        it.current_price, it.current_sale_price = 13559, 7500
    db.commit()
    out = tmp_path / "f.xlsx"
    r = subprocess.run([sys.executable, str(ROOT / "scripts" / "etalon_cost.py"), "--xlsx", str(out)],
                       capture_output=True, cwd=ROOT, env={**os.environ, "PYTHONIOENCODING": "cp1251"})
    assert r.returncode == 0, r.stderr.decode("utf-8", "replace")
    book = load_workbook(out)
    assert book.sheetnames == ["Артикулы", "Категории", "Товары WB", "Как заполнять"]
    ws = book["Артикулы"]
    h = [c.value for c in ws[1]]
    for name in ("Маржинальность артикула", "ИП 1: маржинальность (если иная)", "ИП 1: новая цена со скидкой"):
        assert name in h
    f = ws.cell(2, h.index("ИП 1: новая цена со скидкой") + 1).value
    assert f.startswith("=IF(") and "ROUNDUP($D2*" in f and "Категории!" in f
    assert book["Категории"].cell(1, 4).value == "Маржинальность категории"


def test_top_cost_only_from_skus_mapped_to_platform_cards(db, tmp_path):
    """Под «2101» в 1С и пиджак за 33,52 $, и свитшот за 8,66 $; на площадке только свитшот —
    пиджак верхней планкой быть не должен."""
    w = F.account(db, "wb", "ИП 1", commission=25)
    F.manual_rate(db, "100")
    F.sku(db, "jacket", "2101-008", "48", "Зелёный", barcodes=["j1"], cost_usd=33.52)   # на площадке нет
    F.sku(db, "sweat", "2101-100", "L", "Коричневый", barcodes=["s1"], cost_usd=8.66)
    F.item(db, w, "s1", "2101-100", external_id="1:1")
    for it in db.query(PlatformItem):
        it.current_price, it.current_sale_price = 9429, 5280
    db.commit()
    out = tmp_path / "m.xlsx"
    r = subprocess.run([sys.executable, str(ROOT / "scripts" / "etalon_cost.py"), "--xlsx", str(out)],
                       capture_output=True, cwd=ROOT, env={**os.environ, "PYTHONIOENCODING": "cp1251"})
    assert r.returncode == 0, r.stderr.decode("utf-8", "replace")
    rows = list(load_workbook(out)["Артикулы"].iter_rows(values_only=True))
    row = dict(zip(rows[0], rows[1]))
    assert row["Артикул"] == "2101" and row["Себестоимость макс, $"] == 8.66
