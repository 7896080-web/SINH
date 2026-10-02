"""Страницы: вход, все разделы, ключи (шифрованно), правила с комиссией, курс,
расчёт → подтверждение, ручные цены через Excel, загрузка файлов 1С."""
import io
from decimal import Decimal

from openpyxl import load_workbook

from priceapp.crypto import decrypt_value
from priceapp.models import Account, ApiCredential, OnecCost, PlatformItem, PlatformRule, PriceChange
from priceapp.platforms import CatalogRow
from tests import factories as f

RULE = {"commission_percent": "25", "base_coef": "2,667", "min_markup_coef": "1,3",
        "round_step": "10", "round_minus": "1", "max_change_percent": "20"}


def test_login_required(db):
    from fastapi.testclient import TestClient
    from priceapp.main import app
    with TestClient(app) as c:
        r = c.get("/prices", follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == "/login"


def test_all_pages_render(client, db):
    acc = f.account(db)
    for url in ("/prices?view=rules", "/prices?view=proposals", "/prices?view=products", "/prices?view=log",
                "/mapping", "/mapping?view=analysis", "/mapping?view=candidates", "/rate", "/api-keys",
                "/diagnostics"):
        r = client.get(url)
        assert r.status_code == 200, url


def test_api_keys_encrypted_and_masked(client, db):
    client.post("/api-keys/new", data={"platform": "ozon", "name": "ООО Озон"})
    acc = db.query(Account).one()
    client.post(f"/api-keys/{acc.id}", data={"name": "ООО Озон", "is_active": "1",
                                             "client_id": "123456", "api_key": "secret-api-key-xyz"})
    creds = {c.field_name: c.encrypted_value for c in db.query(ApiCredential)}
    assert "secret-api-key-xyz" not in creds["api_key"] and decrypt_value(creds["api_key"]) == "secret-api-key-xyz"
    page = client.get("/api-keys").text
    assert "secret-api-key-xyz" not in page and "secr" in page
    # пустое поле ключ не стирает
    client.post(f"/api-keys/{acc.id}", data={"name": "ООО Озон", "is_active": "1", "client_id": "", "api_key": ""})
    db.expire_all()
    assert db.query(ApiCredential).count() == 2


def test_browser_autofilled_login_password_never_becomes_a_platform_key(client, db):
    """Браузер подставляет пароль входа в поле ключа сам; «Сохранить» ради галочки
    «Активен» заменял им ключ Lamoda, и площадка отвечала invalid_client."""
    client.post("/api-keys/new", data={"platform": "lamoda", "name": "Ламода"})
    acc = db.query(Account).one()
    client.post(f"/api-keys/{acc.id}", data={"name": "Ламода", "is_active": "1", "client_id": "real-client-id",
                                             "client_secret": "real-secret", "seller_id": "777"})
    r = client.post(f"/api-keys/{acc.id}", data={"name": "Ламода", "is_active": "1", "client_id": "password1",
                                                 "client_secret": "", "seller_id": ""})
    assert "браузер подставил ваш пароль" in r.text and "«Client ID»" in r.text
    db.expire_all()
    creds = {c.field_name: decrypt_value(c.encrypted_value) for c in db.query(ApiCredential)}
    assert creds["client_id"] == "real-client-id"
    assert 'autocomplete="new-password"' in client.get("/api-keys").text


def test_check_and_catalog_with_fake_client(client, db, monkeypatch):
    from priceapp.routers import accounts as r

    class Fake:
        last_truncated = False

        def test_connection(self):
            return True, "ок"

        def get_catalog(self):
            return [CatalogRow("5:1", "b1", "JN", "Джинсы", "48"), CatalogRow("5:1", "b1", "JN", "дубль", "48")]

    monkeypatch.setattr(r, "CLIENT_FACTORY", lambda platform, creds: Fake())
    acc = f.account(db)
    assert "ок" in client.post(f"/api-keys/{acc.id}/check").text
    assert "баркодов 1" in client.post(f"/api-keys/{acc.id}/catalog").text
    assert db.query(PlatformItem).one().size == "48"
    # счётчик на карточке кабинета — число, а не dict.items из Jinja
    page = client.get("/api-keys").text
    assert "Каталог: 1 баркодов" in page and "built-in method" not in page


def test_rules_save_commission_and_validate(client, db):
    f.account(db, commission=None)
    client.post("/prices/rules/wb", data=RULE)
    db.expire_all()
    rule = db.query(PlatformRule).one()
    assert rule.commission_percent == Decimal("25") and rule.base_coef == Decimal("2.667")
    assert 'value="1,3"' in client.get("/prices?view=rules").text
    r = client.post("/prices/rules/wb", data={**RULE, "commission_percent": "100"})
    assert "не сохранено" in r.text
    r = client.post("/prices/rules/wb", data={**RULE, "min_markup_coef": "3"})
    assert "не сохранено" in r.text


def test_rate_manual_mode(client, db):
    assert "не сохранён" in client.post("/rate/mode", data={"mode": "manual", "manual": "abc"}).text
    client.post("/rate/mode", data={"mode": "manual", "manual": "81,5"})
    assert "81,5000" in client.get("/rate").text


def _priced(client, db):
    acc = f.account(db, commission=None)
    client.post("/prices/rules/wb", data=RULE)
    client.post("/rate/mode", data={"mode": "manual", "manual": "81.5"})
    f.sku(db, "u1", "39681", "L", barcodes=["b1"], cost_usd="16.24", name="Свитшот")
    f.item(db, acc, "b1", "39681-L", external_id="5:1")
    return acc


def test_recalculate_without_rate_refused(client, db):
    f.account(db)
    assert "Курса доллара нет" in client.post("/prices/recalculate", data={}).text


def test_recalculate_approve_flow_and_products_economics(client, db):
    acc = _priced(client, db)
    r = client.post("/prices/recalculate", data={})
    assert "предложено 1" in r.text and "3539" in r.text
    page = client.get(f"/prices?view=products&account_id={acc.id}").text
    for s in ("1 323,56", "3539", "2 654,25", "1 330,69", "2,01"):
        assert s in page, s
    ch = db.query(PriceChange).one()
    client.post("/prices/approve", data={"ids": [str(ch.id)]})
    db.refresh(ch)
    assert ch.status == "approved" and ch.decided_by == "op"


def test_manual_price_below_floor_cannot_be_approved(client, db):
    acc = _priced(client, db)
    client.post(f"/prices/manual/{acc.id}", data={"item_id": "u1", "value": "1500"})
    client.post("/prices/recalculate", data={})
    ch = db.query(PriceChange).one()
    assert ch.block_reason == "floor"
    assert "Не подтверждено" in client.post("/prices/approve", data={"ids": [str(ch.id)], "confirm_large": "true"}).text


def test_excel_roundtrip_manual_price(client, db):
    acc = _priced(client, db)
    r = client.get(f"/prices/export/{acc.id}")
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=headers.index("Наценка по расчётной, ₽") + 1).value == 1330.69
    ws.cell(row=2, column=headers.index("Ручная цена, ₽") + 1, value=3999)
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(f"/prices/import/{acc.id}", files={"file": ("p.xlsx", buf.getvalue())})
    assert "изменено: 1" in r.text
    assert db.query(PriceChange).count() == 0


def test_diagnostics_upload_and_dict_refused_before_ready(client, db):
    r = client.post("/diagnostics/upload", data={"kind": "cost"}, files={"file": ("cost.txt", "u1|16.24\n".encode())})
    assert "строк 1" in r.text and db.get(OnecCost, "u1").cost_usd == Decimal("16.24")
    assert "mark-3" in client.post("/diagnostics/request/dict").text
    assert "поставлено" in client.post("/diagnostics/request/cost").text


def test_health_reports_stale_workers(client):
    r = client.get("/health")
    assert r.status_code == 503 and "onec_exchange" in r.text
