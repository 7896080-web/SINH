"""Акции WB: календарь в базу, наши карточки в обычной акции — цена при входе, «% от эталона для
вхождения», порог, пометка; порог «для акции» до её конца; «Добавить в акцию»; строка во «Внимании»;
Excel. Эталон 2009 ₽ без скидки, пол 1387 ₽ (себестоимость 800 ₽ × 1,3 / 0,75)."""
import io
from datetime import timedelta
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from priceapp import accounts as acc_mod, reference, wb_discount as wd, wb_promo as wp
from priceapp.models import AuditLog, PlatformItem, WbDiscountRule, WbPromotion, WbPromotionItem
from priceapp.timeutils import now_utc
from tests import factories as f

FMT = "%Y-%m-%dT%H:%M:%SZ"


def _setup(db, name="ИП 1", nm="5", art="A1", item="u1", barcode="b1", acc=None, category="Шорты"):
    from priceapp.models import OnecCost
    if acc is None:
        f.manual_rate(db, "80")
        acc = f.account(db, "wb", name, commission=25)
        f.rule(db, acc, base_coef="2.5", min_markup_coef="1.3", max_change_percent=1000)
        acc_mod.set_credential(db, acc, "token", "t")
    f.sku(db, item, art, "S", barcodes=[barcode], cost_usd=None if db.get(OnecCost, item) else "10")
    f.item(db, acc, barcode, art, external_id=f"{nm}:1", size="S")
    row = db.query(PlatformItem).filter_by(barcode=barcode).one()
    row.current_price = row.current_sale_price = 2009
    row.category = category
    db.commit()
    assert reference.record(db, acc, item, 2009, [row])
    db.commit()
    return acc


class _Cal:
    """Поддельный WB: календарь акций и `upload/task` для правила возврата."""
    def __init__(self, promos, items=None, details=None):
        self.data, self.items, self.det = promos, items or {}, details or []
        self.uploads, self.posts = [], []

    def promotions(self, start, end, all_promo=True):
        return self.data

    def promotion_details(self, ids):
        return [d for d in self.det if d["id"] in ids]

    def promotion_nomenclatures(self, pid, in_action):
        return self.items.get((pid, in_action), []), False

    def promotion_upload(self, pid, nms, upload_now):
        self.uploads.append((pid, list(nms), upload_now))
        return {"alreadyExists": False, "uploadID": 77}

    def _post(self, url, body):
        self.posts.append(body)
        return {"error": False}


def _promo(pid=1, kind="regular", start_in=timedelta(days=1), days=10, name="Осенняя распродажа"):
    now = now_utc()
    return {"id": pid, "name": name, "type": kind, "startDateTime": (now + start_in).strftime(FMT),
            "endDateTime": (now + start_in + timedelta(days=days)).strftime(FMT)}


def _item(nm, plan_discount, in_action=False, price=2009):
    return {"id": int(nm), "inAction": in_action, "price": price, "currencyCode": "RUB",
            "planPrice": int(price * (100 - plan_discount) / 100), "discount": 0, "planDiscount": plan_discount}


def _load(db, cal):
    return wp.refresh(db, client_factory=lambda p, c: cal)


def test_refresh_stores_promotions_details_and_items(db):
    _setup(db)
    cal = _Cal([_promo(1), _promo(2, "auto", name="Автоакция")],
               {(1, True): [_item(5, 25, True)], (1, False): [_item(99, 10)]},
               [{"id": 1, "description": "Скидки до 50%", "advantages": ["Плашка", "Буст"], "participationPercentage": 40},
                {"id": 2, "exceptionProductsCount": 12}])
    st = _load(db, cal)
    assert st["accounts"] == 1 and st["promos"] == 2 and st["items"] == 2 and not st["errors"]
    p1 = db.query(WbPromotion).filter_by(promo_id=1).one()
    assert p1.advantages == "Плашка; Буст" and p1.participation == 40 and p1.type == "regular"
    assert db.query(WbPromotion).filter_by(promo_id=2).one().exceptions == 12
    items = {i.nm_id: i for i in db.query(WbPromotionItem)}
    assert items["5"].in_action and items["5"].plan_discount == 25 and not items["99"].in_action
    # товар пропал из списка WB — пропадает и у нас; сбой товаров акции оставляет прежние
    cal.items = {(1, True): [_item(5, 25, True)], (1, False): []}
    _load(db, cal)
    assert {i.nm_id for i in db.query(WbPromotionItem)} == {"5"}

    def boom(pid, in_action):
        raise RuntimeError("401 нет категории токена")
    cal.promotion_nomenclatures = boom
    st = _load(db, cal)
    assert "товары не загружены" in st["errors"][0] and db.query(WbPromotionItem).count() == 1
    assert "товары не загружены" in db.query(WbPromotion).filter_by(promo_id=1).one().items_note


def test_classify_fits_returns_floor_and_unknown(db):
    a = _setup(db)
    _setup(db, acc=a, nm="6", art="A2", item="u2", barcode="b2")
    _setup(db, acc=a, nm="7", art="A3", item="u3", barcode="b3")
    _load(db, _Cal([_promo(1)], {(1, False): [_item(5, 25), _item(6, 35), _item(7, 5), _item(42, 20)]}))
    st = {pc.item.nm_id: pc for pc in wp.classify(db, db.query(WbPromotion).all())}
    assert st["5"].entry == 1507 and st["5"].need == Decimal("25.0") and st["5"].status == wp.RETURNS   # порог 10%
    assert st["6"].status == wp.FLOOR and "ниже пола" in st["6"].note                                # скидка 35%
    assert st["7"].status == wp.FITS and st["7"].need == Decimal("5.0")
    assert st["42"].status == wp.UNKNOWN
    wd.set_threshold(db, wd.CATEGORY, "Шорты", Decimal(30), "op")
    db.commit()
    st = {pc.item.nm_id: pc for pc in wp.classify(db, db.query(WbPromotion).all())}
    assert st["5"].status == wp.FITS and st["5"].thr == Decimal(30)


def test_promo_threshold_lasts_until_promo_end_then_returns(db):
    a = _setup(db)
    wd.set_threshold(db, wd.ARTICLE, "A1", Decimal(12), "op", account_id=a.id)
    db.commit()
    _load(db, _Cal([_promo(1)], {(1, False): [_item(5, 25)]}))
    promo = db.query(WbPromotion).one()
    pcs = wp.classify(db, [promo])
    assert wp.set_threshold(db, pcs, Decimal(26), "op") == 1
    db.commit()
    rule = db.query(WbDiscountRule).filter_by(kind="article", key="A1", account_id=a.id).one()
    assert rule.max_disc == Decimal(26) and rule.prev_max_disc == Decimal(12) and rule.until == promo.end_at
    assert wp.classify(db, [promo])[0].status == wp.FITS
    # вторая акция позже — порог больший, срок поздний, «прежний» тот, что был до первой
    later = promo.end_at + timedelta(days=5)
    wd.set_temp_threshold(db, "A1", Decimal(20), later, "op", a.id, 2, "Вторая")
    db.commit()
    db.refresh(rule)
    assert rule.max_disc == Decimal(26) and rule.until == later and rule.prev_max_disc == Decimal(12)
    # срок прошёл — правило возврата перед решениями возвращает прежний порог
    assert wd.expire_temp(db, later + timedelta(minutes=1)) == 1
    db.refresh(rule)
    assert rule.max_disc == Decimal(12) and rule.until is None
    assert db.query(AuditLog).filter_by(action="wb_promo_threshold_back").count() == 1
    # правка руками делает временный порог постоянным
    wd.set_temp_threshold(db, "A1", Decimal(26), later, "op", a.id, 1, "Осенняя")
    db.commit()
    wd.set_threshold(db, wd.ARTICLE, "A1", Decimal(26), "op", account_id=a.id)
    db.commit()
    db.refresh(rule)
    assert rule.until is None and rule.max_disc == Decimal(26)


def test_reference_run_expires_promo_thresholds_first(db):
    a = _setup(db)
    wd.set_temp_threshold(db, "A1", Decimal(30), now_utc() - timedelta(minutes=5), "op", a.id, 1, "Прошла")
    db.commit()
    row = db.query(PlatformItem).filter_by(barcode="b1").one()
    row.current_price = row.current_sale_price = 1500                # −25%: при пороге 30% держалось бы
    db.commit()
    cal = _Cal([])
    st = reference.run(db, client_factory=lambda p, c: cal)
    assert st["restored"] == 1 and cal.posts == [{"data": [{"nmID": 5, "price": 2009, "discount": 0}]}]


def test_promo_end_moved_by_wb_extends_the_threshold(db):
    a = _setup(db)
    cal = _Cal([_promo(1, days=10)], {(1, False): [_item(5, 25)]})
    _load(db, cal)
    promo = db.query(WbPromotion).one()
    wp.set_threshold(db, wp.classify(db, [promo]), Decimal(26), "op")
    db.commit()
    cal.data = [_promo(1, days=20)]
    _load(db, cal)
    rule = db.query(WbDiscountRule).filter_by(key="A1").one()
    assert rule.until == db.query(WbPromotion).one().end_at and a.id


def test_add_sends_only_fitting_cards_from_promo_start(db):
    a = _setup(db)
    _setup(db, acc=a, nm="6", art="A2", item="u2", barcode="b2")
    _setup(db, acc=a, nm="7", art="A3", item="u3", barcode="b3")
    cal = _Cal([_promo(1)], {(1, True): [_item(7, 5, True)], (1, False): [_item(5, 25), _item(6, 35)]})
    _load(db, cal)
    wd.set_threshold(db, wd.CATEGORY, "Шорты", Decimal(30), "op")
    db.commit()
    promo = db.query(WbPromotion).one()
    st = wp.add(db, wp.classify(db, [promo]), "op", client_factory=lambda p, c: cal)
    assert cal.uploads == [(1, [5], False)]                          # акция ещё не началась — со старта
    assert st["sent"] == 1 and st["in"] == 1 and "ниже пола" in st["skipped"][0]
    assert db.query(WbPromotionItem).filter_by(nm_id="5").one().added_at is not None
    # повтор — «отправлено, ждём WB», второй раз не шлём
    st = wp.add(db, wp.classify(db, [promo]), "op", client_factory=lambda p, c: cal)
    assert st["added"] == 1 and len(cal.uploads) == 1
    assert db.query(AuditLog).filter_by(action="wb_promo_add").count() == 1


def test_attention_line_before_start_when_many_in_action_cards_will_return(db):
    a = _setup(db)
    for i, nm in enumerate(("6", "7")):
        _setup(db, acc=a, nm=nm, art=f"A{i + 2}", item=f"u{i + 2}", barcode=f"b{i + 2}")
    items = [_item(n, 25, True) for n in ("5", "6", "7")]
    _load(db, _Cal([_promo(1, start_in=timedelta(hours=10))], {(1, True): items}))
    got = wp.attention_items(db)
    assert len(got) == 1 and got[0][3] == 3 and "вернётся эталон" in got[0][1] and "flt=returns" in got[0][2]
    wd.set_threshold(db, wd.CATEGORY, "Шорты", Decimal(30), "op")
    db.commit()
    assert wp.attention_items(db) == []


def test_page_bulk_add_export_and_import(client, db):
    a = _setup(db)
    _setup(db, acc=a, nm="6", art="A2", item="u2", barcode="b2", category="Джемперы")
    cal = _Cal([_promo(1), _promo(2, "auto", name="Автоакция")], {(1, False): [_item(5, 25), _item(6, 5)]},
               [{"id": 2, "exceptionProductsCount": 4}])
    _load(db, cal)
    html = client.get("/wb-promo").text
    assert "Осенняя распродажа" in html and "автоакция" in html and "исключено до старта: 4" in html
    html = client.get(f"/wb-promo?promo=1&acc={a.id}").text
    assert "% от эталона для вхождения" in html and "вернётся эталон — 1" in html and "уложится в порог — 1" in html
    accs = acc_mod.client_for
    try:
        acc_mod.client_for = lambda db_, acc_, cf=None: cal
        r = client.post("/wb-promo/bulk", data={"acc": str(a.id), "promo": "1", "sel": ["1:c:Шорты"],
                                                "action": "add", "value": "26"})
    finally:
        acc_mod.client_for = accs
    assert "Порог для акции 26%" in r.text and "отправлено в акцию карточек 1 из 1" in r.text
    assert cal.uploads == [(1, [5], False)]
    rule = db.query(WbDiscountRule).filter_by(key="A1").one()
    assert rule.max_disc == Decimal(26) and rule.until is not None
    # Excel: выгрузка разделами, правка «да» в «Добавить в акцию»
    r = client.get(f"/wb-promo/export?promo=1&acc={a.id}")
    rows = [list(x) for x in load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True)]
    assert rows[0] == wp.HEAD and ("Категория", "Джемперы") in {(x[1], x[2]) for x in rows[1:]}
    wb = Workbook()
    wb.active.append(wp.HEAD)
    row = [None] * len(wp.HEAD)
    row[1], row[2], row[6], row[wp.HEAD.index(wp.COL_ADD)] = "Артикул", "A2", 1, "да"
    wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    try:
        acc_mod.client_for = lambda db_, acc_, cf=None: cal
        r = client.post(f"/wb-promo/import?acc={a.id}", files={"file": ("p.xlsx", buf.getvalue())})
    finally:
        acc_mod.client_for = accs
    assert cal.uploads[-1] == (1, [6], False) and "отправлено в акцию карточек 1 из 1" in r.text
    html = client.get(f"/wb-promo?promo=1&acc={a.id}").text
    assert "История" in html and "добавлено карточек 1" in html
    assert client.get("/wb-promo/history").status_code == 200
    # порог «для акции» виден и на «Пороге скидки WB»
    html = client.get(f"/wb-discount?acc={a.id}&by=article").text
    assert "акция «Осенняя распродажа»" in html


def test_add_with_new_threshold_under_returns_filter_still_adds(client, db):
    """Отбор «вернётся эталон» + новый порог + «Добавить в акцию»: после порога карточки уже
    «уложится» и из отбора выпадают — добавить их всё равно надо."""
    a = _setup(db)
    cal = _Cal([_promo(1)], {(1, False): [_item(5, 25)]})
    _load(db, cal)
    accs = acc_mod.client_for
    try:
        acc_mod.client_for = lambda db_, acc_, cf=None: cal
        r = client.post("/wb-promo/save", data={"acc": str(a.id), "promo": "1", "flt": "returns", "by": "article",
                                                "sel": "1:a:A1", "thr": "26", "action": "add"})
    finally:
        acc_mod.client_for = accs
    assert cal.uploads == [(1, [5], False)] and "отправлено в акцию карточек 1 из 1" in r.text


def test_all_promos_mode_page_bulk_and_one_file_for_two_promos(client, db):
    """Акция не выбрана — все обычные акции разделами; в одном файле правка двух акций сразу."""
    a = _setup(db)
    _setup(db, acc=a, nm="6", art="A2", item="u2", barcode="b2", category="Джемперы")
    cal = _Cal([_promo(1), _promo(3, name="Чёрная пятница", start_in=timedelta(days=5))],
               {(1, False): [_item(5, 25), _item(6, 5)], (3, False): [_item(5, 15)]})
    _load(db, cal)
    html = client.get("/wb-promo").text
    assert "Акция «Осенняя распродажа»" in html and "Акция «Чёрная пятница»" in html
    assert "Правка через Excel" in html and "Сейчас ниже эталона" in html and "по умолчанию" in html
    # выгрузка без выбора акции — обе акции своими разделами, «ID акции» у каждой строки
    rows = [list(x) for x in load_workbook(io.BytesIO(client.get("/wb-promo/export").content)).active.iter_rows(values_only=True)]
    assert rows[0] == wp.HEAD
    heads = [r[0] for r in rows[1:] if r[1] is None]
    assert any("ОСЕННЯЯ РАСПРОДАЖА" in h for h in heads) and any("ЧЁРНАЯ ПЯТНИЦА" in h for h in heads)
    assert {r[6] for r in rows[1:] if r[1]} == {1, 3}
    # один файл: A1 в акции 1 — порог 26 и «да»; A1 в акции 3 — «да». Порог 26 «для акции» 1 кончится
    # раньше акции 3 — для неё действует прежний 10%, и −15% в него не укладывается: не добавляем
    wb = Workbook()
    wb.active.append(wp.HEAD)
    for pid, thr, add in ((1, 26, "да"), (3, None, "да")):
        row = [None] * len(wp.HEAD)
        row[1], row[2], row[6] = "Артикул", "A1", pid
        row[wp.HEAD.index(wp.COL_NEW)], row[wp.HEAD.index(wp.COL_ADD)] = thr, add
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    accs = acc_mod.client_for
    try:
        acc_mod.client_for = lambda db_, acc_, cf=None: cal
        r = client.post("/wb-promo/import", files={"file": ("p.xlsx", buf.getvalue())})
    finally:
        acc_mod.client_for = accs
    assert cal.uploads == [(1, [5], False)]
    assert "Чёрная пятница" in r.text and "вернётся эталон" in r.text      # во второй — не уложилась, сказано
    # порог «для акции» виден у строки другой акции как источник
    html = client.get(f"/wb-promo?acc={a.id}").text
    assert "для акции до" in html and "кончится раньше этой" in html
    # массовая правка по всему отбору — по всем акциям
    try:
        acc_mod.client_for = lambda db_, acc_, cf=None: cal
        client.post("/wb-promo/bulk", data={"acc": str(a.id), "q": "A2", "all_filtered": "1", "action": "add"})
    finally:
        acc_mod.client_for = accs
    assert cal.uploads[-1] == (1, [6], False)


# --- срок участия ---------------------------------------------------------------------------------

def test_parse_leave_is_end_of_local_day_and_checks():
    from datetime import date, datetime
    from priceapp.timeutils import local_date_of
    at = wp.parse_leave("2026-10-25")
    assert local_date_of(at) == date(2026, 10, 25) and local_date_of(at + timedelta(seconds=2)) == date(2026, 10, 26)
    assert wp.parse_leave("25.10.2026") == at == wp.parse_leave(datetime(2026, 10, 25))
    assert wp.parse_leave("") is None
    try:
        wp.parse_leave("завтра")
        assert False
    except ValueError:
        pass
    promo = WbPromotion(start_at=now_utc() + timedelta(days=1), end_at=now_utc() + timedelta(days=10))
    assert wp.leave_error(promo, now_utc() - timedelta(hours=1)) == "срок уже прошёл"
    assert "раньше старта" in wp.leave_error(promo, now_utc() + timedelta(hours=5))
    assert "кончится сама" in wp.leave_error(promo, now_utc() + timedelta(days=11))
    assert wp.leave_error(promo, now_utc() + timedelta(days=4)) == ""


def test_add_with_term_and_promo_threshold_ends_with_the_term(db):
    a = _setup(db)
    cal = _Cal([_promo(1)], {(1, False): [_item(5, 25)]})
    _load(db, cal)
    promo = db.query(WbPromotion).one()
    leave = now_utc() + timedelta(days=4)
    wp.set_threshold(db, wp.classify(db, [promo]), Decimal(26), "op", leave)
    db.commit()
    # без срока порог «до 4-го дня» для акции «до 11-го» не в счёт — карточка вылетела бы раньше
    assert wp.classify(db, [promo])[0].status == wp.RETURNS
    st = wp.add(db, wp.classify(db, [promo], leave_at=leave), "op", client_factory=lambda p, c: cal, leave_at=leave)
    assert st["sent"] == 1
    it = db.query(WbPromotionItem).one()
    assert it.leave_at == leave and it.left_at is None
    assert db.query(WbDiscountRule).filter_by(key="A1").one().until == leave and a.id


def test_set_leave_for_cards_in_promo_shortens_the_promo_threshold(db):
    a = _setup(db)
    _setup(db, acc=a, nm="6", art="A2", item="u2", barcode="b2")
    _load(db, _Cal([_promo(1)], {(1, True): [_item(5, 25, True)], (1, False): [_item(6, 5)]}))
    promo = db.query(WbPromotion).one()
    pcs = wp.classify(db, [promo])
    wp.set_threshold(db, pcs, Decimal(26), "op")
    db.commit()
    leave = now_utc() + timedelta(days=3)
    n, bad = wp.set_leave(db, pcs, leave, "op")
    db.commit()
    assert n == 1 and not bad                                         # A2 не в акции — срок не нужен
    assert db.query(WbPromotionItem).filter_by(nm_id="5").one().leave_at == leave
    assert db.query(WbPromotionItem).filter_by(nm_id="6").one().leave_at is None
    assert db.query(WbDiscountRule).filter_by(key="A1").one().until == leave
    n, bad = wp.set_leave(db, pcs, promo.end_at + timedelta(days=1), "op")
    assert n == 0 and "кончится сама" in bad[0]


def test_leave_due_raises_price_to_full_etalon_and_marks_left(db):
    a = _setup(db)
    _setup(db, acc=a, nm="6", art="A2", item="u2", barcode="b2")
    cal = _Cal([_promo(1, start_in=-timedelta(days=1))], {(1, True): [_item(5, 25, True), _item(6, 25, True)]})
    _load(db, cal)
    for nm in ("b1", "b2"):                                           # в акции: 2009 при скидке 25% = 1507
        row = db.query(PlatformItem).filter_by(barcode=nm).one()
        row.current_sale_price = 1507
    row.current_sale_price = 2009                                     # A2 уже вернулся к эталону
    for it in db.query(WbPromotionItem):
        it.leave_at = now_utc() - timedelta(minutes=1)
    db.commit()
    st = wp.leave_due(db, client_factory=lambda p, c: cal)
    assert st["raised"] == 1 and st["same"] == 1
    assert cal.posts == [{"data": [{"nmID": 5, "price": 2009, "discount": 0}]}]
    assert all(i.left_at is not None for i in db.query(WbPromotionItem))
    from priceapp.models import PriceChange
    ch = db.query(PriceChange).filter_by(source="wb_promo_leave").one()
    assert ch.status == "sent" and "по сроку" in ch.note
    assert wp.leave_due(db, client_factory=lambda p, c: cal)["due"] == 0       # второй раз не шлём
    assert db.query(AuditLog).filter_by(action="wb_promo_leave").count() == 1


def test_leave_due_retries_when_wb_refuses_and_names_cards_without_etalon(db):
    _setup(db)
    _load(db, _Cal([_promo(1, start_in=-timedelta(days=1))], {(1, True): [_item(5, 25, True), _item(42, 20, True)]}))
    row = db.query(PlatformItem).filter_by(barcode="b1").one()
    row.current_sale_price = 1507
    for it in db.query(WbPromotionItem):
        it.leave_at = now_utc() - timedelta(minutes=1)
    db.commit()

    class Down(_Cal):
        def _post(self, url, body):
            raise RuntimeError("WB недоступен")
    st = wp.leave_due(db, client_factory=lambda p, c: Down([]))
    assert "WB недоступен" in st["errors"][-1]
    assert db.query(WbPromotionItem).filter_by(nm_id="5").one().left_at is None    # повтор в следующий раз
    assert db.query(WbPromotionItem).filter_by(nm_id="42").one().left_at is not None
    assert "эталона нет" in db.query(AuditLog).filter_by(action="wb_promo_leave").one().details
    cal = _Cal([])
    assert wp.leave_due(db, client_factory=lambda p, c: cal)["raised"] == 1


def test_page_row_term_and_excel_term(client, db):
    a = _setup(db)
    _setup(db, acc=a, nm="6", art="A2", item="u2", barcode="b2")
    _load(db, _Cal([_promo(1)], {(1, True): [_item(5, 25, True)], (1, False): [_item(6, 5)]}))
    day = (now_utc() + timedelta(days=3)).strftime("%Y-%m-%d")
    r = client.post("/wb-promo/save", data={"acc": str(a.id), "promo": "1", "by": "article", "sel": "1:a:A1",
                                            "leave": day, "action": "set"})
    assert "Срок участия до" in r.text and "карточек 1" in r.text
    assert db.query(WbPromotionItem).filter_by(nm_id="5").one().leave_at == wp.parse_leave(day)
    html = client.get(f"/wb-promo?promo=1&acc={a.id}&by=article&flt=leave").text
    assert "участвует до" in html and "A1" in html
    # Excel: «-» снимает срок
    wb = Workbook()
    wb.active.append(wp.HEAD)
    row = [None] * len(wp.HEAD)
    row[1], row[2], row[6], row[wp.HEAD.index(wp.COL_LEAVE)] = "Артикул", "A1", 1, "-"
    wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(f"/wb-promo/import?acc={a.id}", files={"file": ("p.xlsx", buf.getvalue())})
    assert "срок участия изменён у карточек — 1" in r.text
    db.expire_all()
    assert db.query(WbPromotionItem).filter_by(nm_id="5").one().leave_at is None


def test_job_runs_leave_due(db):
    from priceapp.models import WorkerHeartbeat
    from priceapp.workers import jobs
    jobs.job_wb_promo_leave(client_factory=lambda p, c: _Cal([]))
    assert db.get(WorkerHeartbeat, "wb_promo_leave").last_success


def test_import_threshold_term_and_add_in_one_row(client, db):
    """Строка файла: порог 26, «участвовать до» через 4 дня и «да» — порог до срока, карточка
    добавлена со сроком (порог «до срока» для неё в счёт: дальше она в акции не стоит)."""
    a = _setup(db)
    cal = _Cal([_promo(1)], {(1, False): [_item(5, 25)]})
    _load(db, cal)
    day = (now_utc() + timedelta(days=4)).strftime("%d.%m.%Y")
    wb = Workbook()
    wb.active.append(wp.HEAD)
    row = [None] * len(wp.HEAD)
    row[1], row[2], row[6] = "Артикул", "A1", 1
    row[wp.HEAD.index(wp.COL_NEW)], row[wp.HEAD.index(wp.COL_LEAVE)], row[wp.HEAD.index(wp.COL_ADD)] = 26, day, "да"
    wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    accs = acc_mod.client_for
    try:
        acc_mod.client_for = lambda db_, acc_, cf=None: cal
        r = client.post(f"/wb-promo/import?acc={a.id}", files={"file": ("p.xlsx", buf.getvalue())})
    finally:
        acc_mod.client_for = accs
    assert cal.uploads == [(1, [5], False)] and "отправлено в акцию карточек 1 из 1" in r.text
    db.expire_all()
    leave = wp.parse_leave(day)
    assert db.query(WbPromotionItem).one().leave_at == leave
    assert db.query(WbDiscountRule).filter_by(key="A1").one().until == leave


def _ended(db, pid=1):
    """Акция `pid` закончилась минуту назад (товары загружены, пока она шла)."""
    p = db.query(WbPromotion).filter_by(promo_id=pid).one()
    p.end_at = now_utc() - timedelta(minutes=1)
    db.commit()


def test_promo_end_returns_etalon_to_cards_that_stayed_in_promo(db):
    """Решение 18.10: WB после акции цену не возвращает, а «уложившуюся в порог» карточку правило возврата
    не трогает — эталон ей возвращает конец акции. Уже стоящую по эталону и не участвовавшую не трогаем."""
    a = _setup(db)
    _setup(db, acc=a, nm="6", art="A2", item="u2", barcode="b2")
    _setup(db, acc=a, nm="7", art="A3", item="u3", barcode="b3")
    cal = _Cal([_promo(1, start_in=-timedelta(days=3))],
               {(1, True): [_item(5, 8, True), _item(6, 8, True)], (1, False): [_item(7, 8)]})
    _load(db, cal)
    for bc, sale in (("b1", 1848), ("b2", 2009), ("b3", 1848)):    # 5 — по акционной, 6 — уже эталон, 7 — не в акции
        db.query(PlatformItem).filter_by(barcode=bc).one().current_sale_price = sale
    db.commit()
    assert wp.end_due(db, client_factory=lambda p, c: cal)["due"] == 0      # акция ещё идёт
    _ended(db)
    st = wp.end_due(db, client_factory=lambda p, c: cal)
    assert st["raised"] == 1 and st["same"] == 1
    assert cal.posts == [{"data": [{"nmID": 5, "price": 2009, "discount": 0}]}]
    from priceapp.models import PriceChange
    ch = db.query(PriceChange).filter_by(source="wb_promo_end").one()
    assert ch.status == "sent" and "закончилась" in ch.note and "полный эталон" in ch.note
    its = {i.nm_id: i for i in db.query(WbPromotionItem)}
    assert its["5"].end_note == "эталон возвращён" and its["6"].end_note == "цена уже не ниже эталона"
    assert its["7"].end_done_at is None
    assert db.query(PlatformItem).filter_by(barcode="b1").one().current_sale_price == 2009
    assert wp.end_due(db, client_factory=lambda p, c: cal)["due"] == 0      # второй раз не шлём
    assert db.query(AuditLog).filter_by(action="wb_promo_end").count() == 1


def test_promo_end_waits_for_other_promo_and_keeps_minus_n(db):
    """Карточку, записанную ещё в идущую акцию, конец первой не трогает: подъём выкинул бы её оттуда.
    «Эталон − N%» (решение человека с «Порога скидки WB») конец акции не отменяет."""
    from priceapp.models import ReferencePrice
    a = _setup(db)
    _setup(db, acc=a, nm="6", art="A2", item="u2", barcode="b2")
    cal = _Cal([_promo(1, start_in=-timedelta(days=3)), _promo(2, start_in=-timedelta(days=1), name="Вторая")],
               {(1, True): [_item(5, 20, True), _item(6, 20, True)], (2, True): [_item(5, 20, True)]})
    _load(db, cal)
    for bc in ("b1", "b2"):
        db.query(PlatformItem).filter_by(barcode=bc).one().current_sale_price = 1607
    db.query(ReferencePrice).filter_by(item_id="u2").one().held_pct = Decimal(15)
    db.commit()
    _ended(db, 1)
    st = wp.end_due(db, client_factory=lambda p, c: cal)
    assert st["busy"] == 1 and st["raised"] == 1
    assert cal.posts == [{"data": [{"nmID": 6, "price": 1708, "discount": 0}]}]      # 2009 − 15%
    first = {i.nm_id: i for i in db.query(WbPromotionItem).filter_by(promo_id=1)}
    assert "ещё в акции «Вторая»" in first["5"].end_note
    assert db.query(ReferencePrice).filter_by(item_id="u2").one().held_pct == Decimal(15)
    from priceapp.models import PriceChange
    assert "эталон − 15%" in db.query(PriceChange).filter_by(source="wb_promo_end").one().note
    _ended(db, 2)                                                    # кончилась и вторая — теперь эталон
    st = wp.end_due(db, client_factory=lambda p, c: cal)
    assert st["raised"] == 1 and cal.posts[-1] == {"data": [{"nmID": 5, "price": 2009, "discount": 0}]}


def test_promo_end_retries_when_wb_refuses_and_page_shows_result(client, db):
    _setup(db)
    _load(db, _Cal([_promo(1, start_in=-timedelta(days=3))], {(1, True): [_item(5, 8, True)]}))
    db.query(PlatformItem).filter_by(barcode="b1").one().current_sale_price = 1848
    db.commit()
    _ended(db)

    class Down(_Cal):
        def _post(self, url, body):
            raise RuntimeError("WB недоступен")
    st = wp.end_due(db, client_factory=lambda p, c: Down([]))
    assert "WB недоступен" in st["errors"][-1]
    assert db.query(WbPromotionItem).one().end_done_at is None                # повтор в следующий раз
    html = client.get("/wb-promo").text
    assert "Закончившиеся акции" in html and "ждёт возврата эталона — 1" in html
    from priceapp.models import WorkerHeartbeat
    from priceapp.workers import jobs
    jobs.job_wb_promo_leave(client_factory=lambda p, c: _Cal([]))
    hb = db.get(WorkerHeartbeat, "wb_promo_leave")
    db.refresh(hb)
    assert hb.last_success
    db.expire_all()
    assert db.query(WbPromotionItem).one().end_note == "эталон возвращён"
    assert "эталон возвращён — 1" in client.get("/wb-promo").text
