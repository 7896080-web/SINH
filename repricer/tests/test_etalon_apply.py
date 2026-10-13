"""Новый эталон WB по заполненной форме маржинальности (scripts/etalon_apply.py)."""
import importlib
import os
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

from priceapp.database import engine
from priceapp.models import PlatformItem, ReferenceDiscount
from tests import factories as F

ROOT = Path(__file__).resolve().parent.parent
ENV = {**os.environ, "PYTHONIOENCODING": "cp1251"}


def _setup(db, tmp_path):
    w = F.account(db, "wb", "ИП 1", commission=48)
    F.manual_rate(db, "100")
    F.sku(db, "s1", "256250-015", "M", "Бежевый", barcodes=["a1"], cost_usd=12.74)
    F.sku(db, "s2", "256250-030", "M", "Черный", barcodes=["b1"], cost_usd=13.53)
    F.sku(db, "s3", "999", "M", "Синий", barcodes=["c1"], cost_usd=10)
    F.item(db, w, "a1", "256250-015", external_id="111:1")
    F.item(db, w, "b1", "256250-030", external_id="222:1")
    F.item(db, w, "c1", "999", external_id="333:1")
    prices = {"a1": (13559, 7457), "b1": (12089, 6770), "c1": (10000, 3000)}   # 45%, 44%, 70%
    for it in db.query(PlatformItem):
        it.current_price, it.current_sale_price = prices[it.barcode]
    db.commit()
    form = tmp_path / "form.xlsx"
    r = subprocess.run([sys.executable, str(ROOT / "scripts" / "etalon_cost.py"), "--xlsx", str(form)],
                       capture_output=True, cwd=ROOT, env=ENV)
    assert r.returncode == 0, r.stderr.decode("utf-8", "replace")
    book = load_workbook(form)
    ws = book["Артикулы"]
    h = [c.value for c in ws[1]]
    for n in range(2, ws.max_row + 1):
        if ws.cell(n, 1).value == "256250":
            ws.cell(n, h.index("ИП 1: маржинальность (если иная)") + 1).value = 3
    wc = book["Категории"]
    for n in range(2, wc.max_row + 1):
        wc.cell(n, 4).value = 2.8                      # категория — 999 возьмёт её
    book.save(form)
    return w, form


def _mod():
    sys.path.insert(0, str(ROOT / "scripts"))
    return importlib.import_module("etalon_apply")


def test_plan_targets_sale_price_and_keeps_discount_in_corridor(db, tmp_path):
    w, form = _setup(db, tmp_path)
    m = _mod()
    arts, cats, cabs = m.read_form(form)
    live, _ = m.wrd.load(Path(str(engine.url).replace("sqlite:///", "")))
    cards = {c.nm: c for c in m.build_plan(db, arts, cats, live, 35, 50)}
    # 256250: себестоимость артикула 13,53 $ x 100 = 1353 руб; 1353 x 3 / 0,52 = 7805,8 -> 7806
    assert cards["111"].target == 7806 and cards["111"].disc == 45          # скидка в коридоре — та же
    assert cards["111"].new_list == 14193                                    # 7806 / 0,55 = 14192,7 -> вверх
    assert cards["222"].target == 7806 and cards["222"].disc == 44
    # 999: категория 2,8; 1000 x 2,8 / 0,52 = 5384,6 -> 5385; скидка 70% вне коридора -> 50%
    assert cards["333"].target == 5385 and cards["333"].disc == 50 and cards["333"].new_list == 10770
    for c in cards.values():
        assert c.new_list * (100 - c.disc) / 100 >= c.target - 0.01


def test_send_puts_list_and_discount_and_records_discount_etalon(db, tmp_path):
    w, form = _setup(db, tmp_path)
    m = _mod()
    arts, cats, _ = m.read_form(form)
    live, _ = m.wrd.load(Path(str(engine.url).replace("sqlite:///", "")))
    cards = m.build_plan(db, arts, cats, live, 35, 50)
    calls = []

    class Fake:
        def _post(self, url, body):
            calls.append(body)
            return {"error": False}

    res = m.wrd.send_plan(db, cards, factory=lambda p, c: Fake())
    assert res["ИП 1"] == "отправлено карточек: 3"
    assert {d["nmID"]: (d["price"], d["discount"]) for d in calls[0]["data"]} == {
        111: (14193, 45), 222: (13940, 44), 333: (10770, 50)}
    assert {r.nm_id: r.discount for r in db.query(ReferenceDiscount)} == {"111": 45, "222": 44, "333": 50}


def test_script_shows_plan_and_sends_nothing_without_send(db, tmp_path):
    w, form = _setup(db, tmp_path)
    out = tmp_path / "plan.xlsx"
    r = subprocess.run([sys.executable, str(ROOT / "scripts" / "etalon_apply.py"), "--form", str(form),
                        "--xlsx", str(out)], capture_output=True, cwd=ROOT, env=ENV)
    assert r.returncode == 0, r.stderr.decode("utf-8", "replace")
    text = r.stdout.decode("cp1251")
    assert "Карточек к отправке: 3" in text and "Ничего не отправлено" in text
    assert load_workbook(out).sheetnames == ["ИП 1"]


def test_card_with_two_1c_articles_uses_average_cost_and_skips_different_margins(db, tmp_path):
    w = F.account(db, "wb", "ИП 1", commission=48)
    F.manual_rate(db, "100")
    F.sku(db, "a", "2403", "M", barcodes=["a1"], cost_usd=13.3)
    F.sku(db, "b", "2425 b (2403)", "L", barcodes=["b1"], cost_usd=14.16)
    F.item(db, w, "a1", "2403", external_id="777:1"); F.item(db, w, "b1", "2403", external_id="777:2")
    for it in db.query(PlatformItem):
        it.current_price, it.current_sale_price = 9690, 4845          # 50%
    db.commit()
    m = _mod()
    live, _ = m.wrd.load(Path(str(engine.url).replace("sqlite:///", "")))
    arts = {"2403": {"category": "Д", "cost_rub": 1330.0, "art": None, "over": {"ИП 1": 2.8}},
            "2425 b (2403)": {"category": "Д", "cost_rub": 1416.0, "art": None, "over": {"ИП 1": 2.8}}}
    (c,) = m.build_plan(db, arts, {}, live, 35, 50)
    # средняя (1330 + 1416) / 2 = 1373; 1373 x 2,8 / 0,52 = 7393,08 -> 7394
    assert c.averaged and c.cost_rub == 1373 and c.target == 7394 and not c.skip
    arts["2425 b (2403)"]["over"]["ИП 1"] = 3.0
    (c,) = m.build_plan(db, arts, {}, live, 35, 50)
    assert c.skip.startswith("у артикулов карточки разная маржинальность")


def test_adopt_refuses_on_mismatch_then_records_etalon_and_enables_rule(db, tmp_path, capsys):
    from priceapp import reference, settings
    from priceapp.models import ReferencePrice
    w, form = _setup(db, tmp_path)
    settings.put(db, settings.REF_ENABLED, "0")              # как после --send
    db.commit()
    m = _mod()
    after = str(engine.url).replace("sqlite:///", "")
    assert m.main(["--form", str(form), "--adopt", "--after", after]) == 1     # WB держит ещё старое
    assert "Эталон НЕ записан" in capsys.readouterr().out
    db.expire_all()
    assert reference.enabled(db) is False and db.query(ReferencePrice).count() == 0
    # WB применил план
    plan = {"111": (14193, 7806), "222": (13940, 7806), "333": (10770, 5385)}
    for it in db.query(PlatformItem):
        it.current_price, it.current_sale_price = plan[it.external_id.split(":")[0]]
    db.commit()
    assert m.main(["--form", str(form), "--adopt", "--after", after]) == 0
    db.expire_all()
    assert reference.enabled(db) is True
    assert {r.ref_price for r in db.query(ReferencePrice)} == {14193, 13940, 10770}
    assert {r.nm_id: r.discount for r in db.query(ReferenceDiscount)} == {"111": 45, "222": 44, "333": 50}


def test_big_drop_goes_in_steps_and_adopt_waits_for_the_end(db, tmp_path, capsys):
    w = F.account(db, "wb", "ИП 1", commission=50)
    F.manual_rate(db, "100")
    F.sku(db, "s", "3164M", "M", barcodes=["a1"], cost_usd=100)
    F.item(db, w, "a1", "3164M", external_id="180413586:1")
    for it in db.query(PlatformItem):
        it.current_price, it.current_sale_price = 472479, 236240           # 50%
    db.commit()
    m = _mod()
    after = Path(str(engine.url).replace("sqlite:///", ""))
    arts = {"3164M": {"category": "К", "cost_rub": 10000.0, "art": 1.6, "over": {}}}   # цель 32000
    live, _ = m.wrd.load(after)
    (c,) = m.build_plan(db, arts, {}, live, 35, 50, 33)
    # шаг: не ниже 236240 x 0,67 = 158281; цель 32000 — потом
    assert c.step and c.final_target == 32000 and c.target == 158281 and c.disc == 50
    assert (236240 - c.new_list * 0.5) / 236240 <= 0.33 + 1e-9
    # WB применил шаг — следующий шаг от новой цены
    for it in db.query(PlatformItem):
        it.current_price, it.current_sale_price = c.new_list, c.new_list // 2
    db.commit()
    live, _ = m.wrd.load(after)
    (c2,) = m.build_plan(db, arts, {}, live, 35, 50, 33)
    assert c2.step and c2.target < c.target
    # до конца шагов эталон не пишется
    from openpyxl import Workbook
    b = Workbook()
    ws = b.active
    ws.title = "Артикулы"
    ws.append(["Артикул", "Категория", "Себестоимость макс, руб", "Маржинальность артикула"])
    ws.append(["3164M", "К", 10000, 1.6])
    b.create_sheet("Категории").append(["Категория", "a", "b", "Маржинальность категории"])
    f = tmp_path / "f.xlsx"
    b.save(f)
    assert m.main(["--form", str(f), "--adopt", "--after", str(after)]) == 1
    assert "по шагам 1" in capsys.readouterr().out
