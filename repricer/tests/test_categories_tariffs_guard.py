"""Категории и тарифы площадок, наценка на категорию, диапазон безопасности акций,
запрос текущих цен при запуске программы, справка по типам цен."""
import io
from decimal import Decimal

from openpyxl import load_workbook

from priceapp import accounts as acc_mod, dispatch, guard, platforms
from priceapp.models import Account, CategoryTarget, PlatformItem, PriceChange, ProductPrice
from priceapp.platforms import CatalogRow, CurrentPrice, Tariff
from tests import factories as f

RULE = {"commission_percent": "25", "base_coef": "2,5", "min_markup_coef": "1,3",
        "round_step": "10", "round_minus": "1", "max_change_percent": "20", "commission_extra": "0"}


# --- разбор ответов площадок -----------------------------------------------------------

def test_wb_subject_is_category_and_tariffs_by_subject():
    rows = platforms.parse_wb_cards({"cards": [{"nmID": 5, "vendorCode": "39681", "subjectID": 160,
                                                "subjectName": "Свитшоты",
                                                "sizes": [{"chrtID": 1, "techSize": "L", "skus": ["b1"]}]}]})
    assert (rows[0].category, rows[0].category_id) == ("Свитшоты", "160")
    t = platforms.parse_wb_tariffs({"report": [{"subjectID": 160, "subjectName": "Свитшоты", "kgvpMarketplace": 15.5,
                                                "paidStorageKgvp": 19.5, "kgvpSupplier": 12.5}]})
    assert t == {"160": Tariff(15.5, 19.5)}


def test_ozon_prices_carry_tariff_and_seller_promo_price():
    got = platforms.parse_ozon_prices({"items": [
        {"offer_id": "A-1", "price": {"price": "4699.0000", "old_price": "6999", "marketing_seller_price": "4199"},
         "commissions": {"sales_percent_fbs": 18, "sales_percent_fbo": 20}},
        {"offer_id": "A-2", "price": {"price": "1000", "marketing_seller_price": "1000"}, "commissions": {}}]})
    assert got["A-1"] == CurrentPrice(4699, 4199, "", 18.0, 20.0)      # old_price на выручку не влияет
    assert got["A-2"] == CurrentPrice(1000, 1000, "", None, None)
    rows = platforms.parse_ozon_info({"items": [{"id": 7, "offer_id": "A-1", "name": "Свитшот", "type_id": 970,
                                                 "barcodes": ["b1"]}]})
    names = platforms.parse_ozon_type_names({"result": [{"category_name": "Одежда", "children": [
        {"category_name": "Толстовки", "children": [{"type_id": 970, "type_name": "Свитшот", "children": []}]}]}]})
    assert rows[0].category_id == "970" and names == {"970": "Свитшот"}


def test_lamoda_category_is_lowest_russian_level():
    rows = platforms.parse_lamoda_nomenclatures({"nomenclatures": [{
        "parentSku": "MP1", "sku": "MP1S", "barcode": "b1", "externalParentSku": "39681", "name": "Свитшот",
        "categoryLevels": [{"level": 1, "name": "CLOTHES", "language": "EN"}, {"level": 1, "name": "Одежда", "language": "RU"},
                           {"level": 2, "name": "Свитшоты", "language": "RU", "id": 5521}]}]})
    assert (rows[0].category, rows[0].category_id) == ("Свитшоты", "5521")


class _KitSession:
    def __init__(self):
        self.headers = {}

    def get(self, url, params=None, timeout=0):
        page = (params or {}).get("page", 1)
        if url.endswith("/v1/variants"):
            data = {"variants": [{"id": "v1", "barcode": "b1", "sku": "39681-L", "name": "Свитшот", "product_id": "p1"}]
                    if page == 1 else [], "total_count": 1}
        elif url.endswith("/v1/categories"):
            assert params["status"] == "ACTIVE"                       # без status Kit отвечает 400
            data = {"categories": [{"id": 11, "title": "Свитшоты"}] if page == 1 else [], "total_count": 1}
        else:
            data = {"products": [{"id": "p1", "category_ids": [11]}] if page == 1 else []}
        return _Resp(data)


class _Resp:
    def __init__(self, data):
        self._data, self.content, self.status_code = data, b"x", 200

    def raise_for_status(self):
        pass

    def json(self):
        return self._data


def test_kit_catalog_gets_category_from_product():
    rows = platforms.KitClient("t", session=_KitSession()).get_catalog()
    assert (rows[0].category, rows[0].category_id) == ("Свитшоты", "11")


class FakeWb:
    """Каталог с предметами, тарифы и текущие цены — без сети."""
    def __init__(self, tariffs=None, prices=None):
        self.last_truncated = False
        self.tariffs = tariffs if tariffs is not None else {"160": Tariff(15.5, 19.5)}
        self.prices = prices or {}

    def get_catalog(self):
        return [CatalogRow("5:1", "b1", "39681-L", "Свитшот", "L", "Свитшоты", "160"),
                CatalogRow("6:1", "b3", "4033", "Джемпер", "XL", "Джемперы", "161")]

    def get_tariffs(self):
        return self.tariffs

    def price_key(self, item):
        return item.external_id.split(":")[0]

    def get_prices(self):
        return self.prices


def test_load_catalog_keeps_category_and_tariff(db):
    acc = f.account(db)
    acc_mod.load_catalog(db, acc, FakeWb())
    items = {i.barcode: i for i in db.query(PlatformItem)}
    assert items["b1"].category == "Свитшоты" and items["b1"].tariff_fbs == Decimal("15.50")
    assert items["b3"].tariff_fbs is None                         # предмета нет в таблице — тариф неизвестен

    class Broken(FakeWb):
        def get_tariffs(self):
            raise RuntimeError("429")
    acc_mod.load_catalog(db, acc, Broken())                       # тарифы не пришли — каталог всё равно
    db.expire_all()
    assert db.query(PlatformItem).filter_by(barcode="b1").one().tariff_fbs == Decimal("15.50")
    assert "тарифы комиссий не загружены" in db.get(Account, acc.id).catalog_note


# --- расчёт с тарифом и надбавкой -----------------------------------------------------

def _wb(client, db, **rule):
    f.manual_rate(db)
    a1 = f.account(db, "wb", "ИП Яворская")
    a2 = f.account(db, "wb", "ИП Ребрик")
    client.post("/prices/rules/wb", data={**RULE, **rule})
    f.sku(db, "u1", "39681", "L", color="GRI", barcodes=["b1"], cost_usd="16.24", name="Свитшот")
    f.sku(db, "u2", "39681", "M", color="GRI", barcodes=["b2"], cost_usd="16.00", name="Свитшот")
    f.sku(db, "u3", "4033", "XL", color="SIYAH", barcodes=["b3"], cost_usd="12.40", name="Джемпер")
    for a in (a1, a2):
        for bc, ext, cat in (("b1", "5:1", "Свитшоты"), ("b2", "5:2", "Свитшоты"), ("b3", "6:1", "Джемперы")):
            db.add(PlatformItem(account_id=a.id, barcode=bc, article=bc, external_id=ext, category=cat,
                                tariff_fbs=Decimal("15.5") if cat == "Свитшоты" else None))
    db.commit()
    return a1, a2


def test_commission_is_tariff_plus_extra_on_pages(client, db):
    a1, _ = _wb(client, db, commission_extra="3")
    page = client.get("/prices?view=rules").text
    assert "Надбавка к комиссии" in page and "известен у 4 из 6 позиций" in page
    t = client.get("/sku-prices?scope=wb&by=category").text.split("<table", 1)[1]
    assert "Свитшоты" in t and "18,50" in t and "по тарифу" in t        # 15,5 + 3
    assert "28,00" in t and "из правила" in t                             # Джемперы: 25 + 3


def test_category_margin_set_and_sent(client, db):
    a1, a2 = _wb(client, db)
    r = client.post("/sku-prices/coef", data={"scope": "wb", "by": "category", "action": "set", "kind": "margin",
                                              "value": "1,8", "arts": ["Свитшоты"]})
    assert "категорий 1, изменено 1" in r.text
    row = db.query(CategoryTarget).one()
    assert (row.category, row.account_id, row.kind, row.value) == ("Свитшоты", 0, "margin", Decimal("1.8"))
    # 1323,56 × 1,8 / 0,845 = 2819,4 → вверх до 10, минус 1 → 2829; 1304 × 1,8 / 0,845 = 2777,8 → 2779
    t = client.get("/sku-prices?scope=wb").text.split("<table", 1)[1]
    assert "2779–2829" in t and "категория, площадка" in t
    # в кабинете своя маржинальность — только для него
    client.post("/sku-prices/coef", data={"scope": f"a{a2.id}", "by": "category", "row": "Свитшоты",
                                          "kind": "margin", "value": "2"})
    r = client.post("/sku-prices/send", data={"scope": "wb", "by": "category", "arts": ["Свитшоты"]})
    assert "передано на отправку 4 цен по 1 категориям" in r.text
    got = {(c.account_id, c.item_id): c for c in db.query(PriceChange).filter_by(status="approved")}
    assert got[(a1.id, "u1")].new_price == 2829 and got[(a2.id, "u1")].new_price > 2829
    assert got[(a1.id, "u1")].commission_percent == Decimal("15.5") and "маржинальность 1,8" in got[(a1.id, "u1")].note


def test_category_filter_and_export_keep_page_filters(client, db):
    _wb(client, db)
    t = client.get("/sku-prices?scope=wb&cat=Джемперы").text.split("<table", 1)[1]
    assert "4033" in t and "39681" not in t
    ws = load_workbook(io.BytesIO(client.get("/sku-prices/export?scope=wb&cat=Джемперы").content)).active
    arts = [ws.cell(row=i, column=3).value for i in range(2, ws.max_row + 1)]
    assert arts == ["4033"]
    ws = load_workbook(io.BytesIO(client.get("/sku-prices/export?scope=wb&by=category").content)).active
    h = [c.value for c in ws[1]]
    assert "Категория" in h and "Наценка" in h and ws.max_row == 3


def test_category_import_empty_changes_nothing(client, db):
    _wb(client, db)
    book = load_workbook(io.BytesIO(client.get("/sku-prices/export?scope=wb&by=category").content))
    ws = book.active
    h = [c.value for c in ws[1]]
    rows = {ws.cell(row=i, column=h.index("Категория") + 1).value: i for i in range(2, ws.max_row + 1)}
    ws.cell(row=rows["Свитшоты"], column=h.index("Наценка") + 1, value=1.9)
    ws.cell(row=rows["Свитшоты"], column=h.index("Вид наценки") + 1, value="маржинальность")
    buf = io.BytesIO()
    book.save(buf)
    r = client.post("/sku-prices/import", files={"file": ("k.xlsx", buf.getvalue())})
    assert "Наценок изменено: 1" in r.text
    assert [(c.category, c.kind, c.value) for c in db.query(CategoryTarget)] == [("Свитшоты", "margin", Decimal("1.9"))]


# --- пол со скидкой продавца при отправке ---------------------------------------------

def test_dispatch_floor_counts_seller_discount(db):
    f.manual_rate(db)
    acc = f.account(db, commission=25)
    f.rule(db, acc)
    f.sku(db, "u1", "39681", barcodes=["b1"], cost_usd="16.24")
    db.add(PlatformItem(account_id=acc.id, barcode="b1", external_id="5:1", current_price=4000,
                        current_sale_price=2800))                    # скидка продавца 30%
    # 2649 без скидки — маржинальность 1,50 (выше пола), со скидкой 30% — 1,05
    db.add(PriceChange(item_id="u1", account_id=acc.id, barcode="b1", new_price=2649, status="approved"))
    db.commit()

    class Client:
        def push_prices(self, items):
            raise AssertionError("ниже пола со скидкой уходить не должно")
    dispatch.run_account(db, acc, Client())
    ch = db.query(PriceChange).one()
    assert ch.status == "blocked" and "скидка продавца 30%" in ch.note


# --- диапазон безопасности ------------------------------------------------------------

def _prices(db, acc, **by_barcode):
    for it in db.query(PlatformItem).filter_by(account_id=acc.id):
        if it.barcode in by_barcode:
            it.current_price, it.current_sale_price = by_barcode[it.barcode]
    db.commit()


def test_guard_restores_when_platform_lowered_price_and_flags_eaten(client, db):
    a1, a2 = _wb(client, db)
    r = client.post(f"/prices/guard/{a1.id}", data={"guard_min_margin": "1,5", "guard_max_margin": "3"})
    assert "диапазон безопасности сохранён" in r.text
    # b1: площадка опустила цену до 2000 (наша 3309) — маржинальность 2000 × 0,845 / 1323,56 = 1,28
    # b3: цена наша (2529), но скидка 40% съела маржинальность — ценой не исправить
    # b2: 9000 — маржинальность выше «до»
    _prices(db, a1, b1=(2000, 2000), b3=(2529, 1517), b2=(9000, 9000))
    st = guard.run(db)
    assert (st["restored"], st["eaten"], st["above"]) == (1, 1, 1)
    ch = db.query(PriceChange).filter_by(status="approved").one()
    assert (ch.account_id, ch.item_id, ch.new_price, ch.source, ch.old_price) == (a1.id, "u1", 3309, "guard", 2000)
    assert guard.run(db)["restored"] == 0                         # уже в очереди — второй раз не ставим
    page = client.get("/attention").text
    assert "съедает скидка продавца или акция" in page and f"account_id={a1.id}&amp;flt=guard_eaten" in page
    t = client.get(f"/prices?view=products&account_id={a1.id}&flt=guard_above").text.split("<table", 1)[1]
    assert "39681" in t and "4033" not in t
    assert not any(c.account_id == a2.id for c in db.query(PriceChange))   # у кабинета без диапазона — ничего


def test_guard_never_below_floor_and_validation(client, db):
    a1, _ = _wb(client, db)
    # умолчание 1,5 при комиссии 25% — маржинальность 1,13, ниже пола 1,3: правило не сохраняется
    assert "не сохранено" in client.post("/prices/rules/wb", data={**RULE, "base_coef": "1,5"}).text
    assert "«от» больше «до»" in client.post(f"/prices/guard/{a1.id}", data={"guard_min_margin": "3",
                                                                              "guard_max_margin": "2"}).text
    client.post(f"/prices/guard/{a1.id}", data={"guard_min_margin": "1,5", "guard_max_margin": ""})
    client.post("/sku-prices/coef", data={"scope": f"a{a1.id}", "action": "set", "kind": "coef", "value": "1,2",
                                          "arts": ["39681"]})
    _prices(db, a1, b1=(1000, 1000))
    st = guard.run(db)
    assert st["blocked"] == 1 and st["restored"] == 0 and db.query(PriceChange).count() == 0


def test_guard_excel_and_run_button(client, db):
    a1, a2 = _wb(client, db)
    book = load_workbook(io.BytesIO(client.get("/prices/guard-export").content))
    ws = book.active
    rows = {ws.cell(row=i, column=2).value: i for i in range(2, ws.max_row + 1)}
    ws.cell(row=rows["ИП Ребрик"], column=4, value=1.4)
    ws.cell(row=rows["ИП Ребрик"], column=5, value=3.5)
    buf = io.BytesIO()
    book.save(buf)
    assert "Диапазонов изменено: 1" in client.post("/prices/guard-import", files={"file": ("g.xlsx", buf.getvalue())}).text
    db.expire_all()
    assert (db.get(Account, a2.id).guard_min_margin, db.get(Account, a2.id).guard_max_margin) == \
        (Decimal("1.4"), Decimal("3.5"))
    assert db.get(Account, a1.id).guard_min_margin is None
    assert "Диапазоны безопасности (1 каб.)" in client.post("/prices/guard-run").text


# --- запрос цен при запуске -----------------------------------------------------------

def test_startup_refresh_always_requests_prices_then_checks_guard(db, monkeypatch):
    from priceapp.workers import jobs
    from priceapp.timeutils import now_utc
    f.manual_rate(db)
    acc = f.account(db)
    f.rule(db, acc)
    acc_mod.set_credential(db, acc, "token", "t")
    acc.catalog_loaded_at = acc.prices_loaded_at = now_utc()       # всё свежее — суточный прогон бы не пошёл
    acc.guard_min_margin = Decimal("1.5")
    db.commit()
    calls = []

    class Client(FakeWb):
        def get_prices(self):
            calls.append("prices")
            return {}
    monkeypatch.setattr(jobs.guard, "run", lambda db_: calls.append("guard") or {"accounts": 1, "restored": 0,
                                                                                    "eaten": 0, "blocked": 0,
                                                                                    "unpriced": 0, "above": 0})
    jobs.job_daily_refresh(lambda p, c: Client())
    assert calls == []
    jobs.job_daily_refresh(lambda p, c: Client(), force_prices=True)
    assert calls == ["prices", "guard"]


# --- справка ------------------------------------------------------------------------------

def test_price_types_help_page(client, db):
    page = client.get("/help/prices").text
    assert "Типы цен" in page and "СПП" in page and "marketing_seller_price" in page and "Диапазон безопасности" in page
    assert 'href="/help/prices"' in client.get("/sku-prices").text


def test_mapping_export_keeps_status_filter(client, db):
    acc = f.account(db)
    f.sku(db, "u1", "39681", "L", barcodes=["b1"])
    f.item(db, acc, "b1", "39681-L", external_id="5:1")
    f.item(db, acc, "4607001200000", "NEW", external_id="7:1")
    page = client.get(f"/mapping?view=status&account_id={acc.id}&status=not_in_1c").text
    assert f"/mapping/export/{acc.id}?status=not_in_1c" in page
    ws = load_workbook(io.BytesIO(client.get(f"/mapping/export/{acc.id}?status=not_in_1c").content)).active
    assert [str(ws.cell(row=i, column=1).value) for i in range(2, ws.max_row + 1)] == ["4607001200000"]
