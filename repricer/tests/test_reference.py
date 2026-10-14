"""Эталонная цена кабинетов WB: запись при отправке, возврат при падении «к получению» больше
чем на 10%, пересчёт при сдвиге курса больше чем на 2%, опрос WB по расписанию."""
from decimal import Decimal

from priceapp import accounts as acc_mod, dispatch, reference, settings
from priceapp.models import PlatformItem, PriceChange, ProductPrice, ReferencePrice
from tests import factories as f

COST = "10"             # $; курс 80 → себестоимость 800 ₽, цена по правилу ×2,5 = 2000 (с округлением)


def _setup(db, platform="wb", rate="80", current=(2500, 2500)):
    f.manual_rate(db, rate)
    acc = f.account(db, platform=platform, commission=25)
    f.rule(db, acc, base_coef="2.5", min_markup_coef="1.3", max_change_percent=1000)
    f.sku(db, "u1", "A1", "S", barcodes=["b1"], cost_usd=COST)
    f.item(db, acc, "b1", "A1", external_id="5:1", size="S")
    row = db.query(PlatformItem).filter_by(barcode="b1").one()
    row.current_price, row.current_sale_price = current
    db.commit()
    return acc


def _cur(db, price, sale=None):
    row = db.query(PlatformItem).filter_by(barcode="b1").one()
    row.current_price, row.current_sale_price = price, sale if sale is not None else price
    db.commit()


def _ref(db, acc, price=2009):
    rows = db.query(PlatformItem).filter_by(account_id=acc.id).all()
    assert reference.record(db, acc, "u1", price, rows)
    db.commit()
    return db.query(ReferencePrice).filter_by(account_id=acc.id, item_id="u1").one()


def test_record_stores_payout_and_rate(db):
    acc = _setup(db, current=(2500, 2000))                       # скидка продавца 20%
    ref = _ref(db, acc, 2009)
    assert ref.ref_price == 2009 and ref.usd_rub == Decimal("80.0000") and ref.discount == Decimal("0.2000")
    assert ref.ref_payout == Decimal(str(2009 * 0.8 * 0.75)).quantize(Decimal("0.01"))      # 1205,40
    # повторная запись обновляет ту же строку
    ref2 = _ref(db, acc, 2100)
    assert db.query(ReferencePrice).count() == 1 and ref2.ref_price == 2100


def test_non_wb_has_no_reference(db):
    acc = _setup(db, platform="ozon")
    rows = db.query(PlatformItem).all()
    assert reference.record(db, acc, "u1", 2009, rows) is False and db.query(ReferencePrice).count() == 0


class _Push:
    def push_prices(self, items):
        return {"ok": [i.barcode for i in items], "errors": [], "sent_prices": {}}


def test_dispatch_records_reference_after_successful_send(db):
    acc = _setup(db)
    db.add(PriceChange(item_id="u1", account_id=acc.id, barcode="b1", new_price=2009, status="approved", decided_by="op"))
    db.commit()
    dispatch.run_account(db, acc, _Push())
    ref = db.query(ReferencePrice).one()
    assert ref.ref_price == 2009 and ref.source == "sent" and ref.ref_payout == Decimal("1506.75")
    assert db.query(ProductPrice).one().last_sent_price == 2009


def test_adopt_sent_records_prices_sent_before_the_rule(db):
    acc = _setup(db)
    db.add(ProductPrice(item_id="u1", account_id=acc.id, last_sent_price=2009))
    db.commit()
    st = reference.adopt_sent(db)
    assert st["recorded"] == 1 and db.query(ReferencePrice).one().source == "adopt"


class _Wb:
    def __init__(self):
        self.calls = []

    def _post(self, url, body):
        self.calls.append(body)
        return {"error": False}


def _run(db, wb=None):
    wb = wb or _Wb()
    acc_ids = [a.id for a in db.query(__import__("priceapp.models", fromlist=["Account"]).Account)]
    for i in acc_ids:
        a = db.get(__import__("priceapp.models", fromlist=["Account"]).Account, i)
        if not db.query(__import__("priceapp.models", fromlist=["ApiCredential"]).ApiCredential).filter_by(
                account_id=a.id).first():
            acc_mod.set_credential(db, a, "token", "t")
    db.commit()
    return reference.run(db, client_factory=lambda p, c: wb), wb


def test_restore_when_sale_price_drops_more_than_10_percent(db):
    """Эталон — цена после скидки продавца; возврат — эталонной парой одним заданием WB."""
    acc = _setup(db)
    _ref(db, acc, 2009)                                           # скидка 0 — эталон 2009, порог 1808
    _cur(db, 1950)                                                # падение 3% — в порядке
    st, wb = _run(db)
    assert st["restored"] == 0 and wb.calls == []
    _cur(db, 1700)                                                # падение 15,4%
    st, wb = _run(db)
    assert st["restored"] == 1 and st["eaten"] == 0
    assert wb.calls == [{"data": [{"nmID": 5, "price": 2009, "discount": 0}]}]
    # мимо очереди — одним заданием; в «Журнале отправок» — одна строка «отправлено»
    assert [c.status for c in db.query(PriceChange)] == ["sent"]
    db.expire_all()
    row = db.query(PlatformItem).filter_by(barcode="b1").one()
    assert (row.current_price, row.current_sale_price) == (2009, 2009)   # Kit в этом опросе видит эталон
    st, wb = _run(db)
    assert st["restored"] == 0 and wb.calls == []                 # уже на месте


def test_discount_drop_by_platform_is_returned_too(db):
    """Автоакция: цена до скидки та же, скидку подняла площадка — возвращается и скидка."""
    acc = _setup(db, current=(2500, 2000))                        # скидка 20% — эталон 2009 x 0,8 = 1607
    _ref(db, acc, 2009)
    _cur(db, 2009, sale=1100)                                     # площадка поставила скидку ~45%
    st, wb = _run(db)
    assert st["restored"] == 1 and wb.calls == [{"data": [{"nmID": 5, "price": 2009, "discount": 20}]}]


def test_return_raising_price_more_than_50_percent_is_attention_not_a_send(db):
    """13.10: «возврат» эталона 472 479 при цене 158 281 — ошибочный эталон, не занижение."""
    acc = _setup(db)
    _ref(db, acc, 2009)
    _cur(db, 1000)                                                # возврат подняли бы на 101%
    st, wb = _run(db)
    assert st["restored"] == 0 and st["eaten"] == 1 and wb.calls == []
    from priceapp.models import AuditLog
    details = db.query(AuditLog).filter_by(action="reference_big").one().details
    assert "НЕ возвращено" in details and "на 50%" in details               # не «5E+1%»
    settings.put(db, settings.REF_MAX_RAISE_PERCENT, "150")
    db.commit()
    assert _run(db)[0]["restored"] == 1


def test_never_lowers_price_and_keeps_human_decision(db):
    acc = _setup(db)
    _ref(db, acc, 2009)
    _cur(db, 3000)                                                # площадка подняла — не трогаем
    st, wb = _run(db)
    assert st["restored"] == 0 and wb.calls == []
    _cur(db, 1700)
    db.add(PriceChange(item_id="u1", account_id=acc.id, barcode="b1", new_price=5000, status="approved", decided_by="op"))
    db.commit()
    st, wb = _run(db)
    assert wb.calls == [] and [(c.new_price, c.status) for c in db.query(PriceChange)] == [(5000, "approved")]


def test_rate_move_over_2_percent_recomputes_reference(db):
    acc = _setup(db)
    ref = _ref(db, acc, 2009)
    f.manual_rate(db, "81.2")                                     # +1,5% — порога нет
    assert reference.rate_moved(db) is False and reference.run(db)["revalued"] == 0
    assert db.query(ReferencePrice).one().ref_price == 2009
    f.manual_rate(db, "82.4")                                     # +3%
    assert reference.rate_moved(db) is True
    st = reference.run(db)
    db.expire_all()
    ref = db.query(ReferencePrice).one()
    assert st["revalued"] == 1 and ref.usd_rub == Decimal("82.4000")
    assert ref.ref_price == 2069 and ref.source == "rate"        # 10 $ × 82,4 × 2,5 = 2060 → вверх до шага 10, минус 1 = 2069
    assert st["restored"] == 0                                    # цена площадки (2500) не ниже новой эталонной


def test_rate_move_then_first_rule_applies(db):
    acc = _setup(db)
    _ref(db, acc, 2009)
    _cur(db, 1700)                                                # площадка уже занизила
    f.manual_rate(db, "82.4")
    st, wb = _run(db)
    assert st["revalued"] == 1 and st["restored"] == 1
    assert wb.calls[0]["data"][0]["price"] == 2069                # возвращается уже ПЕРЕСЧИТАННАЯ эталонная


def test_rate_fall_lowers_reference_but_sends_nothing(db):
    acc = _setup(db)
    _ref(db, acc, 2009)
    f.manual_rate(db, "77.6")                                     # −3%
    st = reference.run(db)
    db.expire_all()
    assert st["revalued"] == 1 and db.query(ReferencePrice).one().ref_price < 2009 and st["restored"] == 0
    assert db.query(PriceChange).count() == 0


def test_rule_can_be_disabled_and_thresholds_changed(db):
    acc = _setup(db)
    _ref(db, acc, 2009)
    _cur(db, 1700)
    settings.put(db, settings.REF_ENABLED, "0")
    db.commit()
    assert reference.run(db)["disabled"] is True and db.query(PriceChange).count() == 0
    settings.put(db, settings.REF_ENABLED, "1")
    settings.put(db, settings.REF_DROP_PERCENT, "20")             # падение 15,4% < 20% — не возвращаем
    db.commit()
    assert reference.run(db)["restored"] == 0


class _FakeWb:
    last_truncated = False

    def get_prices(self):
        return {}

    def price_key(self, item):
        return (item.external_id or "").split(":")[0]


def test_wb_poll_loads_prices_then_runs_rule_and_skips_fresh_prices(db, monkeypatch):
    from priceapp.timeutils import now_utc
    from priceapp.workers import jobs
    acc = _setup(db)
    acc_mod.set_credential(db, acc, "token", "t")
    db.commit()
    calls = []
    monkeypatch.setattr(jobs.accounts, "load_prices", lambda db_, a, c: calls.append("prices") or {"truncated": False})
    monkeypatch.setattr(jobs.reference, "run", lambda db_, client_factory=None: calls.append("rule") or
                        {"accounts": 1, "revalued": 0, "restored": 0, "eaten": 0, "disabled": False})
    jobs.job_wb_poll(lambda p, c: _FakeWb())
    assert calls == ["prices", "rule"]
    calls.clear()
    acc.prices_loaded_at = now_utc()                              # цены только что загружены суточным обновлением
    db.commit()
    jobs.job_wb_poll(lambda p, c: _FakeWb())
    assert calls == ["rule"]
    assert db.get(__import__("priceapp.models", fromlist=["WorkerHeartbeat"]).WorkerHeartbeat, "wb_poll").last_success


def test_poll_now_flag_is_read_once(db):
    from priceapp.workers import jobs
    assert jobs.wb_poll_requested() is False
    settings.put(db, settings.REF_POLL_NOW, "1")
    db.commit()
    assert jobs.wb_poll_requested() is True and jobs.wb_poll_requested() is False


def test_heavy_loop_polls_wb_at_start_and_every_6_hours():
    from priceapp.workers import background
    assert background.WB_POLL_EVERY == 6 * 3600 and background.WB_POLL_EVERY * 4 == 24 * 3600


def test_etalon_copy_is_saved_on_change_and_not_repeated(db):
    from openpyxl import load_workbook
    from priceapp import config
    for p in config.ETALON_DIR.glob("etalon_*.xlsx") if config.ETALON_DIR.exists() else []:
        p.unlink()
    settings.put(db, settings.REF_SNAPSHOT_HASH, "")
    db.commit()
    acc = _setup(db, current=(2500, 2000))
    _ref(db, acc, 2009)                                             # запись эталона
    p1 = reference.snapshot_if_dirty(db, "тест")
    assert p1 is not None and p1.exists()
    rows = list(load_workbook(p1).active.iter_rows(values_only=True))
    got = dict(zip(rows[0], rows[1]))
    assert got["Артикул WB"] == 5 and got["Цена до скидки, руб"] == 2009 and got["Скидка продавца, %"] == 20
    assert got["ЭТАЛОН: цена после скидки продавца, руб"] == 1607
    assert reference.snapshot_if_dirty(db, "тест") is None          # не менялся — копии нет
    _ref(db, acc, 2009)                                             # записали то же самое
    assert reference.snapshot_if_dirty(db, "тест") is None          # содержимое то же — копии нет
    _ref(db, acc, 2100)
    p2 = reference.snapshot_if_dirty(db, "тест")
    assert p2 is not None and p2 != p1 and len(list(config.ETALON_DIR.glob("etalon_*.xlsx"))) == 2


def test_dispatch_send_saves_etalon_copy(db):
    from priceapp import config
    settings.put(db, settings.REF_SNAPSHOT_HASH, "")
    db.commit()
    before = set(config.ETALON_DIR.glob("etalon_*.xlsx")) if config.ETALON_DIR.exists() else set()
    acc = _setup(db)
    acc_mod.set_credential(db, acc, "token", "t")
    db.add(PriceChange(item_id="u1", account_id=acc.id, barcode="b1", new_price=2009, status="approved", decided_by="op"))
    db.commit()
    dispatch.run(db, lambda p, c: _Push())
    after = set(config.ETALON_DIR.glob("etalon_*.xlsx"))
    new = after - before
    assert len(new) == 1 and "отправка" in next(iter(new)).name
