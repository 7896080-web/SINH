"""Наценки по категориям от эталона WB (scripts/etalon_margins.py): только чтение,
числа сходятся с тем, что потом посчитает `pricing.decide`."""
import io
import os
import subprocess
import sys
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from priceapp.models import CategoryTarget, PlatformItem, PriceChange
from tests import factories as F

ROOT = Path(__file__).resolve().parent.parent


def _setup(db):
    w1 = F.account(db, "wb", "ИП 1", commission=25)
    w2 = F.account(db, "wb", "ИП 2", commission=25)
    kit = F.account(db, "kit", "КИТ", commission=10)
    lam = F.account(db, "lamoda", "Ламода", commission=30)
    F.manual_rate(db, "100")
    F.sku(db, "a1", "100", "M", barcodes=["b1"], cost_usd=10)      # себестоимость 1000 ₽
    F.sku(db, "a2", "200", "M", barcodes=["b2"], cost_usd=20)      # 2000 ₽
    for a, sale in ((w1, 4000), (w2, 6000)):                       # средняя 5000 -> эталон 3450
        F.item(db, a, "b1", "100")
    F.item(db, w1, "b2", "200")
    for a in (kit, lam):
        F.item(db, a, "b1", "100"); F.item(db, a, "b2", "200")
    prices = {(w1.id, "b1"): 4000, (w2.id, "b1"): 6000, (w1.id, "b2"): 10000}
    for it in db.query(PlatformItem):
        it.category = "Куртки"
        if (it.account_id, it.barcode) in prices:
            it.current_price = prices[(it.account_id, it.barcode)] * 2
            it.current_sale_price = prices[(it.account_id, it.barcode)]
    db.commit()
    return kit, lam


def _run(tmp_path, *extra):
    out = tmp_path / "e.xlsx"
    env = {**os.environ, "PYTHONIOENCODING": "cp1251"}
    r = subprocess.run([sys.executable, str(ROOT / "scripts" / "etalon_margins.py"), "--xlsx", str(out), *extra],
                       capture_output=True, env=env, cwd=ROOT)
    assert r.returncode == 0, r.stderr.decode("utf-8", "replace")
    return r.stdout.decode("cp1251"), load_workbook(out)


def test_margins_follow_the_etalon_logic(db, tmp_path):
    _setup(db)
    text, wb = _run(tmp_path)
    assert "по 2 артикулам" in text
    items = {(r[0], r[2]): r for r in list(wb["Товары"].iter_rows(values_only=True))[1:]}
    # товар 100: эталон = (4000 + 6000) / 2 × 0,69 = 3450
    kit = items[("Яндекс KIT", "a1")]
    assert kit[7] == 3450 and kit[8] == 3105                       # Kit = эталон − 10%
    assert abs(kit[10] - 3105 * 0.9 / 1000) < 0.001                  # m = цена × (1 − 10%) / 1000
    lam = items[("Lamoda", "a1")]
    assert lam[8] == 3450 and abs(lam[10] - 3450 * 0.7 / 1000) < 0.001
    # товар 200: эталон 10000 × 0,69 = 6900
    assert abs(items[("Lamoda", "a2")][10] - 6900 * 0.7 / 2000) < 0.001


def _import(client, wb):
    buf = io.BytesIO()
    wb.save(buf)
    return client.post("/sku-prices/import", files={"file": ("e.xlsx", io.BytesIO(buf.getvalue()), "x")})


def test_article_import_puts_every_article_exactly_on_its_etalon(client, db, tmp_path):
    from priceapp.models import ArticleCoef
    from priceapp.routers.sku_prices import build, pick_scope
    _setup(db)
    text, wb = _run(tmp_path)
    assert "наценки артикулов: мин - медиана - макс" in text
    assert wb.active.title == "Импорт"
    assert [c.value for c in wb["Импорт"][1]] == ["Где (код)", "Где", "Категория", "Артикул", "Вид наценки", "Наценка"]
    arts = list(wb["Артикулы"].iter_rows(values_only=True))
    assert arts[0][2] == "Артикул" and len(arts) == 5                 # 2 артикула × Kit и Ламода
    r = _import(client, wb)
    assert "Наценок изменено: 4" in r.text
    assert {(c.platform, c.article, c.account_id, c.kind) for c in db.query(ArticleCoef)} == {
        ("kit", "100", 0, "margin"), ("kit", "200", 0, "margin"),
        ("lamoda", "100", 0, "margin"), ("lamoda", "200", 0, "margin")}
    db.expire_all()
    for scope, art, want in (("lamoda", "100", 3450), ("lamoda", "200", 6900), ("kit", "100", 3105),
                             ("kit", "200", 6210)):
        price = build(db, pick_scope(db, scope), art=art)[0]["rows"][0]["cells"][0]["price"]
        assert want <= price < want * 1.002, (scope, art, price)       # вверх до шага округления


def test_category_level_is_an_option(client, db, tmp_path):
    _setup(db)
    _, wb = _run(tmp_path, "--level", "category")
    assert [c.value for c in wb["Импорт"][1]] == ["Где (код)", "Где", "Категория", "Вид наценки", "Наценка"]
    assert "Наценок изменено: 2" in _import(client, wb).text
    assert {(c.platform, c.category) for c in db.query(CategoryTarget)} == {("kit", "Куртки"), ("lamoda", "Куртки")}


def test_script_changes_nothing(db, tmp_path):
    _setup(db)
    before = (db.query(CategoryTarget).count(), db.query(PriceChange).count())
    _run(tmp_path)
    db.expire_all()
    assert (db.query(CategoryTarget).count(), db.query(PriceChange).count()) == before


def test_percents_are_parameters(db, tmp_path):
    _setup(db)
    _, wb = _run(tmp_path, "--wb-discount", "0", "--kit", "0")
    items = {(r[0], r[2]): r for r in list(wb["Товары"].iter_rows(values_only=True))[1:]}
    assert items[("Яндекс KIT", "a1")][8] == 5000
