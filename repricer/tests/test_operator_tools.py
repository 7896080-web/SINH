"""Инструменты оператора: «Внимание», цена карточки, сводка и подтверждение всего
отбора, сравнение площадок, «что будет, если», скачок курса, история и откат,
сохранённые отборы, статус и минимальная цена площадки."""
import io
from datetime import datetime, timezone
from decimal import Decimal

from openpyxl import load_workbook

from priceapp import accounts as acc_mod, overview, platforms
from priceapp.models import PlatformItem, PlatformRule, PriceChange, ProductPrice, SavedFilter
from priceapp.platforms import CurrentPrice
from priceapp.pricing import get_rule
from tests import factories as f

RULE = {"commission_percent": "25", "base_coef": "2,667", "min_markup_coef": "1,3",
        "round_step": "10", "round_minus": "1", "max_change_percent": "20"}


def _wb(client, db, names=("ИП Яворская",)):
    f.manual_rate(db)
    accs = [f.account(db, "wb", n) for n in names]
    client.post("/prices/rules/wb", data=RULE)
    # два размера ОДНОЙ карточки (nmID 5) с разной себестоимостью и третий товар
    f.sku(db, "u1", "39681", "L", barcodes=["b1"], cost_usd="16.24", name="Свитшот")
    f.sku(db, "u2", "39681", "M", barcodes=["b2"], cost_usd="15.80", name="Свитшот")
    f.sku(db, "u3", "4033", "3XL", barcodes=["b3"], cost_usd="12.40", name="Джемпер")
    for a in accs:
        f.item(db, a, "b1", "39681-L", external_id="5:1")
        f.item(db, a, "b2", "39681-M", external_id="5:2")
        f.item(db, a, "b3", "4033-3XL", external_id="6:1")
    return accs


# --- «Внимание» ---------------------------------------------------------------------

def test_root_and_login_lead_to_attention(client, db):
    assert client.get("/", follow_redirects=False).headers["location"] == "/attention"
    page = client.get("/attention").text
    assert "Курса доллара нет" in page and 'href="/attention"' in page


def test_attention_lists_what_needs_a_decision_with_links(client, db):
    a, = _wb(client, db)
    client.post("/prices/recalculate", data={})
    for it in db.query(PlatformItem).filter_by(barcode="b3"):
        it.current_price = it.current_sale_price = 1000          # 750 к получению < 1010,6 — ниже пола
    db.commit()
    page = client.get("/attention").text
    assert "предложения ждут решения" in page and "/prices?view=proposals&amp;status=proposed" in page
    assert "ниже пола по текущей цене" in page and f"account_id={a.id}&amp;flt=below_floor_now" in page
    assert "не заданы ключи" in page


# --- цена карточки и сводка ---------------------------------------------------------

def test_card_price_is_shown_before_approval(client, db):
    _wb(client, db)
    client.post("/prices/recalculate", data={})
    by_item = {c.item_id: c for c in db.query(PriceChange)}
    assert by_item["u1"].new_price == 3539 and by_item["u2"].new_price == 3439
    info = overview.change_context(db, list(by_item.values()))
    assert info[by_item["u2"].id].card_price == 3539 and info[by_item["u2"].id].card_size == 2
    assert info[by_item["u3"].id].card_size == 1
    s = overview.proposals_summary(db, list(by_item.values()))
    assert (s.total, s.new, s.card_raised) == (3, 3, 1)
    page = client.get("/prices?view=proposals").text
    assert "уйдёт 3539" in page and "Итог по отбору — 3 шт" in page


def test_approve_whole_filter_not_only_visible_rows(client, db):
    _wb(client, db)
    client.post("/prices/recalculate", data={})
    r = client.post("/prices/approve", data={"all_filtered": "1", "q": "39681", "status": ""})
    assert "Подтверждено: 2" in r.text
    assert {c.item_id for c in db.query(PriceChange).filter_by(status="approved")} == {"u1", "u2"}
    r = client.post("/prices/reject", data={"all_filtered": "1", "q": "4033"})
    assert "Отклонено: 1" in r.text


# --- сравнение площадок -------------------------------------------------------------

def test_compare_one_row_per_sku_across_cabinets(client, db):
    a1, a2 = _wb(client, db, ("ИП Яворская", "ИП Ребрик"))
    for it in db.query(PlatformItem):
        it.current_price = it.current_sale_price = 3000 if it.account_id == a1.id else 4000
    db.commit()
    page = client.get("/prices?view=compare").text
    assert "ИП Яворская" in page and "ИП Ребрик" in page and "33%" in page
    assert page.count("/prices/history/") >= 6
    r = client.get("/prices/compare-export?flt=spread")
    ws = load_workbook(io.BytesIO(r.content)).active
    headers = [c.value for c in ws[1]]
    assert "ИП Ребрик: текущая, ₽" in headers and ws.max_row == 4


# --- «что будет, если» --------------------------------------------------------------

def test_rule_preview_counts_without_saving(client, db):
    _wb(client, db)
    r = client.post("/prices/rules/wb/preview", data={**RULE, "base_coef": "2,5"})
    assert "Если сохранить" in r.text and "Wildberries (1 каб.)" in r.text and "изменится 3" in r.text
    assert "ниже 3" in r.text and "Ничего не сохранено" in r.text
    db.expire_all()
    assert get_rule(db, "wb").base_coef == Decimal("2.667")


# --- скачок курса -------------------------------------------------------------------

def test_rate_shift_warns_until_recalculated(client, db):
    _wb(client, db)
    client.post("/prices/recalculate", data={})
    client.post("/rate/mode", data={"mode": "manual", "manual": "85"})
    assert "Курс ушёл на 4.3%" in client.get("/prices?view=proposals").text
    assert "Курс изменился на 4.3%" in client.get("/attention").text
    client.post("/rate/alert", data={"alert": "5"})
    assert "Курс ушёл" not in client.get("/prices?view=proposals").text
    assert "не принят" in client.post("/rate/alert", data={"alert": "abc"}).text
    client.post("/rate/alert", data={"alert": "2"})
    client.post("/prices/recalculate", data={})
    assert "Курс ушёл" not in client.get("/prices?view=proposals").text


# --- история и возврат прежней цены -------------------------------------------------

def test_history_and_rollback_make_an_ordinary_proposal(client, db):
    a, = _wb(client, db)
    old = PriceChange(item_id="u1", account_id=a.id, barcode="b1", new_price=3399, status="sent",
                      sent_at=datetime(2026, 9, 1, tzinfo=timezone.utc).replace(tzinfo=None))
    db.add(old)
    db.add(ProductPrice(item_id="u1", account_id=a.id, last_sent_price=3539))
    db.commit()
    page = client.get(f"/prices/history/{a.id}/u1").text
    assert "3399" in page and "Вернуть эту цену" in page
    r = client.post(f"/prices/rollback/{old.id}")
    assert "Предложение вернуть 3399" in r.text
    new = db.query(PriceChange).filter_by(source="rollback").one()
    assert new.status == "proposed" and new.new_price == 3399 and new.old_price == 3539
    # то, что площадка не принимала, вернуть нельзя
    assert "только цену, которую площадка приняла" in client.post(f"/prices/rollback/{new.id}").text


def test_rollback_below_floor_is_blocked(client, db):
    a, = _wb(client, db)
    old = PriceChange(item_id="u1", account_id=a.id, barcode="b1", new_price=1500, status="sent")
    db.add(old)
    db.commit()
    r = client.post(f"/prices/rollback/{old.id}")
    assert "заблокировано" in r.text
    assert db.query(PriceChange).filter_by(source="rollback").one().block_reason == "floor"


# --- сохранённые отборы -------------------------------------------------------------

def test_saved_filters(client, db):
    _wb(client, db)
    url = "/prices?view=products&flt=no_cost"
    client.post("/filters/save", data={"name": "без себестоимости", "url": url})
    page = client.get("/prices?view=proposals").text
    assert "без себестоимости" in page and 'href="/prices?view=products&amp;flt=no_cost"' in page
    client.post("/filters/save", data={"name": "чужой", "url": "//evil.example/prices"})
    client.post("/filters/save", data={"name": "чужой2", "url": "https://evil.example"})
    assert [x.name for x in db.query(SavedFilter)] == ["без себестоимости"]
    client.post(f"/filters/{db.query(SavedFilter).one().id}/delete")
    assert db.query(SavedFilter).count() == 0


# --- статус и минимальная цена площадки (Lamoda) -------------------------------------

def test_lamoda_min_prices_matched_by_category_and_brand():
    minimal = [{"categoryName": "CLOTHES", "subcategoryName": "SWEATSHIRTS",
                "prices": [{"country": "RU", "price": {"amount": 99900, "currency": "RUB"}}]},
               {"categoryName": "CLOTHES", "subcategoryName": "SWEATSHIRTS", "brand": "Nero",
                "prices": [{"country": "RU", "price": {"amount": 149900, "currency": "RUB"}}]},
               {"categoryName": "SHOES", "subcategoryName": "BOOTS",
                "prices": [{"country": "KZ", "price": {"amount": 1, "currency": "KZT"}}]}]
    lv = [{"level": 1, "language": "EN", "name": "Clothes"}, {"level": 2, "language": "EN", "name": "Sweatshirts"},
          {"level": 1, "language": "RU", "name": "Одежда"}]
    noms = [{"parentSku": "MP1", "brand": "NERO", "categoryLevels": lv},
            {"parentSku": "MP2", "brand": "Other", "categoryLevels": lv},
            {"parentSku": "MP3", "brand": "x", "categoryLevels": [{"level": 1, "language": "EN", "name": "SHOES"}]}]
    assert platforms.parse_lamoda_min_prices(minimal, noms) == {"MP1": 1499, "MP2": 999}


def test_lamoda_worst_status_of_sizes_wins():
    data = {"nomenclatures": [
        {"parentSku": "MP1", "sellValues": [{"country": "RU", "price": {"amount": 300000, "currency": "RUB"},
                                             "priceUpdateStatus": "OK"}]},
        {"parentSku": "MP1", "sellValues": [{"country": "RU", "price": {"amount": 290000, "currency": "RUB"},
                                             "priceUpdateStatus": "QUARANTINE"}]}]}
    assert platforms.parse_lamoda_sell_values(data)["MP1"] == CurrentPrice(3000, 3000, "QUARANTINE")


class _Lam:
    last_truncated = False

    def price_key(self, item):
        return platforms.lamoda_parent(item.external_id)

    def get_prices(self):
        return {"MP1": CurrentPrice(3000, 2700, "QUARANTINE")}

    def get_min_prices(self):
        return {"MP1": 5000}


def test_platform_status_and_min_price_reach_products_and_proposals(client, db):
    f.manual_rate(db)
    acc = f.account(db, "lamoda", "Ламода", commission=35)
    client.post("/prices/rules/lamoda", data={**RULE, "commission_percent": "35"})
    f.sku(db, "u1", "39681", "L", barcodes=["b1"], cost_usd="16.24")
    f.item(db, acc, "b1", "39681", external_id="MP1:MP1R36")
    acc_mod.load_prices(db, acc, _Lam())
    it = db.query(PlatformItem).one()
    assert (it.price_status, it.min_price, it.current_price) == ("QUARANTINE", 5000, 3000)
    page = client.get(f"/prices?view=products&account_id={acc.id}&flt=platform_status").text
    assert "на карантине у площадки" in page and "ниже мин. площадки 5000" in page
    client.post("/prices/recalculate", data={})
    ch = db.query(PriceChange).one()
    assert "ниже минимальной цены площадки 5000" in ch.note
    assert ch.status == "proposed"            # только предупреждение, не блок
    assert "карантине" in client.get("/attention").text


def test_all_new_pages_render(client, db):
    a, = _wb(client, db)
    for url in ("/attention", "/prices?view=compare", f"/prices/history/{a.id}/u1", "/rate"):
        assert client.get(url).status_code == 200, url
