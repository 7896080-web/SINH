"""Цены от эталона WB через скидку продавца (scripts/etalon_prices.py): только чтение."""
import os
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

from priceapp.models import ArticleCoef, PlatformItem, PriceChange, ProductPrice
from tests import factories as F

ROOT = Path(__file__).resolve().parent.parent


def _setup(db):
    w1 = F.account(db, "wb", "ИП 1", commission=25)
    w2 = F.account(db, "wb", "ИП 2", commission=25)
    kit = F.account(db, "kit", "КИТ", commission=10)
    lam = F.account(db, "lamoda", "Ламода", commission=30)
    F.manual_rate(db, "100")
    F.sku(db, "a1", "100", "M", barcodes=["b1"], cost_usd=10)
    F.sku(db, "a3", "300", "M", barcodes=["b3"], cost_usd=10)      # нет на WB
    F.item(db, w1, "b1", "100"); F.item(db, w2, "b1", "100")
    F.item(db, kit, "b1", "100"); F.item(db, lam, "b1", "100"); F.item(db, kit, "b3", "300")
    prices = {(w1.id, "b1"): (8000, 4000), (w2.id, "b1"): (9000, 6000),   # эталон (4000+6000)/2 × 0,69 = 3450
              (kit.id, "b1"): (6000, 4000), (lam.id, "b1"): (5000, 4500), (kit.id, "b3"): (3000, 2000)}
    for it in db.query(PlatformItem):
        it.current_price, it.current_sale_price = prices[(it.account_id, it.barcode)]
    db.commit()


def _run(tmp_path, *extra):
    out = tmp_path / "p.xlsx"
    env = {**os.environ, "PYTHONIOENCODING": "cp1251"}
    r = subprocess.run([sys.executable, str(ROOT / "scripts" / "etalon_prices.py"), "--xlsx", str(out), *extra],
                       capture_output=True, env=env, cwd=ROOT)
    assert r.returncode == 0, r.stderr.decode("utf-8", "replace")
    ws = load_workbook(out)["Kit и Ламода"]
    rows = list(ws.iter_rows(values_only=True))
    return r.stdout.decode("cp1251"), [dict(zip(rows[0], x)) for x in rows[1:]]


def test_target_is_reached_by_discount_within_corridor(db, tmp_path):
    _setup(db)
    text, rows = _run(tmp_path)
    assert "НИЧЕГО НЕ ОТПРАВЛЯЕТСЯ" in text
    kit = next(r for r in rows if r["Площадка"] == "Яндекс KIT" and r["Артикул"] == "100")
    # Kit: цель 3450 × 0,9 = 3105; при цене до скидки 6000 скидка 48,25% — в коридоре 35–50%
    assert kit["Цель со скидкой, ₽"] == 3105 and kit["Что сделать"] == "сменить скидку"
    assert kit["Новая цена до скидки, ₽"] == 6000 and kit["Новая скидка, %"] == 48.2
    assert round(6000 * (1 - kit["Новая скидка, %"] / 100)) in (3105, 3108)


def test_out_of_corridor_moves_the_list_price_to_the_nearest_bound(db, tmp_path):
    _setup(db)
    _, rows = _run(tmp_path)
    lam = next(r for r in rows if r["Площадка"] == "Lamoda")
    # Ламода: цель 3450; при 5000 нужна скидка 31% < 35% — цена до скидки 3450 / 0,65 = 5308, скидка 35%
    assert lam["Цель со скидкой, ₽"] == 3450 and lam["Что сделать"] == "поднять цену до скидки"
    assert lam["Новая цена до скидки, ₽"] == 5308 and lam["Новая скидка, %"] == 35.0
    assert lam["Изменение цены после скидки продавца, %"] == round((3450 - 4500) * 100 / 4500, 1)


def test_article_without_wb_has_no_etalon_and_nothing_changes(db, tmp_path):
    _setup(db)
    before = (db.query(PriceChange).count(), db.query(ArticleCoef).count(), db.query(ProductPrice).count())
    _, rows = _run(tmp_path)
    assert next(r for r in rows if r["Артикул"] == "300")["Что сделать"] == "нет эталона (нет на WB)"
    db.expire_all()
    assert (db.query(PriceChange).count(), db.query(ArticleCoef).count(), db.query(ProductPrice).count()) == before


def test_wb_sheet_shows_each_cabinet_against_the_average(db, tmp_path):
    _setup(db)
    text, _ = _run(tmp_path)
    assert "цена со скидкой продавца НЕ меняется" in text
    wb = load_workbook(tmp_path / "p.xlsx")["WB"]
    rows = list(wb.iter_rows(values_only=True))
    h = rows[0]
    got = {r[h.index("Кабинет")]: r for r in rows[1:]}
    # средняя (4000 + 6000) / 2 = 5000, эталон 3450; ИП 1 на 20% ниже средней, ИП 2 на 20% выше
    assert got["ИП 1"][h.index("Средняя со скидкой по WB, ₽")] == 5000
    assert got["ИП 1"][h.index("Кабинет к средней, %")] == -20.0
    assert got["ИП 2"][h.index("Кабинет к средней, %")] == 20.0
    assert got["ИП 2"][h.index("Эталон (−31%), ₽")] == 3450
    assert got["ИП 1"][h.index("Сейчас: скидка, %")] == 50.0


def test_wb_keeps_the_sale_price_and_moves_list_and_discount_into_corridor(db, tmp_path):
    _setup(db)
    _run(tmp_path)
    wb = load_workbook(tmp_path / "p.xlsx")["WB"]
    rows = list(wb.iter_rows(values_only=True))
    h = rows[0]
    got = {r[h.index("Кабинет")]: r for r in rows[1:]}
    # ИП 1: 8000 / 4000 — скидка 50%, уже в коридоре: ничего не меняем
    assert got["ИП 1"][h.index("Что сделать")] == "скидка уже в коридоре"
    assert got["ИП 1"][h.index("Новая цена до скидки, ₽")] == 8000
    # ИП 2: 9000 / 6000 — скидка 33% < 35%: цена со скидкой та же 6000, скидка 35%, до скидки
    # 6000 / 0,65 = 9230,77 — ВНИЗ, до 9230: с округлением вверх цена со скидкой стала бы 6000,15
    r = got["ИП 2"]
    assert r[h.index("Что сделать")] == "поднять цену до скидки"
    assert r[h.index("Новая цена до скидки, ₽")] == 9230 and r[h.index("Новая скидка, %")] == 35.0
    assert 9230 * 0.65 <= 6000
    assert r[h.index("Цена со скидкой (не меняется), ₽")] == 6000


def test_wb_card_with_different_sale_prices_per_size_is_marked_for_per_size_only(db, tmp_path):
    """Цена на WB уходит карточкой — наибольшая на все размеры. Размер, у которого цена со
    скидкой ниже, при такой отправке подорожал бы, поэтому карточка помечается."""
    w = F.account(db, "wb", "ИП 1", commission=25)
    F.manual_rate(db, "100")
    F.sku(db, "s", "500", "S", barcodes=["bs"], cost_usd=10)
    F.sku(db, "l", "500", "L", barcodes=["bl"], cost_usd=10)
    F.sku(db, "m", "600", "M", barcodes=["bm"], cost_usd=10)
    F.sku(db, "x", "600", "XL", barcodes=["bx"], cost_usd=10)
    F.item(db, w, "bs", "500", external_id="111:1"); F.item(db, w, "bl", "500", external_id="111:2")
    F.item(db, w, "bm", "600", external_id="222:1"); F.item(db, w, "bx", "600", external_id="222:2")
    prices = {"bs": (6000, 3000), "bl": (7000, 3500), "bm": (6000, 3000), "bx": (6000, 3000)}
    for it in db.query(PlatformItem):
        it.current_price, it.current_sale_price = prices[it.barcode]
    db.commit()
    text, _ = _run(tmp_path)
    assert "1  карточек: цены со скидкой по размерам разные" in text
    wb = load_workbook(tmp_path / "p.xlsx")["WB"]
    rows = list(wb.iter_rows(values_only=True))
    h = rows[0]
    mark = {r[h.index("ID_1С")]: r[h.index("Карточка")] for r in rows[1:]}
    assert mark["s"] and mark["l"] and not mark["m"] and not mark["x"]


def test_lamoda_only_uses_recorded_wb_etalon(db, tmp_path):
    """Эталон WB — записанная эталонная цена после скидки продавца, а не нынешняя."""
    from decimal import Decimal
    from priceapp.models import Account, ReferencePrice
    _setup(db)
    w1 = db.query(Account).filter_by(name="ИП 1").one()
    w2 = db.query(Account).filter_by(name="ИП 2").one()
    # эталон: ИП 1 — 6000 / 50% = 3000, ИП 2 — 10000 / 30% = 7000 -> средняя 5000, Ламода 5000 x 0,69 = 3450
    db.add(ReferencePrice(account_id=w1.id, item_id="a1", ref_price=6000, discount=Decimal("0.5")))
    db.add(ReferencePrice(account_id=w2.id, item_id="a1", ref_price=10000, discount=Decimal("0.3")))
    db.commit()
    out = tmp_path / "l.xlsx"
    env = {**os.environ, "PYTHONIOENCODING": "cp1251"}
    r = subprocess.run([sys.executable, str(ROOT / "scripts" / "etalon_prices.py"), "--only", "lamoda",
                        "--xlsx", str(out)], capture_output=True, env=env, cwd=ROOT)
    assert r.returncode == 0, r.stderr.decode("utf-8", "replace")
    text = r.stdout.decode("cp1251")
    assert "из записанного эталона — 2 SKU" in text and "Яндекс KIT" not in text
    book = load_workbook(out)
    assert book.sheetnames == ["Сводка", "Ламода"]
    rows = list(book["Ламода"].iter_rows(values_only=True))
    lam = dict(zip(rows[0], rows[1]))
    assert lam["Эталон WB после скидки продавца (средняя), ₽"] == 5000 and lam["Цель со скидкой, ₽"] == 3450
