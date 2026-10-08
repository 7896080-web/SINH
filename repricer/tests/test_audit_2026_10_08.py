"""Находки аудита 08.10 (`АУДИТ_2026-10-08.md`): каждая закрыта тестом, который
воспроизводил дефект до правки."""
import ast
import io
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import event

from priceapp.models import ArticleCoef, PriceChange, ProductPrice
from tests import factories as F

ROUTERS = Path(__file__).resolve().parent.parent / "priceapp" / "routers"


def _wb_two(db, n=2):
    a1 = F.account(db, "wb", "ИП 1")
    a2 = F.account(db, "wb", "ИП 2")
    F.rule(db, a1)
    F.manual_rate(db)
    for i in range(n):
        F.sku(db, f"i{i}", f"{100 + i}", "M", barcodes=[f"20{i:04d}"], cost_usd=10)
        for a in (a1, a2):
            F.item(db, a, f"20{i:04d}", f"{100 + i}")
    return a1, a2


def _rows(content):
    ws = load_workbook(io.BytesIO(content)).active
    return [[c.value for c in r] for r in ws.iter_rows()]


# --- п. 1: обработчик с синхронной работой не занимает цикл событий ------------------

def test_no_route_handler_is_async():
    """`async def` обработчик выполняется В цикле событий: его синхронная работа с
    базой замораживала весь интерфейс (6,5 с на «Справку» во время «Установить»).
    Обычный `def` FastAPI уводит в пул потоков. Форма — зависимостью `posted_form`."""
    bad = []
    for path in ROUTERS.glob("*.py"):
        for node in ast.walk(ast.parse(path.read_text(encoding="utf-8"))):
            if isinstance(node, ast.AsyncFunctionDef) and any(
                    isinstance(d, ast.Call) and getattr(d.func, "attr", "") in ("get", "post")
                    for d in node.decorator_list):
                bad.append(f"{path.name}:{node.lineno} {node.name}")
    assert not bad, bad


# --- п. 2: массовая передача пишет порциями -------------------------------------------

def test_send_commits_in_portions(client, db, monkeypatch):
    from priceapp.database import SessionLocal
    from priceapp.routers import sku_prices
    a1, a2 = _wb_two(db, n=5)
    monkeypatch.setattr(sku_prices, "COMMIT_EVERY", 2)
    commits = []
    listener = lambda s: commits.append(1)  # noqa: E731
    event.listen(SessionLocal, "after_commit", listener)
    try:
        client.post("/sku-prices/send", data={"scope": "wb", "all_filtered": "1", "confirm_large": "1"},
                    follow_redirects=False)
    finally:
        event.remove(SessionLocal, "after_commit", listener)
    assert db.query(PriceChange).filter(PriceChange.status == "approved").count() == 10
    assert len(commits) >= 5          # 10 цен порциями по 2, а не одной транзакцией


def test_send_still_supersedes_earlier_decision(client, db):
    a1, a2 = _wb_two(db, n=1)
    db.add(PriceChange(item_id="i0", account_id=a1.id, barcode="200000", new_price=1, status="approved"))
    db.commit()
    client.post("/sku-prices/send", data={"scope": f"a{a1.id}", "all_filtered": "1", "confirm_large": "1"})
    db.expire_all()
    got = sorted((c.new_price, c.status) for c in db.query(PriceChange).filter_by(account_id=a1.id))
    assert got[0] == (1, "rejected") and got[1][1] == "approved"


def test_recalculate_rejects_old_proposals_and_adds_new(client, db):
    from priceapp.pricing import recalculate_account
    a1, _ = _wb_two(db, n=3)
    db.add(PriceChange(item_id="i0", account_id=a1.id, barcode="200000", new_price=1, status="proposed"))
    db.commit()
    st = recalculate_account(db, a1)
    db.expire_all()
    assert db.query(PriceChange).filter_by(account_id=a1.id, status="rejected").count() == 1
    assert db.query(PriceChange).filter_by(account_id=a1.id, status="proposed").count() == st.proposed == 3


# --- п. 3: круг «выгрузил — загрузил» в кабинете ничего не меняет ---------------------

def test_cabinet_roundtrip_does_not_copy_platform_markup_into_cabinet(client, db):
    a1, _ = _wb_two(db)
    F.coef(db, "100", "wb", "3")                       # наценка артикула на ПЛОЩАДКЕ
    content = client.get(f"/sku-prices/export?scope=a{a1.id}").content
    rows = _rows(content)
    h = rows[0]
    row = next(r for r in rows[1:] if r[h.index("Артикул")] == "100")
    assert row[h.index("Наценка")] is None                         # не «своё» для кабинета
    assert "площадка" in row[h.index("Откуда наценка")]
    client.post("/sku-prices/import", files={"file": ("f.xlsx", io.BytesIO(content), "x")})
    db.expire_all()
    assert [(c.article, c.account_id) for c in db.query(ArticleCoef)] == [("100", 0)]


def test_platform_export_carries_platform_markup(client, db):
    a1, _ = _wb_two(db)
    F.coef(db, "100", "wb", "3")
    rows = _rows(client.get("/sku-prices/export?scope=wb").content)
    h = rows[0]
    row = next(r for r in rows[1:] if r[h.index("Артикул")] == "100")
    assert row[h.index("Наценка")] == 3


def test_cabinet_row_input_is_empty_for_inherited_markup(client, db):
    a1, _ = _wb_two(db)
    F.coef(db, "100", "wb", "3")
    html = client.get(f"/sku-prices?scope=a{a1.id}").text
    assert 'name="value" value="3"' not in html.replace("\n", " ").replace("  ", " ")


# --- п. 4: ручные цены не ложатся в чужой кабинет -------------------------------------

def test_manual_prices_file_of_one_cabinet_is_refused_in_another(client, db):
    a1, a2 = _wb_two(db)
    client.post(f"/prices/manual/{a1.id}", data={"item_id": "i0", "value": "1999"})
    content = client.get(f"/prices/export/{a1.id}").content
    r = client.post(f"/prices/import/{a2.id}", files={"file": ("f.xlsx", io.BytesIO(content), "x")})
    assert "из другого кабинета" in r.text and "ИП 1" in r.text
    assert db.query(ProductPrice).filter(ProductPrice.account_id == a2.id,
                                         ProductPrice.manual_price.isnot(None)).count() == 0


def test_manual_prices_file_without_cabinet_column_is_refused(client, db):
    from openpyxl import Workbook
    a1, _ = _wb_two(db)
    wb = Workbook()
    wb.active.append(["ID_1С", "Ручная цена, ₽"])
    wb.active.append(["i0", 1999])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(f"/prices/import/{a1.id}", files={"file": ("f.xlsx", io.BytesIO(buf.getvalue()), "x")})
    assert "Кабинет (код)" in r.text
    assert db.query(ProductPrice).filter(ProductPrice.manual_price.isnot(None)).count() == 0


# --- п. 6: расчёт «Цен товаров» не повторяется, пока данные не менялись ---------------

def _count_computes(monkeypatch):
    from priceapp.routers import sku_prices
    calls = []
    real = sku_prices._compute_rows
    monkeypatch.setattr(sku_prices, "_compute_rows", lambda db, sc: calls.append(1) or real(db, sc))
    return calls


def test_rows_are_reused_between_requests_without_changes(client, db, monkeypatch):
    _wb_two(db)
    calls = _count_computes(monkeypatch)
    client.get("/sku-prices?scope=wb")
    client.get("/sku-prices?scope=wb&art=100")          # другой отбор — тот же расчёт
    client.get("/sku-prices/export?scope=wb")
    assert len(calls) == 1


def test_any_price_affecting_commit_invalidates_rows(client, db, monkeypatch):
    from priceapp import settings
    a1, _ = _wb_two(db)
    calls = _count_computes(monkeypatch)
    client.get("/sku-prices?scope=wb")
    settings.put(db, "attention_dirty", "1")             # «Внимание» на цену не влияет
    db.commit()
    client.get("/sku-prices?scope=wb")
    assert len(calls) == 1
    settings.put(db, settings.RATE_MANUAL, "90")          # курс — влияет
    db.commit()
    client.get("/sku-prices?scope=wb")
    assert len(calls) == 2
    pp = ProductPrice(item_id="i0", account_id=a1.id, manual_price=1234)   # как фоновый поток
    db.add(pp)
    db.commit()
    assert "1234" in client.get("/sku-prices?scope=wb").text
    assert len(calls) == 3


def test_row_ok_computes_once(client, db, monkeypatch):
    _wb_two(db)
    client.get("/sku-prices?scope=wb")
    calls = _count_computes(monkeypatch)
    r = client.post("/sku-prices/coef", data={"scope": "wb", "row": "100", "value": "3", "kind": "coef"})
    assert r.status_code == 200 and len(calls) == 1      # POST из кэша, страница после — заново
    db.expire_all()
    assert db.query(ArticleCoef).filter_by(article="100").one().coef == 3


# --- п. 5: отметки переживают смену отбора --------------------------------------------

def test_marked_rows_outside_current_filter_are_still_acted_on(client, db):
    """Отметка, сделанная при другом отборе, уезжает скрытым полем
    (static/selection.js) — и сервер обязан её найти, а не молча выкинуть, потому
    что её нет в текущем отборе."""
    _wb_two(db, n=3)
    r = client.post("/sku-prices/coef", data={"scope": "wb", "art": "102", "action": "set", "kind": "coef",
                                              "value": "3", "arts": ["100", "101"]})
    assert "артикулов 2, изменено 2" in r.text
    db.expire_all()
    assert sorted(c.article for c in db.query(ArticleCoef)) == ["100", "101"]


def test_pages_carry_selection_hooks(client, db):
    a1, _ = _wb_two(db)
    html = client.get("/sku-prices?scope=wb").text
    assert 'data-sel-key="sku:wb:"' in html and 'class="sel-count' in html and "selConfirm(" in html
    assert "/static/selection.js" in html
    html = client.get(f"/prices?view=products&account_id={a1.id}").text
    assert f'data-sel-key="prod:{a1.id}"' in html and "selConfirm(" in html
    js = client.get("/static/selection.js").text
    assert "offscreen" in js and "sessionStorage" in js
