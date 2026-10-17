"""Порог скидки от эталона WB (страница «Порог скидки WB»): порог по категории и артикулу, на все
кабинеты или на один; «Отправить сейчас» — цена «эталон − N%»; дальше обычный сценарий: ниже
порога — возврат ПОЛНОГО эталона, курс держит сдвиг. Пример владельца 17.10: порог на все 10%,
на «Шорты» 20%, сейчас «эталон − 15%»."""
import io
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from priceapp import accounts as acc_mod, etalon, reference, settings, wb_discount as wd
from priceapp.models import AuditLog, PlatformItem, PriceChange, ReferencePrice, WbDiscountRule
from tests import factories as f


class _Wb:
    def __init__(self):
        self.calls = []

    def _post(self, url, body):
        self.calls.append(body)
        return {"error": False}


def _setup(db, category="Шорты", name="ИП 1", nm="5", art="A1", item="u1", barcode="b1", acc=None, cost="10"):
    """Эталон 2009 ₽ без скидки; цена на WB сейчас — эталон. Пол: 800 ₽ × 1,3 / 0,75 = 1387 ₽."""
    if acc is None:
        f.manual_rate(db, "80")
        acc = f.account(db, "wb", name, commission=25)
        f.rule(db, acc, base_coef="2.5", min_markup_coef="1.3", max_change_percent=1000)
        acc_mod.set_credential(db, acc, "token", "t")
    from priceapp.models import OnecCost
    f.sku(db, item, art, "S", barcodes=[barcode], cost_usd=None if db.get(OnecCost, item) else cost)
    f.item(db, acc, barcode, art, external_id=f"{nm}:1", size="S")
    row = db.query(PlatformItem).filter_by(barcode=barcode).one()
    row.current_price = row.current_sale_price = 2009
    row.category = category
    db.commit()
    assert reference.record(db, acc, item, 2009, [row])
    db.commit()
    return acc


def _cur(db, sale, barcode="b1"):
    row = db.query(PlatformItem).filter_by(barcode=barcode).one()
    row.current_price = row.current_sale_price = sale
    db.commit()


def _run(db):
    wb = _Wb()
    st = reference.run(db, client_factory=lambda p, c: wb)
    return st, wb


def _ref(db, item="u1"):
    db.expire_all()
    return db.query(ReferencePrice).filter_by(item_id=item).one()


def test_threshold_order_article_over_category_cabinet_over_all(db):
    a1 = _setup(db)
    a2 = f.account(db, "wb", "ИП 2")
    t = wd.Thresholds(db)
    assert t.for_article("A1", "Шорты", a1.id) == (Decimal(10), "по умолчанию")
    wd.set_threshold(db, wd.CATEGORY, "Шорты", Decimal(20), "op")
    wd.set_threshold(db, wd.CATEGORY, "шорты ", Decimal(25), "op", account_id=a2.id)   # регистр и пробел — та же
    wd.set_threshold(db, wd.ARTICLE, "A1", Decimal(30), "op", account_id=a2.id)
    db.commit()
    t = wd.Thresholds(db)
    assert t.for_article("A1", "Шорты", a1.id) == (Decimal(20), "категория")
    assert t.for_article("A9", "Шорты", a2.id) == (Decimal(25), "категория, кабинет")
    assert t.for_article("A1", "Шорты", a2.id) == (Decimal(30), "артикул, кабинет")
    assert t.for_item("u1", "Шорты", a2.id)[0] == Decimal(30)       # SKU 1С -> артикул
    wd.set_threshold(db, wd.ARTICLE, "A1", Decimal(12), "op")
    db.commit()
    assert wd.Thresholds(db).for_article("A1", "Шорты", a1.id) == (Decimal(12), "артикул")
    hist = [h.details for h in db.query(AuditLog).filter_by(action="wb_threshold")]
    assert len(hist) == 4 and "порог общий -> 20" in hist[0]


def test_category_threshold_lets_promo_stay_and_returns_full_etalon_beyond_it(db):
    _setup(db)
    _cur(db, 1700)                          # −15,4%: при общем пороге 10% — вернули бы эталон
    wd.set_threshold(db, wd.CATEGORY, "Шорты", Decimal(20), "op")
    db.commit()
    st, wb = _run(db)
    assert st["restored"] == 0 and wb.calls == []
    _cur(db, 1500)                          # −25% — ниже порога 20%
    st, wb = _run(db)
    assert st["restored"] == 1 and wb.calls == [{"data": [{"nmID": 5, "price": 2009, "discount": 0}]}]


def test_cabinet_threshold_only_in_its_cabinet(db):
    a1 = _setup(db)
    a2 = f.account(db, "wb", "ИП 2", commission=25)
    acc_mod.set_credential(db, a2, "token", "t")
    _setup(db, acc=a2, nm="7", art="A2", item="u2", barcode="b2")
    wd.set_threshold(db, wd.CATEGORY, "Шорты", Decimal(20), "op", account_id=a1.id)
    db.commit()
    _cur(db, 1700, "b1")
    _cur(db, 1700, "b2")
    st, wb = _run(db)
    assert st["restored"] == 1 and wb.calls == [{"data": [{"nmID": 7, "price": 2009, "discount": 0}]}]


def test_send_now_sets_etalon_minus_n_and_standard_scenario_follows(db):
    """Пример владельца: порог 20%, «эталон − 15%» — уходит 1708; эталон не меняется; акция до
    −20% держится; ниже — полный эталон, сдвиг снят."""
    a = _setup(db)
    wd.set_threshold(db, wd.CATEGORY, "Шорты", Decimal(20), "op")
    db.commit()
    wb = _Wb()
    st = wd.send_now(db, wd.CATEGORY, "Шорты", Decimal(15), "op", client_factory=lambda p, c: wb)
    assert st["sent"] == 1 and wb.calls == [{"data": [{"nmID": 5, "price": 1708, "discount": 0}]}]
    ref = _ref(db)
    assert ref.ref_price == 2009 and ref.held_pct == Decimal("15.00")        # эталон тот же
    ch = db.query(PriceChange).filter_by(source="wb_now").one()
    assert ch.status == "sent" and ch.new_price == 1708
    assert etalon.etalons(db)["A1"].sale == 2009                             # последователи — за эталоном
    rule = db.query(WbDiscountRule).filter_by(kind="category").one()
    assert rule.now_pct == Decimal(15) and rule.now_by == "op" and rule.max_disc == Decimal(20)
    assert db.query(AuditLog).filter_by(action="wb_send_now").count() == 1
    # опрос: на WB 1708 (−15%) — в пределах порога, ничего не уходит
    st, wb = _run(db)
    assert st["restored"] == 0 and wb.calls == []
    _cur(db, 1550)                                                          # акция −22,8%
    st, wb = _run(db)
    assert wb.calls == [{"data": [{"nmID": 5, "price": 2009, "discount": 0}]}]
    assert _ref(db).held_pct is None
    assert a.id


def test_send_now_refuses_n_not_below_threshold_floor_and_big_step(db):
    _setup(db)
    wb = _Wb()
    st = wd.send_now(db, wd.CATEGORY, "Шорты", Decimal(10), "op", client_factory=lambda p, c: wb)
    assert st["sent"] == 0 and wb.calls == [] and "меньше порога" in st["skipped"][0]
    wd.set_threshold(db, wd.CATEGORY, "Шорты", Decimal(50), "op")
    db.commit()
    st = wd.send_now(db, wd.CATEGORY, "Шорты", Decimal(40), "op", client_factory=lambda p, c: wb)
    assert st["sent"] == 0 and "за раз" in st["skipped"][0]                 # 2009 -> 1205: больше 33%
    _cur(db, 1500)
    st = wd.send_now(db, wd.CATEGORY, "Шорты", Decimal(32), "op", client_factory=lambda p, c: wb)
    assert st["sent"] == 0 and "ниже пола" in st["skipped"][0]              # 1366 < 1387
    assert wb.calls == [] and _ref(db).held_pct is None


def test_rate_revaluation_keeps_the_shift(db):
    _setup(db)
    wd.set_threshold(db, wd.ARTICLE, "A1", Decimal(20), "op")
    db.commit()
    wd.send_now(db, wd.ARTICLE, "A1", Decimal(15), "op", client_factory=lambda p, c: _Wb())
    f.manual_rate(db, "84")                                                 # +5%
    st, wb = _run(db)
    ref = _ref(db)
    assert ref.ref_price == 2109 and ref.held_pct == Decimal("15.00")
    assert wb.calls == [{"data": [{"nmID": 5, "price": 1793, "discount": 0}]}]   # 2109 × 0,85


def test_new_etalon_by_hand_or_by_queue_drops_the_shift(db):
    a = _setup(db)
    wd.set_threshold(db, wd.ARTICLE, "A1", Decimal(20), "op")
    db.commit()
    wd.send_now(db, wd.ARTICLE, "A1", Decimal(15), "op", client_factory=lambda p, c: _Wb())
    rows = db.query(PlatformItem).all()
    reference.record(db, a, "u1", 2100, rows)
    db.commit()
    assert _ref(db).held_pct is None
    wd.send_now(db, wd.ARTICLE, "A1", Decimal(15), "op", client_factory=lambda p, c: _Wb())
    assert _ref(db).held_pct == Decimal("15.00")
    etalon.set_etalon(db, "A1", 2200, "op", client_factory=lambda p, c: _Wb())
    assert _ref(db).held_pct is None and _ref(db).ref_price > 2100


def test_page_bulk_and_history(client, db):
    a = _setup(db)
    _setup(db, acc=a, nm="6", art="A2", item="u2", barcode="b2", category="Джемперы")
    html = client.get("/wb-discount").text
    assert "Порог скидки WB" in html and "Шорты" in html and "Джемперы" in html
    r = client.post("/wb-discount/bulk", data={"by": "category", "sel": ["c:Шорты"], "action": "set", "value": "20"})
    assert r.status_code == 200 and "порог 20%" in r.text
    assert wd.Thresholds(db).for_category("Шорты") == (Decimal(20), "категория")
    # весь отбор — только то, что отобрано (поиск «Джем»)
    client.post("/wb-discount/bulk", data={"by": "category", "q": "Джем", "all_filtered": "1", "action": "set",
                                           "value": "25"})
    db.expire_all()
    t = wd.Thresholds(db)
    assert t.for_category("Джемперы")[0] == Decimal(25) and t.for_category("Шорты")[0] == Decimal(20)
    # отбор по кабинету: порог — этому кабинету
    client.post("/wb-discount/bulk", data={"acc": str(a.id), "by": "article", "sel": ["a:A1"], "action": "set",
                                           "value": "30"})
    db.expire_all()
    assert wd.Thresholds(db).for_article("A1", "Шорты", a.id) == (Decimal(30), "артикул, кабинет")
    html = client.get(f"/wb-discount?acc={a.id}&by=article").text
    assert "артикул, кабинет" in html
    # «Отправить сейчас» строкой таблицы
    accs = acc_mod.client_for
    try:
        wb = _Wb()
        acc_mod.client_for = lambda db_, acc_, cf=None: wb
        r = client.post("/wb-discount/save", data={"by": "category", "kind": "category", "key": "Шорты",
                                                   "thr": "20", "now": "15", "action": "send"})
    finally:
        acc_mod.client_for = accs
    assert "отправлено карточек 1 из 1" in r.text and wb.calls[0]["data"][0]["price"] == 1708
    html = client.get("/wb-discount").text
    assert "История изменений" in html and "эталон − 15%" in html
    hist = client.get("/wb-discount/history")
    assert hist.status_code == 200
    rows = list(load_workbook(io.BytesIO(hist.content)).active.iter_rows(values_only=True))
    assert rows[0][:4] == tuple(wd.HISTORY_HEAD) and len(rows) >= 5


def _three(db):
    """ИП 1: A1 и A2 — «Шорты», A3 — «Джемперы»; ИП 2: A1 — общий артикул двух кабинетов."""
    a1 = _setup(db, nm="5", art="A1", item="u1", barcode="b1")
    _setup(db, acc=a1, nm="6", art="A2", item="u2", barcode="b2")
    _setup(db, acc=a1, nm="7", art="A3", item="u3", barcode="b3", category="Джемперы")
    a2 = f.account(db, "wb", "ИП 2", commission=25)
    acc_mod.set_credential(db, a2, "token", "t")
    _setup(db, acc=a2, nm="8", art="A1", item="u1", barcode="b8")
    return a1, a2


def _sheet(content):
    ws = load_workbook(io.BytesIO(content)).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _file(rows):
    """Файл для импорта: [(тип, ключ, где, новый порог, эталон −)] в колонках выгрузки."""
    wb = Workbook()
    out = wb.active
    out.append(wd.HEAD)
    for kind, key, where, thr, now in rows:
        row = [None] * len(wd.HEAD)
        row[wd.HEAD.index(wd.COL_KIND)], row[wd.HEAD.index(wd.COL_KEY)] = kind, key
        row[wd.HEAD.index(wd.COL_WHERE)] = where
        row[wd.HEAD.index(wd.COL_NEW)], row[wd.HEAD.index(wd.COL_NOW)] = thr, now
        out.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_export_lists_shared_articles_then_categories_with_their_articles(client, db):
    _three(db)
    rows = _sheet(client.get("/wb-discount/export").content)
    assert rows[0] == wd.HEAD
    kinds = [(r[1], r[2]) for r in rows[1:]]
    assert rows[1][0].startswith("ОБЩИЕ АРТИКУЛЫ") and rows[1][1] is None          # заголовок раздела
    assert kinds[1:] == [("Артикул", "A1"), ("Категория", "Джемперы"), ("Артикул", "A3"),
                         ("Категория", "Шорты"), ("Артикул", "A2")]
    shared = rows[2]
    assert shared[wd.HEAD.index("Кабинеты")] == "ИП 1, ИП 2" and shared[0] == wd.SHARED_TITLE
    # править — в выделенных колонках, и выгружаются они ПУСТЫМИ: что вписали, то и меняется
    assert all(r[wd.HEAD.index(wd.COL_NEW)] is None and r[wd.HEAD.index(wd.COL_NOW)] is None for r in rows[1:])
    html = client.get("/wb-discount").text
    assert "Общие артикулы — есть в нескольких кабинетах" in html and "кабинеты: ИП 1, ИП 2" in html


def test_new_category_value_goes_to_all_its_articles_but_not_to_shared(db):
    a1, a2 = _three(db)
    wd.set_threshold(db, wd.ARTICLE, "A2", Decimal(30), "op")
    wd.set_threshold(db, wd.ARTICLE, "A2", Decimal(12), "op", account_id=a1.id)
    wd.set_threshold(db, wd.CATEGORY, "Шорты", Decimal(15), "op", account_id=a1.id)
    wd.set_threshold(db, wd.ARTICLE, "A1", Decimal(25), "op")                  # общий — свой
    db.commit()
    changed, cleared = wd.set_category(db, "Шорты", Decimal(20), "op")
    db.commit()
    assert changed and cleared == 3
    t = wd.Thresholds(db)
    assert t.for_article("A2", "Шорты", a1.id) == (Decimal(20), "категория")
    assert t.for_article("A1", "Шорты", a1.id) == (Decimal(25), "артикул")       # общий список — свой
    assert any("каскад категории Шорты" in h.object for h in db.query(AuditLog).filter_by(action="wb_threshold"))
    # «-» у категории снимает только её порог — исключения артикулов остаются
    wd.set_threshold(db, wd.ARTICLE, "A2", Decimal(11), "op")
    db.commit()
    assert wd.set_category(db, "Шорты", None, "op") == (True, 0)
    assert wd.Thresholds(db).for_article("A2", "Шорты", a1.id) == (Decimal(11), "артикул")
    # раздел кабинета снимает свои пороги артикулов только в этом кабинете
    wd.set_threshold(db, wd.ARTICLE, "A2", Decimal(13), "op", account_id=a1.id)
    db.commit()
    assert wd.set_category(db, "Шорты", Decimal(18), "op", account_id=a1.id) == (True, 1)
    assert wd.Thresholds(db).for_article("A2", "Шорты", a1.id) == (Decimal(11), "артикул")   # на все кабинеты — главнее
    assert a2.id


def test_import_category_first_then_article_exception_and_only_changed_rows(client, db):
    a1, _ = _three(db)
    wd.set_threshold(db, wd.ARTICLE, "A3", Decimal(33), "op")                  # не трогаем в файле — останется
    wd.set_threshold(db, wd.ARTICLE, "A2", Decimal(30), "op")
    db.commit()
    content = _file([("Артикул", "A2", "Все кабинеты WB", 12, None),          # исключение — строкой НИЖЕ категории
                     ("Категория", "Шорты", "Все кабинеты WB", 20, None),
                     ("Категория", "Джемперы", "Все кабинеты WB", None, None)])
    r = client.post("/wb-discount/import", files={"file": ("p.xlsx", content)})
    assert r.status_code == 200 and "у категорий — 1, у артикулов — 1" in r.text
    db.expire_all()
    t = wd.Thresholds(db)
    assert t.for_category("Шорты")[0] == Decimal(20)
    assert t.for_article("A2", "Шорты", a1.id) == (Decimal(12), "артикул")      # каскад, затем исключение
    assert t.for_article("A3", "Джемперы", a1.id) == (Decimal(33), "артикул")   # не меняли
    # файл старого образца (rp42: порог в нём стоял заполненным) — отказ, а не каскад по всему каталогу
    wb = Workbook()
    wb.active.append(["Тип", "Категория WB / артикул 1С", "Кабинет", wd.OLD_FORMAT_COL])
    wb.active.append(["Категория", "Шорты", "Все кабинеты WB", 25])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/wb-discount/import", files={"file": ("old.xlsx", buf.getvalue())})
    assert "старого образца" in r.text
    db.expire_all()
    assert wd.Thresholds(db).for_category("Шорты")[0] == Decimal(20)


def test_import_send_now_and_errors(client, db):
    a = _setup(db)
    content = _file([("категория", "шорты", "Все кабинеты WB", 20, 15), ("Артикул", "A1", a.name, 18, None),
                     ("Артикул", "НЕТ", "", 10, None), ("Категория", "Шорты", "Все кабинеты WB", 25, None),
                     ("Артикул", "A1", "ИП Неизвестный", 10, None)])
    fake = _Wb()
    accs = acc_mod.client_for
    try:
        acc_mod.client_for = lambda db_, acc_, cf=None: fake
        r = client.post("/wb-discount/import", files={"file": ("p.xlsx", content)})
    finally:
        acc_mod.client_for = accs
    db.expire_all()
    t = wd.Thresholds(db)
    assert t.for_category("Шорты")[0] == Decimal(20)                           # повтор строки не применён
    assert t.for_article("A1", "Шорты", a.id) == (Decimal(18), "артикул, кабинет")
    # у A1 в кабинете свой порог — «отправить сейчас» категории его не трогает
    assert fake.calls == [] and "свой порог артикула" in r.text
    assert "НЕТ" in r.text and "повторяется" in r.text and "ИП Неизвестный" in r.text
    content = _file([("Артикул", "A1", a.name, "-", None), ("Категория", "Шорты", "Все кабинеты WB", None, 15)])
    try:
        acc_mod.client_for = lambda db_, acc_, cf=None: fake
        client.post("/wb-discount/import", files={"file": ("p.xlsx", content)})
    finally:
        acc_mod.client_for = accs
    db.expire_all()
    assert wd.Thresholds(db).for_article("A1", "Шорты", a.id) == (Decimal(20), "категория")
    assert fake.calls == [{"data": [{"nmID": 5, "price": 1708, "discount": 0}]}]


def test_bulk_does_not_set_articles_covered_by_their_category(client, db):
    a1, _ = _three(db)
    wd.set_threshold(db, wd.ARTICLE, "A2", Decimal(30), "op")
    db.commit()
    client.post("/wb-discount/bulk", data={"sel": ["c:Шорты", "a:A2", "a:A1"], "action": "set", "value": "22"})
    db.expire_all()
    t = wd.Thresholds(db)
    assert t.for_article("A2", "Шорты", a1.id) == (Decimal(22), "категория")    # покрыт категорией, не свой
    assert t.for_article("A1", "Шорты", a1.id) == (Decimal(22), "артикул")      # общий — правится сам
    assert db.query(WbDiscountRule).filter_by(kind="article", key="A2").one().max_disc is None


def test_default_threshold_on_the_page(client, db):
    _setup(db)
    client.post("/wb-discount/default", data={"value": "12"})
    db.expire_all()
    assert settings.get(db, settings.REF_DROP_PERCENT) == "12"
    assert db.query(AuditLog).filter_by(action="wb_threshold_default").count() == 1
    r = client.post("/wb-discount/default", data={"value": "0"})
    assert "больше 0" in r.text and settings.get(db, settings.REF_DROP_PERCENT) == "12"
