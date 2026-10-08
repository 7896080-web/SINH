"""«Цены товаров»: цена = базовая (себестоимость × курс) × коэффициент артикула.

Выбрали площадку или кабинет — задали коэффициент артикулу (сразу всем его
размерам) — увидели новую цену и маржинальность рядом с текущими — передали."""
import io
from decimal import Decimal

from openpyxl import load_workbook

from priceapp import dispatch
from priceapp.models import ArticleCoef, PlatformItem, PriceChange, ProductPrice
from priceapp.pricing import load_inputs, coef_for
from tests import factories as f

RULE = {"commission_percent": "25", "base_coef": "2,5", "min_markup_coef": "1,3",
        "round_step": "10", "round_minus": "1", "max_change_percent": "20"}


def _setup(client, db):
    f.manual_rate(db)
    wb1 = f.account(db, "wb", "ИП Яворская")
    wb2 = f.account(db, "wb", "ИП Ребрик")
    oz = f.account(db, "ozon", "Озон", commission=30)
    client.post("/prices/rules/wb", data=RULE)
    client.post("/prices/rules/ozon", data={**RULE, "commission_percent": "30", "base_coef": "2,8"})
    f.sku(db, "u1", "39681", "L", color="GRI", barcodes=["2000932200000"], cost_usd="16.24", name="Свитшот")
    f.sku(db, "u2", "39681", "M", color="MAVI", barcodes=["2000932200001"], cost_usd="16.00", name="Свитшот")
    f.sku(db, "u3", "4033", "3XL", color="SIYAH", barcodes=["2000932200003"], cost_usd="12.40", name="Джемпер")
    for a in (wb1, wb2):
        f.item(db, a, "2000932200000", "39681-GRI", external_id="5:1")
        f.item(db, a, "2000932200001", "39681-MAVI", external_id="5:2")
        f.item(db, a, "2000932200003", "4033", external_id="6:1")
    f.item(db, oz, "2000932200000", "39681/GRI/L", external_id="o1")
    return wb1, wb2, oz


def _table(page: str) -> str:
    return page.split("<table class=\"art-table\"", 1)[1]


def test_rows_are_articles_base_is_cost_times_rate_and_not_editable(client, db):
    _setup(client, db)
    page = client.get("/sku-prices?scope=wb").text
    t = _table(page)
    # 39681 — одна строка на два размера; 16,24 × 81,5 = 1323,56 и 16,00 × 81,5 = 1304,00
    assert t.count('name="arts" value="39681"') == 1 and "2 разм." in t
    assert "1 304,00–1 323,56" in t
    # по умолчанию WB × 2,5: 1304 → 3269, 1323,56 → 3309; маржинальность 1,88
    assert "3269–3309" in t and "1,88" in t
    assert 'name="base"' not in page                          # базовая руками не правится
    assert "Wildberries — все кабинеты (2)" in page and "Wildberries: ИП Ребрик" in page
    # размеры раскрываются, коэффициент у них общий с артикулом
    t = _table(client.get("/sku-prices?scope=wb&by=sku").text)
    assert "MAVI M" in t and "GRI L" in t


def test_filters_by_article_barcode_size_keep_whole_article(client, db):
    _setup(client, db)
    t = _table(client.get("/sku-prices?scope=wb&art=4033").text)
    assert "4033" in t and "39681" not in t
    t = _table(client.get("/sku-prices?scope=wb&barcode=2200001").text)
    assert 'value="39681"' in t and "4033" not in t
    t = _table(client.get("/sku-prices?scope=wb&size=l").text)
    assert 'value="39681"' in t and "1 разм." in t


def test_coef_for_platform_reaches_every_size_and_every_cabinet(client, db):
    wb1, wb2, _ = _setup(client, db)
    r = client.post("/sku-prices/coef", data={"scope": "wb", "action": "set", "value": "2,2", "arts": ["39681"]})
    assert "изменено 1" in r.text
    assert [(c.article, c.platform, c.account_id, c.coef) for c in db.query(ArticleCoef)] == \
        [("39681", "wb", 0, Decimal("2.2"))]
    # 1323,56 × 2,2 = 2911,8 → 2919; 1304 × 2,2 = 2868,8 → 2869
    assert "2869–2919" in _table(r.text)
    client.post("/prices/recalculate", data={})
    got = {(c.account_id, c.item_id): c.new_price for c in db.query(PriceChange).filter_by(status="proposed")}
    assert got[(wb1.id, "u1")] == got[(wb2.id, "u1")] == 2919 and got[(wb1.id, "u2")] == 2869
    assert got[(wb1.id, "u3")] == 2529                        # 4033 — по умолчанию 2,5


def test_cabinet_coef_beats_platform_and_row_clear_returns_default(client, db):
    wb1, wb2, _ = _setup(client, db)
    client.post("/sku-prices/coef", data={"scope": "wb", "action": "set", "value": "2,2", "all_filtered": "1"})
    client.post("/sku-prices/coef", data={"scope": f"a{wb2.id}", "row": "39681", "value": "3"})
    inp = load_inputs(db, "wb")
    assert coef_for(inp, "u1", wb2.id) == (Decimal("3"), "cabinet")
    assert coef_for(inp, "u1", wb1.id) == (Decimal("2.2"), "article")
    page = client.get("/sku-prices?scope=wb").text
    assert "своя у: ИП Ребрик" in page
    # пустое поле в строке — снять: снова по умолчанию площадки
    client.post("/sku-prices/coef", data={"scope": f"a{wb2.id}", "row": "39681", "value": ""})
    assert coef_for(load_inputs(db, "wb"), "u1", wb2.id) == (Decimal("2.2"), "article")


def test_mult_and_validation(client, db):
    _setup(client, db)
    client.post("/sku-prices/coef", data={"scope": "wb", "action": "set", "value": "2", "arts": ["39681"]})
    client.post("/sku-prices/coef", data={"scope": "wb", "action": "mult", "value": "1,1", "arts": ["39681"]})
    assert db.query(ArticleCoef).one().coef == Decimal("2.2")
    assert "значение" in client.post("/sku-prices/coef", data={"scope": "wb", "action": "set", "value": "abc",
                                                               "arts": ["39681"]}).text
    assert "Ничего не отмечено" in client.post("/sku-prices/coef", data={"scope": "wb", "action": "clear"}).text


def test_margin_uses_platform_discount_and_floor_badge(client, db):
    wb1, wb2, _ = _setup(client, db)
    for it in db.query(PlatformItem).filter(PlatformItem.barcode == "2000932200003"):
        it.current_price, it.current_sale_price = 4000, 3000          # скидка продавца 25%
    db.commit()
    t = _table(client.get("/sku-prices?scope=wb&art=4033").text)
    # новая 2529 со скидкой 25% → 1896,75 × 0,75 / 1010,60 = 1,41; текущая 3000 × 0,75 / 1010,6 = 2,23
    assert "скидка 25%" in t and "1,41" in t and "2,23" in t
    client.post("/sku-prices/coef", data={"scope": "wb", "action": "set", "value": "2", "arts": ["4033"]})
    # 2029 × 0,75 × 0,75 / 1010,6 = 1,13 — ниже пола 1,3 (а без скидки было бы 1,51)
    t = _table(client.get("/sku-prices?scope=wb&flt=floor").text)
    assert "4033" in t and "ниже пола" in t and "39681" not in t


def test_send_queues_approved_for_every_cabinet_and_respects_limits(client, db):
    wb1, wb2, oz = _setup(client, db)
    db.add(ProductPrice(item_id="u2", account_id=wb1.id, last_sent_price=3269))   # уже стоит
    db.add(ProductPrice(item_id="u1", account_id=wb2.id, last_sent_price=2000))   # +65% — больше лимита
    db.add(PriceChange(item_id="u1", account_id=wb1.id, barcode="2000932200000", new_price=1, status="proposed"))
    db.commit()
    r = client.post("/sku-prices/send", data={"scope": "wb", "arts": ["39681"]})
    assert "передано на отправку 2 цен" in r.text and "уже стоит на площадке — 1" in r.text
    assert "изменение больше лимита — 1" in r.text
    q = {(c.account_id, c.item_id): c for c in db.query(PriceChange).filter_by(status="approved")}
    assert set(q) == {(wb1.id, "u1"), (wb2.id, "u2")}
    assert q[(wb1.id, "u1")].new_price == 3309 and q[(wb1.id, "u1")].decided_by == "op"
    assert q[(wb1.id, "u1")].markup_coef == Decimal("1.88") and "× 2,5" in q[(wb1.id, "u1")].note
    assert db.query(PriceChange).filter_by(status="proposed").count() == 0         # вытеснено
    assert all(c.account.platform == "wb" for c in q.values())                    # Озон не тронут
    r = client.post("/sku-prices/send", data={"scope": "wb", "arts": ["39681"], "confirm_large": "1"})
    assert "передано на отправку 3 цен" in r.text       # большой шаг с галочкой + повтор двух прежних

    class Ok:
        def push_prices(self, items):
            return {"ok": [i.barcode for i in items], "errors": [], "sent_prices": {}}
    dispatch.run_account(db, wb1, Ok())
    assert db.query(ProductPrice).filter_by(account_id=wb1.id, item_id="u1").one().last_sent_price == 3309


def test_send_never_below_floor(client, db):
    _setup(client, db)
    client.post("/sku-prices/coef", data={"scope": "wb", "action": "set", "value": "1,2", "arts": ["4033"]})
    r = client.post("/sku-prices/send", data={"scope": "wb", "arts": ["4033"], "confirm_large": "1"})
    assert "ниже пола, не отправлено — 2" in r.text and db.query(PriceChange).count() == 0


def test_excel_roundtrip_empty_changes_nothing(client, db):
    wb1, wb2, _ = _setup(client, db)
    client.post("/sku-prices/coef", data={"scope": "wb", "action": "set", "value": "2,2", "arts": ["39681"]})
    r = client.get("/sku-prices/export?scope=wb")
    book = load_workbook(io.BytesIO(r.content))
    ws = book.active
    h = [c.value for c in ws[1]]
    assert {"Где (код)", "Артикул", "Вид наценки", "Наценка", "Новая маржинальность (мин.)", "Текущая цена, ₽"} <= set(h)
    rows = {ws.cell(row=i, column=h.index("Артикул") + 1).value: i for i in range(2, ws.max_row + 1)}
    col = h.index("Наценка") + 1
    ws.cell(row=rows["39681"], column=col, value=None)                 # не трогать
    ws.cell(row=rows["4033"], column=col, value=2.7)
    ws.append([f"a{wb2.id}", "", "4033"] + [None] * (col - 4) + ["3"])   # свой у кабинета
    ws.append(["nope", "", "4033"] + [None] * (col - 4) + ["3"])
    buf = io.BytesIO()
    book.save(buf)
    r = client.post("/sku-prices/import", files={"file": ("k.xlsx", buf.getvalue())})
    assert "Наценок изменено: 2" in r.text and "«nope» — нет такой площадки" in r.text
    got = {(c.article, c.account_id): c.coef for c in db.query(ArticleCoef)}
    assert got == {("39681", 0): Decimal("2.2"), ("4033", 0): Decimal("2.7"), ("4033", wb2.id): Decimal("3")}
    ws.cell(row=rows["39681"], column=col, value="-")                  # снять
    buf = io.BytesIO()
    book.save(buf)
    client.post("/sku-prices/import", files={"file": ("k.xlsx", buf.getvalue())})
    assert db.query(ArticleCoef).filter_by(article="39681").count() == 0


def test_nav_has_page(client, db):
    assert 'href="/sku-prices"' in client.get("/attention").text


def test_terminal_platform_refusal_closes_at_once(db):
    class Refuses:
        def push_prices(self, items):
            return {"ok": [], "sent_prices": {},
                    "errors": [{"detail": "Kit: INVALID_PRICE", "terminal": True, "items": [i.barcode for i in items]}]}
    f.manual_rate(db)
    acc = f.account(db, "kit", "КИТ", commission=20)
    f.sku(db, "u1", "39681", barcodes=["b1"], cost_usd="16.24")
    f.item(db, acc, "b1", "39681", external_id="v1")
    db.add(PriceChange(item_id="u1", account_id=acc.id, barcode="b1", new_price=9999, status="approved"))
    db.commit()
    dispatch.run_account(db, acc, Refuses())
    ch = db.query(PriceChange).one()
    assert ch.status == "error" and ch.attempts == 1 and "INVALID_PRICE" in ch.last_error
