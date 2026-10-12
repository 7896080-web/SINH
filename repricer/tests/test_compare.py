"""Проверка одинаковых товаров по площадкам («Сравнение площадок»)."""
import io

from openpyxl import load_workbook

from tests import factories as F


def _setup(db):
    w1 = F.account(db, "wb", "ИП 1")
    w2 = F.account(db, "wb", "ИП 2")
    oz = F.account(db, "ozon", "Озон")
    F.rule(db, w1)
    F.rule(db, oz)
    F.manual_rate(db)
    F.sku(db, "same", "100", "M", barcodes=["2001"], cost_usd=10)     # одинаково везде
    F.sku(db, "diff", "200", "M", barcodes=["2002"], cost_usd=10)     # у ИП 2 своя наценка
    F.sku(db, "solo", "300", "M", barcodes=["2003"], cost_usd=10)     # только в одном кабинете
    for a in (w1, w2):
        F.item(db, a, "2001", "100"); F.item(db, a, "2002", "200")
    F.item(db, w1, "2003", "300")
    F.item(db, oz, "2001", "100")
    F.coef(db, "200", "wb", "3", account_id=w2.id)
    return w1, w2, oz


def test_summary_finds_cabinets_of_one_platform_that_differ(client, db):
    _setup(db)
    html = client.get("/prices?view=compare").text
    assert "Проверка одинаковых товаров" in html
    assert "общий SKU 1С хотя бы в двух кабинетах" in html and "): 2." in html      # solo не сравнивается
    assert "кабинеты одной площадки: разная наценка" in html
    rows = client.get("/prices?view=compare&flt=cabinets_differ").text.split('<table class="econ"', 1)[1]
    assert "200" in rows and ">100 " not in rows
    assert "× 3 (артикул, кабинет)" in rows and "× 2,5 (по умолчанию площадки)" in rows


def test_same_markup_everywhere_is_not_an_issue(client, db):
    _setup(db)
    rows = client.get("/prices?view=compare&flt=issues").text.split('<table class="econ"', 1)[1]
    assert rows.count("<tr>") == 2          # заголовок + товар 200


def test_compare_export_has_markup_and_issues(client, db):
    w1, w2, oz = _setup(db)
    ws = load_workbook(io.BytesIO(client.get("/prices/compare-export?flt=issues").content)).active
    h = [c.value for c in ws[1]]
    assert "ИП 2: наценка" in h and "Расхождения" in h and "Озон: маржин. новая" in h
    row = [c.value for c in ws[2]]
    assert row[h.index("ID_1С")] == "diff"
    assert row[h.index("ИП 2: наценка")] == "× 3 (артикул, кабинет)"
    assert "кабинеты одной площадки" in row[h.index("Расхождения")]
    assert isinstance(row[h.index("ИП 1: новая, ₽")], int)


def test_check_script_runs_read_only_on_cp1251_console(db, tmp_path):
    """Скрипт для офисного ПК: печатает сводку, пишет файл, базу не меняет. Консоль
    боевой машины — cp1251, вывод не должен обрываться."""
    import os
    import subprocess
    import sys
    from pathlib import Path
    from priceapp.models import ArticleCoef, PriceChange
    _setup(db)
    before = (db.query(ArticleCoef).count(), db.query(PriceChange).count())
    out = tmp_path / "check.xlsx"
    root = Path(__file__).resolve().parent.parent
    env = {**os.environ, "PYTHONIOENCODING": "cp1251"}
    r = subprocess.run([sys.executable, str(root / "scripts" / "check_same_items.py"), "--xlsx", str(out)],
                       capture_output=True, env=env, cwd=root)
    text = r.stdout.decode("cp1251")
    assert r.returncode == 0, r.stderr.decode("utf-8", "replace")
    assert "Одинаковых товаров (общий SKU 1С хотя бы в двух кабинетах): 2" in text
    assert "кабинеты одной площадки" in text and "x 3 (артикул, кабинет)" in text and "?" not in text
    ws = load_workbook(out).active
    assert ws.max_row == 2                      # заголовок + товар с расхождением
    db.expire_all()
    assert (db.query(ArticleCoef).count(), db.query(PriceChange).count()) == before
