"""Возврат скидки WB по копии базы (scripts/wb_restore_discount.py)."""
import importlib.util
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

from priceapp.database import engine
from priceapp.models import AuditLog, PlatformItem
from tests import factories as F

ROOT = Path(__file__).resolve().parent.parent
spec = importlib.util.spec_from_file_location("wbrd", ROOT / "scripts" / "wb_restore_discount.py")
wbrd = importlib.util.module_from_spec(spec)
sys.modules["wbrd"] = wbrd
spec.loader.exec_module(wbrd)


def _set(db, prices):
    for it in db.query(PlatformItem):
        it.current_price, it.current_sale_price = prices[it.barcode]
    db.commit()


def _two_dbs(db, tmp_path):
    w = F.account(db, "wb", "ИП 1")
    F.item(db, w, "a1", "100", external_id="111:1")
    F.item(db, w, "a2", "100", external_id="111:2")     # тот же nmID — одна карточка
    F.item(db, w, "b1", "200", external_id="222:1")
    F.item(db, w, "c1", "300", external_id="333:1")
    F.item(db, w, "d1", "400", external_id="444:1")
    _set(db, {"a1": (10000, 5000), "a2": (10000, 5000), "b1": (8000, 4800), "c1": (6000, 3000), "d1": (5000, 2500)})
    engine.dispose()
    before = tmp_path / "before.db"
    shutil.copy(Path(str(engine.url).replace("sqlite:///", "")), before)
    # WB сменил скидку у 111 (50 -> 70) и 222 (40 -> 55); у 333 сменилась и цена; 444 не тронут
    _set(db, {"a1": (10000, 3000), "a2": (10000, 3000), "b1": (8000, 3600), "c1": (7000, 2100), "d1": (5000, 2500)})
    return w, before


def test_plan_restores_only_the_discount_and_skips_cards_whose_price_moved(db, tmp_path):
    _, before = _two_dbs(db, tmp_path)
    after = Path(str(engine.url).replace("sqlite:///", ""))
    restore, moved = wbrd.plan(wbrd.load(before)[0], wbrd.load(after)[0])
    assert [(b.nm, b.discount, a.discount) for b, a in restore] == [("111", 50, 70), ("222", 40, 55)]
    assert [b.nm for b, _ in moved] == ["333"]


def test_report_only_by_default_and_file_per_cabinet(db, tmp_path):
    _, before = _two_dbs(db, tmp_path)
    out = tmp_path / "d.xlsx"
    r = subprocess.run([sys.executable, str(ROOT / "scripts" / "wb_restore_discount.py"), "--before", str(before),
                        "--xlsx", str(out)], capture_output=True, cwd=ROOT,
                       env={**__import__("os").environ, "PYTHONIOENCODING": "cp1251"})
    assert r.returncode == 0, r.stderr.decode("utf-8", "replace")
    text = r.stdout.decode("cp1251")
    assert "вернуть: 2 карточек" in text and "Ничего не отправлено" in text
    rows = list(load_workbook(out)["ИП 1"].iter_rows(values_only=True))
    got = {x[0]: x for x in rows[1:]}
    assert got[111][5] == 50 and got[222][5] == 40 and got[333][5] is None and 444 not in got


def test_send_puts_only_nm_and_discount(db, tmp_path):
    w, before = _two_dbs(db, tmp_path)
    after = Path(str(engine.url).replace("sqlite:///", ""))
    restore, _ = wbrd.plan(wbrd.load(before)[0], wbrd.load(after)[0])
    calls = []

    class Fake:
        def _post(self, url, body):
            calls.append((url, body))
            return {"data": {"id": 1}, "error": False}

    res = wbrd.send(db, restore, factory=lambda platform, creds: Fake())
    assert res["ИП 1"].startswith("скидка возвращена: 2")
    url, body = calls[0]
    assert url.endswith("/api/v2/upload/task")
    assert body == {"data": [{"nmID": 111, "discount": 50}, {"nmID": 222, "discount": 40}]}
    assert db.query(AuditLog).filter(AuditLog.action == "wb_restore_discount").count() == 1
    from priceapp.models import ReferenceDiscount      # отправленная скидка — эталонная для проверки скидки
    assert {r.nm_id: r.discount for r in db.query(ReferenceDiscount)} == {"111": 50, "222": 40}


def test_from_report_file_maps_1c_sku_to_card_and_sends_nothing_without_send(db, tmp_path):
    """«Как было» — из отчёта эталона: кабинет + ID_1С -> nmID по каталогу кабинета."""
    from openpyxl import Workbook
    w = F.account(db, "wb", "ИП 1")
    F.sku(db, "s1", "100", "M", barcodes=["a1"]); F.sku(db, "s2", "100", "L", barcodes=["a2"])
    F.sku(db, "s3", "200", "M", barcodes=["b1"])
    F.item(db, w, "a1", "100", external_id="111:1"); F.item(db, w, "a2", "100", external_id="111:2")
    F.item(db, w, "b1", "200", external_id="222:1")
    # сейчас на WB: у 111 скидка 70%, у 222 не менялась
    _set(db, {"a1": (10000, 3000), "a2": (10000, 3000), "b1": (8000, 4800)})
    book = Workbook()
    ws = book.active
    ws.title = "WB"
    ws.append(["Кабинет", "ID_1С", "Сейчас: цена до скидки, ₽", "Сейчас: цена со скидкой, ₽", "Сейчас: скидка, %"])
    ws.append(["ИП 1", "s1", 10000, 5000, 50]); ws.append(["ИП 1", "s2", 10000, 5000, 50])
    ws.append(["ИП 1", "s3", 8000, 4800, 40]); ws.append(["Чужой", "s3", 1, 1, 0])
    f = tmp_path / "et.xlsx"
    book.save(f)
    out = tmp_path / "d.xlsx"
    r = subprocess.run([sys.executable, str(ROOT / "scripts" / "wb_restore_discount.py"), "--from-xlsx", str(f),
                        "--xlsx", str(out)], capture_output=True, cwd=ROOT,
                       env={**__import__("os").environ, "PYTHONIOENCODING": "cp1251"})
    assert r.returncode == 0, r.stderr.decode("utf-8", "replace")
    text = r.stdout.decode("cp1251")
    assert "вернуть: 1 карточек" in text and "Ничего не отправлено" in text
    rows = list(load_workbook(out)["ИП 1"].iter_rows(values_only=True))
    assert [(x[0], x[4], x[5]) for x in rows[1:]] == [(111, 70, 50)]


def test_corridor_never_raises_the_sale_price():
    for lst in range(1000, 20000, 37):
        for sale in range(int(lst * 0.2), lst, max(1, lst // 53)):
            nl, d = wbrd.corridor(lst, sale, 35, 50)
            assert 35 <= d <= 50
            assert nl * (100 - d) / 100 <= sale + 0.5


def test_plan_sends_list_and_discount_per_card_and_skips_unsafe_cards(db, tmp_path):
    from openpyxl import Workbook
    w = F.account(db, "wb", "ИП 1")
    for sku_id, bc, nm in (("s1", "a1", "111:1"), ("s2", "a2", "111:2"), ("s3", "b1", "222:1"),
                           ("s4", "c1", "333:1"), ("s5", "c2", "333:2"), ("s6", "d1", "444:1")):
        F.sku(db, sku_id, "A" + nm[:3], barcodes=[bc])
        F.item(db, w, bc, "A" + nm[:3], external_id=nm)
    # сейчас на WB (WB сменил скидки); у 444 сменилась цена до скидки
    _set(db, {"a1": (10000, 3000), "a2": (10000, 3000), "b1": (9000, 6000), "c1": (6000, 3000),
              "c2": (6000, 3000), "d1": (7000, 3500)})
    book = Workbook()
    ws = book.active
    ws.title = "WB"
    ws.append(["Кабинет", "ID_1С", "Сейчас: цена до скидки, ₽", "Сейчас: цена со скидкой, ₽"])
    ws.append(["ИП 1", "s1", 10000, 5000]); ws.append(["ИП 1", "s2", 10000, 5000])   # 50% — в коридоре
    ws.append(["ИП 1", "s3", 9000, 6000])                                             # 33% — ниже коридора
    ws.append(["ИП 1", "s4", 6000, 3000]); ws.append(["ИП 1", "s5", 6000, 2500])       # размеры разные
    ws.append(["ИП 1", "s6", 5000, 2500])                                             # цена сменилась
    f = tmp_path / "et.xlsx"
    book.save(f)
    after = wbrd.load(Path(str(engine.url).replace("sqlite:///", "")))[0]
    cards = {c.nm: c for c in wbrd.load_plan(f, db, after, 35, 50)}
    assert (cards["111"].new_list, cards["111"].disc) == (10000, 50)
    assert (cards["222"].new_list, cards["222"].disc) == (9230, 35)
    assert cards["333"].skip.startswith("цены со скидкой у размеров разные")
    assert cards["444"].skip.startswith("цена до скидки на WB сменилась")
    calls = []

    class Fake:
        def _post(self, url, body):
            calls.append(body)
            return {"error": False}

    res = wbrd.send_plan(db, list(cards.values()), factory=lambda platform, creds: Fake())
    assert res["ИП 1"] == "отправлено карточек: 2"
    assert calls == [{"data": [{"nmID": 111, "price": 10000, "discount": 50},
                               {"nmID": 222, "price": 9230, "discount": 35}]}]
    from priceapp.models import ReferenceDiscount
    assert {r.nm_id: r.discount for r in db.query(ReferenceDiscount)} == {"111": 50, "222": 35}


def test_verify_compares_plan_with_loaded_wb_prices(db, tmp_path, capsys):
    from openpyxl import Workbook
    w = F.account(db, "wb", "ИП 1")
    F.sku(db, "s1", "A1", barcodes=["a1"]); F.sku(db, "s3", "A2", barcodes=["b1"])
    F.item(db, w, "a1", "A1", external_id="111:1"); F.item(db, w, "b1", "A2", external_id="222:1")
    # WB уже держит план у 111 (10000, 50%), у 222 — нет (план 9230 / 35%, а там 9000 / 33%)
    _set(db, {"a1": (10000, 5000), "b1": (9000, 6000)})
    book = Workbook()
    ws = book.active
    ws.title = "WB"
    ws.append(["Кабинет", "ID_1С", "Сейчас: цена до скидки, ₽", "Сейчас: цена со скидкой, ₽"])
    ws.append(["ИП 1", "s1", 10000, 5000]); ws.append(["ИП 1", "s3", 9000, 6000])
    f = tmp_path / "et.xlsx"
    book.save(f)
    after_path = Path(str(engine.url).replace("sqlite:///", ""))
    assert wbrd.main(["--plan-xlsx", str(f), "--verify", "--after", str(after_path)]) == 0
    out = capsys.readouterr().out
    assert "совпадает с планом: 1; НЕ совпадает: 1" in out and "222" in out and "план 9230, скидка 35%" in out


def test_adopt_current_makes_loaded_wb_prices_and_discounts_the_new_etalon(db):
    from priceapp import discount_guard, reference
    from priceapp.models import ReferenceDiscount, ReferencePrice
    w = F.account(db, "wb", "ИП 1")
    F.manual_rate(db, "100")
    F.sku(db, "s1", "A1", barcodes=["a1"], cost_usd=10)
    F.item(db, w, "a1", "A1", external_id="111:1")
    _set(db, {"a1": (13710, 6855)})
    db.add(ReferencePrice(account_id=w.id, item_id="s1", ref_price=13989))   # старый эталон
    db.commit()
    st = reference.adopt_current(db)
    discount_guard.adopt(db)
    db.expire_all()
    ref = db.query(ReferencePrice).one()
    assert st["recorded"] == 1 and ref.ref_price == 13710 and ref.source == "adopt"
    assert {r.nm_id: r.discount for r in db.query(ReferenceDiscount)} == {"111": 50}
