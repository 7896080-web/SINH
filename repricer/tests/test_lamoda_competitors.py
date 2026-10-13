"""Файл Ламоды «цены конкурентов» по мэппингу (scripts/lamoda_competitors.py)."""
import importlib
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

from priceapp.models import PlatformItem
from tests import factories as F

ROOT = Path(__file__).resolve().parent.parent


def _template(path):
    b = Workbook()
    ws = b.active
    ws.append(["Артикул Ламода", "Площадка конкурента", "Артикул или ссылка на площадке конкурента",
               "Уникальный товар", None, None, "Инструкция"])
    b.save(path)


def test_lamoda_card_gets_wb_link_by_mapping_and_unique_when_nowhere(db, tmp_path):
    w1 = F.account(db, "wb", "ИП 1")
    w2 = F.account(db, "wb", "ИП 2")
    oz = F.account(db, "ozon", "Озон")
    lam = F.account(db, "lamoda", "Ламода")
    for i, art, bc in (("a", "100", ["b1"]), ("a2", "100", ["b2"]), ("u", "300", ["b3"]), ("o", "400", ["b4"])):
        F.sku(db, i, art, "M", "Черный", barcodes=bc)
    F.item(db, w1, "b1", external_id="111:1")                     # один размер у ИП 1
    F.item(db, w2, "b1", external_id="222:1"); F.item(db, w2, "b2", external_id="222:2")   # оба у ИП 2
    F.item(db, oz, "b4", external_id="9")
    F.item(db, lam, "b1", external_id="MP1:s1"); F.item(db, lam, "b2", external_id="MP1:s2")
    F.item(db, lam, "b3", external_id="MP3:s1"); F.item(db, lam, "b4", external_id="MP4:s1")
    for it in db.query(PlatformItem):
        it.current_price, it.current_sale_price = 1000, 500
    db.commit()
    tpl, out = tmp_path / "t.xlsx", tmp_path / "o.xlsx"
    _template(tpl)
    sys.path.insert(0, str(ROOT / "scripts"))
    m = importlib.import_module("lamoda_competitors")
    assert m.main(["--template", str(tpl), "--out", str(out)]) == 0
    b = load_workbook(out)
    rows = {r[0]: r[:4] for r in b.worksheets[0].iter_rows(min_row=2, values_only=True) if r[0]}
    assert rows["MP1"] == ("MP1", "Wildberries", "https://www.wildberries.ru/catalog/222/detail.aspx", None)
    assert rows["MP3"] == ("MP3", None, None, "Да")
    assert "MP4" not in rows                                       # только на Ozon — не уникальный
    assert b.worksheets[0].cell(1, 7).value == "Инструкция"        # шаблон Ламоды не тронут
    det = {r[0]: r for r in b["Подробно"].iter_rows(min_row=2, values_only=True)}
    assert det["MP1"][9] == 2 and det["MP1"][11] == "ИП 1 111"
