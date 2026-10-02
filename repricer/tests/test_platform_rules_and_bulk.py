"""Правила площадки (одно на все её кабинеты), цена от другой площадки,
текущие цены с площадок, отбор и массовая правка, Excel на вкладках «Цен»."""
import io
import os
import sqlite3
import subprocess
import sys
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from priceapp import accounts as acc_mod, platforms
from priceapp.models import Account, PlatformItem, PlatformRule, PriceChange, ProductPrice
from priceapp.platforms import CurrentPrice, LamodaClient, PlatformError, build_client
from priceapp.pricing import decide, get_rule, recalculate_account
from tests import factories as f

RATE = Decimal("81.5")
RULE = {"commission_percent": "25", "markup_coef": "2", "min_markup_coef": "1,3",
        "round_step": "10", "round_minus": "1", "max_change_percent": "20"}


def _xlsx(content: bytes):
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    return wb, ws, [c.value for c in ws[1]]


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- расчёт от базовой площадки ----------------------------------------------------

def test_price_from_base_platform_times_coef(db):
    wb = f.account(db, "wb", "ИП Яворская", commission=25)
    f.rule(db, wb)
    oz = f.account(db, "ozon", "Озон", commission=30)
    rule = f.rule(db, oz, markup_coef=1)
    rule.base_platform, rule.base_coef = "wb", Decimal("1.1")
    db.commit()
    # WB: 16,24 × 81,5 = 1323,56; × 2 / 0,75 = 3529,49 → 3539. Ozon: 3539 × 1,1 = 3892,9 → 3899.
    d = decide(Decimal("16.24"), RATE, 30, rule, None, None, get_rule(db, "wb"))
    assert d.new_price == 3899 and d.source == "base"
    assert "3539" in d.note and "1,1" in d.note
    # Пол — по СВОЕЙ комиссии: 3899 × 0,7 = 2729,3 / 1323,56 = 2,06
    assert d.markup_coef == Decimal("2.06") and d.block_reason is None


def test_base_price_still_checked_against_own_floor(db):
    f.account(db, "wb", commission=25)
    wb_rule = f.rule(db, Account(platform="wb"))
    oz_rule = PlatformRule(platform="ozon", commission_percent=30, markup_coef=1, round_step=1, round_minus=0,
                           min_markup_coef=Decimal("2.5"), max_change_percent=20,
                           base_platform="wb", base_coef=Decimal("0.9"))
    d = decide(Decimal("16.24"), RATE, 30, oz_rule, None, None, wb_rule)
    assert d.block_reason == "floor"


def test_unconfigured_base_gives_no_price(db):
    base = PlatformRule(platform="wb", commission_percent=None, markup_coef=2, round_step=1, round_minus=0,
                        min_markup_coef=1, max_change_percent=20)
    oz = PlatformRule(platform="ozon", commission_percent=30, markup_coef=1, round_step=1, round_minus=0,
                      min_markup_coef=1, max_change_percent=20, base_platform="wb", base_coef=Decimal("1.1"))
    d = decide(Decimal("10"), RATE, 30, oz, None, None, base)
    assert d.new_price is None and "базовая площадка" in d.note


def test_one_rule_serves_every_wb_cabinet(db):
    f.manual_rate(db)
    accs = [f.account(db, "wb", n) for n in ("ИП Яворская", "ИП Ребрик", "ИП Караман")]
    f.rule(db, accs[0])
    f.sku(db, "u1", "39681", "L", barcodes=["b1"], cost_usd="16.24")
    for a in accs:
        f.item(db, a, "b1", "39681-L", external_id="5:1")
        recalculate_account(db, a)
    assert {c.new_price for c in db.query(PriceChange)} == {3539}
    assert db.query(PlatformRule).count() == 1


# --- страница правил ---------------------------------------------------------------

def test_rules_page_one_block_per_platform_with_all_cabinets(client, db):
    for n in ("ИП Яворская", "ИП Ребрик", "ИП Караман"):
        f.account(db, "wb", n)
    f.account(db, "ozon", "Озон", commission=None)
    f.account(db, "lamoda", "Ламода", commission=None)
    page = client.get("/prices?view=rules").text
    assert page.count('action="/prices/rules/wb"') == 1
    assert "ИП Яворская, ИП Караман, ИП Ребрик" in page or "ИП Караман, ИП Ребрик, ИП Яворская" in page
    assert 'action="/prices/rules/lamoda"' in page


def test_rule_with_base_and_chain_refused(client, db):
    f.account(db, "wb")
    f.account(db, "ozon", "Озон")
    client.post("/prices/rules/wb", data=RULE)
    r = client.post("/prices/rules/ozon", data={**RULE, "markup_coef": "1", "base_platform": "wb", "base_coef": "1,1"})
    assert "сохранено" in r.text and "не сохранено" not in r.text
    db.expire_all()
    assert get_rule(db, "ozon").base_coef == Decimal("1.1")
    # цепочка: WB от Ozon, который сам от WB
    r = client.post("/prices/rules/wb", data={**RULE, "base_platform": "ozon", "base_coef": "1"})
    assert "цепочки не поддерживаются" in r.text
    r = client.post("/prices/rules/ozon", data={**RULE, "base_platform": "ozon", "base_coef": "1"})
    assert "сама от себя" in r.text
    r = client.post("/prices/rules/ozon", data={**RULE, "base_platform": "wb", "base_coef": ""})
    assert "Коэффициент к базовой" in r.text


def test_rules_excel_roundtrip_empty_cell_changes_nothing(client, db):
    f.account(db, "wb")
    f.account(db, "ozon", "Озон")
    client.post("/prices/rules/wb", data=RULE)
    client.post("/prices/rules/ozon", data={**RULE, "commission_percent": "30"})
    wb, ws, headers = _xlsx(client.get("/prices/rules-export").content)
    rows = {ws.cell(row=i, column=1).value: i for i in range(2, ws.max_row + 1)}
    col = lambda name: headers.index(name) + 1  # noqa: E731
    ws.cell(row=rows["ozon"], column=col("Комиссия, %"), value=None)          # пусто — не трогать
    ws.cell(row=rows["ozon"], column=col("Цена от площадки (код)"), value="wb")
    ws.cell(row=rows["ozon"], column=col("Коэффициент к базовой"), value=1.15)
    ws.cell(row=rows["wb"], column=col("Коэффициент наценки (2 = +100%)"), value=2.5)
    r = client.post("/prices/rules-import", files={"file": ("r.xlsx", _save(wb))})
    assert "Правил изменено: 2" in r.text
    db.expire_all()
    oz = get_rule(db, "ozon")
    assert oz.commission_percent == Decimal("30") and oz.base_platform == "wb" and oz.base_coef == Decimal("1.15")
    assert get_rule(db, "wb").markup_coef == Decimal("2.5")
    # «-» снимает базу
    wb2, ws2, h2 = _xlsx(client.get("/prices/rules-export").content)
    r2 = {ws2.cell(row=i, column=1).value: i for i in range(2, ws2.max_row + 1)}
    ws2.cell(row=r2["ozon"], column=h2.index("Цена от площадки (код)") + 1, value="-")
    client.post("/prices/rules-import", files={"file": ("r.xlsx", _save(wb2))})
    db.expire_all()
    assert get_rule(db, "ozon").base_platform is None and get_rule(db, "ozon").base_coef is None


# --- текущие цены ------------------------------------------------------------------

def test_parse_wb_and_ozon_prices():
    wb = platforms.parse_wb_prices({"data": {"listGoods": [
        {"nmID": 5, "sizes": [{"price": 3000, "discountedPrice": 2400}, {"price": 3200, "discountedPrice": 2560}]},
        {"nmID": 6, "sizes": [{"price": 0}]}]}})
    assert wb == {"5": CurrentPrice(3200, 2560)}
    oz = platforms.parse_ozon_prices({"items": [{"offer_id": "A-1", "price": {"price": "4699.0000"}},
                                                {"offer_id": "A-2", "price": {}}]})
    assert oz == {"A-1": CurrentPrice(4699, 4699)}


class _Resp:
    def __init__(self, code, data):
        self.status_code, self._data, self.text = code, data, str(data)
        self.content = b"x"

    def json(self):
        return self._data

    def raise_for_status(self):
        import requests
        if self.status_code >= 400:
            raise requests.HTTPError(response=self)


class _LamodaSession:
    """Отвечает по адресу, как площадка по спеке; запоминает запросы."""
    def __init__(self, token_code=201, pages=None, sell=None, push=None):
        self.token_code, self.pages, self.sell, self.push = token_code, pages or [], sell or [], push or {}
        self.calls = []

    def post(self, url, json=None, headers=None, timeout=None):
        self.calls.append(("POST", url, json, headers))
        if url.endswith("/v2/auth-token"):
            return _Resp(self.token_code, {"access_token": "tok", "expires_in": 3600}
                         if self.token_code < 400 else {"error": {"code": "UNAUTHORIZED"}})
        return _Resp(200, self.push)

    def get(self, url, params=None, headers=None, timeout=None):
        self.calls.append(("GET", url, params, headers))
        pages = self.pages if url.endswith("/v2/nomenclatures") else self.sell
        page = params.get("page", 1)
        return _Resp(200, {"nomenclatures": pages[page - 1] if page <= len(pages) else [],
                           "meta": {"page": page, "limit": 25, "total": 0, "totalPages": len(pages)}})


def _nom(parent, sku, barcode, size="M"):
    return {"parentSku": parent, "sku": sku, "barcode": barcode, "externalParentSku": "39681",
            "externalSku": "39681-" + size, "externalSize": size, "name": "Свитшот"}


def test_lamoda_token_seller_id_and_catalog_pages():
    s = _LamodaSession(pages=[[_nom("MP1", "MP1S1", "b1", "S")], [_nom("MP1", "MP1S2", "b2", "M"),
                                                                    {"parentSku": "MP2", "barcode": None}]])
    c = LamodaClient("cid", "secret", "242541217", session=s)
    rows = c.get_catalog()
    assert [(r.external_id, r.barcode, r.article, r.size) for r in rows] == \
        [("MP1:MP1S1", "b1", "39681", "S"), ("MP1:MP1S2", "b2", "39681", "M")]
    assert not c.last_truncated
    method, url, body, _ = s.calls[0]
    assert url.endswith("/api/v2/auth-token")
    assert body == {"grant_type": "client_credentials", "client_id": "cid", "client_secret": "secret"}
    gets = [c_ for c_ in s.calls if c_[0] == "GET"]
    assert len(gets) == 2 and gets[0][2] == {"sellerId": "242541217", "country": "RU", "page": 1, "limit": 25}
    assert gets[0][3] == {"Authorization": "Bearer tok"}
    assert sum(1 for c_ in s.calls if c_[1].endswith("auth-token")) == 1      # токен один на клиента


def test_lamoda_connection_check():
    ok, msg = LamodaClient("c", "s", "1", session=_LamodaSession()).test_connection()
    assert ok and "Seller ID принят" in msg
    ok, msg = LamodaClient("c", "bad", "1", session=_LamodaSession(token_code=401)).test_connection()
    assert not ok and "не приняты" in msg


def test_lamoda_current_prices_kopecks_and_sale_window():
    from datetime import datetime, timezone
    data = {"nomenclatures": [
        {"parentSku": "MP1", "sellValues": [
            {"country": "RU", "price": {"amount": 1499000, "currency": "RUB"},
             "salePrice": {"amount": 1249200, "currency": "RUB"},
             "saleStart": "2026-09-01T00:00:00.000Z", "saleEnd": "2026-10-31T00:00:00.000Z"},
            {"country": "KZ", "price": {"amount": 7759444, "currency": "KZT"}}]},
        {"parentSku": "MP2", "sellValues": [
            {"country": "RU", "price": {"amount": 299900, "currency": "RUB"},
             "salePrice": {"amount": 199900, "currency": "RUB"},
             "saleStart": "2025-06-01T00:00:00.000Z", "saleEnd": "2025-06-30T23:59:59.000Z"}]}]}
    got = platforms.parse_lamoda_sell_values(data, now=datetime(2026, 10, 2, tzinfo=timezone.utc))
    assert got == {"MP1": CurrentPrice(14990, 12492), "MP2": CurrentPrice(2999, 2999)}   # скидка MP2 кончилась


def test_lamoda_push_by_parent_in_kopecks_never_forced():
    resp = {"successCount": 1, "errorCount": 1,
            "errors": [{"sku": "MP2", "code": "VALIDATION_ERROR", "messages": ["Validation error"]}],
            "fraudValidationResults": [{"parentSku": "MP3", "invalidValue": "BLACK_PRICE", "status": "WARNING",
                                        "promotions": [{"name": "Осень"}]}]}
    s = _LamodaSession(push=resp)
    c = LamodaClient("c", "s", "777", session=s)
    out = c.push_prices([platforms.PriceItem("b1", 3499, "MP1:MP1S1"), platforms.PriceItem("b1b", 3599, "MP1:MP1S2"),
                         platforms.PriceItem("b2", 1999, "MP2:MP2S1"), platforms.PriceItem("b3", 999, "MP3:x"),
                         platforms.PriceItem("b4", 500, "")])
    body = [c_ for c_ in s.calls if c_[1].endswith("nomenclatures-prices")][0][2]
    assert body["sellerId"] == "777" and body["country"] == "RU" and body["force"] is False
    assert {p["parentSku"]: p["price"] for p in body["prices"]}["MP1"] == {"amount": 359900, "currency": "RUB"}
    assert sorted(out["ok"]) == ["b1", "b1b"] and out["sent_prices"]["b1"] == 3599
    details = " ".join(str(e["detail"]) for e in out["errors"])
    assert "VALIDATION_ERROR" in details and "Осень" in details and "уберите товар из акции" in details
    assert "нет parentSku" in details
    # ошибка без адресата — неуспешна вся пачка
    s2 = _LamodaSession(push={"successCount": 0, "errorCount": 1, "fraudValidationResults": [],
                              "errors": [{"sku": "", "code": "VALIDATION_ERROR", "messages": ["not found"]}]})
    out2 = LamodaClient("c", "s", "1", session=s2).push_prices([platforms.PriceItem("b1", 100, "MP1:1")])
    assert out2["ok"] == [] and "без указания товара" in str(out2["errors"])


def test_lamoda_needs_seller_id():
    try:
        build_client("lamoda", {"client_id": "x", "client_secret": "y"})
        raise AssertionError("без Seller ID клиент собираться не должен")
    except PlatformError as e:
        assert "seller_id" in str(e)
    assert isinstance(build_client("lamoda", {"client_id": "x", "client_secret": "y", "seller_id": "1"}),
                      LamodaClient)


class FakePrices:
    def __init__(self, prices, truncated=False):
        self.prices, self.last_truncated = prices, truncated

    def price_key(self, item):
        return item.external_id.split(":")[0]

    def get_prices(self):
        return self.prices


def test_load_prices_updates_and_keeps_on_truncation(db):
    acc = f.account(db)
    f.item(db, acc, "b1", external_id="5:1")
    f.item(db, acc, "b2", external_id="6:1")
    st = acc_mod.load_prices(db, acc, FakePrices({"5": CurrentPrice(3000, 2400), "6": CurrentPrice(100, 100)}))
    assert st["updated"] == 2
    # неполная выдача: о «6» площадка промолчала — прежняя цена остаётся
    acc_mod.load_prices(db, acc, FakePrices({"5": CurrentPrice(3100, 2500)}, truncated=True))
    items = {i.barcode: i for i in db.query(PlatformItem)}
    assert items["b1"].current_price == 3100 and items["b2"].current_price == 100
    assert "НЕПОЛНАЯ" in db.get(Account, acc.id).prices_note
    # полная выдача без «6» — цена неизвестна
    acc_mod.load_prices(db, acc, FakePrices({"5": CurrentPrice(3100, 2500)}))
    db.expire_all()
    assert db.query(PlatformItem).filter_by(barcode="b2").one().current_price is None


def _setup_products(client, db):
    f.manual_rate(db)
    a1 = f.account(db, "wb", "ИП Яворская")
    a2 = f.account(db, "wb", "ИП Ребрик")
    client.post("/prices/rules/wb", data=RULE)
    f.sku(db, "u1", "39681", "L", barcodes=["b1"], cost_usd="16.24", name="Свитшот")
    f.sku(db, "u2", "4033", "3XL", barcodes=["b2"], cost_usd="12.40", name="Джемпер")
    for a in (a1, a2):
        f.item(db, a, "b1", "39681-L", external_id="5:1")
        f.item(db, a, "b2", "4033-3XL", external_id="6:1")
    return a1, a2


def test_load_current_button_and_markup_by_current(client, db, monkeypatch):
    from priceapp.routers import accounts as r
    a1, a2 = _setup_products(client, db)
    for a in (a1, a2):
        acc_mod.set_credential(db, a, "token", "t")
    f.account(db, "kit", "КИТ")
    db.commit()
    monkeypatch.setattr(r, "CLIENT_FACTORY", lambda platform, creds: FakePrices(
        {"5": CurrentPrice(2000, 1600), "6": CurrentPrice(4000, 4000)}))
    page = client.post("/prices/load-current", data={"account_id": str(a1.id)}).text
    assert "ИП Яворская — 2" in page and "ИП Ребрик — 2" in page and "чтение цен не подключено" in page
    # 1600 × 0,75 = 1200; 1200 − 1323,56 = −123,56; коэфф. 0,91 — ниже пола 1,3
    page = client.get(f"/prices?view=products&account_id={a1.id}&flt=below_floor_now").text
    assert "−123,56" in page or "-123,56" in page
    assert "39681" in page and "4033" not in page.split("<table")[1]


def test_bulk_manual_from_current_whole_filter_all_cabinets(client, db):
    a1, a2 = _setup_products(client, db)
    for a in (a1, a2):
        for it in db.query(PlatformItem).filter_by(account_id=a.id):
            it.current_price = 2000 if it.barcode == "b1" else 3000
    db.commit()
    r = client.post(f"/prices/bulk/{a1.id}", data={"action": "manual_from_current", "value": "1,1",
                                                   "all_filtered": "1", "q": "39681", "scope": "platform"})
    assert "изменено 2" in r.text
    got = {(p.account_id, p.item_id): p.manual_price for p in db.query(ProductPrice)}
    # 2000 × 1,1 = 2200 → вверх до 10, минус 1 → 2209; только отобранный 39681, в обоих ИП
    assert got == {(a1.id, "u1"): 2209, (a2.id, "u1"): 2209}
    # снять — по отмеченным, только в этом кабинете
    client.post(f"/prices/bulk/{a1.id}", data={"action": "clear_manual", "ids": ["u1"]})
    db.expire_all()
    assert {(p.account_id, p.item_id): p.manual_price for p in db.query(ProductPrice)} == \
        {(a1.id, "u1"): None, (a2.id, "u1"): 2209}
    assert "Ничего не отмечено" in client.post(f"/prices/bulk/{a1.id}", data={"action": "clear_manual"}).text
    assert "значение" in client.post(f"/prices/bulk/{a1.id}",
                                     data={"action": "set_manual", "value": "abc", "ids": ["u1"]}).text


def test_products_import_empty_cell_changes_nothing(client, db):
    a1, _ = _setup_products(client, db)
    client.post(f"/prices/manual/{a1.id}", data={"item_id": "u1", "value": "3999"})
    wb, ws, headers = _xlsx(client.get(f"/prices/export/{a1.id}").content)
    col = headers.index("Ручная цена, ₽") + 1
    ids = {ws.cell(row=i, column=1).value: i for i in range(2, ws.max_row + 1)}
    ws.cell(row=ids["u1"], column=col, value=None)       # пусто — ручная 3999 остаётся
    ws.cell(row=ids["u2"], column=col, value=2999)
    r = client.post(f"/prices/import/{a1.id}", files={"file": ("p.xlsx", _save(wb))})
    assert "изменено: 1" in r.text
    got = {p.item_id: p.manual_price for p in db.query(ProductPrice).filter_by(account_id=a1.id)}
    assert got == {"u1": 3999, "u2": 2999}
    ws.cell(row=ids["u1"], column=col, value="-")
    client.post(f"/prices/import/{a1.id}", files={"file": ("p.xlsx", _save(wb))})
    db.expire_all()
    assert db.query(ProductPrice).filter_by(account_id=a1.id, item_id="u1").one().manual_price is None


def test_proposals_excel_decisions_go_through_approve(client, db):
    a1, _ = _setup_products(client, db)
    client.post(f"/prices/manual/{a1.id}", data={"item_id": "u2", "value": "1000"})   # ниже пола
    client.post("/prices/recalculate", data={"account_id": str(a1.id)})
    wb, ws, headers = _xlsx(client.get(f"/prices/changes-export?view=proposals&account_id={a1.id}").content)
    assert "Решение (Да / Нет)" in headers
    dec = headers.index("Решение (Да / Нет)") + 1
    for i in range(2, ws.max_row + 1):
        ws.cell(row=i, column=dec, value="Да")
    r = client.post("/prices/changes-import", files={"file": ("c.xlsx", _save(wb))})
    assert "подтверждено 1" in r.text and "ниже минимальной" in r.text
    statuses = sorted(c.status for c in db.query(PriceChange).filter_by(account_id=a1.id))
    assert statuses == ["approved", "blocked"]
    log = client.get("/prices/changes-export?view=log")
    _, _, lh = _xlsx(log.content)
    assert "Решение (Да / Нет)" not in lh and "Статус" in lh


def test_migration_moves_rules_from_first_cabinet_to_platform(tmp_path):
    here = Path(__file__).resolve().parents[1]
    db = tmp_path / "m.db"
    env = dict(os.environ, DATABASE_URL=f"sqlite:///{db}")
    run = lambda *a: subprocess.run([sys.executable, "-m", "alembic", *a], cwd=here, env=env,  # noqa: E731
                                    capture_output=True, text=True, timeout=120)
    assert run("upgrade", "0001").returncode == 0
    con = sqlite3.connect(db)
    for i, (p, n, c) in enumerate([("wb", "ИП Яворская", 25), ("wb", "ИП Ребрик", 27), ("ozon", "Озон", 30)], 1):
        con.execute("INSERT INTO accounts (id, platform, name, is_active, commission_percent, last_check_message, "
                    "catalog_note, created_at) VALUES (?, ?, ?, 1, ?, '', '', CURRENT_TIMESTAMP)", (i, p, n, c))
        con.execute("INSERT INTO price_rules (account_id, markup_coef, round_step, round_minus, min_markup_coef, "
                    "max_change_percent, updated_at) VALUES (?, ?, 100, 1, 1.3, 20, CURRENT_TIMESTAMP)",
                    (i, 2 + i / 10))
    con.commit()
    con.close()
    r = run("upgrade", "head")
    assert r.returncode == 0, r.stderr
    con = sqlite3.connect(db)
    rules = {row[0]: row[1:] for row in con.execute(
        "SELECT platform, commission_percent, markup_coef FROM platform_rules")}
    con.close()
    assert rules == {"wb": (25, 2.1), "ozon": (30, 2.3)}
    assert run("check").returncode == 0


# --- импорт на «Сопоставлении» ------------------------------------------------------

def test_mapping_status_import_links_only_not_in_1c(client, db):
    from priceapp.models import ManualLink
    acc = f.account(db)
    f.sku(db, "u1", "39681", "L", barcodes=["b1"])
    f.sku(db, "u2", "4033", "3XL", barcodes=["bx"])
    f.item(db, acc, "b1", "39681-L", external_id="5:1")
    f.item(db, acc, "4607001200000", "NEW", external_id="7:1")
    wb, ws, headers = _xlsx(client.get(f"/mapping/export/{acc.id}").content)
    rows = {str(ws.cell(row=i, column=1).value): i for i in range(2, ws.max_row + 1)}
    col = headers.index("ID_1С") + 1
    ws.cell(row=rows["4607001200000"], column=col, value="u2")     # «нет в 1С» -> ручная связь
    ws.cell(row=rows["b1"], column=col, value="u2")                # уже сопоставлен -> отказ
    r = client.post(f"/mapping/import/{acc.id}", files={"file": ("m.xlsx", _save(wb))})
    assert "Ручных связей создано: 1" in r.text and "перепривязка не делается" in r.text
    assert [(m.barcode, m.item_id) for m in db.query(ManualLink)] == [("4607001200000", "u2")]
    assert "Импорт из Excel" in client.get(f"/mapping?view=status&account_id={acc.id}").text


def test_mapping_candidates_export_import_confirms_through_confirm(client, db):
    from priceapp.models import ManualLink
    acc = f.account(db)
    f.sku(db, "u1", "39681", "L", color="GRI", barcodes=["b1"])
    f.sku(db, "u2", "39681", "M", color="GRI", barcodes=["b2"])
    f.item(db, acc, "b1", "39681", external_id="5:1", size="L")
    f.item(db, acc, "n2", "39681", external_id="5:2", size="M")
    wb, ws, headers = _xlsx(client.get(f"/mapping/candidates-export/{acc.id}").content)
    assert ws.max_row >= 2 and "Подтвердить (Да)" in headers
    for i in range(2, ws.max_row + 1):
        ws.cell(row=i, column=headers.index("Подтвердить (Да)") + 1, value="Да")
    r = client.post(f"/mapping/candidates-import/{acc.id}", files={"file": ("c.xlsx", _save(wb))})
    assert "Сопоставлено баркодов: 1" in r.text
    assert [(m.barcode, m.item_id) for m in db.query(ManualLink)] == [("n2", "u2")]
    # повтор того же файла: уже сопоставлено — отказ, второй связи нет
    r = client.post(f"/mapping/candidates-import/{acc.id}", files={"file": ("c.xlsx", _save(wb))})
    assert "уже сопоставлен" in r.text and db.query(ManualLink).count() == 1
