"""Акции WB (страница «Акции WB»): календарь акций кабинетов и наши карточки в них.

Календарь (`refresh`, фоном раз в 3 часа и по кнопке) — по спеке WB (снимок
github.com/ZloyDeDD/wb-api-skill, 08-promotion.yaml; вживую НЕ проверялось): список акций с датами
и типом, подробности, а у ОБЫЧНЫХ акций — товары кабинета: участвует ли, цена и скидка сейчас,
плановая цена (наибольшая цена со скидкой для участия) и рекомендуемая скидка. У автоакций товаров
API не отдаёт — только число исключённых до старта.

По каждой карточке обычной акции (`classify`):
  * цена при входе — цена сейчас × (100 − рекомендуемая скидка) / 100 (её и поставит WB);
  * «% от эталона для вхождения» — на сколько % она ниже эталона;
  * меньше порога карточки — «уложится в порог»: после старта программа её не тронет;
  * не меньше — «вернётся эталон»: ближайший опрос WB после старта вернёт полный эталон, и карточка
    выйдет из акции;
  * ниже пола хотя бы у одного размера — «ниже пола»: в акцию не добавляем.

«Новый порог для акции» (`set_threshold`) — порог артикула в кабинете ДО КОНЦА АКЦИИ: потом
возвращается прежний (`wb_discount.expire_temp`). Постоянным он не бывает намеренно: если WB после
акции сам прежнюю скидку не вернёт, при постоянном пороге цена так и осталась бы акционной, и вернуть
её было бы нечему. Ставится только тем артикулам, чьи карточки есть в этой акции.

«Добавить в акцию» (`add`) — запись на WB: только «уложится в порог», не ниже пола, ещё не в акции.
Со старта акции, если она не началась, иначе сразу. WB применяет не мгновенно — до следующего
обновления календаря строка помечена «отправлено, ждём WB».

«Участвовать до» (`set_leave`, решение владельца 17.10) — срок участия карточки, раньше конца
акции. По его истечении (`leave_due`, фоном раз в 5 минут) программа ставит ПОЛНЫЙ эталон: цена
поднимается, и карточка выходит из акции. Порог «для акции» этой акции кончается в тот же срок.

Конец акции (`end_due`, решение владельца 18.10, тот же проход раз в 5 минут) — карточкам, которые в
ней стояли, возвращается эталон: WB сам прежнюю цену не возвращает, а «уложившуюся в порог» карточку
правило возврата не трогает, и она так и стояла бы по акционной цене. Карточку, записанную ещё в одну
незакончившуюся акцию, не трогаем — эталон ей вернёт конец той акции. Автоакций этот шаг не видит: их
товаров API не отдаёт.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal

from sqlalchemy.orm import Session

from priceapp import audit, wb_discount as wd
from priceapp.models import Account, ApiCredential, WbDiscountRule, WbPromotion, WbPromotionItem
from priceapp.timeutils import now_utc

REGULAR, AUTO = "regular", "auto"
TYPE_LABELS = {REGULAR: "акция", AUTO: "автоакция"}
HORIZON_DAYS = 60              # календарь — на два месяца вперёд
KEEP_ENDED_DAYS = 30           # закончившиеся акции храним месяц, потом удаляем вместе с товарами
ATTENTION_HOURS = 36           # «накануне старта» — для строки во «Внимании»
ATTENTION_MIN = 3              # «заметное число» карточек
ADDED_WAIT = timedelta(hours=24)

FITS, RETURNS, FLOOR, NOETALON, UNKNOWN = "fits", "returns", "floor", "noetalon", "unknown"
STATUS_LABELS = {FITS: "уложится в порог", RETURNS: "вернётся эталон", FLOOR: "ниже пола",
                 NOETALON: "нет эталона", UNKNOWN: "нет в каталоге"}


def parse_dt(raw) -> datetime | None:
    """«2026-10-20T00:00:00Z» -> наивное UTC, как хранятся отметки."""
    if not raw:
        return None
    try:
        from datetime import timezone
        d = datetime.fromisoformat(str(raw).replace("Z", "+00:00"))
        return d.astimezone(timezone.utc).replace(tzinfo=None) if d.tzinfo else d
    except ValueError:
        return None


def _int(v) -> int | None:
    try:
        return int(Decimal(str(v)).quantize(Decimal(1), ROUND_HALF_UP)) if v is not None and v != "" else None
    except Exception:
        return None


# --- календарь -----------------------------------------------------------------------------------

def refresh(db: Session, client_factory=None, now: datetime | None = None) -> dict:
    """Календарь акций всех активных кабинетов WB — в базу. Сбой кабинета не трогает остальные и
    оставляет его прежние данные. Коммитит по кабинету."""
    from priceapp import accounts
    now = now or now_utc()
    st = {"accounts": 0, "promos": 0, "items": 0, "errors": []}
    for a in db.query(Account).filter(Account.platform == "wb", Account.is_active.is_(True)).all():
        if db.query(ApiCredential.id).filter(ApiCredential.account_id == a.id).first() is None:
            continue
        try:
            client = accounts.client_for(db, a, client_factory)
            got = client.promotions(now - timedelta(days=1), now + timedelta(days=HORIZON_DAYS), True)
        except Exception as e:
            st["errors"].append(f"{a.name}: календарь не загружен — {e}"[:300])
            continue
        st["accounts"] += 1
        existing = {p.promo_id: p for p in db.query(WbPromotion).filter(WbPromotion.account_id == a.id)}
        seen = set()
        for d in got:
            pid = _int(d.get("id"))
            if pid is None:
                continue
            seen.add(pid)
            row = existing.get(pid)
            if row is None:
                row = existing[pid] = WbPromotion(account_id=a.id, promo_id=pid)
                db.add(row)
            row.name = str(d.get("name") or "")[:300]
            row.type = AUTO if str(d.get("type") or "") == AUTO else REGULAR
            row.start_at, row.end_at = parse_dt(d.get("startDateTime")), parse_dt(d.get("endDateTime"))
            row.fetched_at = now
        live = [p for pid, p in existing.items() if pid in seen and (p.end_at is None or p.end_at >= now)]
        try:
            for d in client.promotion_details([p.promo_id for p in live]) if live else []:
                p = existing.get(_int(d.get("id")))
                if p is None:
                    continue
                p.description = str(d.get("description") or "")
                p.advantages = "; ".join(str(x) for x in d.get("advantages") or [])
                p.participation = _int(d.get("participationPercentage"))
                p.in_action_total = _int(d.get("inPromoActionTotal"))
                p.exceptions = _int(d.get("exceptionProductsCount"))
        except Exception as e:
            st["errors"].append(f"{a.name}: подробности акций не загружены — {e}"[:300])
        for p in live:
            if p.type != REGULAR:
                continue
            try:
                inside, t1 = client.promotion_nomenclatures(p.promo_id, True)
                outside, t2 = client.promotion_nomenclatures(p.promo_id, False)
            except Exception as e:
                p.items_note = f"товары не загружены: {e}"[:300]
                st["errors"].append(f"{a.name}, акция «{p.name}»: товары не загружены — {e}"[:300])
                continue
            items = {i.nm_id: i for i in db.query(WbPromotionItem).filter(
                WbPromotionItem.account_id == a.id, WbPromotionItem.promo_id == p.promo_id)}
            fresh = set()
            for flag, rows in ((True, inside), (False, outside)):
                for d in rows:
                    nm = str(_int(d.get("id")) or "")
                    if not nm or nm in fresh:
                        continue
                    fresh.add(nm)
                    it = items.get(nm)
                    if it is None:
                        it = items[nm] = WbPromotionItem(account_id=a.id, promo_id=p.promo_id, nm_id=nm)
                        db.add(it)
                    it.in_action = bool(d.get("inAction", flag))
                    it.price, it.plan_price = _int(d.get("price")), _int(d.get("planPrice"))
                    it.discount, it.plan_discount = _int(d.get("discount")), _int(d.get("planDiscount"))
                    it.fetched_at = now
                    if it.in_action:
                        it.added_at = None            # WB показал участие — ждать больше нечего
            for nm, it in items.items():
                if nm not in fresh:
                    db.delete(it)
            p.items_fetched_at = now
            p.items_note = "список товаров обрезан нашим пределом страниц" if (t1 or t2) else ""
            st["items"] += len(fresh)
        old = now - timedelta(days=KEEP_ENDED_DAYS)
        for pid, p in existing.items():
            if pid not in seen and (p.end_at is None or p.end_at < old):
                db.query(WbPromotionItem).filter(WbPromotionItem.account_id == a.id,
                                                 WbPromotionItem.promo_id == pid).delete()
                db.delete(p)
        st["promos"] += len(seen)
        db.commit()
    _follow_promo_ends(db)
    wd.expire_temp(db, now)
    return st


def _follow_promo_ends(db: Session) -> None:
    """WB продлил акцию — порог «для акции» продлевается вместе с ней: иначе он вернулся бы раньше
    конца, и опрос выкинул бы карточку из ещё идущей акции. Коммитит."""
    changed = False
    for r in db.query(WbDiscountRule).filter(WbDiscountRule.until.isnot(None), WbDiscountRule.until_promo_id.isnot(None)):
        p = db.query(WbPromotion).filter(WbPromotion.account_id == r.account_id,
                                         WbPromotion.promo_id == r.until_promo_id).first()
        if p is not None and p.end_at and p.end_at > r.until:
            r.until = p.end_at
            changed = True
    if changed:
        db.commit()


def summary(st: dict) -> str:
    text = f"Календарь акций WB: кабинетов {st['accounts']}, акций {st['promos']}, товаров в обычных акциях {st['items']}"
    return text + (". Ошибки: " + "; ".join(st["errors"]) if st["errors"] else ".")


# --- карточки акции ------------------------------------------------------------------------------

def entry_sale(it: WbPromotionItem) -> int | None:
    """Цена со скидкой при входе в акцию: WB ставит рекомендуемую скидку на нынешнюю цену."""
    if it.price and it.plan_discount is not None:
        return int((Decimal(it.price) * (100 - Decimal(it.plan_discount)) / 100).quantize(Decimal(1), ROUND_HALF_UP))
    return it.plan_price


@dataclass
class PCard:
    promo: WbPromotion
    item: WbPromotionItem
    account: Account
    card: object = None              # wb_discount.Card или None (карточки нет в нашем каталоге)
    entry: int | None = None
    need: Decimal | None = None      # % от эталона для вхождения
    thr: Decimal | None = None
    thr_src: str = ""                # откуда порог: артикул, категория, по умолчанию, «для акции до …»
    status: str = UNKNOWN
    note: str = ""

    @property
    def articles(self) -> set:
        return self.card.articles if self.card is not None else set()

    @property
    def added(self) -> bool:
        return bool(self.item.added_at) and not self.item.in_action and now_utc() - self.item.added_at < ADDED_WAIT

    @property
    def leaving(self) -> bool:
        """Есть срок участия, и он ещё не прошёл."""
        return self.item.leave_at is not None and self.item.left_at is None

    @property
    def left(self) -> bool:
        return self.item.left_at is not None


def classify(db: Session, promos: list[WbPromotion], all_cards=None, leave_at=None) -> list[PCard]:
    """Каждая карточка обычных акций `promos`: цена при входе, «% от эталона», порог, пометка.
    `leave_at` — срок участия, с которым карточки собираются добавить (у уже стоящих — свой)."""
    from priceapp import etalon, rates
    from priceapp.models import OnecCost
    from priceapp.pricing import get_rule, item_facts
    regular = [p for p in promos if p.type == REGULAR]
    if not regular:
        return []
    all_cards = all_cards if all_cards is not None else wd.cards(db)
    by_nm = {(c.account.id, c.nm): c for c in all_cards}
    accs = {a.id: a for a in db.query(Account)}
    limits = wd.Thresholds(db)
    rule = get_rule(db, "wb")
    rate = rates.current(db)
    costs = {c.item_id: c.cost_usd for c in db.query(OnecCost)}
    temp = {(r.key, r.account_id): r.until for r in db.query(WbDiscountRule).filter(
        WbDiscountRule.kind == wd.ARTICLE, WbDiscountRule.until.isnot(None))}
    after = wd.Thresholds(db, after_temp=True) if temp else limits
    from priceapp.templating import local
    out = []
    for p in regular:
        for it in db.query(WbPromotionItem).filter(WbPromotionItem.account_id == p.account_id,
                                                   WbPromotionItem.promo_id == p.promo_id):
            pc = PCard(p, it, accs.get(p.account_id), by_nm.get((p.account_id, it.nm_id)))
            out.append(pc)
            c = pc.card
            if c is None:
                continue
            pc.entry = entry_sale(it)
            if not c.ref_sale or not pc.entry:
                pc.status = NOETALON
                continue
            pc.need = ((1 - Decimal(pc.entry) / Decimal(c.ref_sale)) * 100).quantize(Decimal("0.1"), ROUND_HALF_UP)
            pc.thr, pc.thr_src = min((limits.for_item(i, c.category, p.account_id) for i in c.item_ids),
                                     key=lambda t: t[0])
            until = next((temp[(a, p.account_id)] for a in c.articles if (a, p.account_id) in temp), None)
            # до какого момента карточка в акции: конец акции или срок участия, если он раньше
            horizon = p.end_at
            lv = leave_at or (it.leave_at if it.left_at is None else None)
            if lv is not None and (horizon is None or lv < horizon):
                horizon = lv
            if until is not None and pc.thr_src.startswith("артикул"):
                if horizon is None or until < horizon:
                    # Временный порог другой акции кончится раньше этой: дальше вернётся прежний, и
                    # опрос выкинул бы карточку из ещё идущей акции. Судим по тому, что будет ПОСЛЕ.
                    pc.thr, src = min((after.for_item(i, c.category, p.account_id) for i in c.item_ids),
                                      key=lambda t: t[0])
                    pc.thr_src = f"{src}: порог для акции до {local(until)} кончится раньше этой"
                else:
                    pc.thr_src = f"для акции до {local(until)}"
            pc.note = etalon.floor_problem(rule, item_facts(rule, c.rows).commission, rate, costs, c.item_ids,
                                           it.price or c.cur_list or 0, it.plan_discount or 0)
            pc.status = FLOOR if pc.note else (FITS if pc.need < pc.thr else RETURNS)
    return out


# --- строки страницы -----------------------------------------------------------------------------

def _fmt(v) -> str:
    return wd.fmt(v)


@dataclass
class PRow:
    kind: str                         # category / article
    key: str
    name: str = ""
    category: str = ""
    shared: bool = False
    cards: list = field(default_factory=list)
    promo: WbPromotion | None = None

    @property
    def sel(self) -> str:
        """Строка — в своей акции: одна и та же категория в двух акциях — две строки."""
        return f"{self.promo.promo_id if self.promo else 0}:{self.kind[0]}:{self.key}"

    @property
    def accounts(self) -> list[str]:
        return sorted({pc.account.name for pc in self.cards if pc.account})

    def count(self, status: str) -> int:
        return sum(1 for pc in self.cards if pc.status == status)

    @property
    def n_in(self) -> int:
        return sum(1 for pc in self.cards if pc.item.in_action)

    @property
    def n_added(self) -> int:
        return sum(1 for pc in self.cards if pc.added)

    def _vals(self, attr):
        return [getattr(pc, attr) for pc in self.cards if getattr(pc, attr) is not None]

    def _range(self, vals) -> str:
        if not vals:
            return ""
        lo, hi = min(vals), max(vals)
        return _fmt(lo) if lo == hi else f"{_fmt(lo)} … {_fmt(hi)}"

    @property
    def etalon(self) -> int | None:
        v = [pc.card.ref_sale for pc in self.cards if pc.card is not None and pc.card.ref_sale]
        return int(sum(v) / len(v)) if v else None

    @property
    def current(self) -> int | None:
        v = [pc.card.cur_sale for pc in self.cards if pc.card is not None and pc.card.cur_sale]
        return int(sum(v) / len(v)) if v else None

    @property
    def entry_text(self) -> str:
        return self._range(self._vals("entry"))

    @property
    def plan_disc_text(self) -> str:
        return self._range([pc.item.plan_discount for pc in self.cards if pc.item.plan_discount is not None])

    @property
    def need_text(self) -> str:
        return self._range(self._vals("need"))

    @property
    def need_max(self) -> Decimal | None:
        v = self._vals("need")
        return max(v) if v else None

    @property
    def thr_text(self) -> str:
        return self._range(self._vals("thr"))

    @property
    def thr_src_text(self) -> str:
        return ", ".join(sorted({pc.thr_src for pc in self.cards if pc.thr_src}))

    @property
    def below_text(self) -> str:
        """На сколько % цена на WB сейчас ниже эталона — как на «Пороге скидки WB»."""
        return self._range([pc.card.below_pct for pc in self.cards
                            if pc.card is not None and pc.card.below_pct is not None])

    @property
    def held_text(self) -> str:
        return ", ".join(f"− {_fmt(h)}%" for h in sorted({pc.card.held for pc in self.cards
                                                          if pc.card is not None and pc.card.held is not None}))

    @property
    def suggest(self) -> int | None:
        """Наименьший целый порог, при котором уложатся все карточки строки."""
        n = self.need_max
        return int(n.to_integral_value(rounding="ROUND_FLOOR")) + 1 if n is not None and n >= 0 else None

    @property
    def mark(self) -> str:
        parts = [f"{STATUS_LABELS[s]} — {self.count(s)}" for s in (FITS, RETURNS, FLOOR, NOETALON) if self.count(s)]
        return "; ".join(parts)

    @property
    def leave_text(self) -> str:
        """«участвует до 25.10» по срокам карточек строки."""
        from priceapp.templating import local
        dates = sorted({pc.item.leave_at for pc in self.cards if pc.leaving})
        return ", ".join(local(d, "%d.%m %H:%M") for d in dates)

    @property
    def n_left(self) -> int:
        return sum(1 for pc in self.cards if pc.left)

    @property
    def can_add(self) -> int:
        return sum(1 for pc in self.cards if pc.status == FITS and not pc.item.in_action and not pc.added)


FILTERS = {"": "все", FITS: "уложится в порог", RETURNS: "вернётся эталон", FLOOR: "ниже пола",
           "in": "уже в акции", "out": "не в акции", "added": "отправлено, ждём WB", NOETALON: "нет эталона",
           "leave": "со сроком участия", "left": "вышли по сроку"}


def keep(row: PRow, q: str = "", cat: str = "", flt: str = "") -> bool:
    ql = q.strip().casefold()
    if ql and ql not in row.key.casefold() and ql not in (row.name or "").casefold():
        return False
    if cat and wd.norm(row.category) != wd.norm(cat):
        return False
    if flt in (FITS, RETURNS, FLOOR, NOETALON) and not row.count(flt):
        return False
    if flt == "in" and not row.n_in:
        return False
    if flt == "out" and row.n_in == len(row.cards):
        return False
    if flt == "added" and not row.n_added:
        return False
    if flt == "leave" and not row.leave_text:
        return False
    if flt == "left" and not row.n_left:
        return False
    return True


def rows(db: Session, pcards: list[PCard], account_id: int = 0, promo: WbPromotion | None = None,
         names: dict | None = None) -> tuple[list[PRow], list[PRow], int]:
    """(категории, артикулы, карточек вне каталога) по карточкам ОДНОЙ акции в кабинете (0 — во
    всех). `names` — артикул -> название (эталоны считаются один раз на страницу)."""
    if names is None:
        from priceapp import etalon
        names = {art: et.name for art, et in etalon.etalons(db).items()}
    pcs = [pc for pc in pcards if not account_id or pc.promo.account_id == account_id]
    unknown = sum(1 for pc in pcs if pc.card is None)
    cats: dict[str, PRow] = {}
    arts: dict[str, PRow] = {}
    art_accs: dict[str, set] = {}
    for pc in pcs:
        if pc.card is None:
            continue
        if pc.card.category:
            cats.setdefault(wd.norm(pc.card.category), PRow(wd.CATEGORY, pc.card.category, category=pc.card.category,
                                                            promo=promo or pc.promo))
        for a in pc.articles:
            arts.setdefault(a, PRow(wd.ARTICLE, a, name=names.get(a, ""), category=pc.card.category,
                                    promo=promo or pc.promo)).cards.append(pc)
            art_accs.setdefault(a, set()).add(pc.promo.account_id)
    for a, row in arts.items():
        row.shared = len(art_accs.get(a, ())) > 1
        if not row.shared and row.category:
            c = cats[wd.norm(row.category)]
            c.cards += [pc for pc in row.cards if pc not in c.cards]
    return (sorted([c for c in cats.values() if c.cards], key=lambda r: r.key.casefold()),
            sorted(arts.values(), key=lambda r: r.key), unknown)


def grouped(cats: list[PRow], arts: list[PRow], q: str = "", cat: str = "", flt: str = "") -> list[wd.Entry]:
    """Как на «Пороге скидки WB»: артикулы, чьи карточки в акции в нескольких кабинетах, затем
    каждая категория и её артикулы."""
    out: list[wd.Entry] = []
    shared = [a for a in arts if a.shared and keep(a, q, cat, flt)]
    if shared:
        out.append(wd.Entry("head", f"{wd.SHARED_TITLE} ({len(shared)})", section=wd.SHARED_TITLE))
        out += [wd.Entry("row", row=a, section=wd.SHARED_TITLE) for a in shared]
    kids: dict[str, list[PRow]] = {}
    for a in arts:
        if not a.shared:
            kids.setdefault(wd.norm(a.category) if a.category else "", []).append(a)
    ql = q.strip().casefold()
    for c in cats:
        if cat and wd.norm(c.key) != wd.norm(cat):
            continue
        name_hit = bool(ql) and ql in c.key.casefold()
        mine = [a for a in kids.get(wd.norm(c.key), []) if keep(a, "" if name_hit else q, "", flt)]
        if not (keep(c, q, "", flt) or mine):
            continue
        out.append(wd.Entry("row", row=c, section=c.key))
        out += [wd.Entry("row", row=a, level=1, section=c.key) for a in mine]
    orphans = [a for a in kids.get("", []) if not cat and keep(a, q, "", flt)]
    if orphans:
        out.append(wd.Entry("head", f"{wd.NO_CATEGORY} ({len(orphans)})", section=wd.NO_CATEGORY))
        out += [wd.Entry("row", row=a, section=wd.NO_CATEGORY) for a in orphans]
    return out


def sections(db: Session, pcards: list[PCard], account_id: int, by: str, q: str = "", cat: str = "", flt: str = "",
             names: dict | None = None) -> tuple[list[wd.Entry], list[str], int]:
    """Строки страницы и файла по всем акциям из `pcards`: каждая акция своим разделом (если их
    несколько), внутри — как на «Пороге скидки WB». -> (строки, категории для отбора, вне каталога)."""
    from priceapp.templating import local
    groups: dict[int, list[PCard]] = {}
    for pc in pcards:
        groups.setdefault(pc.promo.promo_id, []).append(pc)
    out, cats_all, unknown = [], set(), 0
    for pid, pcs in groups.items():
        promo = pcs[0].promo
        cats, arts, unk = rows(db, pcs, account_id, promo, names)
        unknown += unk
        cats_all |= {c.key for c in cats}
        if by == "group":
            part = grouped(cats, arts, q, cat, flt)
        else:
            src = cats if by == wd.CATEGORY else arts
            part = [wd.Entry("row", row=r, section=r.key if r.kind == wd.CATEGORY else (r.category or ""))
                    for r in src if keep(r, q, cat, flt)]
        if part and len(groups) > 1:
            out.append(wd.Entry("promo", f"Акция «{promo.name}» · {local(promo.start_at)} — {local(promo.end_at)} · "
                                         f"ID {pid}", section=promo.name))
        out += part
    return out, sorted(cats_all, key=str.casefold), unknown


# --- действия ------------------------------------------------------------------------------------

def set_threshold(db: Session, pcards: list[PCard], value: Decimal, actor: str, leave_at=None) -> int:
    """Порог «для акции» артикулам этих карточек, в их кабинетах, до конца их акции (или до срока
    участия `leave_at`, если он раньше). Не коммитит."""
    n, done = 0, set()
    for pc in pcards:
        if pc.card is None or pc.promo.end_at is None:
            continue
        for art in pc.articles:
            k = (art, pc.promo.account_id, pc.promo.promo_id)
            if k in done:
                continue
            done.add(k)
            until = min(pc.promo.end_at, leave_at) if leave_at else pc.promo.end_at
            n += wd.set_temp_threshold(db, art, value, until, actor, pc.promo.account_id,
                                       pc.promo.promo_id, pc.promo.name)
    return n


def parse_leave(raw):
    """«Участвовать до»: дата (поле страницы — ГГГГ-ММ-ДД, Excel — дата или дд.мм.гггг) -> конец
    этого МЕСТНОГО дня как наивное UTC. Пусто — None; не дата — ValueError."""
    from datetime import date, datetime as dt
    from priceapp.timeutils import local_day_start_utc
    if raw is None or str(raw).strip() == "":
        return None
    if isinstance(raw, dt):
        d = raw.date()
    elif isinstance(raw, date):
        d = raw
    else:
        text = str(raw).strip()
        for fmt_ in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y"):
            try:
                d = dt.strptime(text, fmt_).date()
                break
            except ValueError:
                continue
        else:
            raise ValueError(text)
    return local_day_start_utc(d + timedelta(days=1)) - timedelta(seconds=1)


def leave_error(promo: WbPromotion, leave_at, now=None) -> str:
    now = now or now_utc()
    if leave_at <= now:
        return "срок уже прошёл"
    if promo.start_at and leave_at <= promo.start_at:
        return "срок раньше старта акции"
    if promo.end_at and leave_at >= promo.end_at:
        return "срок не раньше конца акции — она кончится сама"
    return ""


def set_leave(db: Session, pcards: list[PCard], leave_at, actor: str) -> tuple[int, list[str]]:
    """Срок участия карточкам, что в акции или отправлены в неё; None — снять срок (участвует до
    конца акции). Порог «для акции» этой акции укорачивается до срока. Не коммитит.
    -> (сколько карточек, отказы)."""
    from priceapp.templating import local
    n, bad, done = 0, [], set()
    for pc in pcards:
        name = f"{pc.account.name if pc.account else ''} {pc.item.nm_id}"
        if not (pc.item.in_action or pc.item.added_at):
            continue
        if leave_at is not None:
            err = leave_error(pc.promo, leave_at)
            if err:
                bad.append(f"{name}: {err}")
                continue
        if pc.item.leave_at == leave_at and pc.item.left_at is None:
            continue
        pc.item.leave_at, pc.item.leave_by, pc.item.left_at = leave_at, actor, None
        n += 1
        if leave_at is None:
            continue
        for art in pc.articles:
            k = (art, pc.promo.account_id)
            if k in done:
                continue
            done.add(k)
            r = wd.get_rule(db, wd.ARTICLE, art, pc.promo.account_id)
            if r is not None and r.until is not None and r.until_promo_id == pc.promo.promo_id and r.until > leave_at:
                r.until = leave_at
    if n:
        what = f"до {local(leave_at)}" if leave_at else "снят — до конца акции"
        names = sorted({pc.promo.name for pc in pcards})
        audit.log(db, actor, "wb_promo_leave_set", "", f"срок участия {what}: карточек {n}; акции: "
                  + ", ".join(f"«{x}»" for x in names)[:800])
    return n, bad


def leave_due(db: Session, client_factory=None, now: datetime | None = None) -> dict:
    """Срок участия прошёл — на карточку ставится ПОЛНЫЙ эталон (цена до скидки и скидка эталона):
    цена поднимается, и карточка выходит из акции. Карточка без эталона поднять нечем — она
    помечается вышедшей с записью в журнал «уберите в кабинете WB», иначе её пытались бы каждые
    5 минут. Не принял WB — повтор при следующем проходе. Коммитит."""
    now = now or now_utc()
    due = db.query(WbPromotionItem).filter(WbPromotionItem.leave_at.isnot(None), WbPromotionItem.leave_at <= now,
                                           WbPromotionItem.left_at.is_(None)).all()
    st = {"due": len(due), "raised": 0, "same": 0, "errors": []}
    if not due:
        return st
    cards = {(c.account.id, c.nm): c for c in wd.cards(db)}
    names = {(p.account_id, p.promo_id): p.name for p in db.query(WbPromotion)}
    plan: dict[int, dict[str, tuple]] = {}
    for it in due:
        c = cards.get((it.account_id, it.nm_id))
        promo = names.get((it.account_id, it.promo_id), it.promo_id)
        if c is None or not c.ref_list:
            it.left_at = now
            audit.log(db, "программа", "wb_promo_leave", str(it.account_id),
                      f"акция «{promo}», карточка {it.nm_id}: срок участия прошёл, но эталона нет — цену поднять "
                      "нечем, уберите карточку из акции в кабинете WB")
            st["errors"].append(f"{it.nm_id}: нет эталона")
            continue
        if c.cur_list == c.ref_list and c.cur_sale == c.ref_sale:
            it.left_at = now                  # уже на эталоне — подниматься некуда
            st["same"] += 1
            continue
        plan.setdefault(it.account_id, {}).setdefault(c.nm, (c, [], (int(c.ref_list), int(c.ref_disc or 0)), None))[1] \
            .append((it, promo))

    def done(it):
        it.left_at = now

    def note(c, promos, pair):
        return f"выход из акции {promos} по сроку: полный эталон {c.ref_sale} ₽"

    _raise_cards(db, plan, client_factory, st, "wb_promo_leave", note, done,
                 "срок участия прошёл — поставлен полный эталон, карточки выходят из акций")
    db.commit()
    wd.expire_temp(db, now)
    return st


def _raise_cards(db: Session, plan: dict, client_factory, st: dict, source: str, note, done, what: str) -> None:
    """Поставить на WB пары «цена до скидки, скидка» карточкам `plan` = {кабинет: {nmID: (карточка, [(товар
    акции, название акции)], (цена, скидка), сдвиг «эталон − N%» после отправки)}}. Принятое — сразу у нас
    как текущее и как отправленное, строкой журнала цен с источником `source`; `done(товар акции)`
    помечает обработанное. Не принял WB — пометки нет, повтор при следующем проходе. Коммитит по кабинету.
    «Уже установлены» (400 already set) — принято: WB держит ровно то, что мы шлём (так же в `push_prices`)."""
    from priceapp import accounts, etalon
    from priceapp.dispatch import _record_sent
    from priceapp.platforms import WB_PRICES, wb_error_detail
    for acc_id, by_nm in plan.items():
        acc = db.get(Account, acc_id)
        try:
            client = accounts.client_for(db, acc, client_factory)
        except Exception as e:
            st["errors"].append(f"{acc.name}: {e}"[:300])
            continue
        items = list(by_nm.items())
        for start in range(0, len(items), 1000):
            chunk = items[start:start + 1000]
            body = {"data": [{"nmID": int(nm), "price": int(pair[0]), "discount": int(pair[1])}
                             for nm, (c, _, pair, _) in chunk]}
            try:
                data = client._post(f"{WB_PRICES}/api/v2/upload/task", body)
                if isinstance(data, dict) and data.get("error"):
                    raise RuntimeError(data.get("errorText") or "WB вернул error=true")
            except Exception as e:
                detail = wb_error_detail(e)
                if "already set" not in detail.lower():
                    st["errors"].append(f"{acc.name}: {detail}"[:300])
                    continue
            for nm, (c, its, pair, held) in chunk:
                sale = wd.sale_of(pair[0], pair[1])
                wd._mark_held(db, acc_id, c.item_ids, held)
                for i in c.item_ids:
                    _record_sent(db, i, acc_id, int(pair[0]))
                promos = ", ".join(sorted({f"«{p}»" for _, p in its}))
                etalon.journal_sent(db, acc_id, c.item_ids, c.rows[0].barcode if c.rows else "", c.cur_list,
                                    int(pair[0]), source, note(c, promos, pair), "программа")
                for r in c.rows:
                    r.current_price, r.current_sale_price = int(pair[0]), sale
                for it, _ in its:
                    done(it)
                st["raised"] += 1
            audit.log(db, "программа", source, acc.name,
                      f"{what}: {len(chunk)} — "
                      + "; ".join(f"{nm}: -> {wd.sale_of(pair[0], pair[1])}" for nm, (c, _, pair, _) in chunk[:20])[:900])
        db.commit()


END_SOURCE = "wb_promo_end"


def end_due(db: Session, client_factory=None, now: datetime | None = None) -> dict:
    """Акция закончилась — карточкам, которые в ней стояли, вернуть эталон (решение владельца 18.10).

    WB сам прежнюю цену после акции не возвращает. Карточка, которая «уложилась в порог», в нём же и
    остаётся, и правило возврата (`reference.run`) её не трогает: без этого шага она так и стояла бы по
    акционной цене. Берём карточки обычных акций, у которых WB показывал участие или мы сами их туда
    отправили; вышедшие по сроку уже на эталоне. Цель — эталон с учётом «эталон − N%» (`held_pct`):
    сдвиг — решение человека с «Порога скидки WB», и конец акции его не отменяет. Поднимаем, только если
    цена со скидкой сейчас НИЖЕ цели.

    Карточку, которая записана ещё в одну незакончившуюся акцию (идёт или впереди), не трогаем: подъём
    цены выкинул бы её оттуда. Эталон ей вернёт конец ТОЙ акции. У автоакций список товаров API не
    отдаёт — их этот шаг не видит. Каждая карточка акции обрабатывается один раз (`end_done_at`), итог —
    в `end_note`. Не принял WB — повтор при следующем проходе. Коммитит."""
    now = now or now_utc()
    rows = (db.query(WbPromotionItem, WbPromotion)
            .join(WbPromotion, (WbPromotion.account_id == WbPromotionItem.account_id)
                  & (WbPromotion.promo_id == WbPromotionItem.promo_id))
            .filter(WbPromotion.end_at.isnot(None), WbPromotion.end_at <= now,
                    WbPromotionItem.end_done_at.is_(None), WbPromotionItem.left_at.is_(None))
            .all())
    rows = [(it, p) for it, p in rows if it.in_action or it.added_at is not None]
    st = {"due": len(rows), "raised": 0, "same": 0, "busy": 0, "errors": []}
    if not rows:
        return st
    from priceapp.templating import local
    busy = {}
    for it, p in (db.query(WbPromotionItem, WbPromotion)
                  .join(WbPromotion, (WbPromotion.account_id == WbPromotionItem.account_id)
                        & (WbPromotion.promo_id == WbPromotionItem.promo_id))
                  .filter((WbPromotion.end_at.is_(None)) | (WbPromotion.end_at > now),
                          WbPromotionItem.left_at.is_(None))):
        if it.in_action or it.added_at is not None:
            busy.setdefault((it.account_id, it.nm_id), p)
    cards = {(c.account.id, c.nm): c for c in wd.cards(db)}
    plan: dict[int, dict[str, tuple]] = {}
    for it, p in rows:
        other = busy.get((it.account_id, it.nm_id))
        if other is not None:
            it.end_done_at = now
            it.end_note = (f"ещё в акции «{other.name}»"
                           + (f" до {local(other.end_at, '%d.%m %H:%M')}" if other.end_at else "")
                           + " — эталон вернёт её конец")[:200]
            st["busy"] += 1
            continue
        c = cards.get((it.account_id, it.nm_id))
        if c is None or not c.ref_list:
            it.end_done_at, it.end_note = now, "эталона нет — вернуть нечего"
            audit.log(db, "программа", END_SOURCE, str(it.account_id),
                      f"акция «{p.name}» закончилась, карточка {it.nm_id}: эталона нет — цену после акции "
                      "проверьте в кабинете WB")
            st["errors"].append(f"{it.nm_id}: нет эталона")
            continue
        disc = int(c.ref_disc or 0)
        price = wd.held_list(int(c.ref_list), c.held)
        target = wd.sale_of(price, disc)
        if c.cur_sale is not None and c.cur_sale >= target:
            it.end_done_at, it.end_note = now, "цена уже не ниже эталона"
            st["same"] += 1
            continue
        plan.setdefault(it.account_id, {}).setdefault(c.nm, (c, [], (price, disc), c.held))[1].append((it, p.name))

    def done(it):
        it.end_done_at, it.end_note = now, "эталон возвращён"

    def note(c, promos, pair):
        what = (f"эталон − {wd.fmt(c.held)}%" if c.held else "полный эталон")
        return f"акция {promos} закончилась: {what} {wd.sale_of(pair[0], pair[1])} ₽"

    _raise_cards(db, plan, client_factory, st, END_SOURCE, note, done,
                 "акция закончилась — карточкам возвращён эталон")
    db.commit()
    return st


def ended_summary(db: Session, account_id: int = 0, now: datetime | None = None) -> list[dict]:
    """Закончившиеся акции (их хранят `KEEP_ENDED_DAYS`) и что было с их карточками после конца — для
    страницы: без этого возврат эталона виден только в журнале цен."""
    now = now or now_utc()
    q = db.query(WbPromotion).filter(WbPromotion.end_at.isnot(None), WbPromotion.end_at <= now,
                                     WbPromotion.type == REGULAR)
    if account_id:
        q = q.filter(WbPromotion.account_id == account_id)
    names = {a.id: a.name for a in db.query(Account)}
    out = []
    for p in q.order_by(WbPromotion.end_at.desc()):
        its = [i for i in db.query(WbPromotionItem).filter(WbPromotionItem.account_id == p.account_id,
                                                          WbPromotionItem.promo_id == p.promo_id)
               if i.in_action or i.added_at is not None or i.left_at is not None]
        if not its:
            continue
        notes: dict[str, int] = {}
        for i in its:
            key = ("вышли по сроку" if i.left_at is not None else i.end_note or "ждёт возврата эталона")
            notes[key] = notes.get(key, 0) + 1
        out.append({"promo": p, "account": names.get(p.account_id, ""), "cards": len(its),
                    "notes": sorted(notes.items(), key=lambda kv: -kv[1])})
    return out


def add(db: Session, pcards: list[PCard], actor: str, client_factory=None, leave_at=None) -> dict:
    """«Добавить в акцию» — только «уложится в порог», не ниже пола, ещё не в акции и не отправленные
    недавно. По кабинету и акции одним запросом на 1000 карточек. Коммитит."""
    from priceapp import accounts, progress
    st = {"cards": len(pcards), "sent": 0, "skipped": [], "errors": [], "in": 0, "added": 0}
    plan: dict[tuple[int, int], list[PCard]] = {}
    for pc in pcards:
        name = f"{pc.account.name if pc.account else ''} {pc.item.nm_id}"
        if pc.item.in_action:
            st["in"] += 1
        elif pc.added:
            st["added"] += 1
        elif pc.status != FITS:
            st["skipped"].append(f"{name}: {STATUS_LABELS[pc.status]}"
                                 + (f" — {pc.note}" if pc.note else "")
                                 + (f" (нужно {_fmt(pc.need)}%, порог {_fmt(pc.thr)}%)" if pc.status == RETURNS else ""))
        else:
            plan.setdefault((pc.promo.account_id, pc.promo.promo_id), []).append(pc)
    progress.stage("Добавляю карточки в акцию", total=sum(len(g) for g in plan.values()))
    now = now_utc()
    for (acc_id, pid), group in plan.items():
        promo, acc = group[0].promo, group[0].account
        try:
            client = accounts.client_for(db, acc, client_factory)
        except Exception as e:
            st["errors"].append(f"{acc.name}: {e}"[:300])
            continue
        upload_now = promo.start_at is None or promo.start_at <= now
        done = []
        for start in range(0, len(group), 1000):
            chunk = group[start:start + 1000]
            try:
                res = client.promotion_upload(pid, [int(pc.item.nm_id) for pc in chunk], upload_now)
            except Exception as e:
                st["errors"].append(f"{acc.name}, «{promo.name}»: {e}"[:300])
                continue
            progress.step(len(chunk))
            done += chunk
            for pc in chunk:
                pc.item.added_at, pc.item.added_by, pc.item.left_at = now, actor, None
                if leave_at is not None and not leave_error(pc.promo, leave_at):
                    pc.item.leave_at, pc.item.leave_by = leave_at, actor
            from priceapp.templating import local
            audit.log(db, actor, "wb_promo_add", acc.name,
                      f"акция «{promo.name}» ({pid}): добавлено карточек {len(chunk)}, "
                      f"{'сразу' if upload_now else 'со старта акции'}"
                      + (f", участвуют до {local(leave_at)}" if leave_at else "") + f"; ответ WB: {res}"[:1000])
        st["sent"] += len(done)
        db.commit()
    return st


def add_report(what: str, st: dict) -> tuple[str, str]:
    parts = [f"отправлено в акцию карточек {st['sent']} из {st['cards']}"]
    if st["in"]:
        parts.append(f"уже в акции — {st['in']}")
    if st["added"]:
        parts.append(f"уже отправлены, ждём WB — {st['added']}")
    bad = st["skipped"] + st["errors"]
    if bad:
        parts.append(f"не отправлено ({len(bad)}): " + "; ".join(bad[:5]) + (" …" if len(bad) > 5 else ""))
    if st["sent"]:
        parts.append("WB покажет участие при следующем обновлении календаря")
    return f"{what}: " + "; ".join(parts) + ".", ("warn" if bad else "ok")


# --- «Внимание» ----------------------------------------------------------------------------------

def attention_items(db: Session, now: datetime | None = None) -> list[tuple[str, str, str, int]]:
    """(уровень, текст, ссылка, число): обычная акция стартует в ближайшие `ATTENTION_HOURS` часов,
    а у заметного числа её карточек, уже стоящих в акции, плановая цена ниже порога — после старта
    программа вернёт эталон, и они из акции выйдут."""
    from priceapp.templating import local
    now = now or now_utc()
    soon = db.query(WbPromotion).filter(WbPromotion.type == REGULAR, WbPromotion.start_at.isnot(None),
                                        WbPromotion.start_at > now,
                                        WbPromotion.start_at <= now + timedelta(hours=ATTENTION_HOURS)).all()
    if not soon:
        return []
    out = []
    for p in soon:
        bad = [pc for pc in classify(db, [p]) if pc.item.in_action and pc.status == RETURNS]
        if len(bad) >= ATTENTION_MIN:
            acc = db.get(Account, p.account_id)
            out.append(("warn", f"Акция «{p.name}» ({acc.name if acc else ''}) стартует {local(p.start_at)}: у карточек "
                                "в акции плановая цена ниже порога — после старта вернётся эталон, и они выйдут "
                                "из акции. Поднимите порог для акции или уберите их из акции в кабинете WB",
                        f"/wb-promo?promo={p.promo_id}&acc={p.account_id}&flt={RETURNS}", len(bad)))
    return out


# --- Excel ---------------------------------------------------------------------------------------

HEAD = ["Раздел", "Тип", "Категория WB / артикул 1С", "Название", "Категория WB", "Кабинеты", "ID акции", "Акция",
        "Начало", "Конец", "Карточек в акции", "Карточек всего", "Эталон, ₽", "Сейчас на WB, ₽",
        "Сейчас ниже эталона, %", "Стоит «эталон − N%»", "Цена при входе, ₽", "Рекомендуемая скидка, %",
        "% от эталона для вхождения", "Порог сейчас, %", "Откуда порог", "Пометка", "Участвует до",
        "Новый порог для акции, %", "Участвовать до (дата)", "Добавить в акцию (да)"]
COL_KIND, COL_KEY, COL_PROMO, COL_NEW, COL_ADD = "Тип", "Категория WB / артикул 1С", "ID акции", \
    "Новый порог для акции, %", "Добавить в акцию (да)"
COL_LEAVE = "Участвовать до (дата)"
YES = {"да", "yes", "1", "+", "true", "истина"}


def export_xlsx(entries: list[wd.Entry]) -> bytes:
    """Файл для правки — те же разделы, что на странице (по акциям, если их несколько). Правятся три
    колонки (выделены цветом) и выгружаются ПУСТЫМИ — что вписали, то и меняется."""
    import io

    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from priceapp.templating import local
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Акция WB")
    widths = [22, 11, 24, 36, 18, 26, 10, 30, 14, 14, 9, 9, 10, 10, 11, 12, 11, 11, 12, 10, 20, 30, 14, 12, 14, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "D2"
    edit = PatternFill("solid", fgColor="FFF4C2")
    head_fill = PatternFill("solid", fgColor="DCE6F2")
    promo_fill = PatternFill("solid", fgColor="B8CCE4")
    cat_fill = PatternFill("solid", fgColor="EEF2F7")
    edit_cols = {HEAD.index(COL_NEW), HEAD.index(COL_LEAVE), HEAD.index(COL_ADD)}

    def row_of(values, bold=False, fill=None, indent=0, editable=True):
        out = []
        for i, v in enumerate(values):
            c = WriteOnlyCell(ws, value=v)
            if bold:
                c.font = Font(bold=True)
            if i in edit_cols and editable:
                c.fill = edit
            elif fill:
                c.fill = fill
            if indent and i == 2:
                c.alignment = Alignment(indent=indent)
            out.append(c)
        return out

    ws.append(row_of(HEAD, bold=True))
    for e in entries:
        if e.kind in ("head", "promo"):
            ws.append(row_of([e.text.upper()] + [None] * (len(HEAD) - 1), bold=True,
                             fill=promo_fill if e.kind == "promo" else head_fill, editable=False))
            continue
        r, promo = e.row, e.row.promo
        ws.append(row_of([e.section, wd.KIND_LABELS[r.kind], r.key, r.name, r.category if r.kind == wd.ARTICLE else "",
                          ", ".join(r.accounts), promo.promo_id, promo.name, local(promo.start_at), local(promo.end_at),
                          r.n_in, len(r.cards), r.etalon, r.current, r.below_text, r.held_text, r.entry_text,
                          r.plan_disc_text, r.need_text, r.thr_text, r.thr_src_text, r.mark, r.leave_text,
                          None, None, None],
                         bold=r.kind == wd.CATEGORY, fill=cat_fill if r.kind == wd.CATEGORY else None,
                         indent=2 * e.level))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
