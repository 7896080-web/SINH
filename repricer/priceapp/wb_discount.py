"""Порог скидки от эталона на WB (страница «Порог скидки WB»).

Эталон — цена WB ПОСЛЕ скидки продавца (`reference.py`). Акция площадки её опускает, и правило
возврата (`reference.run`) возвращает ПОЛНЫЙ эталон, когда цена ушла ниже эталона больше чем на
ПОРОГ. Порог по умолчанию — `ref_drop_percent` (10%), на все товары; здесь его задают КАТЕГОРИИ WB
(предмет карточки) или АРТИКУЛУ 1С — на все кабинеты WB или на один. Главнее всех артикул в
кабинете, затем артикул на все кабинеты, категория в кабинете, категория на все кабинеты,
умолчание. Порог действует, пока его не изменят руками.

«Отправить сейчас» (`send_now`, решение владельца 17.10) — разовая отправка на карточки WB
артикула или категории цены «эталон − N%», N меньше порога. Цена до скидки = эталонная ×
(100 − N) / 100, скидка продавца — эталонная. Эталон при этом НЕ меняется: площадки-последователи
идут за ним как шли. Дальше — обычный сценарий:
  * цена держится, пока не ушла ниже порога (акция в пределах порога её не трогает);
  * ушла ниже порога — правило возврата ставит ПОЛНЫЙ эталон, сдвиг снимается;
  * эталон пересчитан по курсу — уходит эталон по курсу − тот же N% (`ReferencePrice.held_pct`):
    иначе первый же сдвиг курса молча отменил бы решение человека;
  * смена эталона руками и любая отправка цены по очереди записывают новый эталон — сдвиг снимается.

Каждая правка порога и каждая отправка пишутся в журнал (`HISTORY_ACTIONS`): страница показывает
историю и выгружает её в Excel.

Список (страница и Excel, решение владельца 17.10): сначала ОБЩИЕ артикулы — те, что есть в
нескольких кабинетах WB, — затем каждая категория со списком своих артикулов. Новое значение у
КАТЕГОРИИ ставится на все её артикулы (`set_category`): их собственные пороги снимаются, действует
порог категории. Не трогали категорию — меняются только изменённые артикулы. Общие артикулы
правятся своим списком: каскад категории их своих порогов не лишает (без своего порога на них и
так действует порог категории).
"""
from __future__ import annotations

from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from sqlalchemy.orm import Session

from priceapp import audit, settings
from priceapp.models import Account, PriceChange, PriceChangeStatus, ReferencePrice, WbDiscountRule
from priceapp.timeutils import now_utc

CATEGORY, ARTICLE = "category", "article"
KIND_LABELS = {CATEGORY: "Категория", ARTICLE: "Артикул"}
SOURCE = "wb_now"
MAX_THRESHOLD = Decimal(90)
CLEAR_CELL = "-"
ALL_ACCOUNTS = "Все кабинеты WB"
HISTORY_ACTIONS = ("wb_threshold", "wb_send_now", "wb_threshold_default")


def norm(category: str) -> str:
    return (category or "").strip().casefold()


def parse_kind(raw) -> str | None:
    t = str(raw or "").strip().casefold()
    if t in ("категория", "category", "кат", "к"):
        return CATEGORY
    if t in ("артикул", "article", "арт", "а"):
        return ARTICLE
    return None


def parse_key(raw) -> str:
    """Артикул из Excel приходит и числом: 27643.0 — это «27643»."""
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    return str(raw if raw is not None else "").strip()


def parse_pct(raw) -> Decimal | None:
    """Процент из поля или ячейки: «15», «15,5», «15%». Не число — InvalidOperation."""
    text = str(raw if raw is not None else "").strip().replace(" ", "").replace(",", ".").rstrip("%")
    if text == "":
        return None
    return Decimal(text)


def cell(raw) -> tuple[str, Decimal | None]:
    """('skip'|'clear'|'set'|'bad', значение): пусто — без изменений, «-» — снять."""
    text = str(raw if raw is not None else "").strip()
    if text == "":
        return "skip", None
    if text == CLEAR_CELL:
        return "clear", None
    try:
        return "set", parse_pct(text)
    except (InvalidOperation, ValueError):
        return "bad", None


def threshold_error(v: Decimal) -> str:
    if not Decimal(0) < v <= MAX_THRESHOLD:
        return f"порог — больше 0 и не больше {fmt(MAX_THRESHOLD)}%"
    return ""


def now_error(v: Decimal) -> str:
    if not Decimal(0) <= v < MAX_THRESHOLD:
        return f"«эталон −» — от 0 до {fmt(MAX_THRESHOLD)}% (0 — вернуть полный эталон)"
    return ""


def fmt(v) -> str:
    return "" if v is None else format(Decimal(v).normalize(), "f").replace(".", ",")


def wb_accounts(db: Session) -> list[Account]:
    return db.query(Account).filter(Account.platform == "wb", Account.is_active.is_(True)).order_by(Account.name).all()


def account_name(db: Session, account_id: int) -> str:
    if not account_id:
        return ALL_ACCOUNTS
    a = db.get(Account, account_id)
    return a.name if a is not None else f"кабинет #{account_id}"


# --- порог -----------------------------------------------------------------------------------

class Thresholds:
    """Порог по товару: артикул (кабинет → все) → категория (кабинет → все) → умолчание.
    Строится один раз на проход: правил мало, а спрашивают по каждому SKU."""

    def __init__(self, db: Session):
        from priceapp.pricing import article_map
        from priceapp.reference import drop_percent
        self.default = drop_percent(db)
        self.by_art: dict[tuple[str, int], Decimal] = {}
        self.by_cat: dict[tuple[str, int], Decimal] = {}
        for r in db.query(WbDiscountRule).filter(WbDiscountRule.max_disc.isnot(None)):
            if r.kind == ARTICLE:
                self.by_art[(r.key, r.account_id or 0)] = Decimal(r.max_disc)
            else:
                self.by_cat[(norm(r.key), r.account_id or 0)] = Decimal(r.max_disc)
        self._articles = article_map(db) if self.by_art else {}

    def for_category(self, category: str, account_id: int = 0) -> tuple[Decimal, str]:
        k = norm(category)
        if account_id and (k, account_id) in self.by_cat:
            return self.by_cat[(k, account_id)], "категория, кабинет"
        if (k, 0) in self.by_cat:
            return self.by_cat[(k, 0)], "категория"
        return self.default, "по умолчанию"

    def for_article(self, article: str, category: str = "", account_id: int = 0) -> tuple[Decimal, str]:
        if account_id and (article, account_id) in self.by_art:
            return self.by_art[(article, account_id)], "артикул, кабинет"
        if (article, 0) in self.by_art:
            return self.by_art[(article, 0)], "артикул"
        return self.for_category(category, account_id)

    def for_item(self, item_id: str, category: str = "", account_id: int = 0) -> tuple[Decimal, str]:
        from priceapp.pricing import article_of
        art = article_of(self._articles, item_id) if self.by_art else item_id
        return self.for_article(art, category, account_id)


# --- цена «эталон − N%» ---------------------------------------------------------------------

def held_list(ref_price: int, pct) -> int:
    """Цена до скидки при сдвиге N% от эталонной; скидка продавца остаётся эталонной."""
    if not pct:
        return int(ref_price)
    return int((Decimal(ref_price) * (100 - Decimal(pct)) / 100).quantize(Decimal(1), ROUND_HALF_UP))


def sale_of(price: int, disc: int) -> int:
    return int((Decimal(price) * (100 - Decimal(disc)) / 100).quantize(Decimal(1), ROUND_HALF_UP))


def held_pair(ref: ReferencePrice) -> tuple[int, int]:
    """(цена до скидки, скидка), которая должна стоять на WB по эталону с учётом сдвига."""
    from priceapp.reference import ref_discount_int
    return held_list(int(ref.ref_price), ref.held_pct), ref_discount_int(ref)


def card_category(card) -> str:
    return next((r.category for r in card.rows if getattr(r, "category", "")), "")


@dataclass
class Card:
    """Карточка WB (кабинет + nmID) со всеми её SKU — у одной карточки бывает два артикула 1С."""
    account: Account
    nm: str
    category: str
    articles: set = field(default_factory=set)
    item_ids: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    ref_list: int | None = None
    ref_disc: int | None = None
    ref_sale: int | None = None
    held: Decimal | None = None
    cur_list: int | None = None
    cur_sale: int | None = None

    @property
    def below_pct(self) -> Decimal | None:
        """На сколько % цена после скидки сейчас ниже эталона (минус — выше)."""
        if not self.ref_sale or not self.cur_sale:
            return None
        return ((1 - Decimal(self.cur_sale) / Decimal(self.ref_sale)) * 100).quantize(Decimal("0.1"), ROUND_HALF_UP)


def cards(db: Session, ets=None) -> list[Card]:
    """Все карточки WB с эталоном, без повторов (одна карточка под двумя артикулами — одна)."""
    from priceapp import etalon
    ets = ets if ets is not None else etalon.etalons(db)
    held = {(r.account_id, r.item_id): r.held_pct for r in db.query(ReferencePrice)
            .filter(ReferencePrice.held_pct.isnot(None))}
    out: dict[tuple, Card] = {}
    for art, et in ets.items():
        for c in et.cards:
            k = (c.account.id, c.nm)
            m = out.get(k)
            if m is None:
                m = out[k] = Card(c.account, c.nm, card_category(c))
            m.articles.add(art)
            m.item_ids += [i for i in c.item_ids if i not in m.item_ids]
            m.rows += [r for r in c.rows if r not in m.rows]
            if c.ref_list and (m.ref_list is None or c.ref_list > m.ref_list):
                m.ref_list, m.ref_disc, m.ref_sale = c.ref_list, c.ref_disc, c.ref_sale
            if c.cur_list and (m.cur_list is None or c.cur_list > m.cur_list):
                m.cur_list, m.cur_sale = c.cur_list, c.cur_sale
            for i in c.item_ids:
                if held.get((c.account.id, i)) is not None:
                    m.held = Decimal(held[(c.account.id, i)])
    return list(out.values())


def select(all_cards: list[Card], kind: str, key: str, account_id: int = 0) -> list[Card]:
    if account_id:
        all_cards = [c for c in all_cards if c.account.id == account_id]
    if kind == ARTICLE:
        return [c for c in all_cards if key in c.articles]
    k = norm(key)
    return [c for c in all_cards if norm(c.category) == k]


def shared_articles(all_cards: list[Card]) -> set[str]:
    """Артикулы, которые есть в нескольких кабинетах WB, — общий список."""
    accs: dict[str, set] = {}
    for c in all_cards:
        for a in c.articles:
            accs.setdefault(a, set()).add(c.account.id)
    return {a for a, s in accs.items() if len(s) > 1}


def category_articles(all_cards: list[Card], category: str, account_id: int = 0) -> set[str]:
    """Артикулы из списка категории (без общих) — на них каскадом ставится порог категории."""
    shared = shared_articles(all_cards)
    k = norm(category)
    return {a for c in all_cards if norm(c.category) == k and (not account_id or c.account.id == account_id)
            for a in c.articles if a not in shared}


# --- правила -----------------------------------------------------------------------------------

def get_rule(db: Session, kind: str, key: str, account_id: int = 0) -> WbDiscountRule | None:
    q = db.query(WbDiscountRule).filter(WbDiscountRule.kind == kind, WbDiscountRule.account_id == (account_id or 0))
    if kind == ARTICLE:
        return q.filter(WbDiscountRule.key == key).first()
    return next((r for r in q if norm(r.key) == norm(key)), None)


def set_threshold(db: Session, kind: str, key: str, value: Decimal | None, actor: str,
                  account_id: int = 0, where: str = "") -> bool:
    """Свой порог артикула или категории в кабинете (0 — на все кабинеты); None — снять, тогда
    действует следующий по старшинству. Пишет историю. Возвращает, изменилось ли. Не коммитит."""
    rule = get_rule(db, kind, key, account_id)
    old = Decimal(rule.max_disc) if rule is not None and rule.max_disc is not None else None
    if old == value:
        return False
    if rule is None:
        rule = WbDiscountRule(kind=kind, key=key, account_id=account_id or 0, updated_by=actor)
        db.add(rule)
    rule.max_disc = value
    rule.updated_by = actor
    rule.updated_at = now_utc()
    db.flush()
    audit.log(db, actor, "wb_threshold", (where or account_name(db, account_id))[:100],
              f"{KIND_LABELS[kind].lower()} {key}: порог {fmt(old) or 'общий'} -> {fmt(value) or 'общий'}")
    return True


def set_category(db: Session, key: str, value: Decimal | None, actor: str, account_id: int = 0, where: str = "",
                 all_cards: list[Card] | None = None) -> tuple[bool, int]:
    """Порог категории. НОВОЕ значение (не None) ставится на все артикулы её списка: их свои пороги
    в этом же разделе снимаются, и на них действует порог категории. Раздел «все кабинеты» снимает
    ещё и свои пороги этой категории и её артикулов в отдельных кабинетах — иначе «на все артикулы
    категории» было бы неправдой. «-» (None) — только снять порог самой категории: исключения
    артикулов остаются. Возвращает (изменился ли порог категории, сколько своих порогов снято).
    Не коммитит."""
    changed = set_threshold(db, CATEGORY, key, value, actor, account_id, where)
    if value is None:
        return changed, 0
    arts = category_articles(all_cards if all_cards is not None else cards(db), key, account_id)
    cleared = 0
    q = db.query(WbDiscountRule).filter(WbDiscountRule.max_disc.isnot(None))
    for r in q.filter(WbDiscountRule.kind == ARTICLE).all():
        if r.key in arts and (not account_id or r.account_id == account_id):
            cleared += set_threshold(db, ARTICLE, r.key, None, actor, r.account_id,
                                     f"{account_name(db, r.account_id)} — каскад категории {key}")
    if not account_id:
        for r in q.filter(WbDiscountRule.kind == CATEGORY, WbDiscountRule.account_id != 0).all():
            if norm(r.key) == norm(key):
                cleared += set_threshold(db, CATEGORY, r.key, None, actor, r.account_id,
                                         f"{account_name(db, r.account_id)} — каскад категории {key}")
    return changed, cleared


def set_default(db: Session, value: Decimal, actor: str) -> bool:
    """Порог на все товары (`ref_drop_percent`). Пишет историю. Не коммитит."""
    old = settings.get(db, settings.REF_DROP_PERCENT)
    new = format(value.normalize(), "f")
    if old == new:
        return False
    settings.put(db, settings.REF_DROP_PERCENT, new)
    audit.log(db, actor, "wb_threshold_default", ALL_ACCOUNTS, f"порог на все товары: {old} -> {new}")
    return True


# --- «Отправить сейчас» ----------------------------------------------------------------------

def send_now(db: Session, kind: str, key: str, pct: Decimal, actor: str, client_factory=None,
             all_cards: list[Card] | None = None, account_id: int = 0) -> dict:
    """На карточки WB артикула или категории (в кабинете `account_id`, 0 — во всех) — «эталон −
    pct%» одним заданием WB на карточку. pct = 0 — вернуть полный эталон. Карточку не трогает и
    называет, если: эталона нет, pct не меньше её порога (ближайший опрос вернул бы эталон), в
    очереди подтверждённая человеком цена, снижение больше `wb_max_drop_percent` за раз (WB такое
    не применит), ниже пола любого размера. Коммитит по кабинету."""
    from priceapp import accounts, etalon, progress, rates
    from priceapp.dispatch import _record_sent
    from priceapp.models import OnecCost
    from priceapp.platforms import WB_PRICES
    from priceapp.pricing import get_rule as platform_rule
    from priceapp.pricing import item_facts
    all_cards = all_cards if all_cards is not None else cards(db)
    chosen = select(all_cards, kind, key, account_id)
    limits = Thresholds(db)
    own = 0
    if kind == CATEGORY:
        # категория правит артикулы без своего порога; со своим — правятся своей строкой
        def governed(c: Card) -> bool:
            return any(not limits.for_article(a, c.category, c.account.id)[1].startswith("артикул")
                       for a in c.articles)
        own = sum(1 for c in chosen if not governed(c))
        chosen = [c for c in chosen if governed(c)]
    st = {"cards": len(chosen), "sent": 0, "same": 0, "skipped": [], "errors": [], "own": own}
    if not chosen:
        return st
    rule = platform_rule(db, "wb")
    rate = rates.current(db)
    ids = [i for c in chosen for i in c.item_ids]
    costs = {c.item_id: c.cost_usd for c in db.query(OnecCost).filter(OnecCost.item_id.in_(ids))}
    max_drop = Decimal(settings.get(db, settings.WB_MAX_DROP_PERCENT) or "33")
    busy = {(a, i) for a, i in db.query(PriceChange.account_id, PriceChange.item_id).filter(
        PriceChange.is_test.is_(False), PriceChange.status == PriceChangeStatus.approved.value)}
    held_value = pct if pct else None
    plan: dict[int, list[tuple[Card, int, int, int]]] = {}
    same: list[Card] = []
    for c in chosen:
        name = f"{c.account.name} {c.nm}"
        if not c.ref_list:
            st["skipped"].append(f"{name}: нет эталона")
            continue
        thr = min(limits.for_item(i, c.category, c.account.id)[0] for i in c.item_ids)
        if pct >= thr:
            st["skipped"].append(f"{name}: порог {fmt(thr)}% — «эталон −» должно быть меньше порога, "
                                 "иначе ближайший опрос вернул бы эталон")
            continue
        if any((c.account.id, i) in busy for i in c.item_ids):
            st["skipped"].append(f"{name}: в очереди подтверждённая цена — сначала она")
            continue
        price, d = held_list(c.ref_list, pct), int(c.ref_disc or 0)
        sale = sale_of(price, d)
        if c.cur_list == price and c.cur_sale == sale:
            same.append(c)
            continue
        if c.cur_sale and Decimal(sale) < Decimal(c.cur_sale) * (1 - max_drop / 100):
            st["skipped"].append(f"{name}: {c.cur_sale} -> {sale} ₽ — снижение больше {fmt(max_drop)}% за раз "
                                 "WB не применит; задайте меньший процент")
            continue
        problem = etalon.floor_problem(rule, item_facts(rule, c.rows).commission, rate, costs, c.item_ids, price, d)
        if problem:
            st["skipped"].append(f"{name}: {problem}")
            continue
        plan.setdefault(c.account.id, []).append((c, price, d, sale))
    st["same"] = len(same)
    now = now_utc()
    note = f"эталон − {fmt(pct)}% (порог скидки WB)" if pct else "полный эталон (порог скидки WB)"
    for group in plan.values():
        a = group[0][0].account
        try:
            client = accounts.client_for(db, a, client_factory)
        except Exception as e:
            st["errors"].append(f"{a.name}: {e}"[:300])
            continue
        done = []
        for start in range(0, len(group), 1000):
            chunk = group[start:start + 1000]
            body = {"data": [{"nmID": int(c.nm), "price": price, "discount": d} for c, price, d, _ in chunk]}
            try:
                data = client._post(f"{WB_PRICES}/api/v2/upload/task", body)
                if isinstance(data, dict) and data.get("error"):
                    raise RuntimeError(data.get("errorText") or "WB вернул error=true")
            except Exception as e:
                st["errors"].append(f"{a.name}: {e}"[:300])
                continue
            done += chunk
        for c, price, d, sale in done:
            _mark_held(db, a.id, c.item_ids, held_value)
            for i in c.item_ids:
                _record_sent(db, i, a.id, price)       # цена наша — диапазон безопасности её не «вернёт»
            etalon.journal_sent(db, a.id, c.item_ids, c.rows[0].barcode if c.rows else "", c.cur_list, price,
                                SOURCE, f"{note}: {sale} ₽ после скидки {d}%", actor)
            for r in c.rows:
                r.current_price, r.current_sale_price = price, sale
            c.cur_list, c.cur_sale, c.held = price, sale, held_value
        st["sent"] += len(done)
        if done:
            sample = "; ".join(f"{c.nm}: -> {sale}" for c, _, _, sale in done[:20])
            audit.log(db, actor, "wb_send_now", a.name,
                      f"{KIND_LABELS[kind].lower()} {key}: {note}, карточек {len(done)} — {sample}"[:1000])
        db.commit()
    for c in same:
        _mark_held(db, c.account.id, c.item_ids, held_value)
        c.held = held_value
    rule_row = get_rule(db, kind, key, account_id)
    if rule_row is None:
        rule_row = WbDiscountRule(kind=kind, key=key, account_id=account_id or 0, updated_by=actor)
        db.add(rule_row)
    rule_row.now_pct, rule_row.now_at, rule_row.now_by = pct, now, actor
    if not st["sent"] and (st["skipped"] or st["errors"]):
        # ничего не ушло — в историю всё равно: «нажимал, но WB/пол не дали» тоже ответ на вопрос
        audit.log(db, actor, "wb_send_now", account_name(db, account_id)[:100],
                  f"{KIND_LABELS[kind].lower()} {key}: {note} — не отправлено ни одной карточки: "
                  + "; ".join((st["skipped"] + st["errors"])[:5])[:800])
    db.commit()
    if st["sent"]:
        from priceapp import overview
        overview.mark_dirty(db)
        db.commit()
    return st


def _mark_held(db: Session, account_id: int, item_ids, pct) -> None:
    for ref in db.query(ReferencePrice).filter(ReferencePrice.account_id == account_id,
                                               ReferencePrice.item_id.in_(list(item_ids))):
        ref.held_pct = pct


def merge_reports(sts: list[dict]) -> dict:
    out = {"cards": 0, "sent": 0, "same": 0, "skipped": [], "errors": [], "own": 0}
    for st in sts:
        for k in ("cards", "sent", "same", "own"):
            out[k] += st.get(k, 0)
        out["skipped"] += st["skipped"]
        out["errors"] += st["errors"]
    return out


def report(what: str, pct: Decimal | None, st: dict) -> tuple[str, str]:
    """Текст и уровень флеша по итогам `send_now`; pct None — проценты разные (файл)."""
    if pct is not None:
        what = f"{what}: " + (f"эталон − {fmt(pct)}%" if pct else "полный эталон")
    own = st.get("own", 0)
    if not st["cards"]:
        if own:
            return (f"{what} — не отправлено: у всех {own} карточек свой порог артикула, их правят строкой "
                    "артикула.", "warn")
        return f"{what} — на WB нет карточек с эталоном.", "warn"
    parts = [f"отправлено карточек {st['sent']} из {st['cards']}"]
    if st["same"]:
        parts.append(f"уже стояло так — {st['same']}")
    if own:
        parts.append(f"не трогали {own} — у их артикулов свой порог, правьте строкой артикула")
    bad = st["skipped"] + st["errors"]
    if bad:
        parts.append(f"не отправлено ({len(bad)}): " + "; ".join(bad[:5]) + (" …" if len(bad) > 5 else ""))
    return f"{what} — " + "; ".join(parts) + ".", ("warn" if bad else "ok")


# --- страница и Excel ---------------------------------------------------------------------------

@dataclass
class Row:
    kind: str
    key: str
    name: str = ""
    category: str = ""
    cards: int = 0
    etalon: int | None = None
    current: int | None = None
    below_min: Decimal | None = None
    below_max: Decimal | None = None
    held: list = field(default_factory=list)      # сдвиги «эталон − N%», что стоят сейчас
    sales: list = field(default_factory=list)     # цены после скидки на WB сейчас, по карточкам
    accounts: list = field(default_factory=list)  # кабинеты, где есть этот артикул (категория)
    shared: bool = False                          # артикул есть в нескольких кабинетах — общий список
    articles: int = 0                             # категория: артикулов в её списке (без общих)
    own_articles: int = 0                         # категория: из них со своим порогом здесь
    rule: WbDiscountRule | None = None            # свой порог там, где смотрим (кабинет или все)
    others: list = field(default_factory=list)    # «все кабинеты»: свои пороги отдельных кабинетов
    eff: Decimal | None = None
    eff_src: str = ""

    @property
    def sel(self) -> str:
        return f"{self.kind[0]}:{self.key}"

    @property
    def own(self) -> Decimal | None:
        return Decimal(self.rule.max_disc) if self.rule is not None and self.rule.max_disc is not None else None

    @property
    def below_text(self) -> str:
        if self.below_min is None:
            return ""
        if self.below_min == self.below_max:
            return fmt(self.below_min)
        return f"{fmt(self.below_min)} … {fmt(self.below_max)}"


def table(db: Session, account_id: int = 0, all_cards: list[Card] | None = None) -> tuple[list[Row], list[Row]]:
    """(категории, артикулы) кабинета (0 — всех кабинетов WB): свой порог, действующий порог и что
    стоит на WB."""
    from priceapp import etalon
    ets = etalon.etalons(db)
    all_cards = all_cards if all_cards is not None else cards(db, ets)
    shared = shared_articles(all_cards)
    art_accs: dict[str, set] = {}
    for c in all_cards:
        for a in c.articles:
            art_accs.setdefault(a, set()).add(c.account.name)
    if account_id:
        all_cards = [c for c in all_cards if c.account.id == account_id]
    limits = Thresholds(db)
    names = {a.id: a.name for a in db.query(Account)}
    rules: dict[tuple, WbDiscountRule] = {}
    others: dict[tuple, list] = {}
    for r in db.query(WbDiscountRule):
        k = (r.kind, r.key if r.kind == ARTICLE else norm(r.key))
        if (r.account_id or 0) == (account_id or 0):
            rules[k] = r
        elif not account_id and r.max_disc is not None:
            others.setdefault(k, []).append(f"{names.get(r.account_id, r.account_id)} {fmt(r.max_disc)}%")
    cats: dict[str, Row] = {}
    arts: dict[str, Row] = {}
    art_cats: dict[str, dict[str, int]] = {}
    for c in all_cards:
        if c.category:
            _add(cats.setdefault(norm(c.category), Row(CATEGORY, c.category, category=c.category)), c)
        for art in c.articles:
            row = arts.get(art)
            if row is None:
                et = ets.get(art)
                row = arts[art] = Row(ARTICLE, art, name=et.name if et else "", etalon=et.sale if et else None)
            _add(row, c)
            if c.category:
                art_cats.setdefault(art, {}).setdefault(c.category, 0)
                art_cats[art][c.category] += 1
    # правило на ключ, которого в каталоге нет (карточку убрали, опечатка в файле), — тоже видно
    for (kind, k), r in rules.items():
        if kind == CATEGORY and k not in cats:
            cats[k] = Row(CATEGORY, r.key, category=r.key)
        if kind == ARTICLE and k not in arts:
            arts[k] = Row(ARTICLE, r.key)
    for k, row in cats.items():
        row.rule, row.others = rules.get((CATEGORY, k)), others.get((CATEGORY, k), [])
        row.eff, row.eff_src = limits.for_category(row.key, account_id)
    for art, row in arts.items():
        cs = art_cats.get(art, {})
        row.category = max(cs, key=cs.get) if cs else row.category
        row.rule, row.others = rules.get((ARTICLE, art)), others.get((ARTICLE, art), [])
        row.eff, row.eff_src = limits.for_article(art, row.category, account_id)
        row.current = int(sum(row.sales) / len(row.sales)) if row.sales else None
        row.accounts, row.shared = sorted(art_accs.get(art, ())), art in shared
        if not row.shared and row.category and norm(row.category) in cats:
            c = cats[norm(row.category)]
            c.articles += 1
            c.own_articles += row.own is not None
    return (sorted(cats.values(), key=lambda r: r.key.casefold()),
            sorted(arts.values(), key=lambda r: r.key))


def _add(row: Row, c: Card) -> None:
    row.cards += 1
    if c.cur_sale:
        row.sales.append(c.cur_sale)
    b = c.below_pct
    if b is not None:
        row.below_min = b if row.below_min is None else min(row.below_min, b)
        row.below_max = b if row.below_max is None else max(row.below_max, b)
    if c.held is not None and c.held not in row.held:
        row.held.append(c.held)


FILTERS = {"": "все", "own": "со своим порогом", "held": "стоит «эталон − N%»",
           "below": "цена сейчас ниже эталона", "over": "ниже порога (вернётся эталон)"}


def keep(row: Row, q: str = "", cat: str = "", flt: str = "") -> bool:
    ql = q.strip().casefold()
    if ql and ql not in row.key.casefold() and ql not in (row.name or "").casefold():
        return False
    if cat and norm(row.category) != norm(cat):
        return False
    if flt == "own" and row.own is None:
        return False
    if flt == "held" and not row.held:
        return False
    if flt == "below" and not (row.below_max is not None and row.below_max > 0):
        return False
    if flt == "over" and not (row.below_max is not None and row.eff is not None and row.below_max > row.eff):
        return False
    return True


@dataclass
class Entry:
    """Строка списка: заголовок раздела или строка категории / артикула (level 1 — артикул под категорией)."""
    kind: str                      # head / row
    text: str = ""
    row: Row | None = None
    level: int = 0
    section: str = ""


NO_CATEGORY = "Без категории"
SHARED_TITLE = "Общие артикулы — есть в нескольких кабинетах"


def grouped(cats: list[Row], arts: list[Row], q: str = "", cat: str = "", flt: str = "") -> list[Entry]:
    """Общие артикулы, затем каждая категория и её артикулы. Категория показывается, если подходит
    под отбор сама или хоть один её артикул; совпала по названию с поиском — со всеми артикулами."""
    out: list[Entry] = []
    shared = [a for a in arts if a.shared and keep(a, q, cat, flt)]
    if shared:
        out.append(Entry("head", f"{SHARED_TITLE} ({len(shared)})", section=SHARED_TITLE))
        out += [Entry("row", row=a, section=SHARED_TITLE) for a in shared]
    kids: dict[str, list[Row]] = {}
    for a in arts:
        if not a.shared:
            kids.setdefault(norm(a.category) if a.category else "", []).append(a)
    ql = q.strip().casefold()
    for c in cats:
        if cat and norm(c.key) != norm(cat):
            continue
        name_hit = bool(ql) and ql in c.key.casefold()
        mine = [a for a in kids.get(norm(c.key), []) if keep(a, "" if name_hit else q, "", flt)]
        if not (keep(c, q, "", flt) or mine):
            continue
        out.append(Entry("row", row=c, section=c.key))
        out += [Entry("row", row=a, level=1, section=c.key) for a in mine]
    orphans = [a for a in kids.get("", []) if not cat and keep(a, q, "", flt)]
    if orphans:
        out.append(Entry("head", f"{NO_CATEGORY} ({len(orphans)})", section=NO_CATEGORY))
        out += [Entry("row", row=a, section=NO_CATEGORY) for a in orphans]
    return out


def covered(targets: list[tuple[str, str]], arts: dict[str, Row]) -> list[tuple[str, str]]:
    """Массовая правка: строку артикула, чья категория тоже в правке, отдельно не правим — её
    покрывает каскад категории (иначе каждый артикул получил бы свой порог, равный порогу категории,
    — лишние строки, которые разорвали бы связь с категорией). Общие артикулы правятся всегда сами."""
    cats = {norm(k) for kind, k in targets if kind == CATEGORY}
    out = []
    for kind, k in targets:
        row = arts.get(k) if kind == ARTICLE else None
        if row is not None and not row.shared and row.category and norm(row.category) in cats:
            continue
        out.append((kind, k))
    return out


HEAD = ["Раздел", "Тип", "Категория WB / артикул 1С", "Название", "Категория WB", "Кабинеты", "Где действует",
        "Карточек WB", "Эталон, ₽", "Сейчас на WB, ₽", "Сейчас ниже эталона, %", "Стоит «эталон − N%»",
        "Порог сейчас, %", "Свой порог, %", "Новый порог, %", "Отправить сейчас: эталон −, %"]
COL_KIND, COL_KEY, COL_WHERE, COL_NEW, COL_NOW = HEAD[1], HEAD[2], HEAD[6], HEAD[14], HEAD[15]
OLD_FORMAT_COL = "Порог: макс. скидка от эталона, %"      # выгрузка rp42: там порог стоял заполненным


def export_xlsx(entries: list[Entry], where: str) -> bytes:
    """Файл для правки: те же разделы, что на странице. Правятся две колонки (выделены цветом) и
    выгружаются ПУСТЫМИ — что вписали, то и меняется; остальное справочное."""
    import io

    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Порог скидки WB")
    widths = [22, 11, 26, 40, 20, 30, 18, 9, 10, 10, 12, 12, 20, 10, 12, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "D2"
    edit = PatternFill("solid", fgColor="FFF4C2")
    head_fill = PatternFill("solid", fgColor="DCE6F2")
    cat_fill = PatternFill("solid", fgColor="EEF2F7")
    edit_cols = {HEAD.index(COL_NEW), HEAD.index(COL_NOW)}

    def row_of(values, bold=False, fill=None, indent=0, editable=True):
        out = []
        for i, v in enumerate(values):
            c = WriteOnlyCell(ws, value=v)
            if bold:
                c.font = Font(bold=True)
            if i in edit_cols and editable:
                c.fill = edit
            elif fill:
                c.fill = fill
            if indent and i == 2:
                c.alignment = Alignment(indent=indent)
            out.append(c)
        return out

    ws.append(row_of(HEAD, bold=True))
    for e in entries:
        if e.kind == "head":
            ws.append(row_of([e.text.upper()] + [None] * (len(HEAD) - 1), bold=True, fill=head_fill, editable=False))
            continue
        r = e.row
        ws.append(row_of([e.section, KIND_LABELS[r.kind], r.key, r.name, r.category if r.kind == ARTICLE else "",
                          ", ".join(r.accounts), where, r.cards or None, r.etalon, r.current, r.below_text,
                          ", ".join(fmt(h) for h in sorted(r.held)),
                          f"{fmt(r.eff)} ({r.eff_src})" if r.eff is not None else "",
                          float(r.own) if r.own is not None else None, None, None],
                         bold=r.kind == CATEGORY, fill=cat_fill if r.kind == CATEGORY else None, indent=2 * e.level))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HISTORY_HEAD = ["Когда", "Кто", "Где", "Что"]


def history(db: Session, limit: int = 50) -> list:
    from priceapp.models import AuditLog
    return (db.query(AuditLog).filter(AuditLog.action.in_(HISTORY_ACTIONS))
            .order_by(AuditLog.id.desc()).limit(limit).all())
