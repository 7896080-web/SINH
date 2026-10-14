"""«Утверждение цен»: страница и Excel по кабинету — строка = «карточка площадки + артикул 1С».

Тот же расчёт, что в книге «WB — утверждение цен по карточкам»: вводится маржинальность
артикула в кабинете (к получению / себестоимость); из неё — цена в списке артикула (как у
`pricing.decide`, наценка «маржинальность»), ИТОГОВАЯ цена карточки — НАИБОЛЬШАЯ из цен её
артикулов (так цену ставит `dispatch`), цена продавца и покупателя, коэффициент от базовой,
прибыль на единицу, окупаемость, предупреждения. Эталонная цена (`reference.py`) показана
рядом с текущей ценой площадки.

Экономика сверх комиссии (логистика с отказами `f_logistics`, доля выплаты `k`, цена
покупателя / цена продавца `w`, нетто-продажи) берётся из таблицы `card_econ` — её заполняет
импорт из Excel (колонки «k», «W/P», «Логистика+прочее на ед., ₽», «Нетто-продаж за 7 нед., шт»);
без неё: k = 1 − комиссия, w — 0,594 у WB и 1 у остальных, логистика 0 (Lamoda — 49 ₽ за обработку возврата).
"""
from __future__ import annotations

import io
import math
from decimal import Decimal
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session

from priceapp import mapping, rates, reference, settings
from priceapp.models import (Account, CardEcon, OnecBarcode, OnecCost, ProductPrice, ReferencePrice)
from priceapp.platforms import PLATFORMS, card_key
from priceapp.pricing import (BLOCK_FLOOR, KIND_MARGIN, SRC_CABINET, _dec, article_of, decide_for, item_facts, load_inputs,
                              target_for)

# --- параметры расчёта (страница «Утверждение цен», в Excel — лист «Параметры») -----------------------------
ECON_TAX, ECON_ROI = "econ_tax_pct", "econ_roi_pct"
ECON_EXTRA_ADS, ECON_RENT = "econ_extra_ads_pct", "econ_rent_month"
ECON_RENT_SHARE, ECON_NET_MONTH = "econ_rent_share_pct", "econ_net_month"
ECON_DEFAULTS = {ECON_TAX: "10", ECON_ROI: "30", ECON_EXTRA_ADS: "5", ECON_RENT: "140000",
                 ECON_RENT_SHARE: "100", ECON_NET_MONTH: "958"}
DEFAULT_ADS = "5"          # реклама + хранение кабинета, % от цены продавца (факт — в настройке кабинета)
W_DEFAULT = {"wb": 0.594}  # цена покупателя / цена продавца: скидка WB (СПП) ≈ 41%
# Логистика и прочее на единицу, ₽, когда по карточке нет своей экономики (card_econ). Lamoda: фиксированные 49 ₽ за
# обработку возврата; ставится на каждую проданную единицу (доля возвратов неизвестна — оценка сверху). Меняется
# настройкой `econ_f_<площадка>`.
F_DEFAULT = {"lamoda": "49"}


def f_key(platform: str) -> str:
    return f"econ_f_{platform}"


def ads_key(account_id: int) -> str:
    return f"econ_ads_pct_{account_id}"


def econ_value(db: Session, key: str) -> float:
    raw = (settings.get(db, key) or ECON_DEFAULTS.get(key, "0")).replace(",", ".").strip()
    try:
        return float(raw)
    except ValueError:
        return float(ECON_DEFAULTS.get(key, "0"))


def _f_default(db: Session, platform: str) -> float:
    raw = (settings.get(db, f_key(platform)) or F_DEFAULT.get(platform, "0")).replace(",", ".").strip()
    try:
        return max(float(raw), 0.0)
    except ValueError:
        return float(F_DEFAULT.get(platform, "0"))


def params(db: Session, account: Account) -> dict:
    """Параметры расчёта кабинета; проценты — долями."""
    ads_raw = (settings.get(db, ads_key(account.id)) or DEFAULT_ADS).replace(",", ".")
    try:
        ads = float(ads_raw)
    except ValueError:
        ads = float(DEFAULT_ADS)
    net_month = econ_value(db, ECON_NET_MONTH) or 1
    return {"tax": econ_value(db, ECON_TAX) / 100, "roi": econ_value(db, ECON_ROI) / 100,
            "ads": ads / 100, "ads_pct": ads, "extra_ads": econ_value(db, ECON_EXTRA_ADS) / 100,
            "rent_month": econ_value(db, ECON_RENT), "rent_share": econ_value(db, ECON_RENT_SHARE) / 100,
            "net_month": net_month,
            "f_default": _f_default(db, account.platform),
            "rent_unit": econ_value(db, ECON_RENT) * econ_value(db, ECON_RENT_SHARE) / 100 / net_month}


def save_params(db: Session, account: Account, form) -> list[str]:
    """Сохранить параметры из формы страницы; список ошибок."""
    errors = []
    pairs = [(ECON_TAX, "tax", "налог"), (ECON_ROI, "roi", "целевая окупаемость"), (ECON_EXTRA_ADS, "extra_ads", "доп. реклама"),
             (ECON_RENT, "rent_month", "аренда"), (ECON_RENT_SHARE, "rent_share", "доля аренды"),
             (ECON_NET_MONTH, "net_month", "продаж в месяц"), (ads_key(account.id), "ads", "реклама кабинета")]
    for key, field, label in pairs:
        raw = str(form.get(field) or "").replace(",", ".").replace(" ", "").strip()
        if raw == "":
            continue
        try:
            v = float(raw)
        except ValueError:
            errors.append(f"«{label}»: «{raw}» — не число")
            continue
        if v < 0 or (field in ("tax", "roi", "extra_ads", "rent_share", "ads") and v > 100) or (field == "net_month" and v <= 0):
            errors.append(f"«{label}»: {raw} вне допустимого")
            continue
        settings.put(db, key, str(v).rstrip("0").rstrip(".") if "." in str(v) else str(v))
    return errors


# --- расчёт строк ---------------------------------------------------------------------------------------------

def ceil10(x: float) -> int:
    return int(math.ceil(x / 10 - 1e-9)) * 10 - 1


def _col_stats(vals):
    vals = [v for v in vals if v is not None]
    return (sum(vals) / len(vals), max(vals)) if vals else (None, None)


def compute(db: Session, account: Account) -> dict:
    """Строки «карточка × артикул 1С» кабинета с полным расчётом. {'rows': [...], 'p': параметры, 'rate': курс}."""
    p = params(db, account)
    platform = account.platform
    inp = load_inputs(db, platform)
    rule = inp.rule
    rate = rates.current(db)
    usd = rate.usd_rub if rate else None
    items = mapping.account_items(db, account.id)                 # item_id -> строки каталога кабинета
    item_ids = list(items)
    info: dict[str, OnecBarcode] = {}
    costs: dict[str, OnecCost] = {}
    for i in range(0, len(item_ids), 500):
        chunk = item_ids[i:i + 500]
        for b in db.query(OnecBarcode).filter(OnecBarcode.item_id.in_(chunk)):
            info.setdefault(b.item_id, b)
        for c in db.query(OnecCost).filter(OnecCost.item_id.in_(chunk)):
            costs[c.item_id] = c
    pps = {pp.item_id: pp for pp in db.query(ProductPrice).filter(ProductPrice.account_id == account.id)}
    refs = {r.item_id: r for r in db.query(ReferencePrice).filter(ReferencePrice.account_id == account.id)}
    econ = {(e.card, e.article): e for e in db.query(CardEcon).filter(CardEcon.account_id == account.id)}
    drop = float(reference.drop_percent(db)) / 100

    # карточка -> {артикул 1С -> [(item_id, строки SKU в этой карточке)]}; карточка -> все её строки каталога
    card_rows: dict[str, list] = {}
    card_arts: dict[str, dict[str, list]] = {}
    for item_id, rows in items.items():
        art = article_of(inp.articles, item_id)
        by_card: dict[str, list] = {}
        for r in rows:
            by_card.setdefault(card_key(platform, r) or r.barcode, []).append(r)
        for ck, rs in by_card.items():
            card_rows.setdefault(ck, []).extend(rs)
            card_arts.setdefault(ck, {}).setdefault(art, []).append((item_id, rs))

    out = []
    for ck in sorted(card_arts):
        facts_card = item_facts(rule, card_rows[ck])
        commission = facts_card.commission
        disc = float(facts_card.discount or 0)
        sale = facts_card.sale or facts_card.current
        arts = card_arts[ck]
        per_art = {}
        for art, skus in arts.items():
            lists, cost_rub = [], []
            target = None
            category = ""
            decision = None
            for item_id, rs in skus:
                cost = costs.get(item_id)
                facts = item_facts(rule, rs)
                category = category or facts.category
                d = decide_for(inp, item_id, account.id, cost.cost_usd if cost else None, usd, pps.get(item_id), rs, facts)
                decision = decision or d
                if d.new_price:
                    lists.append(d.new_price)
                if d.cost_rub is not None:
                    cost_rub.append(float(d.cost_rub))
                if target is None:
                    target = target_for(inp, item_id, account.id, facts.category)
            barcodes = [r.barcode for _, rs in skus for r in rs]
            sizes = sorted({(r.size or (info[i].size if i in info else "") or "").strip() for i, rs in skus for r in rs} - {""})
            per_art[art] = {"barcodes": barcodes, "sizes": sizes, "ala": max(lists) if lists else None, "cost_mean": (sum(cost_rub) / len(cost_rub)) if cost_rub else None,
                            "cost_max": max(cost_rub) if cost_rub else None, "target": target, "category": category,
                            "skus": skus, "note": decision.note if decision and not lists else ""}
        card_list = max((v["ala"] for v in per_art.values() if v["ala"]), default=None)
        cur_list = facts_card.current
        for art, v in per_art.items():
            e = econ.get((ck, art))
            tar = float(commission) / 100 if commission is not None else None
            k = float(e.k) if e and e.k is not None else (1 - tar if tar is not None else None)
            w = float(e.w) if e and e.w is not None else W_DEFAULT.get(platform, 1.0)
            f_log = float(e.f_logistics) if e and e.f_logistics is not None else p["f_default"]
            net = int(e.net_units) if e and e.net_units is not None else 0
            how = (e.how if e and e.how else "нет данных по логистике и продажам — загрузите из Excel")
            # Доля «к получению» k взята по продажам артикула ДРУГИХ кабинетов (или всех): у карточки здесь свой предмет
            # WB и своя комиссия по тарифу — k сдвигается на разницу тарифов. Без продаж в этом кабинете её не измерить.
            # Идемпотентно: у скорректированного k «1 − k» уже равно тарифу карточки.
            if (platform == "wb" and e and e.k is not None and tar is not None
                    and not (e.how or "").startswith(("артикул в кабинете", "категория в кабинете"))):
                t_emp = round((1 - k) * 200) / 200                 # комиссия, зашитая в k (шаг тарифов WB — 0,5 п.п.)
                if 0.01 <= abs(t_emp - tar) <= 0.12:               # больше — не разница предметов, а чужие данные
                    how += f" · k сдвинут по тарифу карточки: комиссия {tar * 100:g}% вместо {t_emp * 100:g}%"
                    k = k + (t_emp - tar)
            cm, cx = v["cost_mean"], v["cost_max"]
            row = {"card": ck, "wbart": card_rows[ck][0].article, "art": art, "narts": len(arts), "category": v["category"],
                   "net": net, "cost": cm, "costmax": cx, "sale": sale, "current": cur_list, "tar": tar, "disc": disc,
                   "k": k, "w": w, "F": f_log, "how": how, "nsku": len(v["skus"]),
                   "ala": v["ala"], "target": v["target"], "item_ids": [i for i, _ in v["skus"]], "note": v["note"],
                   "barcodes": v["barcodes"], "sizes": v["sizes"]}
            tgt = v["target"]
            row["margin"] = float(tgt.value) if tgt is not None and tgt.kind == KIND_MARGIN else None
            row["target_text"] = ("" if tgt is None else (f"маржинальность {float(tgt.value):g}" if tgt.kind == KIND_MARGIN
                                                          else f"× {float(tgt.value):g} от базовой"))
            row["target_source"] = tgt.source if tgt is not None else ""
            # цены
            if card_list and k is not None:
                row["al"] = card_list
                row["ap"] = round(card_list * (1 - disc))
                row["bap"] = row["ap"] * w
                row["apa"] = round(v["ala"] * (1 - disc)) if v["ala"] else None
                row["ac"] = math.floor(v["ala"] / cx * 10000) / 10000 if (v["ala"] and cx) else None
                ad = p["ads"] + p["extra_ads"]
                unit = k - p["tax"] * w - ad
                row["apr"] = row["ap"] * unit - f_log - p["rent_unit"] - cm if cm is not None else None
                row["aroi"] = row["apr"] / cm if row["apr"] is not None and cm else None
                row["amg"] = row["apr"] / row["bap"] if row["apr"] is not None and row["bap"] else None
                row["achg"] = row["ap"] / sale - 1 if sale else None
                if cx and unit > 0:
                    row["recp"] = ceil10((cx * (1 + p["roi"]) + f_log + p["rent_unit"]) / unit)
                else:
                    row["recp"] = None
                floor = float(rule.min_markup_coef)
                if row["apr"] is not None and row["apr"] < 0:
                    row["warn"] = "УБЫТОК при этой цене"
                elif cx and row["ap"] * (1 - (tar or 0)) / cx < floor:
                    row["warn"] = "ниже пола репрайсера — будет заблокировано"
                elif row["achg"] is not None and row["achg"] > 0.3:
                    row["warn"] = "рост более +30% — повышать шагами, следить за заказами"
                elif row["aroi"] is not None and row["aroi"] < p["roi"]:
                    row["warn"] = "окупаемость ниже целевой"
                else:
                    row["warn"] = ""
            else:
                for key in ("al", "ap", "bap", "apa", "ac", "apr", "aroi", "amg", "achg", "recp"):
                    row[key] = None
                row["warn"] = v["note"] or "цена не считается"
            # эталон: по SKU артикула берём наименьшее «к получению» против эталонного
            ref_prices = [refs[i].ref_price for i, _ in v["skus"] if i in refs]
            row["ref_price"] = max(ref_prices) if ref_prices else None
            row["ref_status"] = ""
            if ref_prices:
                row["ref_status"] = "в порядке"
                for i, rs in v["skus"]:
                    rf = refs.get(i)
                    cur_payout = (float(sale) * (1 - float(commission) / 100)) if sale and commission is not None else None
                    if rf and rf.ref_payout is not None and cur_payout is not None and cur_payout < float(rf.ref_payout) * (1 - drop):
                        row["ref_status"] = ("будет возвращена" if cur_list and cur_list < rf.ref_price else "съедено скидкой")
                        break
            out.append(row)
    out.sort(key=lambda r: (r["category"] or "", r["art"], r["card"]))
    return {"rows": out, "p": p, "rate": rate, "rule": rule, "account": account}


# --- сохранение маржинальностей из формы и из файла -----------------------------------------------------------

def set_margins(db: Session, user, account: Account, values: dict[str, Decimal | None]) -> tuple[int, list[str]]:
    """values: артикул -> маржинальность (None — снять). Пишет ArticleCoef вида «маржинальность» этого кабинета."""
    from priceapp.routers.sku_prices import set_coef
    changed, errors = 0, []
    for art, v in values.items():
        if v is not None and (v < Decimal("0.5") or v > 100):
            errors.append(f"{art}: маржинальность {v} вне 0,5–100")
            continue
        changed += set_coef(db, user, art, account.platform, account.id, v, KIND_MARGIN)
    return changed, errors


# --- Excel ------------------------------------------------------------------------------------------------------------

HF = PatternFill("solid", fgColor="202545")
HFONT = Font(bold=True, color="FFFFFF")
INP = PatternFill("solid", fgColor="FFF2CC")
GREEN = PatternFill("solid", fgColor="E2EFDA")
BLUE = PatternFill("solid", fgColor="DDEBF7")
GREY = PatternFill("solid", fgColor="EDEDED")
RUB = "#,##0;[Red]-#,##0"
PCT = "0.0%"
GRP = {"b": HF, "g": PatternFill("solid", fgColor="2E6B3A"), "a": PatternFill("solid", fgColor="BF8F00"),
       "o": PatternFill("solid", fgColor="1F4E79"), "s": PatternFill("solid", fgColor="7F7F7F")}
BODY = {"b": None, "g": GREEN, "a": None, "o": BLUE, "s": GREY}

# (ключ, заголовок, ширина, формат, группа) — те же заголовки и порядок, что на листах книги «по карточкам»
COLUMNS = [
    ("cat", "Категория", 16, None, "b"), ("nm", "Карточка WB (nmID)", 13, "0", "b"), ("wbart", "Артикул продавца на WB", 24, None, "b"),
    ("art", "Артикул 1С", 22, None, "b"), ("narts", "Артикулов 1С в карточке", 9, None, "b"), ("net", "Нетто-продаж за 7 нед., шт", 9, None, "b"),
    ("cost", "Себестоимость средняя, ₽", 10, RUB, "b"), ("costmax", "Себестоимость макс. по размерам, ₽", 10, RUB, "b"),
    ("sale", "Цена продавца карточки сейчас, ₽", 11, RUB, "b"), ("bnow", "Цена для покупателя сейчас, ₽", 11, RUB, "b"),
    ("effc", "Наценка артикула в репрайсере сейчас", 14, None, "b"), ("setsrc", "Где задана", 11, None, "b"),
    ("pnow", "Прибыль на ед. сейчас, ₽", 10, RUB, "b"), ("rnow", "Окупаемость сейчас", 10, PCT, "b"),
    ("recp", "Рекомендуемая цена продавца по артикулу, ₽", 12, RUB, "g"), ("recpc", "Рекомендуемая цена КАРТОЧКИ (макс.), ₽", 12, RUB, "g"),
    ("prec", "Прибыль на ед. при рекомендации, ₽", 11, RUB, "g"), ("rrec", "Окупаемость при рекомендации", 11, PCT, "g"),
    ("chg", "Изменение цены (рекомендация)", 11, PCT, "g"), ("sug", "Предложение", 20, None, "g"),
    ("ac", "К ВНЕСЕНИЮ В РЕПРАЙСЕР: коэф. от базовой (по артикулу)", 16, "0.0000", "o"),
    ("ap", "ИТОГОВАЯ цена продавца КАРТОЧКИ, ₽", 13, RUB, "o"), ("bap", "ИТОГОВАЯ цена для покупателя, ₽", 13, RUB, "o"),
    ("am", "Маржинальность (к получению / сс) — ВВОДИТСЯ ВРУЧНУЮ", 16, "0.00", "a"), ("al", "ИТОГОВАЯ цена карточки в списке, ₽", 12, RUB, "o"),
    ("apa", "Цена продавца по этому артикулу (до максимума карточки), ₽", 13, RUB, "o"),
    ("apr", "Прибыль на ед. при цене карточки, ₽", 11, RUB, "o"), ("aroi", "Окупаемость", 10, PCT, "o"), ("amg", "Маржа от цены покупателя", 10, PCT, "o"),
    ("achg", "Изменение цены карточки к текущей", 11, PCT, "o"), ("warn", "Предупреждение", 40, None, "o"), ("note", "Комментарий", 34, None, "a"),
    ("manual", "Ваша ручная маржинальность (сохранена)", 14, "0.00", "a"),
    ("tar", "Тариф WB", 8, "0%", "s"), ("disc", "Скидка продавца", 8, "0%", "s"), ("k", "k", 8, "0.000", "s"), ("w", "W/P", 8, "0.000", "s"),
    ("F", "Логистика+прочее на ед., ₽", 10, RUB, "s"), ("how", "Откуда параметры", 24, None, "s"),
    ("ala", "Цена в списке по артикулу, ₽", 12, RUB, "s"), ("apm", "Цена продавца по ручной маржинальности, ₽", 13, RUB, "s"),
    ("peff", "Цена в расчёте экономики, ₽", 12, RUB, "s"), ("vnow", "Прибыль за 7 нед. сейчас, ₽", 12, RUB, "s"),
    ("vnew", "Прибыль за 7 нед. по утверждённой цене, ₽", 12, RUB, "s"),
    ("refp", "Эталонная цена (последняя отправленная), ₽", 13, RUB, "s"), ("refs", "Статус эталона", 16, None, "s"),
]
LET = {key: CL(i + 1) for i, (key, *_r) in enumerate(COLUMNS)}
HEADER_ROW = 4


def sheet_name(account: Account) -> str:
    n = account.name.replace("ИП ", "").strip() or account.name
    return "".join(ch for ch in n if ch not in '[]:*?/\\')[:31]


def build_workbook(db: Session, accounts_: list[Account]) -> bytes:
    """Книга по кабинетам: «Параметры», «Экономика», лист на каждый кабинет — как книга «по карточкам»."""
    data = {a.id: compute(db, a) for a in accounts_}
    wb = Workbook()
    pm = wb.active
    pm.title = "Параметры"
    pm["A1"] = "Утверждение цен по карточкам — параметры (меняются и в программе, на странице «Утверждение цен»)"
    pm["A1"].font = Font(bold=True, size=14, color="202545")
    p0 = data[accounts_[0].id]["p"]
    floor = float(data[accounts_[0].id]["rule"].min_markup_coef)
    base = data[accounts_[0].id]["rule"].base_coef
    rows = [("Налог, % от цены покупателя", p0["tax"], PCT), ("Целевая окупаемость (ROI), % от себестоимости", p0["roi"], PCT),
            ("Пол репрайсера (мин. маржинальность)", floor, "0.00"),
            ("Наценка по умолчанию в репрайсере", float(base) if base is not None else None, "0.00")]
    for i, a in enumerate(accounts_):
        rows.append((f"Реклама + хранение, % от цены продавца — {sheet_name(a)}", data[a.id]["p"]["ads"], PCT))
    pad = 6 + len(accounts_)
    extra_row = pad
    rows += [("ДОПОЛНИТЕЛЬНАЯ реклама, % от цены продавца", p0["extra_ads"], PCT), ("Аренда склада, ₽ в месяц", p0["rent_month"], RUB),
             ("Доля аренды, относимая на площадку", p0["rent_share"], "0%"), ("Нетто-продаж в месяц, шт", p0["net_month"], "#,##0"),
             ("Аренда на 1 нетто-единицу, ₽", None, "#,##0.0")]
    for i, (n, v, f) in enumerate(rows, 2):
        pm.cell(i, 1, n)
        if v is None and n.startswith("Аренда на 1"):
            cl = pm.cell(i, 2, f"=B{extra_row + 1}*B{extra_row + 2}/B{extra_row + 3}")
        else:
            cl = pm.cell(i, 2, v)
            if v is not None:
                cl.fill = INP
        cl.number_format = f
    pm.column_dimensions["A"].width = 60
    pm.column_dimensions["B"].width = 14
    tx, ro, fl, bs = "Параметры!$B$2", "Параметры!$B$3", "Параметры!$B$4", "Параметры!$B$5"
    extra = f"Параметры!$B${extra_row}"
    rent = f"Параметры!$B${extra_row + 4}"
    ad_cell = {a.id: f"(Параметры!$B${6 + i}+{extra})" for i, a in enumerate(accounts_)}
    ad_only = {a.id: f"Параметры!$B${6 + i}" for i, a in enumerate(accounts_)}
    info = {}
    for acc in accounts_:
        rows_ = data[acc.id]["rows"]
        ws = wb.create_sheet(sheet_name(acc))
        ws["A1"] = f"{acc.name} — утверждение цен по карточкам ({PLATFORMS.get(acc.platform, acc.platform)})"
        ws["A1"].font = Font(bold=True, size=13, color="202545")
        ws["A2"] = ("Жёлтые ячейки — ввод. Итоговая цена карточки = максимум по артикулам карточки. Файл можно загрузить обратно "
                    "на странице «Утверждение цен» (читается колонка «Маржинальность» по артикулу и колонки k, W/P, логистика, продажи).")
        ws["A2"].font = Font(italic=True, color="666666")
        for j, (key, h, wd, f, grp) in enumerate(COLUMNS, 1):
            cl = ws.cell(HEADER_ROW, j, h)
            cl.font = HFONT
            cl.fill = GRP[grp]
            cl.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            ws.column_dimensions[CL(j)].width = wd
        ws.row_dimensions[HEADER_ROW].height = 84
        first, last = HEADER_ROW + 1, HEADER_ROW + len(rows_)
        rng = lambda key: f"${LET[key]}${first}:${LET[key]}${max(last, first)}"  # noqa: E731
        ad = ad_cell[acc.id]
        for n, r in enumerate(rows_, first):
            q = {k_: f"{LET[k_]}{n}" for k_ in LET}
            vals = {"cat": r["category"], "nm": (int(r["card"]) if str(r["card"]).isdigit() else str(r["card"])), "wbart": r["wbart"], "art": r["art"], "narts": r["narts"], "net": r["net"],
                    "cost": r["cost"], "costmax": r["costmax"], "sale": r["sale"], "tar": r["tar"], "disc": r["disc"], "k": r["k"], "w": r["w"],
                    "F": r["F"], "how": r["how"], "effc": r["target_text"], "setsrc": r["target_source"],
                    "refp": r["ref_price"], "refs": r["ref_status"]}
            for key, val in vals.items():
                if val is not None:
                    ws[q[key]] = float(val) if isinstance(val, (Decimal,)) else val
            if r["margin"] is not None:
                ws[q["am"]] = r["margin"]
                ws[q["manual"]] = r["margin"]
            can = r["cost"] is not None and r["costmax"] and r["k"] is not None and r["tar"] is not None
            if can:
                ws[q["bnow"]] = f"={q['sale']}*{q['w']}"
                ws[q["pnow"]] = f"={q['sale']}*({q['k']}-{tx}*{q['w']}-{ad})-{q['F']}-{rent}-{q['cost']}"
                ws[q["rnow"]] = f"={q['pnow']}/{q['cost']}"
                ws[q["recp"]] = f"=CEILING(({q['costmax']}*(1+{ro})+{q['F']}+{rent})/({q['k']}-{tx}*{q['w']}-{ad}),10)-1"
                ws[q["recpc"]] = f"=_xlfn.MAXIFS({rng('recp')},{rng('nm')},{q['nm']})"
                ws[q["prec"]] = f"={q['recpc']}*({q['k']}-{tx}*{q['w']}-{ad})-{q['F']}-{rent}-{q['cost']}"
                ws[q["rrec"]] = f"={q['prec']}/{q['cost']}"
                ws[q["chg"]] = f"={q['recpc']}/{q['sale']}-1"
                ws[q["sug"]] = f'=IF({q["chg"]}>1,"Пересмотреть",IF({q["chg"]}<0,"Оставить текущую","Принять рекомендацию"))'
                raw = f"{q['costmax']}*{q['am']}/((1-{q['tar']})*(1-{q['disc']}))"
                ws[q["ala"]] = f'=IF({q["am"]}="","",CEILING({raw},10)-1+IF(CEILING({raw},10)-1<{raw},10,0))'
                ws[q["al"]] = f'=IF(COUNTIFS({rng("nm")},{q["nm"]},{rng("ala")},">0")=0,"",_xlfn.MAXIFS({rng("ala")},{rng("nm")},{q["nm"]}))'
                ws[q["ap"]] = f'=IF({q["al"]}="","",ROUND({q["al"]}*(1-{q["disc"]}),0))'
                ws[q["bap"]] = f'=IF({q["al"]}="","",{q["ap"]}*{q["w"]})'
                ws[q["apa"]] = f'=IF({q["ala"]}="","",ROUND({q["ala"]}*(1-{q["disc"]}),0))'
                ws[q["ac"]] = f'=IF({q["ala"]}="","",ROUNDDOWN({q["ala"]}/{q["costmax"]},4))'
                ws[q["apr"]] = f'=IF({q["ap"]}="","",{q["ap"]}*({q["k"]}-{tx}*{q["w"]}-{ad})-{q["F"]}-{rent}-{q["cost"]})'
                ws[q["aroi"]] = f'=IF({q["ap"]}="","",{q["apr"]}/{q["cost"]})'
                ws[q["amg"]] = f'=IF({q["ap"]}="","",{q["apr"]}/({q["ap"]}*{q["w"]}))'
                ws[q["achg"]] = f'=IF({q["ap"]}="","",{q["ap"]}/{q["sale"]}-1)'
                rawm = f"{q['costmax']}*{q['manual']}/((1-{q['tar']})*(1-{q['disc']}))"
                ws[q["apm"]] = (f'=IF({q["manual"]}="","",ROUND((CEILING({rawm},10)-1+IF(CEILING({rawm},10)-1<{rawm},10,0))'
                                f'*(1-{q["disc"]}),0))')
                ws[q["peff"]] = f'=IF({q["ap"]}="",{q["sale"]},{q["ap"]})'
                ws[q["vnow"]] = f"={q['net']}*{q['pnow']}"
                ws[q["vnew"]] = f'=IF({q["ap"]}="",{q["net"]}*{q["pnow"]},{q["net"]}*{q["apr"]})'
                ws[q["warn"]] = (f'=IF({q["ap"]}="","",IF({q["apr"]}<0,"УБЫТОК при этой цене",IF({q["ap"]}*(1-{q["tar"]})/{q["costmax"]}<{fl},'
                                 f'"ниже пола репрайсера — будет заблокировано",IF({q["achg"]}>0.3,"рост более +30% — повышать шагами, следить за заказами",'
                                 f'IF({q["aroi"]}<{ro},"окупаемость ниже целевой","")))))')
            else:
                ws[q["warn"]] = r["warn"]
                ws[q["peff"]] = f"={q['sale']}"
        for rr in range(first, last + 1):
            for j, (key, h, wd, f, grp) in enumerate(COLUMNS, 1):
                cl = ws.cell(rr, j)
                if f and f != "@":
                    cl.number_format = f
                if key in ("am", "note", "manual"):
                    cl.fill = INP
                elif BODY[grp] is not None:
                    cl.fill = BODY[grp]
        dv = DataValidation(type="decimal", operator="between", formula1="0.5", formula2="100", allow_blank=True,
                            errorTitle="Маржинальность", error="Число от 0,5 до 100 (как принимает репрайсер)")
        ws.add_data_validation(dv)
        dv.add(f"{LET['am']}{first}:{LET['am']}{max(last, first)}")
        ws.freeze_panes = f"E{first}"
        ws.auto_filter.ref = f"A{HEADER_ROW}:{CL(len(COLUMNS))}{max(last, first)}"
        for key in ("pnow", "apr"):
            ws.conditional_formatting.add(f"{LET[key]}{first}:{LET[key]}{max(last, first)}",
                                          CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))
        ws.conditional_formatting.add(f"{LET['warn']}{first}:{LET['warn']}{max(last, first)}",
                                      FormulaRule(formula=[f'LEFT({LET["warn"]}{first},6)="УБЫТОК"'], font=Font(color="C00000", bold=True)))
        info[acc.id] = (sheet_name(acc), first, last)
    # --- Экономика: по кабинетам и всего
    we = wb.create_sheet("Экономика")
    we["A1"] = "Экономика при заданных маржинальностях (по ценам карточек; 7 недель при прошлом объёме и в пересчёте на месяц)"
    we["A1"].font = Font(bold=True, size=13, color="202545")
    labels = ["Нетто-продаж, шт", "Цена продавца Σ, ₽", "Цена для покупателя Σ, ₽", "К получению (после комиссии и скидки площадки), ₽",
              "− Логистика с отказами, штрафы, приёмка, ₽", "− Реклама и хранение (факт), ₽", "− Дополнительная реклама, ₽", "− Налог, ₽",
              "− Аренда склада, ₽", "− Себестоимость, ₽", "ПРИБЫЛЬ, ₽", "Окупаемость от себестоимости", "Маржа от цены покупателя",
              "Прибыль в месяц, ₽ (×30,44/49)"]
    RR = {lab: 5 + i for i, lab in enumerate(labels)}
    for i, lab in enumerate(labels):
        we.cell(5 + i, 1, lab).font = Font(bold=lab.startswith("ПРИБЫЛЬ"))
    blocks = {}
    col = 2
    for acc in accounts_ + [None]:
        blocks[acc.id if acc else 0] = col
        col += 4
    for acc in accounts_:
        c = blocks[acc.id]
        name, first, last = info[acc.id]
        g_ = lambda key: f"'{name}'!${LET[key]}${first}:${LET[key]}${max(last, first)}"  # noqa: E731
        we.cell(3, c, name).font = Font(bold=True, size=12, color="202545")
        for t, h in enumerate(["Сейчас", "При заданных", "Разница"]):
            cl = we.cell(4, c + t, h)
            cl.fill = HF
            cl.font = HFONT
        for t, pk in enumerate(("sale", "peff")):
            cc = CL(c + t)
            P_ = g_(pk)
            we.cell(RR[labels[0]], c + t, f"=SUM({g_('net')})")
            we.cell(RR[labels[1]], c + t, f"=SUMPRODUCT({g_('net')},{P_})")
            we.cell(RR[labels[2]], c + t, f"=SUMPRODUCT({g_('net')},{P_},{g_('w')})")
            we.cell(RR[labels[3]], c + t, f"=SUMPRODUCT({g_('net')},{P_},{g_('k')})")
            we.cell(RR[labels[4]], c + t, f"=SUMPRODUCT({g_('net')},{g_('F')})")
            we.cell(RR[labels[5]], c + t, f"={cc}{RR[labels[1]]}*{ad_only[acc.id]}")
            we.cell(RR[labels[6]], c + t, f"={cc}{RR[labels[1]]}*{extra}")
            we.cell(RR[labels[7]], c + t, f"={tx}*{cc}{RR[labels[2]]}")
            we.cell(RR[labels[8]], c + t, f"={rent}*{cc}{RR[labels[0]]}")
            we.cell(RR[labels[9]], c + t, f"=SUMPRODUCT({g_('net')},{g_('cost')})")
            we.cell(RR[labels[10]], c + t, f"={cc}{RR[labels[3]]}-SUM({cc}{RR[labels[4]]}:{cc}{RR[labels[9]]})")
            we.cell(RR[labels[11]], c + t, f'=IFERROR({cc}{RR[labels[10]]}/{cc}{RR[labels[9]]},"")')
            we.cell(RR[labels[12]], c + t, f'=IFERROR({cc}{RR[labels[10]]}/{cc}{RR[labels[2]]},"")')
            we.cell(RR[labels[13]], c + t, f"={cc}{RR[labels[10]]}*30.44/49")
    cT = blocks[0]
    we.cell(3, cT, "ВСЕ КАБИНЕТЫ").font = Font(bold=True, size=12, color="202545")
    for t, h in enumerate(["Сейчас", "При заданных", "Разница"]):
        cl = we.cell(4, cT + t, h)
        cl.fill = HF
        cl.font = HFONT
    for t in range(2):
        for lab in labels[:11] + [labels[13]]:
            rr = RR[lab]
            we.cell(rr, cT + t, "=" + "+".join(f"{CL(blocks[a.id] + t)}{rr}" for a in accounts_))
        we.cell(RR[labels[11]], cT + t, f'=IFERROR({CL(cT + t)}{RR[labels[10]]}/{CL(cT + t)}{RR[labels[9]]},"")')
        we.cell(RR[labels[12]], cT + t, f'=IFERROR({CL(cT + t)}{RR[labels[10]]}/{CL(cT + t)}{RR[labels[2]]},"")')
    for bid, c in blocks.items():
        for lab in labels:
            rr = RR[lab]
            we.cell(rr, c + 2, f'=IFERROR({CL(c + 1)}{rr}-{CL(c)}{rr},"")')
            for t in range(3):
                we.cell(rr, c + t).number_format = PCT if lab in (labels[11], labels[12]) else ("#,##0" if lab == labels[0] else RUB)
        for t in range(3):
            we.cell(RR[labels[10]], c + t).fill = GREY
            we.cell(RR[labels[10]], c + t).font = Font(bold=True)
    we.column_dimensions["A"].width = 52
    for c_ in range(2, 2 + 4 * (len(accounts_) + 1)):
        we.column_dimensions[CL(c_)].width = 14
    wb.move_sheet("Экономика", offset=-len(accounts_))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- импорт --------------------------------------------------------------------------------------------------------------

def _num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace(",", ".")
    if s == "" or s == "-":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _find_header(ws):
    for r in range(1, 12):
        titles = [str(c.value).strip() if c.value is not None else "" for c in ws[r]]
        if "Артикул 1С" in titles and any(t.startswith("Маржинальность (к получению") for t in titles):
            return r, {t: i for i, t in enumerate(titles) if t}
    return None, {}


def import_workbook(db: Session, user, data: bytes, accounts_: list[Account], apply_margins: bool = True) -> dict:
    """Загрузка книги этого формата: по листу кабинета — маржинальности артикулов («Маржинальность …») и
    экономика карточки (k, W/P, логистика, продажи, откуда параметры). Лист → кабинет по имени.
    Пустая ячейка маржинальности ничего не меняет, «-» снимает наценку артикула в кабинете."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        return {"errors": [f"Файл не читается как .xlsx ({type(e).__name__})."], "changed": 0, "econ": 0, "sheets": 0}
    by_sheet = {}
    for a in accounts_:
        by_sheet[sheet_name(a).lower()] = a
        by_sheet[a.name.lower()] = a
    res = {"changed": 0, "econ": 0, "sheets": 0, "errors": [], "conflicts": 0, "skipped": 0}
    for ws in wb.worksheets:
        acc = by_sheet.get(ws.title.lower())
        if acc is None:
            continue
        hr, idx = _find_header(ws)
        if hr is None:
            res["errors"].append(f"лист «{ws.title}»: нет колонок «Артикул 1С» и «Маржинальность…» — выгрузите файл со страницы")
            continue
        res["sheets"] += 1
        c_art = idx["Артикул 1С"]
        c_m = next(i for t, i in idx.items() if t.startswith("Маржинальность (к получению"))
        c_card = idx.get("Карточка WB (nmID)")
        c_net, c_k, c_w = idx.get("Нетто-продаж за 7 нед., шт"), idx.get("k"), idx.get("W/P")
        c_f, c_how = idx.get("Логистика+прочее на ед., ₽"), idx.get("Откуда параметры")
        c_eff, c_src = idx.get("Наценка артикула в репрайсере сейчас"), idx.get("Где задана")
        margins: dict[str, list] = {}
        clears = set()
        econ_rows = []
        for r in ws.iter_rows(min_row=hr + 1, values_only=True):
            if not r or len(r) <= c_art or not r[c_art]:
                continue
            art = str(r[c_art]).strip()
            raw = r[c_m] if len(r) > c_m else None
            if isinstance(raw, str) and raw.strip() == "-":
                clears.add(art)
            else:
                v = _num(raw)
                # Выгрузка кладёт в колонку ДЕЙСТВУЮЩУЮ маржинальность, в том числе унаследованную
                # от категории или площадки. Если человек её не трогал (то же число, задано не на
                # артикуле кабинета) — не делаем её своей: иначе правка категории перестала бы
                # действовать на артикул молча (аудит 14.10, тот же класс, что п.3 аудита 08.10).
                inherited = (c_src is not None and c_eff is not None and len(r) > max(c_src, c_eff)
                             and str(r[c_src] or "") != SRC_CABINET and v is not None
                             and str(r[c_eff] or "") == f"маржинальность {float(v):g}")
                if v is not None and not inherited:
                    margins.setdefault(art, []).append(v)
            if c_card is not None and c_k is not None and r[c_card] is not None:
                econ_rows.append((str(r[c_card]).strip().split(".")[0], art, _num(r[c_net]) if c_net is not None else None,
                                  _num(r[c_k]), _num(r[c_w]) if c_w is not None else None,
                                  _num(r[c_f]) if c_f is not None else None,
                                  (str(r[c_how]) if c_how is not None and r[c_how] else "")))
        values: dict[str, Decimal | None] = {}
        for art, vals in margins.items():
            if len(set(round(v, 4) for v in vals)) > 1:
                res["conflicts"] += 1
            values[art] = Decimal(str(round(max(vals), 4)))
        for art in clears:
            values.setdefault(art, None)
        ch, errs = set_margins(db, user, acc, values) if apply_margins else (0, [])
        res["changed"] += ch
        res["errors"] += [f"{acc.name}: {e}" for e in errs]
        existing = {(e.card, e.article): e for e in db.query(CardEcon).filter(CardEcon.account_id == acc.id)}
        for card, art, net, k, w, f_log, how in econ_rows:
            e = existing.get((card, art))
            if e is None:
                e = CardEcon(account_id=acc.id, card=card, article=art)
                db.add(e)
                existing[(card, art)] = e
            e.net_units = int(net) if net is not None else 0
            e.k, e.w, e.f_logistics, e.how = k, w, f_log, how[:200]
            res["econ"] += 1
        db.flush()
    db.commit()
    return res
