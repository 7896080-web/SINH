"""Сопоставление каталога площадки с товарами 1С — по правилам sync_admin.

1. **Ключ — штрихкод.** Баркод площадки сопоставлен, если он есть в справочнике
   1С (`OnecBarcode`) у ровно ОДНОГО SKU (размер-цвет). У двух SKU — нарушение
   «неоднозначно»: какой товар имеется в виду, сказать нельзя, и цену по нему
   не считаем.
2. **Пул размер-цвета площадки.** Строки каталога с одним `external_id` — это
   один размер на площадке с несколькими баркодами. Если сопоставленные баркоды
   пула ведут к ОДНОМУ SKU 1С, остальные баркоды пула относятся к нему же
   (так работает автопривязка sync_admin). Ведут к разным — пул не мержим.
0. **Мэппинг Sync Admin** (`SyncBarcode`, снимок выгрузки «Мэппинг»): привязки, которые
   Sync Admin сделал мимо справочника 1С. Файлом (`excel_import`) и переподвязкой
   (`excel_repoint`) — решение человека, они главнее справочника, как и в Sync Admin.
   Догадка по пулу (`pool_match`) — только для баркода, которого в справочнике нет:
   справочник 1С в Sync Admin перебивает догадку, и здесь так же.
3. **Подтверждённые человеком связи** (`ManualLink`, страница «Сопоставление»
   по правилам артикулов) — для баркодов, которых в справочнике 1С нет вовсе.
   Существующую связь по баркоду справочника они не перебивают никогда.

Справочник 1С — снимок `barcodes_*.txt` (формат sync_admin): по команде
`BARCODE_DICT` или файлом вручную. Снимок заменяется целиком.
"""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field

from sqlalchemy.orm import Session

from priceapp import settings
from priceapp.models import ManualLink, OnecBarcode, PlatformItem, SyncBarcode
from priceapp.timeutils import now_utc

BARCODE_DICT_FIELDS = 6

STATUS_LABELS = {
    "sync": "по мэппингу Sync Admin",
    "ok": "сопоставлен по баркоду",
    "pool": "по пулу размера",
    "manual": "подтверждён вручную",
    "not_in_1c": "нет в 1С",
    "ambiguous": "баркод у нескольких SKU 1С",
}
MAPPED = ("ok", "pool", "manual", "sync")
SYNC_GUESS = "pool_match"          # догадка Sync Admin по пулу — справочник 1С её перебивает


class MappingError(ValueError):
    pass


# --- Справочник 1С ---------------------------------------------------------------

def _split_from_the_right(line: str, fields: int) -> list[str] | None:
    """Как в sync_admin: «|» внутри наименования режется справа."""
    parts = line.split("|")
    if len(parts) < fields:
        return None
    if len(parts) == fields:
        return parts
    tail = fields - 3
    return [parts[0], parts[1], "|".join(parts[2:-tail])] + parts[-tail:]


def parse_barcode_dict(text: str) -> list[dict]:
    """`uid|артикул|наименование|баркод|размер|цвет` — строка на баркод."""
    rows = []
    for raw in text.splitlines():
        line = raw.rstrip("\r\n")
        if not line.strip():
            continue
        p = _split_from_the_right(line, BARCODE_DICT_FIELDS)
        if p is None:
            p = line.split("|")
            if len(p) < 4:
                continue
        item, barcode = p[0].strip(), p[3].strip()
        if not item or not barcode:
            continue
        rows.append({"item_id": item, "article": p[1].strip(), "name": p[2].strip(),
                     "barcode": barcode, "size": p[4].strip() if len(p) > 4 else "",
                     "color": p[5].strip() if len(p) > 5 else ""})
    return rows


def decode_upload(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise MappingError("файл не читается ни как UTF-8, ни как cp1251")


def load_dictionary(db: Session, rows: list[dict], source: str) -> dict:
    """Заменить снимок справочника. Не коммитит. Пустой разбор — отказ: обрезанный
    файл иначе стёр бы снимок, и весь каталог стал бы «нет в 1С»."""
    if not rows:
        raise MappingError("в файле нет ни одной строки справочника "
                           "(ждём uid|артикул|наименование|баркод|размер|цвет)")
    seen: set[tuple[str, str]] = set()
    unique = []
    for r in rows:
        key = (r["barcode"], r["item_id"])
        if key in seen:
            continue
        seen.add(key)
        unique.append({k: (r[k] or "")[:lim] for k, lim in
                       (("barcode", 64), ("item_id", 64), ("article", 200), ("name", 500),
                        ("size", 100), ("color", 100))})
    db.query(OnecBarcode).delete(synchronize_session=False)
    db.bulk_insert_mappings(OnecBarcode, unique)
    settings.put(db, settings.DICT_LOADED_AT, now_utc().isoformat(timespec="seconds"))
    settings.put(db, settings.DICT_ROWS, str(len(unique)))
    return {"rows": len(unique), "items": len({i for _, i in seen}), "source": source}


# --- Сопоставление ------------------------------------------------------------------

@dataclass
class MapRow:
    item: PlatformItem
    status: str
    item_id: str = ""
    others: list[str] = field(default_factory=list)   # SKU 1С при «неоднозначно»

    @property
    def label(self) -> str:
        return STATUS_LABELS[self.status]


def owners_of(db: Session, barcodes) -> dict[str, set[str]]:
    out: dict[str, set[str]] = defaultdict(set)
    barcodes = list(set(barcodes))
    for i in range(0, len(barcodes), 500):
        for b in db.query(OnecBarcode.barcode, OnecBarcode.item_id).filter(
                OnecBarcode.barcode.in_(barcodes[i:i + 500])):
            out[b.barcode].add(b.item_id)
    return out


def sync_of(db: Session, barcodes) -> dict[str, tuple[str, str]]:
    """баркод -> (SKU 1С, источник) по мэппингу Sync Admin."""
    out: dict[str, tuple[str, str]] = {}
    barcodes = list(set(barcodes))
    for i in range(0, len(barcodes), 500):
        for b in db.query(SyncBarcode.barcode, SyncBarcode.item_id, SyncBarcode.source).filter(
                SyncBarcode.barcode.in_(barcodes[i:i + 500])):
            out[b.barcode] = (b.item_id, b.source)
    return out


SYNC_HEAD = ("ID_1С", "Баркод")
SYNC_SOURCES = {"excel_import", "excel_repoint", SYNC_GUESS}


def parse_sync_mapping(data: bytes) -> list[dict]:
    """Выгрузка «Мэппинг» Sync Admin (мэппинг_баркодов.xlsx): ID_1С, Баркод, Артикул,
    Наименование, Размер, Цвет, Источник, Добавлен. Берём только привязки, сделанные в
    Sync Admin (`SYNC_SOURCES`): строки «выгрузка из 1С» / 1c_dict / 1c_export — тот же
    справочник 1С, что у нас, — пропускаются."""
    import io

    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise MappingError(f"файл не читается как Excel: {e}") from None
    ws = book.worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = [str(c or "").strip() for c in next(it, ())]
    if not all(h in head for h in SYNC_HEAD):
        raise MappingError("это не выгрузка «Мэппинг» Sync Admin: нет колонок «ID_1С» и «Баркод»")
    ix = {h: head.index(h) for h in head if h}
    def cell(r, name):
        i = ix.get(name)
        return str(r[i]).strip() if i is not None and i < len(r) and r[i] is not None else ""
    rows = []
    for r in it:
        src = cell(r, "Источник")
        item, barcode = cell(r, "ID_1С"), cell(r, "Баркод")
        if not item or not barcode or src not in SYNC_SOURCES:
            continue
        rows.append({"item_id": item[:64], "barcode": barcode[:64], "article": cell(r, "Артикул")[:200],
                     "name": cell(r, "Наименование")[:500], "size": cell(r, "Размер")[:100],
                     "color": cell(r, "Цвет")[:100], "source": src})
    return rows


def load_sync_mapping(db: Session, rows: list[dict]) -> dict:
    """Заменить снимок мэппинга Sync Admin. Не коммитит. Пустой — отказ: иначе пустой файл
    стёр бы все ручные привязки. Баркод в Sync Admin один на товар — повтор берём первым."""
    if not rows:
        raise MappingError("в файлах нет привязок Sync Admin (источник excel_import / excel_repoint / "
                           "pool_match) — выгрузите «Мэппинг» с отбором по источнику")
    seen: dict[str, dict] = {}
    for r in rows:
        seen.setdefault(r["barcode"], r)
    db.query(SyncBarcode).delete(synchronize_session=False)
    now = now_utc()
    db.bulk_insert_mappings(SyncBarcode, [{**r, "loaded_at": now} for r in seen.values()])
    forget(db)
    by = defaultdict(int)
    for r in seen.values():
        by[r["source"]] += 1
    return {"rows": len(seen), "by_source": dict(by)}


def build(db: Session, account_id: int) -> list[MapRow]:
    """Статус каждой строки каталога кабинета. Сборка дорогая (на кабинете в
    150 тыс. строк — секунды), поэтому в пределах транзакции считается один раз
    (`database.session_cache`); коммит, откат или `forget` кэш сбрасывают."""
    from priceapp.database import session_cache
    cache = session_cache(db)
    if ("build", account_id) not in cache:
        cache[("build", account_id)] = _build(db, account_id)
    return cache[("build", account_id)]


def forget(db: Session) -> None:
    """Сбросить кэш сборки — после записи связей в ТОЙ ЖЕ транзакции."""
    from priceapp.database import session_cache
    session_cache(db).clear()


def _build(db: Session, account_id: int) -> list[MapRow]:
    items = (db.query(PlatformItem).filter(PlatformItem.account_id == account_id)
             .order_by(PlatformItem.article, PlatformItem.size, PlatformItem.barcode).all())
    owners = owners_of(db, [i.barcode for i in items])
    manual = {m.barcode: m.item_id for m in
              db.query(ManualLink).filter(ManualLink.account_id == account_id)}
    sync = sync_of(db, [i.barcode for i in items])

    rows: list[MapRow] = []
    for it in items:
        own = owners.get(it.barcode, set())
        s = sync.get(it.barcode)
        if s is not None and s[1] != SYNC_GUESS:
            rows.append(MapRow(it, "sync", s[0]))
        elif len(own) == 1:
            rows.append(MapRow(it, "ok", next(iter(own))))
        elif len(own) > 1:
            rows.append(MapRow(it, "ambiguous", others=sorted(own)))
        elif it.barcode in manual:
            rows.append(MapRow(it, "manual", manual[it.barcode]))
        elif s is not None:
            rows.append(MapRow(it, "sync", s[0]))
        else:
            rows.append(MapRow(it, "not_in_1c"))

    # Пул размера на площадке: сопоставленные соседи ведут к одному SKU — остальные к нему же.
    pools: dict[str, list[MapRow]] = defaultdict(list)
    for r in rows:
        if r.item.external_id:
            pools[r.item.external_id].append(r)
    for pool in pools.values():
        targets = {r.item_id for r in pool if r.status in ("ok", "manual", "sync")}
        if len(targets) != 1 or any(r.status == "ambiguous" for r in pool):
            continue
        target = next(iter(targets))
        for r in pool:
            if r.status == "not_in_1c":
                r.status, r.item_id = "pool", target
    return rows


def account_items(db: Session, account_id: int) -> dict[str, list[PlatformItem]]:
    """SKU 1С -> строки каталога кабинета, которые на него ведут (для расчёта цены)."""
    from priceapp.database import session_cache
    cache = session_cache(db)
    key = ("items", account_id)
    if key not in cache:
        out: dict[str, list[PlatformItem]] = defaultdict(list)
        for r in build(db, account_id):
            if r.status in MAPPED:
                out[r.item_id].append(r.item)
        cache[key] = dict(out)
    return cache[key]


def link_manually(db: Session, account_id: int, pairs: list[tuple[str, str]], actor: str = "",
                  source: str = "excel") -> tuple[int, dict]:
    """Ручные связи баркод -> SKU 1С (файлом со страницы «Сопоставление»).

    Связывается ТОЛЬКО баркод со статусом «нет в 1С» и только с SKU, который есть
    в справочнике 1С. Уже сопоставленный (по баркоду, по пулу, вручную) не
    перепривязывается никогда: связь меняет, с какой себестоимости считается цена,
    и старый файл, залитый по ошибке, переразнёс бы каталог. Возвращает
    (создано, {причина: n})."""
    status = {r.item.barcode: r.status for r in build(db, account_id)}
    known = {r[0] for r in db.query(OnecBarcode.item_id).distinct()}
    created, refused = 0, defaultdict(int)
    for barcode, item_id in pairs:
        st = status.get(barcode)
        if st is None:
            refused["баркода нет в каталоге кабинета"] += 1
        elif st != "not_in_1c":
            refused[f"баркод уже {STATUS_LABELS[st]} — перепривязка не делается"] += 1
        elif item_id not in known:
            refused["ID_1С нет в справочнике 1С"] += 1
        else:
            db.add(ManualLink(account_id=account_id, barcode=barcode, item_id=item_id,
                              source=source, created_by=actor))
            from priceapp.database import session_cache
            session_cache(db).clear()   # сборка в этой же транзакции должна увидеть связь
            status[barcode] = "manual"     # повтор баркода в файле — не второй INSERT
            created += 1
    db.flush()
    return created, dict(refused)


def counts(rows: list[MapRow]) -> dict[str, int]:
    c = {k: 0 for k in STATUS_LABELS}
    for r in rows:
        c[r.status] += 1
    return c
