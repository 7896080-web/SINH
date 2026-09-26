"""Старт учёта задним числом: выписки за прошлые месяцы, предложения, свод за период."""
import io
import os
import sqlite3

from openpyxl import load_workbook

from conftest import CHAT, png, TODAY, payment, answer
from finance.flow import _parse_period_arg
from finance.storage import Storage
from test_sverka import PDF, op, statement


def start_period(flow, db, card_index=0):
    [picker] = flow.on_command(CHAT, "sverka")
    flow.on_button(CHAT, "s:m:period")
    return flow.on_button(CHAT, f"s:c:{db.cards()[card_index].id}")


def sug(d, amount, desc, cat):
    return {**op(d, amount, desc=desc), "business_category": cat}


def test_picker_covers_last_12_months(env):
    db, rec, flow = env
    [picker] = flow.on_command(CHAT, "sverka")
    data = [d for row in picker.buttons for _, d in row]
    assert data[0] == "s:m:2025-10" and "s:m:2026-01" in data and data[-2] == "s:m:2026-09"
    assert data[-1] == "s:m:period"
    labels = [label for row in picker.buttons for label, _ in row]
    assert "Янв 26" in labels


def test_multi_month_statement_split_by_month(env):
    db, rec, flow = env
    [wait] = start_period(flow, db)
    assert "разложатся по месяцам" in wait.text
    rec.statements.append(statement([
        op("2026-01-05", "1000", "in", "Зарплата"),
        op("2026-01-10", "300", desc="Магазин"),
        op("2026-02-03", "400", desc="Кафе"),
        op("2026-03-15", "500", desc="Аптека"),
    ], total_in="1000", total_out="1200"))
    [got] = flow.on_files(CHAT, [PDF], "")
    assert "Принято операций: 4" in got.text
    assert "январь 2026: 2" in got.text and "март 2026: 1" in got.text
    assert "Итоги из документа" not in got.text  # итоги за 3 месяца на месяц не кладём
    card_id = db.cards()[0].id
    jan = db.statement(card_id, "2026-01")
    assert jan[0]["total_out"] is None and len(jan[1]) == 2
    assert len(db.statement(card_id, "2026-03")[1]) == 1
    rec.payments.append(payment(amount="400", date="2026-02-03"))
    replies = flow.on_command(CHAT, "done")
    texts = "\n".join(r.text for r in replies)
    assert texts.count("💳 Сбер") == 3
    assert "/vypiska" in texts and "Итог" in str(replies[-1].buttons)
    assert "s:itog:2026-01..2026-03" in str(replies[-1].buttons)


def test_totals_only_needs_a_month_in_period_mode(env):
    db, rec, flow = env
    start_period(flow, db)
    rec.statements.append(statement([], total_in="100", total_out="50"))
    [r] = flow.on_text(CHAT, "пришло 100 ушло 50")
    assert "только за конкретный месяц" in r.text


def test_single_month_file_in_period_mode_keeps_totals(env):
    db, rec, flow = env
    start_period(flow, db)
    rec.statements.append(statement([op("2026-04-02", "100")], total_out="150"))
    flow.on_files(CHAT, [PDF], "")
    assert db.statement(db.cards()[0].id, "2026-04")[0]["total_out"] == 15000


def test_suggestions_accept_all_and_notbiz(env):
    db, rec, flow = env
    start_period(flow, db)
    rec.statements.append(statement([
        sug("2026-01-12", "2500", "CDEK", "Логистика и доставка"),
        sug("2026-02-20", "7000", "YANDEX DIRECT", "Реклама и продвижение"),
        sug("2026-02-21", "900", "OZON", "Упаковка и расходники"),
        sug("2026-02-22", "650", "Пятёрочка", ""),
        {**op("2026-02-23", "5000", desc="на карту *2222", own=True),
         "business_category": "Прочее"},  # свой перевод не предлагаем
    ]))
    flow.on_files(CHAT, [PDF], "")
    replies = flow.on_command(CHAT, "done")
    [s] = [r for r in replies if r.text.startswith("🔎")]
    assert "Похоже на бизнес-расходы: 3 на 10 400,00 ₽" in s.text
    assert "Пятёрочка" not in s.text and "s:acc" in str(s.buttons)

    ozon = next(ln for m in ("2026-02",) for ln in db.statement(db.cards()[0].id, m)[1]
                if ln.description == "OZON")
    [left] = flow.on_command(CHAT, "notbiz", str(ozon.id))
    assert "Осталось 2" in left.text

    [done] = flow.on_button(CHAT, f"s:acc:{db.get_state(CHAT)['review_id']}")
    assert done.text.startswith("✅ Записано как бизнес: 2 на 9 500,00 ₽")
    assert "Реклама и продвижение: 7 000,00 ₽" in done.text
    assert "review" not in db.get_state(CHAT)
    assert sorted(e.amount for e in db.expenses("2026-02")) == [700000]
    # Записанные строки теперь сопоставлены и больше не предлагаются.
    [again] = flow.on_command(CHAT, "vypiska", "2026-02")
    assert "YANDEX" not in again.text and "OZON" in again.text
    assert flow.on_command(CHAT, "biz", "все")[0].text.startswith("Нет предложенных")


def test_accept_asks_category_only_when_not_suggested(env):
    db, rec, flow = env
    start_period(flow, db)
    rec.statements.append(statement([sug("2026-03-01", "100", "ИП Иванов", ""),
                                     sug("2026-03-02", "200", "CDEK", "Логистика и доставка")]))
    flow.on_files(CHAT, [PDF], "")
    flow.on_command(CHAT, "done")
    lines = db.statement(db.cards()[0].id, "2026-03")[1]
    replies = flow.on_command(CHAT, "biz", " ".join(str(ln.id) for ln in lines))
    assert "статья" in replies[-1].text.lower()  # первая без статьи — спрашиваем
    replies = answer(flow, f"d:cat:{db.category_id('Подрядчики и зарплата')}")
    [summary] = replies
    assert summary.text.startswith("✅ Записано как бизнес: 2 на 300,00 ₽")
    assert "Подрядчики и зарплата: 100,00 ₽" in summary.text


def test_old_screenshot_filed_by_payment_month(env, tmp_path):
    db, rec, flow = env
    rec.payments.append(payment(date="2026-01-17"))
    flow.on_files(CHAT, [png()], "")
    e = db.expenses("2026-01")[0]
    assert os.path.basename(os.path.dirname(e.receipt_path)) == "2026-01"
    assert os.path.exists(e.receipt_path)
    assert not os.listdir(os.path.join(flow.receipts_dir, "2026-09"))


def test_period_itog_business_by_category_and_personal(env):
    db, rec, flow = env
    for d, amount, cat in (("2026-01-10", "1000", "Реклама и продвижение"),
                           ("2026-02-10", "2000", "Логистика и доставка"),
                           ("2026-05-10", "500", "Реклама и продвижение")):
        rec.payments.append(payment(date=d, amount=amount, category=cat))
        flow.on_files(CHAT, [png()], "")
    # Выписка по Сберу за февраль: ушло 10 000, из них 2 000 записано как бизнес.
    flow.on_command(CHAT, "sverka", "2026-02")
    flow.on_button(CHAT, f"s:c:{db.cards()[0].id}")
    rec.statements.append(statement([op("2026-02-10", "2000"), op("2026-02-11", "8000")]))
    flow.on_files(CHAT, [PDF], "")
    flow.on_command(CHAT, "done")

    [r] = flow.on_command(CHAT, "itog", "2026")
    assert "январь 2026 — сентябрь 2026" in r.text
    assert "янв 26: 1 000,00 ₽ · 0,00 ₽  (выписок нет)" in r.text
    assert "фев 26: 2 000,00 ₽ · 8 000,00 ₽  (нет выписки: Тинькофф, Альфа)" in r.text
    assert "💼 Ушло на бизнес: 3 500,00 ₽" in r.text
    assert "• Логистика и доставка: 2 000,00 ₽" in r.text
    assert "• Реклама и продвижение: 1 500,00 ₽" in r.text
    assert "🏠 Личные расходы: 8 000,00 ₽" in r.text
    assert "должен" not in r.text and "вернул" not in r.text
    wb = load_workbook(io.BytesIO(r.file[1]))
    assert wb.sheetnames == ["По месяцам", "По статьям", "Бизнес-расходы"]
    assert wb["По месяцам"].max_row == 3 + 9 * 3 + 3
    cats = wb["По статьям"]
    assert cats.cell(cats.max_row, 1).value == "Итого бизнес"
    assert cats.cell(cats.max_row, 11).value == 3500

    [feb] = flow.on_command(CHAT, "itog", "2026-02")
    assert "💼 Ушло на бизнес: 2 000,00 ₽" in feb.text and "🏠 Личные расходы: 8 000,00 ₽" in feb.text


def test_parse_period_arg():
    assert _parse_period_arg("2026", TODAY) == ("2026-01", "2026-09")
    assert _parse_period_arg("2025", TODAY) == ("2025-01", "2025-12")
    assert _parse_period_arg("2027", TODAY) is None
    assert _parse_period_arg("2026-03..2026-01", TODAY) is None
    assert _parse_period_arg("2026-01 .. 2026-03", TODAY) == ("2026-01", "2026-03")
    assert _parse_period_arg("март", TODAY) is None


def test_fix_command(env):
    db, rec, flow = env
    rec.payments.append(payment())
    flow.on_files(CHAT, [png()], "")
    eid = db.expenses("2026-09")[0].id
    assert flow.on_command(CHAT, "fix", str(eid))[0].text.startswith("✅ Записано")
    assert "/fix 42" in flow.on_command(CHAT, "fix", "999")[0].text


def test_old_database_gets_new_column(tmp_path):
    path = str(tmp_path / "old.db")
    conn = sqlite3.connect(path)
    conn.executescript("""
        CREATE TABLE statement_lines (id INTEGER PRIMARY KEY, statement_id INTEGER NOT NULL,
            op_date TEXT NOT NULL, amount INTEGER NOT NULL, direction TEXT NOT NULL,
            description TEXT NOT NULL DEFAULT '', own_transfer INTEGER NOT NULL DEFAULT 0);
        INSERT INTO statement_lines VALUES (1, 1, '2026-01-01', 100, 'out', 'x', 0);
    """)
    conn.commit()
    conn.close()
    db = Storage(path)
    row = db.conn.execute("SELECT suggested_category FROM statement_lines").fetchone()
    assert row[0] == ""


def test_identical_purchases_kept_but_resent_file_not_doubled(env):
    db, rec, flow = env
    start_period(flow, db)
    coffee = [op("2026-02-02", "300", desc="Кофейня"), op("2026-02-02", "300", desc="Кофейня")]
    rec.statements += [statement(coffee), statement(coffee), statement(coffee + [coffee[0]])]
    [first] = flow.on_files(CHAT, [(b"%PDF part 1", "application/pdf")], "")
    assert "Принято операций: 2" in first.text
    # Другой файл (скриншот той же страницы) с теми же строками — строки не задвоятся.
    [again] = flow.on_files(CHAT, [(b"%PDF screenshot", "application/pdf")], "")
    assert "Принято операций: 0 (повторы пропущены: 2)" in again.text
    # в новом файле их уже три — добавится одна
    flow.on_files(CHAT, [(b"%PDF part 2", "application/pdf")], "")
    assert len(db.statement(db.cards()[0].id, "2026-02")[1]) == 3
