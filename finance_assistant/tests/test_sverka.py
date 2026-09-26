from openpyxl import load_workbook
import io

from conftest import CHAT, PNG, payment
from finance.reconcile import match, summarize
from finance.storage import Expense, StatementLine

PDF = (b"%PDF-1.4", "application/pdf")


def op(d, amount, direction="out", desc="", own=False):
    return {"date": d, "amount": amount, "direction": direction, "description": desc,
            "own_transfer": own}


def statement(ops, total_in="", total_out="", last4="1111"):
    return {"is_statement": True, "card_last4": last4, "total_in": total_in,
            "total_out": total_out, "operations": ops}


def start_sverka(flow, db, card_index=0, month="2026-09"):
    [ask_month] = flow.on_command(CHAT, "sverka")
    assert f"s:m:{month}" in str(ask_month.buttons)
    [ask_card] = flow.on_button(CHAT, f"s:m:{month}")
    card = db.cards()[card_index]
    return flow.on_button(CHAT, f"s:c:{card.id}")


def record(db, rec, flow, **kw):
    rec.payments.append(payment(**kw))
    flow.on_files(CHAT, [PNG], "")


def test_full_month(env):
    db, rec, flow = env
    record(db, rec, flow, amount="1500", date="2026-09-20")                        # найдётся
    record(db, rec, flow, amount="3000", date="2026-09-10", merchant="VK Реклама",
           category="Реклама и продвижение")                                        # найдётся (+2 дня)
    record(db, rec, flow, amount="999", date="2026-09-11", merchant="Где-то")      # не найдётся
    rec.payments.append(payment(direction="in", amount="2000", date="2026-09-25"))
    [skip] = flow.on_files(CHAT, [PNG], "")
    assert "поступление — не записываю" in skip.text

    start_sverka(flow, db)
    rec.statements.append(statement([
        op("2026-09-01", "100000", "in", "Зарплата"),
        op("2026-09-20", "1500", desc="CDEK"),
        op("2026-09-12", "3000", desc="VK ADS"),
        op("2026-09-15", "5000", desc="Перевод на карту *2222", own=True),
        op("2026-09-18", "4200", desc="Пятёрочка"),
        op("2026-09-25", "2000", "in", "Перевод от ООО"),
        op("2026-08-31", "777", desc="прошлый месяц"),
    ]))
    [got] = flow.on_files(CHAT, [PDF], "")
    assert "Принято операций: 6" in got.text and "Пропущено" in got.text
    rec.statements.append(statement([op("2026-09-18", "4200", desc="Пятёрочка")]))
    [again] = flow.on_files(CHAT, [PNG], "")
    assert "повторы пропущены: 1" in again.text

    card_reply, unmatched, nxt = flow.on_command(CHAT, "done")
    assert "Пришло: 102 000,00 ₽" in card_reply.text
    assert "Ушло:   13 700,00 ₽" in card_reply.text
    assert "на бизнес: 5 499,00 ₽" in card_reply.text
    assert "на личное: 3 201,00 ₽" in card_reply.text  # 13700 - 5000 своих - 5499 бизнес
    assert "не найдено в выписке" in card_reply.text and "999" in card_reply.text
    assert "Пятёрочка" in unmatched.text and "CDEK" not in unmatched.text
    assert "Тинькофф" in nxt.text

    [itog] = flow.on_command(CHAT, "itog", "2026-09")
    assert "💼 Ушло на бизнес: 5 499,00 ₽" in itog.text
    assert "• Реклама и продвижение: 3 000,00 ₽" in itog.text
    assert "• Логистика и доставка: 2 499,00 ₽" in itog.text
    # Личное: Сбер 3 201 по выписке; у Тинькофф и Альфа выписок нет.
    assert "🏠 Личные расходы: 3 201,00 ₽" in itog.text
    assert "без выписки, личное не посчитано: Тинькофф ·2222, Альфа ·3333" in itog.text
    assert "должен" not in itog.text
    name, data = itog.file
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Свод", "Бизнес-расходы", "Не найдено в выписке"]
    assert wb["Бизнес-расходы"].max_row == 4


def test_biz_from_statement_line(env):
    db, rec, flow = env
    start_sverka(flow, db)
    rec.statements.append(statement([op("2026-09-18", "4200", desc="OZON упаковка")]))
    flow.on_files(CHAT, [PDF], "")
    flow.on_command(CHAT, "done")
    line_id = db.statement(db.cards()[0].id, "2026-09")[1][0].id
    [question] = flow.on_command(CHAT, "biz", str(line_id))
    assert "статья" in question.text.lower()
    flow.on_button(CHAT, f"d:cat:{db.category_id('Упаковка и расходники')}")
    s = summarize(db, "2026-09")
    cs = s.cards[0]
    assert cs.business == 420000 and cs.personal == 0 and not cs.missing and not cs.unmatched_out


def test_totals_only_text(env):
    db, rec, flow = env
    start_sverka(flow, db, card_index=1)
    rec.statements.append(statement([], total_in="120000", total_out="95000.50", last4=""))
    [r] = flow.on_text(CHAT, "пришло 120000 ушло 95000,50")
    assert "Итоги из документа" in r.text
    card_reply, *_ = flow.on_command(CHAT, "done")
    assert "Ушло:   95 000,50 ₽" in card_reply.text and "только итоги" in card_reply.text


def test_printed_totals_win_over_line_sum(env):
    db, rec, flow = env
    start_sverka(flow, db)
    rec.statements.append(statement([op("2026-09-02", "100")], total_out="250"))
    flow.on_files(CHAT, [PDF], "")
    assert summarize(db, "2026-09").cards[0].total_out == 25000


def test_wrong_card_warning_and_not_statement(env):
    db, rec, flow = env
    start_sverka(flow, db)
    rec.statements.append(statement([op("2026-09-02", "100")], last4="9999"))
    [r] = flow.on_files(CHAT, [PDF], "")
    assert "…9999" in r.text
    rec.statements.append({**statement([]), "is_statement": False})
    [r] = flow.on_files(CHAT, [PNG], "")
    assert "не похоже на выписку" in r.text


def test_reset_statement(env):
    db, rec, flow = env
    start_sverka(flow, db)
    rec.statements.append(statement([op("2026-09-02", "100")]))
    flow.on_files(CHAT, [PDF], "")
    flow.on_command(CHAT, "done")
    [r] = start_sverka(flow, db)
    assert "Уже загружено строк: 1" in r.text
    flow.on_button(CHAT, "s:reset")
    assert db.statement(db.cards()[0].id, "2026-09") is None


def test_itog_without_statements_and_default_month(env):
    db, rec, flow = env
    record(db, rec, flow)
    [r] = flow.on_command(CHAT, "itog")
    assert "сентябрь 2026" in r.text and "(нет выписки)" in r.text
    assert "💼 Ушло на бизнес: 1 500,00 ₽" in r.text and "🏠 Личные расходы: 0,00 ₽" in r.text


def _e(i, d, amount, kind="expense"):
    return Expense(i, d, amount, kind, "business", 1, "Сбер", None, "", "", "")


def _l(i, d, amount, direction="out"):
    return StatementLine(i, d, amount, direction, "", False)


def test_match_rules():
    exps = [_e(1, "2026-09-10", 100), _e(2, "2026-09-10", 100), _e(3, "2026-09-10", 100),
            _e(4, "2026-09-10", 500)]
    lines = [_l(10, "2026-09-14", 100), _l(11, "2026-09-11", 100), _l(12, "2026-09-09", 100),
             _l(13, "2026-09-10", 500, "in")]
    pairs, missing = match(exps, lines)
    assert pairs == {1: 11, 2: 12}  # 14-е — дальше 3 дней; зачисление расходу не пара
    assert [e.id for e in missing] == [3, 4]


def test_bad_command_args(env):
    db, rec, flow = env
    assert "2026-09" in flow.on_command(CHAT, "itog", "сентябрь")[0].text
    assert "/biz 12" in flow.on_command(CHAT, "biz", "")[0].text
    assert "Не нашёл" in flow.on_command(CHAT, "biz", "999")[0].text
    assert "не идёт" in flow.on_command(CHAT, "done")[0].text
