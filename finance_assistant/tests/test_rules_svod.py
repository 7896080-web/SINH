"""Привязка получателей к статьям и сводная за период."""
import io
from datetime import date

from openpyxl import load_workbook

from conftest import CHAT, TODAY, payment, png
from finance.pivot import buckets, parse_period
from test_business_account import biz, card, kontur  # noqa: F401 — фикстура biz
from test_sverka import statement


def cat(db, name):
    return db.category_id(name)


def test_chosen_category_remembered_for_merchant(env):
    db, rec, flow = env
    yandex = dict(merchant="Сервисы Яндекса (АО Яндекс Банк)", category="Реклама и продвижение",
                  category_confident=False)
    rec.payments += [payment(amount="122000", **yandex),
                     payment(amount="5000", date="2026-09-22",
                             **{**yandex, "merchant": "АО «Сервисы Яндекса Яндекс Банк»",
                                "category": "Прочее"})]
    [q] = flow.on_files(CHAT, [png()], "")
    assert "статья" in q.text.lower()
    flow.on_button(CHAT, f"d:cat:{cat(db, 'Маркетплейсы: комиссии и услуги')}")
    # Второй такой же получатель (другое написание) — без вопроса, по правилу,
    # хотя модель предложила другую статью.
    [saved] = flow.on_files(CHAT, [png()], "")
    assert saved.text.startswith("✅") and "Маркетплейсы" in saved.text
    assert "📌 Статья по правилу" in saved.text
    [r] = flow.on_command(CHAT, "rules")
    assert "Маркетплейсы: комиссии и услуги:" in r.text and "Сервисы Яндекса" in r.text


def test_changing_category_updates_rule_and_offers_past_records(env):
    db, rec, flow = env
    for amount in ("1000", "2000", "3000"):
        rec.payments.append(payment(amount=amount, merchant="VK Реклама",
                                    category="Прочее", category_confident=True))
        flow.on_files(CHAT, [png()], "")
    first = db.expenses("2026-09")[0]
    replies = flow.on_button(CHAT, f"e:cat:{first.id}:{cat(db, 'Реклама и продвижение')}")
    assert "Запомнил: «VK Реклама» → Реклама и продвижение" in replies[1].text
    assert "другая статья: 2 на 5 000,00 ₽" in replies[1].text
    [data] = [d for row in replies[1].buttons for _, d in row]
    [done] = flow.on_button(CHAT, data)
    assert "изменена у записей «VK Реклама»: 2" in done.text
    assert {e.category for e in db.expenses("2026-09")} == {"Реклама и продвижение"}
    # Новые записи этого получателя — сразу в новую статью.
    rec.payments.append(payment(amount="4000", merchant="vk реклама", category="Прочее"))
    [saved] = flow.on_files(CHAT, [png()], "")
    assert "Реклама и продвижение" in saved.text


def test_rules_not_learned_from_bulk_or_generic_names(env):
    db, rec, flow = env
    rec.payments.append(payment(merchant="ИП", category="Прочее", category_confident=False))
    flow.on_files(CHAT, [png()], "")
    flow.on_button(CHAT, f"d:cat:{cat(db, 'Прочее')}")
    assert db.rules() == []  # «ИП» без имени — слишком общее
    [r] = flow.on_command(CHAT, "rules")
    assert "Правил пока нет" in r.text


def test_delete_rule(env):
    db, rec, flow = env
    db.set_rule("СДЭК", cat(db, "Логистика и доставка"))
    rule_id = db.rules()[0]["id"]
    assert flow.on_command(CHAT, "delrule", str(rule_id))[0].text == "Правило удалено."
    assert "/delrule 3" in flow.on_command(CHAT, "delrule", "99")[0].text


def test_rule_applies_to_statement_lines(env):
    db, rec, flow = env
    db.set_rule("CDEK Москва", cat(db, "Логистика и доставка"))
    flow.on_command(CHAT, "sverka", "2026-09")
    flow.on_button(CHAT, f"s:c:{db.cards()[0].id}")
    rec.statements.append(statement([
        {"date": "2026-09-02", "time": "", "amount": "700", "direction": "out",
         "description": "CDEK МОСКВА", "own_transfer": False, "business_category": ""}]))
    flow.on_files(CHAT, [png()], "")
    replies = flow.on_command(CHAT, "done")
    [s] = [r for r in replies if r.text.startswith("🔎")]
    assert "CDEK МОСКВА → Логистика и доставка" in s.text


def test_parse_period_and_columns():
    today = TODAY  # суббота 26.09.2026
    week = parse_period("week", today)
    assert (week.start, week.end, week.columns) == (date(2026, 9, 21), date(2026, 9, 27), "day")
    prev = parse_period("prevmonth", today)
    assert (prev.start, prev.end, prev.columns) == (date(2026, 8, 1), date(2026, 8, 31), "week")
    assert parse_period("2026", today).columns == "month"
    custom = parse_period("01.09.2026-15.09.2026", today)
    assert (custom.start, custom.end) == (date(2026, 9, 1), date(2026, 9, 15))
    assert [b[0] for b in buckets(custom)] == ["01.09–06.09", "07.09–13.09", "14.09–15.09"]
    assert parse_period("15.09.2026-01.09.2026", today) is None
    assert parse_period("вчера", today) is None
    assert parse_period("2026-02", today).title == "февраль 2026"


def test_svod_buttons_and_errors(env):
    db, rec, flow = env
    [r] = flow.on_command(CHAT, "svod")
    assert "v:week" in str(r.buttons) and "v:year" in str(r.buttons)
    assert "Не понял период" in flow.on_command(CHAT, "svod", "когда-то")[0].text
    [empty] = flow.on_button(CHAT, "v:prevweek")
    assert "Бизнес-расходов за этот период нет" in empty.text


def test_svod_month_layout_and_excel(biz):
    db, rec, flow = biz
    rows = [("2026-09-02", "1000", "СДЭК", "Логистика и доставка"),
            ("2026-09-09", "2500", "СДЭК", "Логистика и доставка"),
            ("2026-09-10", "700", "Деловые Линии", "Логистика и доставка"),
            ("2026-09-16", "12000", "Яндекс Директ", "Реклама и продвижение"),
            ("2026-08-30", "9999", "СДЭК", "Логистика и доставка")]  # другой месяц
    for d, amount, merchant, category in rows:
        rec.payments.append(payment(card_last4="5501", bank="ВТБ", date=d, amount=amount,
                                    merchant=merchant, category=category))
        flow.on_files(CHAT, [png()], "")
    rec.payments.append(payment(card_last4="5501", bank="ВТБ", date="2026-09-03", amount="800",
                                looks_personal=True))
    flow.on_files(CHAT, [png()], "")
    flow.on_button(CHAT, "d:purpose:personal")  # личное — в сводную не попадает
    # Бизнес-счёт: списание по выписке без записи → «Без статьи».
    flow.on_command(CHAT, "sverka", "2026-09")
    flow.on_button(CHAT, f"s:c:{card(db, 'ПСБ').id}")
    rec.statements.append(statement([
        {"date": "2026-09-20", "time": "", "amount": "3000", "direction": "out",
         "description": "Петров П. П.", "own_transfer": False, "business_category": ""}],
        last4=""))
    flow.on_files(CHAT, [png()], "")
    flow.on_command(CHAT, "done")

    [r] = flow.on_command(CHAT, "svod", "2026-09")
    assert "📊 Бизнес-расходы: сентябрь 2026" in r.text
    assert "💼 Всего: 19 200,00 ₽ (5 операций)" in r.text
    assert "▪️ Реклама и продвижение: 12 000,00 ₽ · 62%" in r.text
    assert "▪️ Логистика и доставка: 4 200,00 ₽ · 22%" in r.text
    assert "    СДЭК: 3 500,00 ₽" in r.text and "    Деловые Линии: 700,00 ₽" in r.text
    assert "▪️ Без статьи: 3 000,00 ₽" in r.text
    assert "По неделям:" in r.text and "31.08" not in r.text

    wb = load_workbook(io.BytesIO(r.file[1]))
    assert wb.sheetnames == ["Сводная", "Раскладка", "Данные", "Фильтр по датам"]
    pv = wb["Сводная"]
    header = [c.value for c in pv[3]]
    assert header[0] == "Статья" and header[-2:] == ["Итого", "Доля"] and len(header) == 3 + 5
    assert pv.cell(4, 1).value == "Реклама и продвижение" and pv.cell(4, 7).value == 12000
    assert pv.cell(pv.max_row, 1).value == "Итого" and pv.cell(pv.max_row, 7).value == 19200
    lay = [r[0].value for r in wb["Раскладка"].iter_rows(min_row=2)]
    assert lay[:2] == ["Реклама и продвижение", "    Яндекс Директ"]
    data = wb["Данные"]
    assert data.max_row == 6 and "Operations" in data.tables
    flt = wb["Фильтр по датам"]
    assert flt["B1"].value.date() == date(2026, 9, 1)
    assert flt["B6"].value.startswith("=SUMIFS(Данные!$G$2:$G$6,")

    [year] = flow.on_command(CHAT, "svod", "2026")
    assert "💼 Всего: 29 199,00 ₽ (6 операций)" in year.text and "По месяцам:" in year.text
    assert "авг 26: 9 999,00 ₽" in year.text
