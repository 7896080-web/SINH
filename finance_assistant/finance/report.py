"""Тексты для чата и Excel-файлы свода.

Итог всегда в двух суммах: «Ушло на бизнес» (по статьям) и «Личные расходы».
"""

import io
from datetime import date

from .money import format_amount
from .reconcile import CardSummary, MonthSummary
from .storage import BUSINESS, BUSINESS_ACCOUNT, REIMBURSEMENT, Expense, StatementLine

MONTHS = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август",
          "сентябрь", "октябрь", "ноябрь", "декабрь"]


def rub(kopecks: int | None) -> str:
    return "—" if kopecks is None else f"{format_amount(kopecks)} ₽"


def month_name(month: str) -> str:
    year, mon = month.split("-")
    return f"{MONTHS[int(mon) - 1]} {year}"


def _short(month: str) -> str:
    year, mon = month.split("-")
    return f"{MONTHS[int(mon) - 1][:3]} {year[2:]}"


def short_date(iso: str) -> str:
    d = date.fromisoformat(iso)
    return f"{d.day:02d}.{d.month:02d}"


def _what(e: Expense) -> str:
    if e.kind == REIMBURSEMENT:  # запись прежней версии — в итоги не входит
        return "перевод между своими счетами"
    if e.purpose == BUSINESS:
        return f"бизнес · {e.category or 'без статьи'}"
    return "личное (с бизнес-счёта)" if e.card_kind == BUSINESS_ACCOUNT else "личное"


def expense_line(e: Expense) -> str:
    who = f" · {e.merchant}" if e.merchant else ""
    return f"{short_date(e.op_date)}  {rub(e.amount)}  {e.card}  {_what(e)}{who}"


def expense_card(e: Expense) -> str:
    what = _what(e)
    lines = [f"✅ Записано №{e.id}", f"{rub(e.amount)} · {short_date(e.op_date)} · {e.card}",
             what[0].upper() + what[1:]]
    if e.merchant or e.description:
        lines.append(" — ".join(x for x in (e.merchant, e.description) if x))
    return "\n".join(lines)


def _categories(out: list[str], by_category: dict[str, int], indent: str = "  "):
    out += [f"{indent}• {name}: {rub(amount)}" for name, amount in by_category.items()]


def card_text(cs: CardSummary, month: str) -> str:
    icon = "🏦" if cs.card.is_business else "💳"
    out = [f"{icon} {cs.card.label} — {month_name(month)}"]
    if cs.has_statement:
        out.append(f"Пришло: {rub(cs.total_in)}")
        out.append(f"Ушло:   {rub(cs.total_out)}")
        if cs.own_out:
            label = "переведено вам" if cs.card.is_business else "переводы на свои счета"
            out.append(f"  {label}: {rub(cs.own_out)}")
    else:
        out.append("Выписка не загружена — есть только записанное вами:")
    out.append(f"  на бизнес: {rub(cs.business)}")
    _categories(out, cs.categories, "    ")
    if cs.card.is_business:
        if cs.personal_marked:
            out.append(f"  на личное: {rub(cs.personal_marked)}")
    else:
        out.append(f"  на личное: {rub(cs.personal)}")
    if cs.has_statement and not cs.lines_checked:
        out.append("(по строкам не сверялось — в выписке были только итоги)")
    if cs.missing:
        out.append("")
        out.append("⚠️ Записано, но не найдено в выписке (проверьте карту или дату):")
        out += [f"  №{e.id} {expense_line(e)}" for e in cs.missing]
    return "\n".join(out)


def _line(ln: StatementLine) -> str:
    # Время важно: в выписке банка «Россия» у всех оплат одно описание
    # («Оплата по QR-коду через СБП»), различить их можно только так.
    when = short_date(ln.op_date) + (f" {ln.op_time}" if ln.op_time else "")
    return f"#{ln.id}  {when}  {rub(ln.amount)}  {ln.description}"


def unmatched_text(lines: list[StatementLine]) -> str:
    out = ["Списания из выписки, которые сейчас считаются личными.",
           "Если среди них есть бизнес — отправьте /biz и номера, например: /biz 12 15", ""]
    return "\n".join(out + [_line(ln) for ln in lines])


def suggestions_text(lines: list[StatementLine], title: str = "🔎 Похоже на бизнес-расходы") -> str:
    out = [f"{title}: {len(lines)} на {rub(sum(ln.amount for ln in lines))}",
           "Проверьте список. Лишние уберите: /notbiz 12 15, недостающие добавьте: /biz 7",
           ""]
    out += [f"{_line(ln)} → {ln.suggested_category or 'Прочее'}" for ln in lines]
    return "\n".join(out)


def _totals(out: list[str], business: int, by_category: dict[str, int], personal: int,
            incomplete: list[str]):
    out.append(f"💼 Ушло на бизнес: {rub(business)}")
    _categories(out, by_category)
    out.append(f"🏠 Личные расходы: {rub(personal)}")
    if incomplete:
        out.append(f"  (без выписки, личное не посчитано: {', '.join(incomplete)})")


def month_text(s: MonthSummary) -> str:
    out = [f"📊 Свод за {month_name(s.month)}", ""]
    for cs in s.cards:
        icon = "🏦" if cs.card.is_business else "💳"
        state = "" if cs.has_statement else "  (нет выписки)"
        out.append(f"{icon} {cs.card.label}{state}")
        if cs.has_statement:
            out.append(f"  пришло {rub(cs.total_in)} · ушло {rub(cs.total_out)}")
        out.append(f"  бизнес {rub(cs.business)} · личное {rub(cs.personal)}")
        if cs.missing:
            out.append(f"  ⚠️ не найдено в выписке: {len(cs.missing)}")
    out.append("")
    _totals(out, s.business, s.by_category, s.personal,
            [c.label for c in s.personal_incomplete])
    return "\n".join(out)


def _period_name(summaries: list[MonthSummary]) -> str:
    return f"{month_name(summaries[0].month)} — {month_name(summaries[-1].month)}"


def _merge_categories(summaries: list[MonthSummary]) -> dict[str, int]:
    merged: dict[str, int] = {}
    for s in summaries:
        for name, amount in s.by_category.items():
            merged[name] = merged.get(name, 0) + amount
    return dict(sorted(merged.items(), key=lambda kv: -kv[1]))


def period_text(summaries: list[MonthSummary]) -> str:
    out = [f"📊 Свод: {_period_name(summaries)}", "", "По месяцам (бизнес · личное):"]
    for s in summaries:
        gap = ""
        if s.personal_incomplete:
            gap = ("  (выписок нет)" if len(s.personal_incomplete) == len(s.personal_cards)
                   else f"  (нет выписки: {', '.join(c.name for c in s.personal_incomplete)})")
        out.append(f"  {_short(s.month)}: {rub(s.business)} · {rub(s.personal)}{gap}")
    out.append("")
    out.append("По картам за период:")
    for i, cs0 in enumerate(summaries[0].cards):
        per = [s.cards[i] for s in summaries]
        icon = "🏦" if cs0.card.is_business else "💳"
        with_stmt = [c for c in per if c.has_statement]
        line = (f"  {icon} {cs0.card.label}: бизнес {rub(sum(c.business for c in per))}"
                f" · личное {rub(sum(c.personal or 0 for c in per))}")
        if with_stmt:
            line += (f" · пришло {rub(sum(c.total_in or 0 for c in with_stmt))}"
                     f" · ушло {rub(sum(c.total_out or 0 for c in with_stmt))}")
        missing = sum(len(c.missing) for c in per)
        if missing:
            line += f" · ⚠️ не найдено в выписках: {missing}"
        out.append(line)
    out.append("")
    incomplete = sorted({c.label for s in summaries for c in s.personal_incomplete})
    _totals(out, sum(s.business for s in summaries), _merge_categories(summaries),
            sum(s.personal for s in summaries), incomplete)
    return "\n".join(out)


def _num(kopecks: int | None):
    return None if kopecks is None else kopecks / 100


def _bold_row(ws, row: int):
    from openpyxl.styles import Font
    for cell in ws[row]:
        cell.font = Font(bold=True)


def _finish(wb) -> bytes:
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "# ##0.00"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _expenses_sheet(wb, expenses: list[Expense]):
    ops = wb.create_sheet("Бизнес-расходы")
    ops.append(["№", "Дата", "Сумма", "Карта", "Статья", "Получатель", "Что оплачено", "Чек"])
    _bold_row(ops, 1)
    for e in expenses:
        ops.append([e.id, e.op_date, _num(e.amount), e.card, e.category or "", e.merchant,
                    e.description, e.receipt_path])
    for col, width in zip("ABCDEFGH", (6, 12, 14, 18, 30, 26, 40, 30)):
        ops.column_dimensions[col].width = width


def _missing_sheet(wb, missing: list[Expense]):
    if missing:
        m = wb.create_sheet("Не найдено в выписке")
        m.append(["№", "Дата", "Сумма", "Карта", "Получатель"])
        for e in missing:
            m.append([e.id, e.op_date, _num(e.amount), e.card, e.merchant])


CARD_HEADER = ["Карта", "Пришло", "Ушло", "Переводы на свои счета", "На бизнес", "На личное",
               "Не найдено в выписке"]


def _card_row(cs: CardSummary) -> list:
    return [cs.card.label, _num(cs.total_in), _num(cs.total_out), _num(cs.own_out),
            _num(cs.business), _num(cs.personal), len(cs.missing)]


def month_xlsx(s: MonthSummary) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Свод"
    ws.append([f"Свод за {month_name(s.month)}"])
    _bold_row(ws, 1)
    ws.append([])
    ws.append(CARD_HEADER)
    _bold_row(ws, 3)
    for cs in s.cards:
        ws.append(_card_row(cs))
    ws.append([])
    ws.append(["Ушло на бизнес", _num(s.business)])
    _bold_row(ws, ws.max_row)
    for name, amount in s.by_category.items():
        ws.append([f"  {name}", _num(amount)])
    ws.append(["Личные расходы", _num(s.personal)])
    _bold_row(ws, ws.max_row)
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 16
    _expenses_sheet(wb, s.business_expenses)
    _missing_sheet(wb, [e for cs in s.cards for e in cs.missing])
    return _finish(wb)


def period_xlsx(summaries: list[MonthSummary]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "По месяцам"
    ws.append([f"Свод: {_period_name(summaries)}"])
    _bold_row(ws, 1)
    ws.append([])
    ws.append(["Месяц"] + CARD_HEADER)
    _bold_row(ws, 3)
    for s in summaries:
        for cs in s.cards:
            ws.append([s.month] + _card_row(cs))
    ws.append([])
    ws.append(["Ушло на бизнес", "", "", "", "", _num(sum(s.business for s in summaries))])
    ws.append(["Личные расходы", "", "", "", "", "", _num(sum(s.personal for s in summaries))])
    for col, width in zip("ABCDEFGH", (14, 30, 14, 14, 16, 14, 14, 12)):
        ws.column_dimensions[col].width = width

    cats = wb.create_sheet("По статьям")
    cats.append(["Статья"] + [s.month for s in summaries] + ["Итого"])
    _bold_row(cats, 1)
    for name, total in _merge_categories(summaries).items():
        cats.append([name] + [_num(s.by_category.get(name, 0)) for s in summaries] + [_num(total)])
    cats.append(["Итого бизнес"] + [_num(s.business) for s in summaries]
                + [_num(sum(s.business for s in summaries))])
    _bold_row(cats, cats.max_row)
    cats.column_dimensions["A"].width = 34

    _expenses_sheet(wb, [e for s in summaries for e in s.business_expenses])
    _missing_sheet(wb, [e for s in summaries for cs in s.cards for e in cs.missing])
    return _finish(wb)
