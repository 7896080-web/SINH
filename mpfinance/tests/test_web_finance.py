"""Программа через HTTP: страницы, загрузка отчёта, отправка в 1С,
себестоимость, настройки, кабинеты и ключи, диагностика."""
import io
from datetime import date

from openpyxl import load_workbook

from finapp.finance import values as v
from finapp.models import (ApiCredential, AuditLog, Barcode, OnecTask, OneCUploadStatus,
                           PlatformAccount, ProductCost, SalesReport, SalesReportLine)
from tests.finance_factories import make_product, make_report, make_xlsx, wb_file, wb_row

PAGES = ["/finance", "/finance/reports", "/finance/unit-economics", "/finance/costs", "/api-keys",
         "/diagnostics"]


def _upload(client, data, platform="wb", **form):
    return client.post("/finance/reports/upload", data={"platform": platform, **form},
                       files={"file": ("report.xlsx", data)}, follow_redirects=False)


def test_pages_require_login(anon_client):
    for url in PAGES:
        response = anon_client.get(url, follow_redirects=False)
        assert response.status_code == 303 and response.headers["location"] == "/login"


def test_empty_pages_render(client):
    for url in PAGES:
        assert client.get(url).status_code == 200, url


def test_upload_redirects_to_report_and_pages_show_it(client, db):
    make_product(db)
    response = _upload(client, wb_file([wb_row()]))
    assert response.status_code == 303
    report = db.query(SalesReport).one()
    assert response.headers["location"] == f"/finance/reports/{report.id}"

    page = client.get(f"/finance/reports/{report.id}").text
    assert "Сверка денег" in page and "Загрузка продаж в 1С" in page
    assert "Wildberries_Склад_FBO" in page

    period = "date_from=2026-09-01&date_to=2026-09-30"
    dashboard = client.get(f"/finance?{period}").text
    assert "Wildberries" in dashboard and "revenue-chart" in dashboard
    ue = client.get(f"/finance/unit-economics?{period}").text
    assert "A-1" in ue
    assert client.get(f"/finance/unit-economics?{period}&group_by=sku&allocate=true").status_code == 200
    assert "без себестоимости" in ue


def test_bad_upload_shows_reason_and_saves_nothing(client, db):
    response = _upload(client, make_xlsx(["Что-то"], [["1"]]))
    assert response.headers["location"] == "/finance/reports"
    page = client.get("/finance/reports").text
    assert "Отчёт не загружен" in page and "«Что-то»" in page
    assert db.query(SalesReport).count() == 0


def test_bad_period_in_form_is_refused(client, db):
    _upload(client, wb_file([wb_row()]), period_from="вчера", period_to="2026-09-07")
    assert db.query(SalesReport).count() == 0


def test_payouts_form_and_money_check(client, db):
    _upload(client, wb_file([wb_row()]))
    report = db.query(SalesReport).one()
    client.post(f"/finance/reports/{report.id}/payouts",
                          data={"declared": "2 100,00", "bank": "2000", "note": "удержали долг"})
    db.refresh(report)
    assert (report.declared_payout_kop, report.bank_payout_kop) == (210000, 200000)
    page = client.get(f"/finance/reports/{report.id}").text
    assert "разница" in page and "удержали долг" in page

    client.post(f"/finance/reports/{report.id}/payouts", data={"declared": "abc"})
    db.refresh(report)
    assert report.declared_payout_kop == 210000       # мусор не затирает введённое


def test_send_to_1c_then_delete_is_blocked_until_answer(client, db, epf_ready):
    make_product(db)
    _upload(client, wb_file([wb_row()]))
    report = db.query(SalesReport).one()

    client.post(f"/finance/reports/{report.id}/send-1c")
    db.refresh(report)
    assert report.onec_status == OneCUploadStatus.pending
    assert db.query(AuditLog).filter(AuditLog.action == "sales_report_onec_requested").count() == 1

    client.post(f"/finance/reports/{report.id}/delete")
    assert db.query(SalesReport).count() == 1     # в пути — удалять нельзя

    report.onec_status = OneCUploadStatus.done
    db.commit()
    client.post(f"/finance/reports/{report.id}/delete")
    assert db.query(SalesReport).count() == 0
    assert db.query(SalesReportLine).count() == 0


def test_send_with_unmatched_goods_shows_reason(client, db):
    _upload(client, wb_file([wb_row(**{"Баркод": "999"})]))
    report = db.query(SalesReport).one()
    client.post(f"/finance/reports/{report.id}/send-1c")
    db.refresh(report)
    assert report.onec_status == OneCUploadStatus.none
    assert "без товара 1С" in client.get(f"/finance/reports/{report.id}").text


def test_remap_after_mapping(client, db):
    _upload(client, wb_file([wb_row()]))
    report = db.query(SalesReport).one()
    assert report.unmatched_lines == 1
    make_product(db)
    client.post(f"/finance/reports/{report.id}/remap")
    db.refresh(report)
    assert report.unmatched_lines == 0


def test_unified_export_roundtrips_through_import(client, db):
    """Отчёт в единой форме можно загрузить обратно как отчёт любой площадки —
    те же суммы до копейки."""
    _upload(client, wb_file([wb_row(), wb_row(**{
        "Тип документа": "Возврат", "Обоснование для оплаты": "Возврат", "Srid": "r"})]))
    original = db.query(SalesReport).one()
    exported = client.get(f"/finance/reports/{original.id}/export").content

    ws = load_workbook(io.BytesIO(exported)).active
    headers = [c.value for c in ws[1]]
    assert headers[5:5 + len(v.UNIFIED_COLUMNS)] == [c.header for c in v.UNIFIED_COLUMNS]

    _upload(client, exported, platform="ozon")
    copy = db.query(SalesReport).filter(SalesReport.platform == "ozon").one()
    assert copy.source_format == "unified"
    totals = lambda r: sorted((l.operation, l.revenue_kop, l.payout_kop) for l in r.lines)
    assert totals(copy) == totals(original)


def test_exports_are_xlsx(client, db):
    make_product(db)
    make_report(db, lines=[{"operation": v.SALE, "uid_1c": "u1", "article": "A-1",
                                "operation_date": date(2026, 9, 2), "revenue": 100000,
                                "payout": 80000, "commission": 20000}])
    for url in ("/finance/unit-economics/export?date_from=2026-09-01&date_to=2026-09-30",
                "/finance/costs/export", "/finance/reports/template"):
        response = client.get(url)
        assert response.status_code == 200
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml"), url
    rows = list(load_workbook(io.BytesIO(client.get(
        "/finance/unit-economics/export?date_from=2026-09-01&date_to=2026-09-30").content)).active.values)
    assert rows[-1][1] == "ИТОГО"


def test_costs_import_by_uid_barcode_and_article(client, db):
    make_product(db, uid="u1", article="A-1", barcodes=("111",))
    make_product(db, uid="u2", article="A-1", size="48", barcodes=("222",))
    make_product(db, uid="u3", article="B-2", barcodes=("333",))
    data = make_xlsx(["ID_1С", "Артикул", "Баркод", "Себестоимость", "Действует с"], [
        ["u3", "", "", "1 200,50", "01.09.2026"],
        ["", "", "222", 900, None],
        ["", "A-1", "", 800, None],           # всем размерам артикула (u1 и u2, у u2 — новая цена)
        ["", "", "", 100, None],              # не найти товар — ошибка, не падение
        ["u1", "", "", "", None],             # пустая себестоимость — пропуск
    ])
    response = client.post("/finance/costs/import",
                                     files={"file": ("costs.xlsx", data)}, follow_redirects=False)
    assert response.status_code == 303
    costs = {(c.uid_1c, c.valid_from): c.cost_minor for c in db.query(ProductCost)}
    assert costs[("u3", date(2026, 9, 1))] == 120050
    assert costs[("u1", date(2000, 1, 1))] == 80000
    assert costs[("u2", date(2000, 1, 1))] == 80000
    page = client.get("/finance/costs").text
    assert "Ошибок: 1" in page


def test_settings_validation_and_save(client, db):
    response = client.post("/finance/settings", data={
        "FIN_TAX_MODE": "revenue", "FIN_TAX_RATE": "200", "FIN_EXTRA_COST_PER_UNIT": "0",
        "FIN_WAREHOUSE_WB": "W", "FIN_WAREHOUSE_OZON": "O", "FIN_WAREHOUSE_KIT": "K",
        "FIN_WAREHOUSE_LAMODA": ""}, follow_redirects=True)
    assert "не сохранены" in response.text

    client.post("/finance/settings", data={
        "FIN_TAX_MODE": "profit", "FIN_TAX_RATE": "15", "FIN_EXTRA_COST_PER_UNIT": "35,5",
        "FIN_USD_RATE_MANUAL": "92,5", "FIN_KIT_CRM_COMMISSION": "4,2", "FIN_KIT_SALES_SOURCE": "api",
        "FIN_WAREHOUSE_WB": "Wildberries_Склад_FBO", "FIN_WAREHOUSE_OZON": "OZON_Склад",
        "FIN_WAREHOUSE_KIT": "Сайт AWER", "FIN_WAREHOUSE_LAMODA": "Lamaoda_Склад"})
    from finapp.finance.settings import load_settings
    settings = load_settings(db)
    assert settings.tax_mode == "profit" and str(settings.tax_rate) == "15"
    assert settings.extra_cost_kop == 3550
    assert settings.warehouses["lamoda"] == "Lamaoda_Склад"
    assert str(settings.usd_rate_manual) == "92.5"
    assert str(settings.kit_crm_commission) == "4.2" and settings.kit_sales_source == "api"
    for bad in ({"FIN_KIT_CRM_COMMISSION": "150"}, {"FIN_KIT_SALES_SOURCE": "почта"}):
        response = client.post("/finance/settings", data={"FIN_TAX_MODE": "revenue", "FIN_TAX_RATE": "6",
                                                          **bad}, follow_redirects=True)
        assert "не сохранены" in response.text


def test_missing_report_redirects(client):
    response = client.get("/finance/reports/999", follow_redirects=False)
    assert response.status_code == 303


def test_accounts_page_creates_account_and_keeps_key_secret(client, db, sync_db):
    sync_db.account(7, "wb", "ИП Яворская")
    client.post("/api-keys/new", data={"platform": "wb", "name": "ИП Яворская"})
    account = db.query(PlatformAccount).one()
    client.post(f"/api-keys/{account.id}", data={"name": "ИП Яворская", "is_active": "1",
                                                  "finance_token": "secret-token-123456",
                                                  "sync_account_id": "7"})
    db.refresh(account)
    assert account.sync_account_id == 7 and account.is_active
    row = db.query(ApiCredential).one()
    assert row.field_name == "finance_token" and "secret-token" not in row.encrypted_value
    page = client.get("/api-keys").text
    assert "secr" in page and "secret-token-123456" not in page   # только маска
    # Пустое поле при сохранении ключ не стирает.
    client.post(f"/api-keys/{account.id}", data={"name": "ИП Яворская", "is_active": "1"})
    assert db.query(ApiCredential).count() == 1


def test_password_autofilled_into_key_field_is_not_saved(client, db):
    client.post("/api-keys/new", data={"platform": "kit", "name": "Сайт"})
    account = db.query(PlatformAccount).one()
    client.post(f"/api-keys/{account.id}", data={"name": "Сайт", "token": "password1"})
    assert db.query(ApiCredential).count() == 0


def test_diagnostics_ping_and_manual_dictionary_upload(client, db):
    # Обмен с 1С по умолчанию выключен: проверка версии не ставится.
    client.post("/diagnostics/request/FIN_PING")
    assert db.query(OnecTask).count() == 0
    client.post("/diagnostics/onec-enabled", data={"enabled": "1"})
    client.post("/diagnostics/request/FIN_PING")
    assert db.query(OnecTask).one().command == "FIN_PING"
    # До ответа mark-4 остальные команды не ставятся.
    client.post("/diagnostics/request/BARCODE_DICT")
    assert db.query(OnecTask).count() == 1
    client.post("/diagnostics/upload", data={"kind": "dict"},
                files={"file": ("barcodes.txt", "u1|A-1|Джинсы|111|46|синий\n".encode("cp1251"))})
    assert db.query(Barcode).count() == 1
    assert "Справочник загружен" in client.get("/diagnostics").text


def test_costs_request_waits_for_mark4(client, db):
    response = client.post("/finance/costs/request-1c", follow_redirects=True)
    assert "mark-4" in response.text
    assert db.query(OnecTask).count() == 0


def test_health_without_login(anon_client):
    response = anon_client.get("/health")
    assert response.status_code == 503                  # фоновые задания в тестах не запущены
    assert "onec_exchange" in " ".join(response.json()["stale"])
