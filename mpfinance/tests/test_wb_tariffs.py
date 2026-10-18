"""Текущая комиссия WB по категории (`finance/tariffs.py`) — по API, как в «Репрайсере».

Тариф берётся из `GET common-api…/api/v1/tariffs/commission`: FBS —
`kgvpMarketplace`, FBO — `paidStorageKgvp`. Предмет карточки — из детализации
отчёта WB («Предмет» в файле, `subjectName` в API). В юнит-экономике у позиции
WB — предмет, тариф FBS сейчас (и прежний, если WB его сменил) и сколько WB
удержал на деле."""
from datetime import date, timedelta

import pytest

from finapp import settings
from finapp.finance import fetch, tariffs
from finapp.finance.formats import parse_records
from finapp.finance.importer import import_report
from finapp.finance.tariffs import fetch_wb        # настоящий: автофикстура подменяет его в модуле
from finapp.models import CategoryTariff, WbCard
from finapp.timeutils import now_utc
from tests.finance_factories import make_account, make_product, make_xlsx, wb_row, WB_HEADERS
from tests.test_finance_api_fetch import FakeResponse, FakeSession, _cred, _wb_line

ANSWER = {"report": [
    {"kgvpBooking": 23.5, "kgvpMarketplace": 25.5, "kgvpPickup": 23.5, "kgvpSupplier": 22,
     "kgvpSupplierExpress": 3, "paidStorageKgvp": 24.5, "parentID": 1, "parentName": "Одежда",
     "subjectID": 11, "subjectName": "Джинсы"},
    {"kgvpMarketplace": 27, "paidStorageKgvp": 26, "parentName": "Одежда", "subjectID": 12,
     "subjectName": "Свитшоты"},
    {"kgvpMarketplace": 150, "paidStorageKgvp": "abc", "subjectID": 13, "subjectName": "Брак"},
    {"subjectName": "без id"},
]}


def test_answer_is_read_like_the_repricer_reads_it():
    rows = {r["category_id"]: r for r in tariffs.parse_wb_tariffs(ANSWER)}
    assert set(rows) == {"11", "12", "13"}
    assert (rows["11"]["fbs_bp"], rows["11"]["fbo_bp"], rows["11"]["category_name"]) == (2550, 2450, "Джинсы")
    assert (rows["13"]["fbs_bp"], rows["13"]["fbo_bp"]) == (None, None)        # вне 0–100 — не тариф
    with pytest.raises(tariffs.TariffError, match="разметки"):
        tariffs.parse_wb_tariffs({"data": []})


def test_request_goes_to_common_api_with_the_token_as_is():
    session = FakeSession({tariffs.WB_TARIFFS_PATH: ANSWER})
    assert len(fetch_wb("fin", session)) == 3
    assert session.headers["Authorization"] == "fin"
    assert session.calls == [("GET", tariffs.WB_TARIFFS_PATH, {"locale": "ru"})]
    refusing = FakeSession({"*": FakeResponse(401, {"title": "unauthorized"})})
    with pytest.raises(tariffs.TariffError, match="401 — токен не подошёл"):
        fetch_wb("bad", refusing)


def _wb_with_token(db):
    account = make_account(db)
    _cred(db, account, "finance_token", "fin")
    return account


def _serve(monkeypatch, answer):
    monkeypatch.setattr(tariffs, "fetch_wb", lambda token, session=None: tariffs.parse_wb_tariffs(answer))


def test_refresh_stores_tariffs_and_remembers_a_change(db, monkeypatch):
    assert "нет активного кабинета WB" in tariffs.refresh(db)
    assert settings.get(db, settings.TARIFFS_OK_AT) == ""
    _wb_with_token(db)
    _serve(monkeypatch, ANSWER)
    assert tariffs.refresh(db).startswith("предметов 3, новых 3")
    raised = {"report": [dict(ANSWER["report"][0], kgvpMarketplace=26.5)]}
    _serve(monkeypatch, raised)
    assert "сменился тариф FBS: 1" in tariffs.refresh(db)
    jeans = db.get(CategoryTariff, ("wb", "11"))
    assert (jeans.fbs_bp, jeans.prev_fbs_bp) == (2650, 2550) and jeans.changed_at is not None
    assert db.get(CategoryTariff, ("wb", "12")).fbs_bp == 2700       # пропавший из ответа не удалён
    assert "В базе предметов: 3" in tariffs.status_text(db)


def test_daily_job_asks_once_a_day_and_a_failure_does_not_count(db, monkeypatch):
    from finapp.workers import jobs
    _wb_with_token(db)
    jobs.job_daily()                                               # сеть «выключена» автофикстурой
    assert settings.get(db, settings.TARIFFS_OK_AT) == ""
    assert "сеть в тестах выключена" in settings.get(db, settings.TARIFFS_RESULT)
    _serve(monkeypatch, ANSWER)
    jobs.job_daily()
    db.expire_all()
    assert db.query(CategoryTariff).count() == 3
    asked = []
    monkeypatch.setattr(tariffs, "fetch_wb", lambda token, session=None: asked.append(token) or [])
    assert tariffs.refresh_if_due(db) is None and asked == []      # сутки не прошли
    settings.put(db, settings.TARIFFS_OK_AT, (now_utc() - timedelta(hours=21)).isoformat(timespec="seconds"))
    db.commit()
    tariffs.refresh_if_due(db)
    assert asked == ["fin"]


def test_report_file_remembers_the_subject_of_each_wb_card(db):
    data = make_xlsx(WB_HEADERS, [wb_row(), wb_row(**{"Код номенклатуры": 777, "Предмет": "Свитшоты",
                                                      "Название": "Свитшот базовый", "Srid": "srid-2"})])
    outcome = import_report(db, platform="wb", file_bytes=data, filename="wb.xlsx", user="t",
                            report_number="900", period_from=date(2026, 9, 1), period_to=date(2026, 9, 7))
    assert outcome.errors == []
    assert {c.nm_id: c.subject_name for c in db.query(WbCard)} == {"123456": "Джинсы", "777": "Свитшоты"}
    assert sorted(l.name for l in outcome.report.lines) == ["Джинсы прямые", "Свитшот базовый"]


def test_api_row_gives_subject_and_keeps_title_as_name():
    line = parse_records([_wb_line(subjectName="Джинсы")], "wb").lines[0]
    assert (line["subject"], line["name"], line["platform_sku"]) == ("Джинсы", "Джинсы", "123")
    line = parse_records([_wb_line(subjectName="Джинсы", title="Джинсы прямые")], "wb").lines[0]
    assert (line["subject"], line["name"]) == ("Джинсы", "Джинсы прямые")


def _sold(db, account):
    make_product(db)
    data = make_xlsx(WB_HEADERS, [wb_row(), wb_row(**{"Srid": "srid-2"})])
    import_report(db, platform="wb", file_bytes=data, filename="wb.xlsx", user="t", account_id=account.id,
                  report_number="901", period_from=date(2026, 9, 1), period_to=date(2026, 9, 7))
    db.commit()


def test_unit_economics_shows_current_fbs_next_to_what_wb_kept(client, db, monkeypatch):
    account = _wb_with_token(db)
    _sold(db, account)
    _serve(monkeypatch, ANSWER)
    tariffs.refresh(db)
    _serve(monkeypatch, {"report": [dict(ANSWER["report"][0], kgvpMarketplace=26.5)]})
    tariffs.refresh(db)
    page = client.get("/finance/unit-economics?date_from=2026-09-01&date_to=2026-09-30").text
    assert "FBS сейчас <b>26,5 %</b> (было 25,5 %)" in page
    assert "WB удержал 16 %" in page                    # (2500 − 2100) / 2500
    assert "Джинсы" in page

    from io import BytesIO
    from openpyxl import load_workbook
    raw = client.get("/finance/unit-economics/export?date_from=2026-09-01&date_to=2026-09-30").content
    rows = list(load_workbook(BytesIO(raw), read_only=True).active.iter_rows(values_only=True))
    head = list(rows[0])
    first = dict(zip(head, rows[1]))
    assert first["Предмет WB"] == "Джинсы" and first["Комиссия WB FBS сейчас, %"] == 26.5
    assert first["Комиссия WB FBO сейчас, %"] == 24.5 and first["Удержано WB с продаж, %"] == 16
    assert first["Комиссия WB FBS до смены, %"] == 25.5


def test_two_subjects_with_one_name_and_different_rates_are_not_guessed(db, monkeypatch):
    account = _wb_with_token(db)
    _sold(db, account)
    _serve(monkeypatch, {"report": [dict(ANSWER["report"][0]),
                                    dict(ANSWER["report"][0], subjectID=99, kgvpMarketplace=10)]})
    tariffs.refresh(db)
    book = tariffs.TariffBook(db, ["123456"])
    assert book.subject_for(["123456"]) == "Джинсы" and book.for_cards(["123456"]) is None


def test_diagnostics_has_status_and_button(client, db, monkeypatch):
    page = client.get("/diagnostics").text
    assert "Тарифы WB (комиссия по категориям)" in page and "ещё не загружались" in page
    _wb_with_token(db)
    _serve(monkeypatch, ANSWER)
    response = client.post("/diagnostics/wb-tariffs", follow_redirects=True)
    assert "Тарифы WB: предметов 3" in response.text


def test_migration_creates_tariff_tables(tmp_path):
    import sqlite3
    from tests.test_cost_file import _alembic
    url = f"sqlite:///{tmp_path / 'm.db'}"
    _alembic(url, "upgrade", "head")
    conn = sqlite3.connect(tmp_path / "m.db")
    conn.execute("SELECT platform, category_id, fbs_bp, fbo_bp, prev_fbs_bp FROM category_tariffs").fetchall()
    conn.execute("SELECT nm_id, subject_name FROM wb_cards").fetchall()
    conn.close()
    _alembic(url, "downgrade", "0007")
    _alembic(url, "upgrade", "head")
