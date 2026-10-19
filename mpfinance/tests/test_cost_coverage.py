"""Себестоимость считается по ПРОДАННЫМ товарам, а не по всему справочнику 1С.

Справочник приходит из Sync Admin целиком — 152 тысячи позиций, почти все на
площадках не продаются. «Задана у 68 042 из 152 415» ничего не говорило о
расчёте, а список «без себестоимости» состоял из ненужных строк."""
import io
from datetime import date

import openpyxl

from finapp.finance import values as v
from finapp.finance.importer import resolve_uids
from finapp.models import ProductCost
from tests.finance_factories import make_product, make_report


def _catalog(db):
    make_product(db, uid="sold-cost", article="27643", barcodes=("2000000000011",))
    make_product(db, uid="sold-none", article="36446", barcodes=("2000000000028",))
    make_product(db, uid="unsold", article="1532", barcodes=("2000000000035",))
    db.add(ProductCost(uid_1c="sold-cost", valid_from=date(2000, 1, 1), cost_minor=150000, currency="RUB",
                       source="excel", updated_by="admin"))
    db.commit()
    report = make_report(db, key="SR-cov", lines=[
        {"operation": v.SALE, "quantity": 1, "revenue": 460000, "payout": 400000, "barcode": "2000000000011"},
        {"operation": v.SALE, "quantity": 3, "revenue": 900000, "payout": 800000, "barcode": "2000000000028"},
        # Еженедельный отчёт Lamoda: баркода нет, товар 1С не определить.
        {"operation": v.SALE, "quantity": 2, "revenue": 500000, "payout": 300000, "barcode": None},
    ])
    resolve_uids(db, report)
    db.commit()
    return report


def test_costs_page_counts_sold_products_and_lists_only_sold_without_cost(client, db):
    _catalog(db)
    page = client.get("/finance/costs").text
    assert "1 из 2</b> проданных товаров" in page and "В справочнике 1С всего 3 товаров" in page
    missing = client.get("/finance/costs?missing=true").text
    assert "36446" in missing and "1532" not in missing.split("<tbody>")[1]


def test_export_follows_the_filter(client, db):
    _catalog(db)
    data = client.get("/finance/costs/export?missing=true").content
    rows = list(openpyxl.load_workbook(io.BytesIO(data)).active.iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == ["sold-none"]
    everything = client.get("/finance/costs/export").content
    assert len(list(openpyxl.load_workbook(io.BytesIO(everything)).active.iter_rows(min_row=2))) == 3


def test_dashboard_says_where_pieces_without_cost_come_from(client, db):
    report = _catalog(db)
    page = client.get(f"/finance?date_from={report.period_from}&date_to={report.period_to}").text
    assert "Продано 5 шт без себестоимости" in page
    assert "Из них 2 шт — товар 1С не определён" in page and "Ещё 3 шт у товара не задана себестоимость" in page
    assert "Из них 1 — без баркода в самом отчёте" in page
    assert "без себестоимости: 3 из 3 шт" in page


def test_size_u_is_not_listed_or_counted(client, db):
    """Размера «U» на площадках нет (владелец, 19.10.2026) — ни в списке, ни в счётчиках."""
    _catalog(db)
    make_product(db, uid="u-size", article="99999 U-ONLY", size="U", barcodes=("2000000000042",))
    make_product(db, uid="u-sold", article="99998 U-SOLD", size=" u ", barcodes=("2000000000059",))
    report = make_report(db, key="SR-u", lines=[
        {"operation": v.SALE, "quantity": 1, "revenue": 300000, "payout": 250000, "barcode": "2000000000059"}])
    resolve_uids(db, report)
    db.commit()
    from finapp.finance.costs import sold_without_cost
    assert sold_without_cost(db) == 1                                    # только «36446» 46, не U
    page = client.get("/finance/costs").text.split("<tbody>")[1]
    assert "99999 U-ONLY" not in page and "99998 U-SOLD" not in page and "27643" in page
    assert "1 из 2</b> проданных товаров" in client.get("/finance/costs").text
    assert "В справочнике 1С всего 3 товаров" in client.get("/finance/costs").text


def test_new_sold_without_cost_asks_1c_again_before_a_day_passes(db, epf_ready):
    from datetime import timedelta

    from finapp import settings
    from finapp.models import OnecTask
    from finapp.timeutils import now_utc
    from finapp.workers import jobs
    _catalog(db)                                                          # «36446» продан без цены
    settings.put(db, settings.CATALOG_LOADED_AT, now_utc().isoformat(timespec="seconds"))

    def asked(loaded_hours_ago, missing_after_load):
        db.query(OnecTask).delete()
        settings.put(db, settings.COSTS_LOADED_AT,
                     (now_utc() - timedelta(hours=loaded_hours_ago)).isoformat(timespec="seconds"))
        settings.put(db, settings.COSTS_MISSING_AT_LOAD, str(missing_after_load))
        db.commit()
        jobs.job_daily()
        db.expire_all()
        return [t.command for t in db.query(OnecTask)] == ["EXPORT_COST_PRICES"]

    assert asked(4, 0)                    # новый без цены, выгрузке больше трёх часов — спросить
    assert not asked(1, 0)                # выгрузка свежая — подождать
    assert not asked(4, 1)                # тот же товар был без цены и после выгрузки — 1С его не знает
    assert asked(25, 1)                   # раз в сутки — как всегда


def test_cost_load_remembers_what_stayed_without_cost(db):
    from finapp import settings
    from finapp.finance import costs
    _catalog(db)
    costs.import_costs(db, {"sold-cost": 1900})
    assert settings.get(db, settings.COSTS_MISSING_AT_LOAD) == "1"
